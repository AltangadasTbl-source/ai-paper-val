#!/usr/bin/env python3
"""Construction only. No review-time Python; 2025 only; never overwrite outputs."""
import argparse, hashlib, json, os, tempfile
from pathlib import Path
from copy import copy
from urllib.parse import unquote
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as letter

HERE=Path(__file__).resolve().parent
(HERE/".tmp").mkdir(exist_ok=True)
tempfile.tempdir=str(HERE/".tmp")
ROOT=HERE.parents[2]
BASE=ROOT/'human_adjudication/previous/human_adjudication_v_2_1_5_3'
ANNUAL=BASE/'human_adjudication_v_2_1_5_3_2025.xlsx'
SOURCES=['Candidates','Evidence','Links','Relations']
OPTS={'D':['retain','exclude','discuss'], 'E':['statistical_results','report_internal_inconsistency','plan_material','protocol_report'], 'F':['no','primary_done'], 'H':['confirmed_issue','probable_issue','not_issue','indeterminate'], 'I':['substantive','minor_presentation','indeterminate'], 'J':['direct','possible','unlikely','unknown'], 'K':['substantive_inconsistency','minor_presentation','non_inconsistency','indeterminate'], 'L':['substantive_difference','minor_presentation','version_timing','non_inconsistency','indeterminate']}
NAVY='173F57'; YELLOW='FFF2CC'; BLUE='EAF3FC'; GRAY='EDF2F5'
def q(s): return "'"+s.replace("'","''")+"'"
def lit(s): return '"'+str(s).replace('"','""')+'"'
def valid(cell,values): return 'OR('+','.join(f'{cell}={lit(v)}' for v in values)+')'
def active(col,d,e,potential=False):
    branch=f'{e}="statistical_results"' if col in 'HIJ' else (f'OR({e}="report_internal_inconsistency",{e}="plan_material")' if col=='K' else f'{e}="protocol_report"')
    disp=f'OR({d}="retain",{d}="")' if potential else f'{d}="retain"'
    return f'AND({disp},{branch})'
def preserve(expr): return f'IF({expr}="","",{expr})'
def lookup(sheet,col,key,n,keycol='A'):
    return f'INDEX({q(sheet)}!${col}$2:${col}${n},MATCH({key},{q(sheet)}!${keycol}$2:${keycol}${n},0))'
def local_dir(value):
    p=(HERE/value).resolve()
    if not p.is_relative_to(HERE) or p==HERE or p.is_relative_to(HERE/'release'): raise ValueError('Use a construction subdirectory inside 2025, outside release.')
    return p

def table(w,name,rows,widths=None,height=48):
    if name in w: del w[name]
    s=w.create_sheet(name)
    for row in rows: s.append(row)
    for row in s:
        for c in row:
            c.font=Font(name='Calibri',size=11,color=NAVY)
            c.alignment=Alignment(wrap_text=True,vertical='center')
            c.protection=Protection(locked=False,hidden=False)
            if c.row==1: c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
            elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F4F8FA')
    for r in range(1,s.max_row+1): s.row_dimensions[r].height=42 if r==1 else height
    for i in range(1,s.max_column+1):s.column_dimensions[letter(i)].width=(widths or {}).get(i,28)
    s.freeze_panes='B2';s.sheet_view.showGridLines=False
    if name not in ['Start','Setup','Checks','Manifest','Rule_defaults']:s.auto_filter.ref=s.dimensions
    return s

def hint(s,r,sheets):
    a,d,e,f=[f'{c}{r}' for c in 'ADEF']
    count='+'.join(f'COUNTIF({q(t.title)}!$A$2:$A${t.max_row},{a})' for t in sheets)
    key=f'OR({a}="",({count})<>1,COUNTIFS(Candidates!$G$2:$G$132,{a},Candidates!$K$2:$K$132,{lit(s.title[2:])})<>1)'
    # CLEAN plus NBSP substitution catches whitespace-only pasted identities as well as ordinary spaces.
    ident='LEN(TRIM(CLEAN(SUBSTITUTE(Setup!$B$2,CHAR(160)," "))))=0'
    stale=[];errors=[f'NOT({valid(d,OPTS["D"])})',f'NOT({valid(e,OPTS["E"])})']
    for col in 'HIJKL':
        cell=f'{col}{r}';act=active(col,d,e)
        stale.append(f'AND(NOT({act}),{cell}<>"",{cell}<>"not_applicable")')
        opt=f'{cell}<>"",' if col in 'JL' else ''
        errors.append(f'AND({act},{opt}NOT({valid(cell,OPTS[col])}))')
    return f'=IFERROR(IF({ident},"CHECK REVIEWER ID",IF({key},"CHECK KEY",IF(OR({",".join(stale)}),"CLEAR INACTIVE VALUES",IF(OR({",".join(errors)}),"CHECK FIELDS",IF(AND({f}<>"",NOT({valid(f,OPTS["F"])})),"CHECK COMPLETE",IF({f}="primary_done","PRIMARY READY","PRIMARY DRAFT")))))),"CHECK INPUT / KEY")'

START=[['Step','Instructions'],['Release','v.2.3.1 | 2025 | 15 papers | 131 candidates | Independent primary review'],
['Coordinator','Assign two different people to the two role files. Distribute only the corresponding file to each person. Roles are labels, not person IDs. This standalone file cannot verify the identity in another file.'],
['1. Begin','Save As a working copy in Excel. Enter your own reviewer ID (or identifying name) in Setup B2 and optional full name in B3. Setup B4 identifies this file\'s role; B5 is the cohort year.'],
['2. Read','Use Papers links to open each C_ sheet. Candidate-key links open embedded Evidence. Filter Candidates, Evidence, Links and Relations by key for full context. Click source labels in Links to open original PDFs. Keep the source-folder structure when using relative links.'],
['3. Assess','Choose D=retain, exclude, or discuss. Confirm or change the AI-suggested domain in E. Answer applicable H-L, using Codebook definitions. R holds optional English notes. Yellow = common/required; blue = optional; gray = inactive. M-Q are empty spacers.'],
['4. Applicability','For retain: statistical_results requires H and I; J is optional. report_internal_inconsistency and plan_material require K. protocol_report has no extra required judgment; L is optional. Exclude/discuss have no active H-L judgments.'],
['5. Defaults','Untouched H-L formulas display blank for a potentially applicable field while D is blank or retain; otherwise not_applicable. Choosing an answer replaces its formula. Blank means unanswered, with no inherited answer. Optional active fields may stay blank. Active fields cannot use not_applicable.'],
['6. Changes','Before changing a completed row, manually set F=no, make changes, review again, then mark primary_done. Excel formulas cannot reset an entered flag. Domain/disposition changes do not erase literal answers. CLEAR INACTIVE VALUES: delete stale inactive answers or set them to not_applicable; answer newly active required fields.'],
['7. Finish each row','After manual evidence review, set F=primary_done. G must show PRIMARY READY. Invalid pasted values are checked as well as dropdown choices. D and a valid E are required even for exclude/discuss. Completion alone does not establish scientific accuracy.'],
['8. Inspect and deliver','Inspect G and Checks; filter Decisions status for unfinished rows. When Checks shows PRIMARY READY, save the completed workbook and deliver this file. Review, calculations and delivery use Excel alone. Keep normal Excel backup copies; never regenerate a completed review file.'],
['Uncertainty','Discuss is a completed independent response when otherwise valid. It records unresolved scientific uncertainty and initiates no shared workflow. Checks displays discuss totals separately; completing data entry does not resolve uncertainty.'],
['Results','Decisions contains only your assessment. Unfinished and discuss outcomes stay blank. P_R/P_P retained records are descriptive and never temporal. Paper_counts shows independent completion separately and withholds outcome counts until every row is ready and none is discuss.'],
['Editing and recovery','All cells, including unused cells, are unlocked; there are no protected ranges, passwords or unlock steps. Edit Setup B2:B3 and C_ D/E/F/H-L/R. Preserve keys, source text and formulas; recover accidental changes with Undo or a normal backup.'],
['Navigation and sorting','Sort or filter complete records only, including all columns A:R on C_ sheets. Do not sort a single input column or edit candidate keys. Results match by candidate key, including the one-candidate paper. Use embedded source filters if a source sheet has been sorted after following a link.'],
['Hints','CHECK REVIEWER ID: enter an identity. CHECK KEY: restore the unique baseline key in its paper. CHECK FIELDS: check D/E and active H-L options. CHECK COMPLETE: use no or primary_done. PRIMARY DRAFT: finish manual review. CHECK INPUT / KEY: inspect pasted errors and references.'],
['Validation status','Construction static and formula-engine checks are documented in the accompanying VALIDATION.md. Interactive desktop Excel smoke testing and visual rendering are pending. Excel recalculates on opening.']]

USAGE='''# 2025 independent primary review — v.2.3.1

Assign two different people and distribute only the corresponding R_A or R_B file. Each file has the same 15 papers / 131 candidates and starts with no human answers.

1. Open your file in Excel and Save As a working copy. Enter your ID/name in Setup B2 (optional full name B3).
2. Use Papers to open each C_ sheet, and candidate-key links, Candidates, Evidence and Links to inspect sources.
3. Choose D, confirm the editable AI suggestion E, answer applicable H–L, and add optional English notes in R. Read Codebook for option meanings.
4. After manual review set F=primary_done; inspect G and Checks. Before changing a completed row, reset F=no, review again, then mark primary_done. Flags do not reset automatically. Clear stale inactive answers after domain/disposition changes.
5. When Checks says PRIMARY READY, save and deliver this workbook. Keep normal Excel backups. All calculations and review actions use Excel alone.

Discuss can complete independent entry but remains unresolved. Discuss outcomes are blank; Paper_counts withholds counts for any paper with unfinished or discuss rows. P_R/P_P retained records stay descriptive.

H–L start with applicability formulas. Selecting an answer replaces that formula; blank is unanswered. Active optional fields may be blank; not_applicable is only for inactive fields. No cells or sheets are protected, including unused cells. Preserve source keys and calculation cells; use Undo/backups after accidental edits.

Source text is embedded. Original-document links are relative to the repository source folders; retain that layout for clickable PDFs. Formula calculations have no external dependencies. The workbook includes all review instructions.

Desktop Excel smoke testing and visual inspection are pending; see ../VALIDATION.md for actual checks and exact manual tests. Construction scripts are pre-release records only. Do not rebuild or overwrite completed review files. Future corrections require a separately named new version and review of changes.
'''

def build(role,out):
    w=load_workbook(ANNUAL)
    candidates=[dict(zip(next(w['Candidates'].values),r)) for r in list(w['Candidates'].values)[1:]]
    sheets=[s for s in w if s.title.startswith('C_')]
    assert len(sheets)==15 and len(candidates)==131
    for n in ['Start','Setup','Checks','Primary_snapshot','Decisions','Paper_counts','Protocol_plan_table','Codebook','Manifest']:del w[n]
    w.defined_names.clear()
    start=table(w,'Start',START,{1:29,2:115},88)
    start['B2']=START[1][1]+f' | {role}'
    setup=table(w,'Setup',[['Field','Value'],['reviewerId / identifying name',None],['reviewerName (optional)',None],['reviewerRole',role],['year',2025]],{1:35,2:70},42)
    for row in [2,3]:setup.cell(row,2).fill=PatternFill('solid',fgColor=YELLOW)
    checks=table(w,'Checks',[['Check','Result'],['Reviewer role','=Setup!B4'],['Candidates',131],['Primary ready rows','=COUNTIF(Decisions!G2:G132,"PRIMARY READY")'],['Unfinished rows','=B3-B4'],['Discuss / unresolved rows','=COUNTIF(Decisions!K2:K132,"discuss")'],['Workflow status','=IF(B4=B3,"PRIMARY READY","ACTION NEEDED")'],['Interpretation','Completion means independent entry is finished. Discuss remains unresolved; affected paper outcome counts are withheld.'],['Next action','Filter Decisions status; return to C_ inputs. When primary-ready, save and deliver this file.'],['Ready discuss rows','=COUNTIFS(Decisions!G2:G132,"PRIMARY READY",Decisions!K2:K132,"discuss")']],{1:35,2:105},60)
    for s in sheets:
        s.tables.clear() # Empty spacer headers must remain genuinely empty; use a whole-record AutoFilter.
        s.auto_filter.ref=f'A1:R{s.max_row}'
        s.data_validations=DataValidationList();s.conditional_formatting=ConditionalFormattingList()
        s['E1']='Review domain (AI suggestion; editable)'
        for col in 'MNOPQ':
            s.column_dimensions[col].width=3;s.column_dimensions[col].hidden=True
            for r in range(1,s.max_row+1):
                c=s[f'{col}{r}'];c.value=None;c.comment=None;c.hyperlink=None;c.fill=PatternFill('solid',fgColor=GRAY)
        for col,width in {'E':35,'H':29,'K':36,'L':28}.items():s.column_dimensions[col].width=width
        for col,options in OPTS.items():
            dv=DataValidation(type='list',formula1=lit(','.join(options+(['not_applicable'] if col in 'HIJKL' else []))),allow_blank=True)
            dv.errorTitle='Use an allowed option';dv.error='Choose an option from the list; see Codebook.';dv.showErrorMessage=True;dv.errorStyle='stop';dv.showDropDown=False
            s.add_data_validation(dv);dv.add(f'{col}2:{col}{s.max_row}')
        for r in range(2,s.max_row+1):
            s[f'D{r}']=None;s[f'F{r}']='no';s[f'R{r}']=None
            s[f'G{r}']=hint(s,r,sheets)
            for col in 'HIJKL':s[f'{col}{r}']=f'=IF({active(col,f"D{r}",f"E{r}",True)},"","not_applicable")'
        s.conditional_formatting.add(f'G2:G{s.max_row}',FormulaRule(formula=['G2<>"PRIMARY READY"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color='9C0006')))
        for col in 'HIJKL':s.conditional_formatting.add(f'{col}2:{col}{s.max_row}',FormulaRule(formula=[f'NOT({active(col,"$D2","$E2",True)})'],fill=PatternFill('solid',fgColor=GRAY),font=Font(color='687783')))
    headers=['year','paper','key','relationship','reviewerRole','reviewerId','status','outcome','aiDisposition','aiDomain','disposition','domain','statisticalIssueStatus','substantiveStatus','downstreamImpact','internalStatus','protocolInterpretation','notes','reviewerName']
    dec=table(w,'Decisions',[headers]+[[None]*len(headers) for c in candidates],{1:12,2:27,3:44,7:28,8:28,18:55},72)
    for r,c in enumerate(candidates,2):
        s=w['C_'+c['package']]; src=lambda col:lookup(s.title,col,f'C{r}',s.max_row)
        for j,val in enumerate([2025,c['package'],c['key'],c['primary'],'=Setup!$B$4','='+preserve('Setup!$B$2'),f'=IFERROR({src("G")},"CHECK KEY")'],1):dec.cell(r,j,val)
        dec[f'H{r}']=f'=IF(OR(G{r}<>"PRIMARY READY",K{r}="discuss"),"",IF(K{r}<>"retain","not_eligible",IF(D{r}="P_R","protocol_descriptive",IF(D{r}="P_P","plan_descriptive",IF(AND(OR(D{r}="S_C",D{r}="S_W"),L{r}="statistical_results",OR(M{r}="confirmed_issue",M{r}="probable_issue"),N{r}="substantive"),"primary_temporal",IF(AND(OR(D{r}="S_C",D{r}="S_W"),L{r}="report_internal_inconsistency",P{r}="substantive_inconsistency"),"secondary_temporal","not_eligible"))))))'
        dec[f'I{r}']=c['ai_disposition'];dec[f'J{r}']=c['ai_domain']
        for j,col in enumerate(['D','E','H','I','J','K','L','R'],11):
            expr=preserve(src(col))
            if col in 'HIJKL':expr=f'IF({active(col,f"K{r}",f"L{r}")},{expr},"not_applicable")'
            dec.cell(r,j,'=IFERROR('+expr+',"")')
        dec[f'S{r}']='='+preserve('Setup!$B$3')
    outcomes=['primary_temporal','secondary_temporal','protocol_descriptive','plan_descriptive','not_eligible']
    pc=table(w,'Paper_counts',[['year','paper','candidates','ready','unfinished','discuss_unresolved','independent_review_complete','outcome_counts_available']+outcomes]+[[2025,s.title[2:],s.max_row-1]+[None]*10 for s in sheets],{1:12,2:28,7:29,8:29},55)
    for r in range(2,pc.max_row+1):
        match=f'Decisions!$B$2:$B$132,B{r}'
        for col,formula in {'D':f'COUNTIFS({match},Decisions!$G$2:$G$132,"PRIMARY READY")','E':f'C{r}-D{r}','F':f'COUNTIFS({match},Decisions!$K$2:$K$132,"discuss")','G':f'IF(E{r}=0,"yes","no")','H':f'IF(AND(E{r}=0,F{r}=0),"yes","no")'}.items():pc[f'{col}{r}']='='+formula
        for j,o in enumerate(outcomes,9):pc.cell(r,j,f'=IF(H{r}<>"yes","",COUNTIFS({match},Decisions!$H$2:$H$132,"{o}"))')
    plans=[c for c in candidates if c['primary'] in ['P_R','P_P']]
    pp=table(w,'Protocol_plan_table',[['year','paper','key','relationship','title','statement','status','disposition','domain','interpretation','internalStatus','notes','outcome','reviewerRole','reviewerId']]+[[2025,c['package'],c['key'],c['primary'],c['title'],c['statement']]+[None]*9 for c in plans],{1:12,2:27,3:44,5:60,6:90,12:55},105)
    for r in range(2,pp.max_row+1):
        for j,col in enumerate(['G','K','L','Q','P','R','H','E','F'],7):pp.cell(r,j,'=IFERROR('+preserve(lookup('Decisions',col,f'C{r}',132,'C'))+',"")')
    # Keep methodological field/option definitions, replacing only obsolete workflow text.
    old=list(load_workbook(ANNUAL)['Codebook'].values);rows=[];skip=False
    for row in old:
        row=list(row)
        if row[0] and str(row[0]).startswith('5  |'):skip=True;continue
        if row[0] and str(row[0]).startswith('6  |'):skip=False;row[0]='5  |  Candidate context'
        if skip or row[4]=='primarySummary' or row[1]=='checker_done':continue
        if row[0]=='Working method':row[2]='Independent primary review';row[3]='Blank = no answer; indeterminate/unknown = uncertain; not_applicable = inactive. Clear obsolete answers after changing domain or disposition.'
        if row[4]=='disposition':row[3]='Keep, exclude, or defer as unresolved. Blank means untouched, not discuss. No AI disposition is preselected.'
        if row[4]=='reviewComplete':
            row[2]='Independent primary review'
            row[3]='Use no while drafting; primary_done only after manual review. Reset F=no before changing a completed row, review again, then mark primary_done.'
        if row[4]=='readiness':row[3]='PRIMARY READY means this independent primary assessment passes entry checks. Discuss remains unresolved. Checks does not establish scientific accuracy or prevent edits to source/formula cells.'
        if row[4]=='locations':row[2]='Source reference';row[3]='Where to inspect the claim. Use Candidates, Evidence and Links for embedded context and original documents.'
        if row[1]=='no':row[3]='Draft. Manual review is not yet confirmed complete; AI suggestions or partial entries do not count as completion.'
        rows.append(row)
    table(w,'Codebook',rows,{1:34,2:29,3:40,4:96,5:29},96)
    table(w,'Rule_defaults',[['Domain (retain only)','Required','Optional','Inactive default'],['statistical_results','H and I','J','not_applicable'],['report_internal_inconsistency','K','None','not_applicable'],['plan_material','K','None','not_applicable'],['protocol_report','None','L','not_applicable'],['exclude / discuss','D and E only','R notes','H-L: not_applicable']],{1:36,2:25,3:25,4:28},50)
    man=[['field','value'],['version','v.2.3.1'],['schema','xlsx-adjudication-independent-primary-10'],['year',2025],['reviewerRole',role],['workflow','One independent primary review; all review operations and calculations are internal Excel features.'],['provenance','Annual v.2.1.5.3 cohort and frozen catalog. Sources, relationships, order and AI suggestions preserved. No human answers imported.'],['protection','Protected ranges: NONE. Every cell and default style explicitly unlocked, formulas visible; no sheet or workbook lock or password. Rule_defaults is an editable reference.'],['columns A-C','Candidate key, title, location: baseline contents unchanged.'],['columns D-G','Disposition; editable AI-suggested domain; completion no/primary_done; calculated validation hint.'],['columns H-L','Statistical issue; substantiveness; optional downstream impact; internal/plan inconsistency; optional protocol interpretation.'],['columns M-Q / R','M-Q empty hidden spacers; R optional English notes.'],['paper count gate','v.2.3.1 adaptation: all rows primary-ready AND no discuss before outcome counts; independent completion and unresolved totals are separate.'],['outcomes','P_R/P_P precede temporal rules. Only retained S_C/S_W substantive confirmed/probable statistical issues are primary_temporal; retained substantive report-internal inconsistency is secondary_temporal.'],['source hyperlinks','Frozen href cell text is unchanged. Clickable hyperlink targets are rebased from repository source paths to this output directory. Embedded source tables require no refresh.'],['calculation','Automatic; full recalculation on open. No external formula references, scripts, macros or runtime dependencies.']]
    for f in json.loads((HERE/'baseline_inventory.current.json').read_text())['source_files']:man.append(['source sha256: '+Path(f['path']).name,f['sha256']])
    table(w,'Manifest',man,{1:47,2:113},75)
    # Make actual navigation links; original annual file stored href strings but no hyperlink objects.
    evidence={}
    for r in range(2,w['Evidence'].max_row+1):evidence.setdefault(w['Evidence'].cell(r,1).value,r)
    for s in sheets:
        for r in range(2,s.max_row+1):s.cell(r,1).hyperlink=Hyperlink(ref=f'A{r}',location=f"'Evidence'!A{evidence[s.cell(r,1).value]}",tooltip='Read embedded evidence; filter by candidate key')
    for r in range(2,w['Papers'].max_row+1):w['Papers'].cell(r,2).hyperlink=Hyperlink(ref=f'B{r}',location=q(w['Papers'].cell(r,2).value)+'!A1')
    for name,col in [('Links',5),('Candidates',15)]:
        for r in range(2,w[name].max_row+1):
            cell=w[name].cell(r,col);href=cell.value
            if not href:continue
            path,sep,fragment=href.partition('#')
            target=ROOT/unquote(path.removeprefix('../../'))
            if not target.is_file():raise FileNotFoundError(target)
            rebased=os.path.relpath(target,out)+(sep+fragment if sep else '')
            cell.hyperlink=rebased
            if name=='Links':w[name].cell(r,4).hyperlink=rebased
    # Explicit defaults, including unused cells and serialized legacy styles.
    for s in w:
        s.protection.sheet=False
        for row in s:
            for c in row:c.protection=Protection(locked=False,hidden=False)
        for dims in [s.row_dimensions,s.column_dimensions]:
            for d in dims.values():d.protection=Protection(locked=False,hidden=False)
    unlocked=w._protections.add(Protection(locked=False,hidden=False))
    for st in w._cell_styles:st.protectionId=unlocked
    for st in w._named_styles:st.protection=Protection(locked=False,hidden=False)
    w.security=None
    w.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    order=['Start','Setup','Checks','Papers']+[s.title for s in sheets]+['Decisions','Paper_counts','Protocol_plan_table']+SOURCES+['Codebook','Rule_defaults','Manifest']
    w._sheets=[w[n] for n in order];w.active=0
    dest=out/f'human_adjudication_v_2_3_1_2025_{role}.xlsx'
    if dest.exists():raise FileExistsError('Never overwrite an existing workbook: '+str(dest))
    w.save(dest)
    return dest

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--year',type=int,required=True,choices=[2025]);ap.add_argument('--output-dir',default='staging');args=ap.parse_args()
    out=local_dir(args.output_dir)
    for role in ['R_A','R_B']:
        if (out/f'human_adjudication_v_2_3_1_2025_{role}.xlsx').exists():raise FileExistsError('Output exists; choose a fresh construction directory.')
    # Confirm baseline hashes before constructing either independent file.
    inv=json.loads((HERE/'baseline_inventory.json').read_text())
    for f in inv['source_files']:assert hashlib.sha256(Path(f['path']).read_bytes()).hexdigest()==f['sha256'],f['path']
    out.mkdir(parents=True,exist_ok=True)
    for role in ['R_A','R_B']:print(build(role,out))
    (out/'USAGE.md').write_text(USAGE)
    (HERE/'source_hash_manifest.json').write_text(json.dumps(inv['source_files'],indent=2)+'\n')
    (HERE/'TRANSFORMATION.md').write_text('''# 2025 construction record\n\nRead v.2.1.5.3 USAGE, DESIGN, VALIDATION, actual annual workbook, maintenance builder/codebook and source catalog. Preflight matches the supplied inventory byte-for-byte. Annual cohort, not DOI year, determines membership.\n\nCloned annual layout and immutable annual source cell values independently for each role; replaced all workflow, validation, completion and result formulas. Removed snapshot and legacy review inputs. Added key-based primary-only results, explicit pasted-input/completion checks, discuss-aware paper gates and working navigation/source hyperlinks. Source href text remains immutable; hyperlink relationships are rebased. All styles are explicitly unlocked, including unused defaults; no protected ranges. Added Rule_defaults as an unprotected reference.\n\nPreserved A-C/E defaults, H-L applicability defaults, primary option meanings, navy/wrapped styling, yellow required/blue optional and gray inactive formatting, freeze panes and candidate order. Widened E/H/K/L for codes; blank M-Q are hidden. C_ structured tables replaced by whole-record AutoFilters to permit truly blank spacer headers. No review-time automation. Builder accepts only 2025, confines output to local construction subdirectories, refuses release paths and any existing output workbook.\n''')
if __name__=='__main__':main()
