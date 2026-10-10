#!/usr/bin/env python3
"""Pre-release checks only. All synthetic edits stay in tests/; never alter inputs."""
import argparse, copy, hashlib, json, re, os, sys, time
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from urllib.parse import unquote
import formulas
import numpy as np
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries
from build_v_2_3_1 import HERE, BASE, ANNUAL, SOURCES, OPTS
NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REPORT={'static':[], 'functional':[], 'limitations':['Desktop Excel reviewer smoke test: NOT RUN (desktop Excel unavailable).','Visual rendering and actual Excel dropdown display: NOT RUN (no desktop Excel or LibreOffice).','formulas 1.3.4 is a compatible subset evaluator, not Excel; its tests do not establish exact Excel UI behavior.'], 'protected_ranges':[]}
def check(value,msg,group='static'):
    if not value: raise AssertionError(msg)
    REPORT[group].append(msg)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def getvals(s):return list(s.values)
def statics(path,role,baseline,catalog):
    w=load_workbook(path);papers=[s for s in w if s.title.startswith('C_')]
    keys=[s.cell(r,1).value for s in papers for r in range(2,s.max_row+1)]
    check(len(papers)==15 and len(keys)==len(set(keys))==131,f'{role}: 15 papers / 131 unique candidates')
    check(w['Setup']['B2'].value is None and w['Setup']['B3'].value is None and w['Setup']['B4'].value==role and w['Setup']['B5'].value==2025,f'{role}: independent blank identity, correct role/year')
    for name in SOURCES:
        vals=getvals(w[name]);check(vals==getvals(baseline[name]),f'{role}: {name} annual source values identical')
        ch=list(catalog[name].values);ki=ch[0].index('key');annual=[ch[0]]+[r for r in ch[1:] if str(r[ki]).startswith('2025::')]
        check(vals==annual,f'{role}: {name} matches frozen catalog cohort')
    maxlen=0;fc=0;hyperlinks=0
    for s in w:
        check(not s.protection.sheet,f'{role}: {s.title} unprotected')
        for row in s:
            for c in row:
                assert c.protection.locked is False and c.protection.hidden is False,(s.title,c.coordinate)
                if c.data_type=='f':
                    fc+=1;maxlen=max(maxlen,len(c.value));assert '#REF!' not in c.value and '[' not in c.value
                    for tok in Tokenizer(c.value).items:
                        if tok.type=='OPERAND' and tok.subtype=='RANGE' and '!' in tok.value:
                            ref=tok.value.split('!')[0].strip("'").replace("''", "'");assert ref in w,ref
                elif c.data_type=='e':raise AssertionError((s.title,c.coordinate,c.value))
                if isinstance(c.value,str):
                    assert not re.search(r'checker|Primary_snapshot|FINAL READY|checkStatus|resolutionStatus|Setup!\$?B\$?6',c.value,re.I),(s.title,c.coordinate,c.value)
                if c.hyperlink:
                    hyperlinks+=1;h=c.hyperlink
                    if h.location:
                        sn,coord=h.location.rsplit('!',1);sn=sn.strip("'");assert sn in w;assert w[sn][coord].value is not None
                    elif h.target:
                        assert (path.parent/unquote(h.target.split('#')[0])).resolve().is_file(),h.target
        for dv in s.data_validations.dataValidation:
            assert not any(x in str(dv.formula1).lower() for x in ['checker','snapshot','stage'])
            assert len(dv.formula1 or '')<=255
    for s in papers:
        old=baseline[s.title];assert s.max_column==18 and s.freeze_panes==old.freeze_panes
        assert [[s.cell(r,c).value for c in [1,2,3,5]] for r in range(2,s.max_row+1)]==[[old.cell(r,c).value for c in [1,2,3,5]] for r in range(2,old.max_row+1)]
        assert s.auto_filter.ref==f'A1:R{s.max_row}'
        for row in s:
            for c in row[12:17]:assert c.value is None and c.hyperlink is None and c.comment is None
        for col in 'MNOPQ':assert s.column_dimensions[col].hidden and s.column_dimensions[col].width==3
        for r in range(2,s.max_row+1):
            assert s[f'D{r}'].value is None and s[f'R{r}'].value is None and s[f'F{r}'].value=='no'
            for col in 'HIJKL':assert s[f'{col}{r}'].value==old[f'{col}{r}'].value
            for col in 'DEFHIJKLR':assert not s[f'{col}{r}'].protection.locked
        for dv in s.data_validations.dataValidation:
            assert all(not any(f'{col}{r}' in dv for dv in s.data_validations.dataValidation) for col in 'MNOPQ' for r in range(1,s.max_row+1))
        assert len(s.data_validations.dataValidation)==8
        for col in 'DE FHIJKL'.replace(' ',''):
            dv=next(d for d in s.data_validations.dataValidation if f'{col}2' in d)
            assert json.loads(dv.formula1).split(',')==OPTS[col]+(['not_applicable'] if col in 'HIJKL' else [])
    check(fc>2000 and maxlen<8192,f'{role}: {fc} internal formulas; maximum length {maxlen} < 8192')
    check(hyperlinks>600,f'{role}: {hyperlinks} internal/source hyperlinks resolve from {path.parent.name}')
    check(not w.defined_names and not w._external_links and not w.security,f'{role}: no stale names, external workbook links or structure locks')
    check(w.calculation.calcMode=='auto' and w.calculation.fullCalcOnLoad and w.calculation.forceFullCalc,f'{role}: automatic full recalculation on opening')
    with ZipFile(path) as z:
        assert not any(any(t in n.lower() for t in ['externallinks','vba','connections','querytables']) for n in z.namelist())
        x=ET.fromstring(z.read('xl/styles.xml'))
        for xf in x.findall('m:cellXfs/m:xf',NS)+x.findall('m:cellStyleXfs/m:xf',NS):
            p=xf.find('m:protection',NS);assert p is not None and p.get('locked')=='0' and p.get('hidden')=='0'
        for n in z.namelist():
            if n.startswith('xl/worksheets/') and n.endswith('.xml'):
                assert ET.fromstring(z.read(n)).find('m:sheetProtection',NS) is None
    check(True,f'{role}: serialized styles (including default/unused cells) unlocked; no passwords or hidden formulas')
    check(True,f'{role}: primary layout, blank assessments, defaults, dropdowns and spacers verified')
    return w

def solution_value(sol,sheet,coord):
    suffix=f']{sheet.upper()}\'!{coord.upper()}'
    matches=[v for k,v in sol.items() if isinstance(k,str) and k.endswith(suffix)]
    if len(matches)!=1:raise AssertionError((suffix,len(matches)))
    return matches[0].value[0,0]
def calculate(path):
    print('Calculate',path.name,flush=True);t=time.time()
    model=formulas.ExcelModel().loads(str(path)).finish();sol=model.calculate()
    for k,v in sol.items():
        if hasattr(v,'value'):
            for x in np.asarray(v.value).flat:
                assert not isinstance(x,formulas.tokens.operand.XlError),(k,x)
    print('Calculated in',round(time.time()-t,1),'seconds',flush=True)
    return model,sol

def cache_copy(source,dest,sol):
    # Write evaluation caches into a TEST copy without changing formulas or formatting.
    wb=load_workbook(source);xmlns=NS['m'];ns=lambda t:'{'+xmlns+'}'+t
    with ZipFile(source) as zin,ZipFile(dest,'w',ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            data=zin.read(name)
            m=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',name)
            if m:
                sheet=wb.worksheets[int(m[1])-1].title;xml=ET.fromstring(data)
                for c in xml.findall('.//m:sheetData/m:row/m:c',NS):
                    if c.find('m:f',NS) is None:continue
                    val=solution_value(sol,sheet,c.get('r'));v=c.find('m:v',NS)
                    if v is None:v=ET.SubElement(c,ns('v'))
                    if isinstance(val,(str,np.str_)):c.set('t','str');v.text=str(val)
                    elif isinstance(val,(bool,np.bool_)):c.set('t','b');v.text='1' if val else '0'
                    else:c.set('t','n');v.text=str(val)
                data=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
            zout.writestr(name,data)
    calc=load_workbook(dest,data_only=True);raw=load_workbook(dest)
    check(calc['Checks']['B7'].value==solution_value(sol,'Checks','B7'),'Engine results saved and reopened from formula caches on TEST copy','functional')
    check(raw['Checks']['B7'].data_type=='f','Saved/reopened TEST copy retains formulas','functional')

def clone(w):
    buffer=BytesIO();w.save(buffer);buffer.seek(0);return load_workbook(buffer)

def fill(w):
    w['Setup']['B2']='TEST_REVIEWER';w['Setup']['B3']='TEST ONLY'
    w['Start']['B2']='TEST ONLY — SYNTHETIC INPUTS — NEVER RESEARCH DATA'
    for s in w:
        if not s.title.startswith('C_'):continue
        for r in range(2,s.max_row+1):
            for col in 'HIJKL':s[f'{col}{r}']=None
            s[f'D{r}']='retain';s[f'F{r}']='primary_done'
            e=s[f'E{r}'].value
            if e=='statistical_results':s[f'H{r}']='confirmed_issue';s[f'I{r}']='substantive'
            elif e in ['report_internal_inconsistency','plan_material']:s[f'K{r}']='substantive_inconsistency'
            s[f'R{r}']='TEST '+s[f'A{r}'].value

def outcome(rel,e,h=None,i=None,k=None,d='retain'):
    if d=='discuss':return ''
    if d!='retain':return 'not_eligible'
    if rel=='P_R':return 'protocol_descriptive'
    if rel=='P_P':return 'plan_descriptive'
    if rel in ['S_C','S_W'] and e=='statistical_results' and h in ['confirmed_issue','probable_issue'] and i=='substantive':return 'primary_temporal'
    if rel in ['S_C','S_W'] and e=='report_internal_inconsistency' and k=='substantive_inconsistency':return 'secondary_temporal'
    return 'not_eligible'

def functional(w,testdir):
    source=next(s for s in w if s.title.startswith('C_'))
    for d,e in [('',x) for x in OPTS['E']]+[('retain',x) for x in OPTS['E']]+[('exclude','statistical_results'),('discuss','protocol_report')]:
        for col in 'HIJKL':
            fn=formulas.Parser().ast(source[f'{col}2'].value)[1].compile()
            val=fn(*[{'D2':d,'E2':e}[k] for k in fn.inputs]).item()
            enabled=d in ['', 'retain'] and ((col in 'HIJ' and e=='statistical_results') or (col=='K' and e in ['report_internal_inconsistency','plan_material']) or (col=='L' and e=='protocol_report'))
            assert val==('' if enabled else 'not_applicable'),(col,d,e,val)
        check(True,f'Untouched H-L defaults: disposition {d!r}, domain {e}','functional')
    complete=clone(w);fill(complete);path=testdir/'TEST_all_complete.xlsx';complete.save(path)
    model,sol=calculate(path)
    check(solution_value(sol,'Checks','B4')==131 and solution_value(sol,'Checks','B5')==0 and solution_value(sol,'Checks','B7')=='PRIMARY READY','All 131 candidates primary-ready; 15 papers independently complete','functional')
    crows=list(complete['Candidates'].values);candidates=[dict(zip(crows[0],r)) for r in crows[1:]]
    expected={c['key']:outcome(c['primary'],c['ai_domain'],'confirmed_issue','substantive','substantive_inconsistency') for c in candidates}
    for r,c in enumerate(candidates,2):assert solution_value(sol,'Decisions',f'H{r}')==expected[c['key']]
    for r in range(2,17):
        paper=complete['Paper_counts'].cell(r,2).value
        assert solution_value(sol,'Paper_counts',f'G{r}')=='yes' and solution_value(sol,'Paper_counts',f'H{r}')=='yes'
        for j,o in enumerate(['primary_temporal','secondary_temporal','protocol_descriptive','plan_descriptive','not_eligible'],9):
            from openpyxl.utils import get_column_letter
            assert solution_value(sol,'Paper_counts',f'{get_column_letter(j)}{r}')==sum(expected[c['key']]==o for c in candidates if c['package']==paper)
    check(True,'Every completed paper outcome count reconciles with independently expected results, including one-candidate paper','functional')
    cache_copy(path,testdir/'TEST_all_complete_recalculated.xlsx',sol)
    # Reuse evaluator graph, overriding only identity and selecting the real G formula output.
    s=next(s for s in complete if s.title.startswith('C_'))
    key=next(k for k in sol if isinstance(k,str) and k.endswith("]SETUP'!B2"))
    output=next(k for k in sol if isinstance(k,str) and k.endswith(f"]{s.title.upper()}'!G2"))
    for identity in ['', '   ', '\t\n', '\u00a0']:
        si=model.calculate(inputs={key:[[identity]]},outputs=[output])
        check(solution_value(si,s.title,'G2')=='CHECK REVIEWER ID',f'Identity {identity!r} rejected by actual G formula','functional')
    # Assign independent edge cases to different rows in a single synthetic workbook.
    cases=[]
    def add(name,values,status='PRIMARY READY',out=None):cases.append((name,values,status,out))
    add('statistical confirmed substantive',{})
    add('statistical probable substantive',{'H':'probable_issue'})
    add('statistical minor',{'I':'minor_presentation'},out='not_eligible')
    add('statistical indeterminate issue',{'H':'indeterminate'},out='not_eligible')
    add('statistical indeterminate importance',{'I':'indeterminate'},out='not_eligible')
    add('statistical not issue',{'H':'not_issue'},out='not_eligible')
    for col in ['H','I']:
        for value in [None,'not_applicable','PASTED_INVALID']:
            add(f'{col} required {value}',{col:value},'CHECK FIELDS','')
    for col in ['J']:
        for value in [None,'direct','possible','unlikely','unknown']:add(f'optional {col} {value}',{col:value})
        for value in ['not_applicable','PASTED_INVALID']:add(f'optional invalid {col} {value}',{col:value},'CHECK FIELDS','')
    for domain in ['report_internal_inconsistency','plan_material']:
        base={'E':domain,'H':None,'I':None,'K':'substantive_inconsistency'}
        add(domain,base)
        for value in [None,'not_applicable','PASTED_INVALID']:add(f'{domain} required {value}',dict(base,K=value),'CHECK FIELDS','')
        for value in ['minor_presentation','non_inconsistency','indeterminate']:add(f'{domain} {value}',dict(base,K=value))
    base={'E':'protocol_report','H':None,'I':None}
    for value in [None]+OPTS['L']:add(f'protocol optional {value}',dict(base,L=value))
    for value in ['not_applicable','PASTED_INVALID']:add(f'protocol invalid {value}',dict(base,L=value),'CHECK FIELDS','')
    for d in ['exclude','discuss']:
        add(d,{'D':d,'H':None,'I':None},out='not_eligible' if d=='exclude' else '')
        add(d+' stale',{'D':d},'CLEAR INACTIVE VALUES','')
        add(d+' missing domain',{'D':d,'E':None,'H':None,'I':None},'CHECK FIELDS','')
    for col in ['D','E']:
        for value in [None,'PASTED_INVALID']:add(f'invalid {col} {value}',{col:value,'H':None,'I':None},'CHECK FIELDS','')
    add('domain changed stale',{'E':'protocol_report'},'CLEAR INACTIVE VALUES','')
    add('domain changed cleared',base)
    add('optional cleared',{'J':None})
    add('completion manually reset',{'F':'no'},'PRIMARY DRAFT','')
    add('completion blank',{'F':None},'PRIMARY DRAFT','')
    add('forbidden completion',{'F':'checker_done'},'CHECK COMPLETE','')
    add('pasted completion',{'F':'yes'},'CHECK COMPLETE','')
    add('missing key',{'A':None},'CHECK KEY','')
    add('duplicate key',{},'CHECK KEY','')
    add('duplicate counterpart',{},'CHECK KEY','')
    # Explicit relationship precedence despite human domain changes.
    mixed=clone(complete)
    records={c['key']:c for c in candidates}
    slots=[(s,r) for s in mixed if s.title.startswith('C_') for r in range(2,s.max_row+1)]
    slots.sort(key=lambda sr: records[sr[0].cell(sr[1],1).value]['primary'] in ['P_R','P_P'])
    expectations=[]
    for idx,(name,changes,status,expected_out) in enumerate(cases):
        s,r=slots[idx];original=s[f'A{r}'].value
        for col in 'HIJKL':s[f'{col}{r}']=None
        for col,value in {'D':'retain','E':'statistical_results','F':'primary_done','H':'confirmed_issue','I':'substantive'}.items():s[f'{col}{r}']=value
        if name=='duplicate key':changes={'A':slots[idx+1][0].cell(slots[idx+1][1],1).value}
        for col,value in changes.items():s[f'{col}{r}']=value
        if expected_out is None:
            expected_out=outcome(records[original]['primary'],s[f'E{r}'].value,s[f'H{r}'].value,s[f'I{r}'].value,s[f'K{r}'].value,s[f'D{r}'].value)
        expectations.append((name,s.title,r,original,status,expected_out))
    # Locate descriptive relationships outside assigned rows and change their domain.
    for relationship in ['P_R','P_P']:
        for s,r in slots[len(cases):]:
            key0=s[f'A{r}'].value
            if records[key0]['primary']==relationship:
                for col in 'HIJKL':s[f'{col}{r}']=None
                s[f'E{r}']='statistical_results';s[f'H{r}']='confirmed_issue';s[f'I{r}']='substantive'
                expectations.append((relationship+' precedence',s.title,r,key0,'PRIMARY READY','protocol_descriptive' if relationship=='P_R' else 'plan_descriptive'));break
        else:raise AssertionError('No unused relationship fixture '+relationship)
    path=testdir/'TEST_primary_cases.xlsx';mixed.save(path)
    _,msol=calculate(path)
    for name,s,r,key0,status,expected_out in expectations:
        assert solution_value(msol,s,f'G{r}')==status,(name,solution_value(msol,s,f'G{r}'),status)
        dr=next(i for i,c in enumerate(candidates,2) if c['key']==key0)
        assert solution_value(msol,'Decisions',f'H{dr}')==expected_out,(name,solution_value(msol,'Decisions',f'H{dr}'),expected_out)
        check(True,name+': readiness and outcome','functional')
    check(solution_value(msol,'Checks','B6')==3,'Discuss counted separately, including unfinished discuss rows','functional')
    for r in range(2,17):
        pending=solution_value(msol,'Paper_counts',f'E{r}');discuss=solution_value(msol,'Paper_counts',f'F{r}')
        if pending or discuss:assert solution_value(msol,'Paper_counts',f'I{r}')==''
    check(True,'Paper counts withheld for unfinished or discuss papers','functional')
    # All-ready with discuss: completion yes, counts unavailable. Use the single-candidate paper.
    s=next(s for s in complete if s.title.startswith('C_') and s.max_row==2)
    singlekey=s['A2'].value
    changed=clone(complete);cs=changed[s.title];cs['D2']='discuss'
    for col in 'HIJKL':cs[f'{col}2']=None
    # Reverse all complete C_ rows, translating relative formula references as an Excel whole-record sort does.
    for sheet in changed:
        if not sheet.title.startswith('C_'):continue
        vals=[[(c.value,c.coordinate) for c in row] for row in list(sheet)[1:]]
        for r,row in enumerate(reversed(vals),2):
            for j,(v,origin) in enumerate(row,1):
                dest=sheet.cell(r,j)
                dest.value=Translator(v,origin=origin).translate_formula(dest.coordinate) if isinstance(v,str) and v.startswith('=') else v
    path=testdir/'TEST_sorted_and_discuss.xlsx';changed.save(path)
    _,ssol=calculate(path)
    for dr,c in enumerate(candidates,2):
        assert solution_value(ssol,'Decisions',f'R{dr}')=='TEST '+c['key']
        assert solution_value(ssol,'Decisions',f'H{dr}')==('' if c['key']==singlekey else expected[c['key']])
    check(True,'Whole-record reordering preserves all 131 key-to-assessment/notes/outcome mappings','functional')
    pr=next(r for r in range(2,17) if changed['Paper_counts'].cell(r,2).value==s.title[2:])
    check(solution_value(ssol,'Checks','B7')=='PRIMARY READY' and solution_value(ssol,'Checks','B6')==1,'All-ready annual review can retain one unresolved discuss','functional')
    check(solution_value(ssol,'Paper_counts',f'G{pr}')=='yes' and solution_value(ssol,'Paper_counts',f'H{pr}')=='no' and solution_value(ssol,'Paper_counts',f'I{pr}')=='','One-candidate discuss paper independently complete, counts withheld','functional')
    (testdir/'cases.json').write_text(json.dumps([{'case':n,'sheet':s,'row':r,'original_key':k,'expected_status':st,'expected_outcome':o} for n,s,r,k,st,o in expectations],indent=2)+'\n')


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--year',type=int,required=True,choices=[2025]);ap.add_argument('--input-dir',default='staging');ap.add_argument('--static-only',action='store_true');args=ap.parse_args()
    inputdir=(HERE/args.input_dir).resolve();assert inputdir.is_relative_to(HERE)
    testdir=HERE/'tests';testdir.mkdir(exist_ok=True)
    baseline=load_workbook(ANNUAL);catalog=load_workbook(BASE/'resources/source_catalog_v_2_1_5_3.xlsx',read_only=True)
    before={f['path']:sha(Path(f['path'])) for f in json.loads((HERE/'source_hash_manifest.json').read_text())}
    for f in json.loads((HERE/'source_hash_manifest.json').read_text()):assert before[f['path']]==f['sha256']
    paths=[inputdir/f'human_adjudication_v_2_3_1_2025_{role}.xlsx' for role in ['R_A','R_B']]
    originals={str(p):sha(p) for p in paths}
    wa=statics(paths[0],'R_A',baseline,catalog);wb=statics(paths[1],'R_B',baseline,catalog)
    for name in SOURCES+[s.title for s in wa if s.title.startswith('C_')]:assert getvals(wa[name])==getvals(wb[name])
    # Workbook cells and formulas identical apart from documented role metadata.
    diffs=[]
    for sa in wa:
        sb=wb[sa.title]
        for row in sa:
            for ca in row:
                if ca.value!=sb[ca.coordinate].value:diffs.append((sa.title,ca.coordinate))
    check(diffs==[('Start','B2'),('Setup','B4'),('Manifest','B5')],'Roles differ only in their three role metadata cells')
    sa=next(s for s in wa if s.title.startswith('C_'));sa['D2']='retain'
    wa.save(testdir/'TEST_independence_R_A.xlsx')
    check(sha(paths[1])==originals[str(paths[1])] and wb[sa.title]['D2'].value is None,'Editing and saving an A test copy leaves B unchanged')
    if not args.static_only:
        blank=load_workbook(paths[0]);blankpath=testdir/'TEST_blank.xlsx';blank.save(blankpath)
        _,sol=calculate(blankpath)
        check(solution_value(sol,'Checks','B4')==0 and solution_value(sol,'Checks','B5')==131 and solution_value(sol,'Checks','B6')==0,'Clean template has zero ready, 131 unfinished, zero discuss','functional')
        check(all(solution_value(sol,'Decisions',f'H{r}')=='' for r in range(2,133)),'Every clean-template outcome blank','functional')
        functional(load_workbook(paths[0]),testdir)
    check(all(sha(p)==originals[str(p)] for p in paths),'Validator left input templates unchanged')
    check(all(sha(Path(p))==v for p,v in before.items()),'All inventoried baseline sources unchanged after validation')
    REPORT['engine']='formulas '+formulas.__version__;REPORT['input_hashes']=originals;REPORT['status']='PASS for executed checks; desktop Excel and rendering pending';REPORT['static_only']=args.static_only
    report=HERE/('validation_report.release.json' if args.static_only else 'validation_report.json')
    report.write_text(json.dumps(REPORT,indent=2)+'\n')
    print('PASS',len(REPORT['static']),'static assertions;',len(REPORT['functional']),'functional assertions. Desktop Excel/rendering pending.',flush=True)
if __name__=='__main__':main()
