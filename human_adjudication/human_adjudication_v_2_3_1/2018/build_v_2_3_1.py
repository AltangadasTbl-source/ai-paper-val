#!/usr/bin/env python3
"""2018 only; pre-release construction. Never run on completed review files."""
import argparse, hashlib, json, os, re, tempfile
from pathlib import Path
from copy import copy
from urllib.parse import unquote
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as letter

HERE=Path(__file__).resolve().parent
(HERE/'.tmp').mkdir(exist_ok=True)
tempfile.tempdir=str(HERE/'.tmp')
ROOT=HERE.parents[2]
BASE=ROOT/'human_adjudication/previous/human_adjudication_v_2_1_5_3'
SOURCE=BASE/'human_adjudication_v_2_1_5_3_2018.xlsx'
SCHEMA='xlsx-adjudication-independent-primary-2.3.1'
OPTIONS={'D':['retain','exclude','discuss'],'E':['statistical_results','report_internal_inconsistency','plan_material','protocol_report'],'F':['no','primary_done'],'H':['confirmed_issue','probable_issue','not_issue','indeterminate'],'I':['substantive','minor_presentation','indeterminate'],'J':['direct','possible','unlikely','unknown'],'K':['substantive_inconsistency','minor_presentation','non_inconsistency','indeterminate'],'L':['substantive_difference','minor_presentation','version_timing','non_inconsistency','indeterminate']}
NAVY='173F57'; YELLOW='FFF2CC'; BLUE='EAF3FC'; GRAY='EDF2F5'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def quote(v): return '"'+str(v).replace('"','""')+'"'
def sq(s): return "'"+s.replace("'","''")+"'"
def valid(c,opts): return 'OR('+','.join(f'EXACT({c},{quote(v)})' for v in opts)+')'
def active(col,d,e):
    branch=f'{e}="statistical_results"' if col in 'HIJ' else f'OR({e}="report_internal_inconsistency",{e}="plan_material")' if col=='K' else f'{e}="protocol_report"'
    return f'AND({d}="retain",{branch})'
def lookup(sheet,col,key,n,keycol='A'):
    return f'INDEX({sq(sheet)}!${col}$2:${col}${n+1},MATCH({key},{sq(sheet)}!${keycol}$2:${keycol}${n+1},0))'
def preserve(x): return f'IF({x}="","",{x})'
def hint(r,n,total):
    d,e,f,a=(f'{c}{r}' for c in 'DEFA')
    bad=[f'NOT({valid(d,OPTIONS["D"])})',f'NOT({valid(e,OPTIONS["E"])})']; stale=[]
    for col in 'HIJKL':
        x=f'{col}{r}'; act=active(col,d,e)
        optional=f'{x}<>"",' if col in 'JL' else ''
        bad.append(f'AND({act},{optional}NOT({valid(x,OPTIONS[col])}))')
        stale.append(f'AND(NOT({act}),{x}<>"",NOT(EXACT({x},"not_applicable")))')
    identity='LEN(TRIM(CLEAN(SUBSTITUTE(Setup!$B$2,CHAR(160)," "))))=0'
    # Membership in this paper is verified against immutable construction-time Decisions metadata.
    key=f'OR({a}="",COUNTIF($A$2:$A${n+1},{a})<>1,COUNTIFS(Decisions!$C$2:$C${total+1},{a},Decisions!$B$2:$B${total+1},MID({a},7,LEN({a})-12))<>1)'
    return f'=IFERROR(IF({identity},"CHECK REVIEWER ID",IF({key},"CHECK KEY",IF(OR({",".join(stale)}),"CLEAR INACTIVE VALUES",IF(OR({",".join(bad)}),"CHECK FIELDS",IF(AND({f}<>"",NOT({valid(f,OPTIONS["F"])})),"CHECK COMPLETE",IF(EXACT({f},"primary_done"),"PRIMARY READY","PRIMARY DRAFT")))))),"CHECK INPUT / KEY")'

def table(w,name,rows,widths=None,height=48):
    if name in w:
        idx=w.sheetnames.index(name); del w[name]; s=w.create_sheet(name,idx)
    else: s=w.create_sheet(name)
    for row in rows: s.append(row)
    for row in s:
        for c in row:
            c.font=Font(name='Calibri',size=11,color=NAVY)
            c.alignment=Alignment(wrap_text=True,vertical='center')
            c.protection=Protection(locked=False,hidden=False)
            if c.row==1:
                c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(name='Calibri',size=11,color='FFFFFF',bold=True)
            elif c.row%2==0: c.fill=PatternFill('solid',fgColor='F4F8FA')
        s.row_dimensions[row[0].row].height=42 if row[0].row==1 else height
    for i in range(1,s.max_column+1): s.column_dimensions[letter(i)].width=(widths or [])[i-1] if widths and i<=len(widths) else 28
    s.freeze_panes='B2'; s.sheet_view.showGridLines=False
    return s

START=[['Step','Instructions'],['Release','v.2.3.1 | 2018 | 10 papers / 119 candidates | Independent primary review; Excel only.'],
['Coordinator','Assign two different people to the two role files. Distribute only the corresponding file to each person. Role labels are not person IDs. This standalone workbook cannot verify the other person’s identity.'],
['1. Begin','Save As a working copy in this release folder to preserve relative source links. Enter your own reviewer ID/name in Setup B2/B3. Keep backups using normal Excel saves.'],
['2. Read evidence','Use Papers to open a C_ paper sheet. Click a candidate key to open embedded Evidence; use Candidates, Links and Relations for full source context. Links opens original documents. AI suggestions are not human-confirmed findings.'],
['3. Assess','Choose D (retain/exclude/discuss), confirm or change suggested domain E, answer applicable H–L, and optionally add English notes in R. Yellow = common/required; blue = optional; gray = inactive. See Codebook and Rule_defaults. M–Q are empty hidden spacers.'],
['4. Mark completion','After manual review set F=primary_done. G must show PRIMARY READY. Use no while drafting. Inspect Checks and filter Decisions readiness for outstanding rows.'],
['5. Save and deliver','When Checks shows PRIMARY READY, save, close and deliver this workbook. Decisions, Paper_counts and Protocol_plan_table update inside Excel. Preserve access to original source folders; embedded evidence requires no refresh. No scripts or external calculation services are needed.'],
['Changes after completion','Manually reset F to no BEFORE editing a completed row, review again, then mark primary_done. Excel formulas cannot erase an entered completion flag. Domain/disposition changes never automatically erase literal judgments.'],
['Blank and inactive inputs','Untouched H–L formulas display blank when potentially applicable and not_applicable otherwise. Choosing an answer replaces its formula. Blank means no answer, never an inherited answer. Clear stale inactive answers or choose not_applicable; G flags CLEAR INACTIVE VALUES. Newly active required fields need answers.'],
['Applicability','For retain: statistical_results requires H/I; J optional. report_internal_inconsistency and plan_material require K. protocol_report has optional L and no further required judgments. For exclude/discuss all H–L are inactive. D and a valid E are always required.'],
['Optional fields','Active J/L may be blank; if populated they need a listed active option. not_applicable is invalid for active judgments. R is optional English notes. Dropdowns assist entry; G also checks pasted values.'],
['Unresolved assessments','discuss is a completed independent primary response when all entry checks pass. It records unresolved uncertainty; it does not initiate a shared workflow. Checks displays the discuss total separately. A discuss outcome stays blank.'],
['Results and paper counts','Outcomes stay blank until the row is ready. Retained P_R/P_P are descriptive and never temporal. Paper outcome counts stay blank until ALL its rows are ready AND none is discuss. Review completion and scientific resolution are separate. Indeterminate codes retain baseline classification rules; see Codebook.'],
['Editable workbook','All cells, including unused cells, are unlocked. No sheet or structure protection, passwords or runtime unlock steps. Fixed Rule_defaults are visible references and are also unlocked. Preserve reference cells and formulas; undo accidental changes or restore your backup.'],
['Source navigation','Source text and href cell values are preserved exactly. Clickable link targets are adjusted for this release location. If you Save As elsewhere, keep the same directory depth and source tree or use embedded evidence and locate originals manually.'],
['Validation limits','Entry checks do not establish scientific accuracy. Construction validation and pending desktop Excel recalculation/rendering checks are recorded in the accompanying VALIDATION.md.']]

def build(role,out):
    w=load_workbook(SOURCE)
    del w['Primary_snapshot']; w.defined_names.clear()
    data=list(w['Candidates'].values); candidates=[dict(zip(data[0],r)) for r in data[1:]]
    papers=list(dict.fromkeys(c['package'] for c in candidates)); total=len(candidates)
    assert len(papers)==10 and total==119 and len({c['key'] for c in candidates})==119
    table(w,'Start',START,[32,115],80)
    setup=table(w,'Setup',[['Field','Value'],['reviewerId',None],['reviewerName (optional)',None],['reviewerRole',role],['year',2018]],[34,65],42)
    for addr in ['B2','B3']: setup[addr].fill=PatternFill('solid',fgColor=YELLOW)
    checks=table(w,'Checks',[['Check','Result'],['Reviewer role','=Setup!B4'],['Candidates',total],['Primary ready rows',f'=COUNTIF(Decisions!G2:G{total+1},"PRIMARY READY")'],['Discuss / unresolved rows',f'=COUNTIF(Decisions!I2:I{total+1},"discuss")'],['Unfinished rows','=B3-B4'],['Workflow status','=IF(B4=B3,"PRIMARY READY","ACTION NEEDED")'],['Meaning','PRIMARY READY means independent entry is complete. Discuss remains unresolved and is not counted as absence of error.'],['Next action','Inspect G on paper sheets and readiness in Decisions. Save and deliver this file after review.']],[35,100],55)
    ps=table(w,'Papers',[['Paper','Review worksheet','Candidates']]+[[p,'C_'+p,sum(c['package']==p for c in candidates)] for p in papers],[32,38,18],36)
    for r,p in enumerate(papers,2): ps[f'B{r}'].hyperlink=Hyperlink(ref=f'B{r}',location=f"{sq('C_'+p)}!A1")
    evidence={}
    for row in w['Evidence'].iter_rows(min_row=2): evidence.setdefault(row[0].value,row[0].row)
    for p in papers:
        s=w['C_'+p]; n=s.max_row-1
        s.tables.clear(); s.auto_filter.ref=f'A1:R{n+1}'; s.conditional_formatting._cf_rules.clear(); s.data_validations.dataValidation.clear()
        s['E1']='Review domain (AI suggestion; editable)'; s.row_dimensions[1].height=70
        s.column_dimensions['E'].width=35; s.column_dimensions['D'].width=20; s.column_dimensions['G'].width=29
        for col in 'MNOPQ':
            s.column_dimensions[col].width=3; s.column_dimensions[col].hidden=True
            for r in range(1,n+2): s[f'{col}{r}']=None
        for col,opts in OPTIONS.items():
            choices=opts+(['not_applicable'] if col in 'HIJKL' else [])
            dv=DataValidation(type='list',formula1=quote(','.join(choices)),allow_blank=True)
            dv.showDropDown=False; dv.showErrorMessage=True; dv.errorTitle='Use a listed option'; dv.error='Choose an option; G also checks pasted entries.'; dv.errorStyle='stop'
            s.add_data_validation(dv); dv.add(f'{col}2:{col}{n+1}')
        for r in range(2,n+2):
            s[f'D{r}']=None; s[f'F{r}']='no'; s[f'R{r}']=None; s[f'G{r}']=hint(r,n,total)
            # Verify membership against this sheet's paper, rather than deriving it from editable key text.
            s[f'G{r}']=s[f'G{r}'].value.replace(f'MID(A{r},7,LEN(A{r})-12)',quote(p))
            s[f'A{r}'].hyperlink=Hyperlink(ref=f'A{r}',location=f"'Evidence'!A{evidence[s[f'A{r}'].value]}")
            for col in 'HIJKL':
                act=active(col,f'D{r}',f'E{r}').replace(f'D{r}="retain"',f'OR(D{r}="retain",D{r}="")')
                s[f'{col}{r}']=f'=IF({act},"","not_applicable")'
            for col in 'DEFHIK': s[f'{col}{r}'].fill=PatternFill('solid',fgColor=YELLOW)
            for col in 'JLR': s[f'{col}{r}'].fill=PatternFill('solid',fgColor=BLUE)
        for col in 'HIJKL':
            act=active(col,'$D2','$E2').replace('$D2="retain"','OR($D2="retain",$D2="")')
            s.conditional_formatting.add(f'{col}2:{col}{n+1}',FormulaRule(formula=[f'NOT({act})'],fill=PatternFill('solid',fgColor=GRAY),font=Font(color='687783')))
        s.conditional_formatting.add(f'G2:G{n+1}',FormulaRule(formula=['G2<>"PRIMARY READY"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color='9C0006')))
    headers=['year','paper','candidate_key','relationship','reviewer_role','reviewer_id','readiness','outcome','disposition','domain','statistical_issue','substantiveness','downstream_impact','internal_status','protocol_interpretation','notes','ai_disposition','ai_domain']
    ds=table(w,'Decisions',[headers]+[[2018,c['package'],c['key'],c['primary']]+[None]*14 for c in candidates],[12,27,44,18,20,24,28,28,20,35,26,24,24,30,28,55,22,35],70)
    ds.auto_filter.ref=f'A1:R{total+1}'
    for r,c in enumerate(candidates,2):
        n=sum(x['package']==c['package'] for x in candidates)
        look=lambda col:lookup('C_'+c['package'],col,f'$C{r}',n)
        ds[f'E{r}']='=Setup!$B$4'; ds[f'F{r}']='='+preserve('Setup!$B$2'); ds[f'G{r}']=f'=IFERROR({look("G")},"CHECK KEY")'
        for dest,col in zip('IJKLMNOP','DEHIJKLR'):
            exp=preserve(look(col))
            if col in 'HIJKL': exp=f'IF({active(col,f"I{r}",f"J{r}")},{exp},"not_applicable")'
            ds[f'{dest}{r}']=f'=IFERROR({exp},"")'
        ds[f'Q{r}']=c['ai_disposition']; ds[f'R{r}']=c['ai_domain']
        ds[f'H{r}']=f'=IF(OR(G{r}<>"PRIMARY READY",I{r}="discuss"),"",IF(I{r}<>"retain","not_eligible",IF(D{r}="P_R","protocol_descriptive",IF(D{r}="P_P","plan_descriptive",IF(AND(OR(D{r}="S_C",D{r}="S_W"),J{r}="statistical_results",OR(K{r}="confirmed_issue",K{r}="probable_issue"),L{r}="substantive"),"primary_temporal",IF(AND(OR(D{r}="S_C",D{r}="S_W"),J{r}="report_internal_inconsistency",N{r}="substantive_inconsistency"),"secondary_temporal","not_eligible"))))))'
    counts=table(w,'Paper_counts',[['year','paper','candidates','unfinished','discuss_unresolved','primary_complete','outcomes_releasable','primary_temporal','secondary_temporal','protocol_descriptive','plan_descriptive','not_eligible']]+[[2018,p,sum(c['package']==p for c in candidates)]+[None]*9 for p in papers],[12,30,18,18,24,24,25,25,25,25,25,24])
    for r in range(2,12):
        match=f'Decisions!$B$2:$B${total+1},B{r}'
        counts[f'D{r}']=f'=C{r}-COUNTIFS({match},Decisions!$G$2:$G${total+1},"PRIMARY READY")'
        counts[f'E{r}']=f'=COUNTIFS({match},Decisions!$I$2:$I${total+1},"discuss")'
        counts[f'F{r}']=f'=IF(D{r}=0,"yes","no")'; counts[f'G{r}']=f'=IF(AND(D{r}=0,E{r}=0),"yes","no")'
        for col,outcome in zip('HIJKL',['primary_temporal','secondary_temporal','protocol_descriptive','plan_descriptive','not_eligible']): counts[f'{col}{r}']=f'=IF(G{r}<>"yes","",COUNTIFS({match},Decisions!$H$2:$H${total+1},"{outcome}"))'
    pc=[c for c in candidates if c['primary'] in ['P_R','P_P']]
    plans=table(w,'Protocol_plan_table',[['year','paper','candidate_key','relationship','title','statement','readiness','disposition','protocol_interpretation','internal_status','notes']]+[[2018,c['package'],c['key'],c['primary'],c['title'],c['statement']]+[None]*5 for c in pc],[12,27,44,18,60,90,28,22,28,30,60],100)
    for r in range(2,len(pc)+2):
        for dest,col in zip('GHIJK','GIONP'): plans[f'{dest}{r}']='='+preserve(lookup('Decisions',col,f'C{r}',total,'C'))
    rows=json.loads((BASE/'maintenance/codebook.json').read_text()); clean=[rows[0]]; skip=False
    for row in rows[1:]:
        if str(row[0]).startswith('5  |'): skip=True; continue
        if str(row[0]).startswith('6  |'): skip=False
        if skip or row[4]=='primarySummary' or row[1]=='checker_done': continue
        row=list(row)
        if row[0]=='Working method': row[2]='Independent primary'; row[3]='Blank means no answer, never inherited. Clear obsolete literal answers after branch changes. not_applicable means inactive.'
        if row[4]=='disposition': row[3]=row[3].replace(' Checker edits the existing answer directly.','')
        if row[4]=='reviewComplete':
            row[2]='Independent primary'; row[3]='Use no while drafting; primary_done after manual review. Reset to no before editing a completed row, review again, then mark primary_done.'
        if row[4]=='readiness': row[3]='PRIMARY READY means valid independent primary entry and manual completion, including discuss. Other hints identify missing or invalid input. This does not establish scientific accuracy.'
        if row[4]=='locations': row[2]='Source reference'; row[3]='Read Candidates, Evidence, Links and Relations for context and original documents.'
        clean.append(row)
    table(w,'Codebook',clean,[34,30,40,96,28],100)
    table(w,'Rule_defaults',[['Retained domain','Required','Optional','Inactive'],['statistical_results','H, I','J','K, L'],['report_internal_inconsistency','K','None','H, I, J, L'],['plan_material','K','None','H, I, J, L'],['protocol_report','None','L','H, I, J, K'],['Default inactive value','not_applicable','exclude/discuss','H–L all inactive']],[36,28,28,38])
    manifest=[['field','value'],['version','v.2.3.1'],['schema',SCHEMA],['year',2018],['role',role],['cohort','Annual v.2.1.5.3 baseline; 10 papers / 119 candidates'],['source_sha256',sha(SOURCE)],['catalog_sha256',sha(BASE/'resources/source_catalog_v_2_1_5_3.xlsx')],['workflow','Independent primary assessment; no external workbook or refresh dependencies.'],['protection','Protected ranges: NONE. All styles/cells explicitly unlocked; fixed Rule_defaults also unlocked.'],['C_ columns','A key; B title; C location; D disposition; E suggested editable domain; F completion; G validation; H issue; I substantive; J impact; K internal; L protocol; M–Q empty spacers; R notes.'],['Decisions columns','A–D source identity/relationship; E–F reviewer role/ID; G readiness; H outcome; I–P effective assessment/notes; Q–R unchanged AI suggestions.'],['Paper_counts adaptation','Review complete when all rows ready; outcome counts withheld until all rows ready and none discuss. Discuss outcome blank; count separately.'],['Defaults','H–L retain baseline formula defaults. Input choices replace formulas. Inactive effective fields normalize only after stale inputs are flagged by G.'],['Source links','Original href cell values unchanged. Hyperlink relationship targets resolved to repository sources relative to this output folder.'],['Calculation','Automatic full recalculation on open; no VBA, scripts, external formulas or services.'],['Validation','See VALIDATION.md for actual static/synthetic checks and pending desktop Excel recalculation, rendering and saved/reopened smoke test.']]
    table(w,'Manifest',manifest,[34,110],65)
    # Source value text is immutable; hyperlink metadata provides corrected navigation.
    for name,col in [('Links',5),('Candidates',15)]:
        for row in w[name].iter_rows(min_row=2):
            cell=row[col-1]; href=cell.value
            if href:
                path,sep,fragment=href.partition('#')
                path=path.removeprefix('../../')
                target=ROOT/unquote(path)
                if not target.exists(): raise FileNotFoundError(target)
                cell.hyperlink=os.path.relpath(target,out)+(sep+fragment if sep else '')
    # Explicit unlock includes serialized default style and unused-cell inheritance.
    for style in w._cell_styles: style.protectionId=0
    w._protections[0]=Protection(locked=False,hidden=False)
    for s in w:
        s.protection.sheet=False
        for row in s:
            for c in row: c.protection=Protection(locked=False,hidden=False)
    for ns in w._named_styles: ns.protection=Protection(locked=False,hidden=False)
    w.security=None
    w.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    w.active=0
    dest=out/f'human_adjudication_v_2_3_1_2018_{role}.xlsx'
    if dest.exists(): raise FileExistsError(f'Refusing overwrite: {dest}')
    w.save(dest)
    # openpyxl omits protectionId=0 even when its in-memory protection is unlocked.
    # Set every serialized cell/style xf explicitly, including the unused-cell default.
    temp=dest.with_suffix('.unlock.xlsx')
    with ZipFile(dest) as src, ZipFile(temp,'w',ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content=src.read(item.filename)
            if item.filename=='xl/styles.xml':
                root=ET.fromstring(content); ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                for kind in ['cellXfs','cellStyleXfs']:
                    for xf in root.find(ns+kind):
                        xf.set('applyProtection','1')
                        for old in list(xf.findall(ns+'protection')): xf.remove(old)
                        ET.SubElement(xf,ns+'protection',{'locked':'0','hidden':'0'})
                content=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            dst.writestr(item,content)
    temp.replace(dest)
    return dest

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--year',type=int,required=True,choices=[2018]); ap.add_argument('--output-dir',default='staging'); args=ap.parse_args()
    out=(HERE/args.output_dir).resolve()
    if not out.is_relative_to(HERE) or out==HERE/'release' or (HERE/'release') in out.parents: raise ValueError('Construction output must be inside this year and outside release')
    if any(out.glob('*.xlsx')): raise FileExistsError('Use a new empty staging directory; never overwrite workbooks')
    expected=json.loads((HERE/'baseline_inventory.json').read_text())
    for src in expected['source_files']: assert sha(Path(src['path']))==src['sha256'],src['path']
    out.mkdir(parents=True,exist_ok=True)
    files=[build(role,out) for role in ['R_A','R_B']]
    (HERE/'source_hash_manifest.json').write_text(json.dumps(expected['source_files'],indent=2)+'\n')
    (out/'output_hashes.json').write_text(json.dumps({p.name:sha(p) for p in files},indent=2)+'\n')
    print('\n'.join(str(p) for p in files))
if __name__=='__main__': main()
