"""Construction-only limited interpreter for actual emitted formula strings.
NOT Excel or a certified compatible engine; cannot validate rendering/cached values.
"""
import re
from functools import lru_cache
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries

@lru_cache(None)
def parse(formula):
    ts=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE']; i=0
    def expr(minp=0):
        nonlocal i
        t=ts[i]; i+=1
        if t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            if ts[i].subtype!='CLOSE':
                while True:
                    args.append(expr())
                    if ts[i].type!='SEP': break
                    i+=1
            assert ts[i].subtype=='CLOSE'; i+=1; node=('call',t.value[:-1],args)
        elif t.type=='OPERAND':
            node=('ref',t.value) if t.subtype=='RANGE' else ('value',t.value[1:-1].replace('""','"') if t.subtype=='TEXT' else float(t.value) if t.subtype=='NUMBER' else t.value=='TRUE')
        elif t.type=='PAREN' and t.subtype=='OPEN':
            node=expr(); assert ts[i].subtype=='CLOSE'; i+=1
        else: raise ValueError(('Unsupported token',t.value,t.type))
        while i<len(ts) and ts[i].type=='OPERATOR-INFIX':
            op=ts[i].value; p={'=':1,'<>':1,'>':1,'<':1,'>=':1,'<=':1,'+':2,'-':2,'*':3,'/':3}.get(op)
            if p is None: raise ValueError(op)
            if p<minp: break
            i+=1; node=('op',op,node,expr(p+1))
        return node
    ast=expr(); assert i==len(ts),(formula,i,len(ts)); return ast

def eq(a,b): return a.lower()==b.lower() if isinstance(a,str) and isinstance(b,str) else a==b
class Probe:
    def __init__(self,w): self.w=w; self.cache={}; self.busy=set()
    def cell(self,s,a):
        a=a.replace('$',''); k=(s,a)
        if k in self.cache: return self.cache[k]
        if k in self.busy: raise ValueError('circular reference')
        self.busy.add(k)
        try:
            v=self.w[s][a].value
            v=self.evaluate(parse(v),s) if isinstance(v,str) and v.startswith('=') else '' if v is None else v
            self.cache[k]=v; return v
        finally: self.busy.remove(k)
    def evaluate(self,n,s):
        typ=n[0]
        if typ=='value': return n[1]
        if typ=='ref':
            a=n[1]
            if '!' in a: s,a=a.rsplit('!',1); s=s.strip("'").replace("''", "'")
            if ':' not in a: return self.cell(s,a)
            x1,y1,x2,y2=range_boundaries(a)
            return [self.cell(s,self.w[s].cell(r,c).coordinate) for r in range(y1,y2+1) for c in range(x1,x2+1)]
        ev=lambda x:self.evaluate(x,s)
        if typ=='op':
            op,a,b=n[1],ev(n[2]),ev(n[3])
            if op=='=': return eq(a,b)
            if op=='<>': return not eq(a,b)
            if op=='+': return a+b
            if op=='-': return a-b
            return {'>':lambda:a>b,'<':lambda:a<b,'>=':lambda:a>=b,'<=':lambda:a<=b,'*':lambda:a*b,'/':lambda:a/b}[op]()
        name,args=n[1],n[2]
        if name=='IF': return ev(args[1]) if ev(args[0]) else ev(args[2])
        if name=='IFERROR':
            try: return ev(args[0])
            except (ValueError,IndexError,KeyError,ZeroDivisionError,StopIteration): return ev(args[1])
        v=[ev(x) for x in args]
        if name=='AND': return all(v)
        if name=='OR': return any(v)
        if name=='NOT': return not v[0]
        if name=='EXACT': return str(v[0])==str(v[1])
        if name=='LEN': return len(str(v[0]))
        if name=='TRIM': return re.sub(' +',' ',str(v[0]).strip(' '))
        if name=='CLEAN': return ''.join(c for c in str(v[0]) if ord(c)>=32)
        if name=='SUBSTITUTE': return str(v[0]).replace(str(v[1]),str(v[2]))
        if name=='CHAR': return chr(int(v[0]))
        if name=='INDEX': return v[0][int(v[1])-1]
        if name=='MATCH':
            assert v[2]==0
            return next(i for i,x in enumerate(v[1],1) if eq(x,v[0]))
        if name=='COUNTIF': return sum(eq(x,v[1]) for x in v[0])
        if name=='COUNTIFS': return sum(all(eq(v[j][i],v[j+1]) for j in range(0,len(v),2)) for i in range(len(v[0])))
        raise ValueError(('unsupported function',name))
