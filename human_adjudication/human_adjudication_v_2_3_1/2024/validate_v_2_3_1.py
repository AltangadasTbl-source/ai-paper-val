#!/usr/bin/env python3
"""Construction-only static validation and bounded formula regression probes.
No Excel/compatible calculation engine is present: those checks remain NOT RUN.
"""
import argparse, copy, json, re, sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from build_v_2_3_1 import HERE,BASE,ANNUAL,SOURCES,OPTIONS,SCHEMA,sha,lookup,hint
sys.path.insert(0,str(HERE/'tests'))
from formula_probe import Probe,parse
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RESULTS=[]
def check(name,condition,detail=None):
    RESULTS.append({'check':name,'result':'PASS' if condition else 'FAIL',**({'detail':detail} if detail is not None else {})})
    if not condition:raise AssertionError(name+': '+str(detail))
def values(s):return list(s.values)

def static(path,role,baseline,catalog):
    w=openpyxl.load_workbook(path);cs=[s for s in w if s.title.startswith('C_')];keys=[s.cell(r,1).value for s in cs for r in range(2,s.max_row+1)]
    check(role+' totals',len(cs)==10 and len(keys)==len(set(keys))==91)
    check(role+' baseline key order',keys==[s.cell(r,1).value for s in baseline if s.title.startswith('C_') for r in range(2,s.max_row+1)])
    check(role+' required sheets',set(w.sheetnames)==set(['Start','Setup','Checks','Papers','Decisions','Paper_counts','Protocol_plan_table','Codebook','Rule_defaults','Manifest']+SOURCES+[s.title for s in cs]))
    for name in SOURCES:
        check(role+' frozen '+name,values(w[name])==values(baseline[name]))
        rows=values(catalog[name]);ki=rows[0].index('key')
        check(role+' catalog '+name,values(w[name])==[rows[0]]+[r for r in rows[1:] if r[ki].startswith('2024::')])
    check(role+' blank identity',w['Setup']['B2'].value is None and w['Setup']['B3'].value is None and w['Setup']['B4'].value==role)
    links=0;forms=[]
    for s in w:
        check(role+' unprotected '+s.title,not s.protection.sheet and not s.protection.password)
        for row in s:
            for c in row:
                assert c.protection.locked is False and c.protection.hidden is False
                if c.hyperlink:
                    target=c.hyperlink.target;links+=1
                    if target.startswith('#'):
                        sn,ref=target[1:].split('!');assert sn.strip("'") in w;assert w[sn.strip("'")][ref].value is not None
                    else:assert (path.parent/target.split('#')[0]).resolve().is_file(),target
                if c.data_type=='f':
                    forms.append(c.value);parse(c.value)
                    assert '#REF!' not in c.value and '[' not in c.value and '\\' not in c.value
                    assert len(c.value)<=8192
                    for token in Tokenizer(c.value).items:
                        if token.subtype=='RANGE' and '!' in token.value:assert token.value.rsplit('!',1)[0].strip("'") in w
                if c.data_type=='e':raise AssertionError('Error cell '+s.title+'!'+c.coordinate)
                # Frozen sources legitimately include paths named checkers and narrative words.
                if s.title not in SOURCES and isinstance(c.value,str):assert not re.search(r'checker|snapshot|consensus|FINAL READY|stage switch',c.value,re.I),(s.title,c.coordinate)
                assert ('R_B' if role=='R_A' else 'R_A') not in str(c.value)
    check(role+' links and formulas',True,{'hyperlinks_checked':links,'formulas_parsed':len(forms),'maximum_formula_length':max(map(len,forms))})
    check(role+' no names or external links',not w.defined_names and not w._external_links)
    for s in cs:
        bs=baseline[s.title]
        assert s.max_column==18 and s.freeze_panes==bs.freeze_panes
        for r in range(2,s.max_row+1):
            assert [s.cell(r,c).value for c in [1,2,3,5]]==[bs.cell(r,c).value for c in [1,2,3,5]]
            assert s[f'D{r}'].value is None and s[f'R{r}'].value is None and s[f'F{r}'].value=='no'
            for c in 'HIJKL':assert s[f'{c}{r}'].data_type=='f'
        for c in 'MNOPQ':assert all(s[f'{c}{r}'].value is None for r in range(1,s.max_row+1)) and s.column_dimensions[c].width==3
        assert len(s.data_validations.dataValidation)==8 and len(s.conditional_formatting)==6
        for dv in s.data_validations.dataValidation:
            assert not dv.showDropDown and dv.showErrorMessage
            assert not any(c in str(dv.sqref) for c in 'MNOPQ')
            assert len(dv.formula1)<256
        assert s.auto_filter.ref==f'A1:R{s.max_row}'
    check(role+' blank inputs / layout / dropdowns / conditional formatting',True)
    # Every preserved methodological option definition remains exact (discuss has an appended workflow clarification).
    old=json.loads((BASE/'maintenance/codebook.json').read_text()); retained=[]
    for row in old:
        if row[0] and str(row[0]).startswith('5  |'):break
        if row[2]=='Option' and row[1] not in ['no','discuss']:retained.append(row)
    new=values(w['Codebook'])
    assert all(tuple(row) in new for row in retained)
    check(role+' methodological definitions preserved',True)
    check(role+' auto recalculation',w.calculation.calcMode=='auto' and w.calculation.fullCalcOnLoad and w.calculation.forceFullCalc)
    with ZipFile(path) as z:
        forbidden=[n for n in z.namelist() if any(x in n.lower() for x in ['externallink','connection','querytable','vbaproject'])]
        check(role+' no runtime dependencies',not forbidden,forbidden)
        styles=ET.fromstring(z.read('xl/styles.xml'))
        for family in ['cellXfs','cellStyleXfs']:
            for xf in styles.find('s:'+family,NS):
                p=xf.find('s:protection',NS);assert p is not None and p.get('locked')=='0' and p.get('hidden')=='0'
        for n in z.namelist():
            if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'):assert ET.fromstring(z.read(n)).find('s:sheetProtection',NS) is None
        assert ET.fromstring(z.read('xl/workbook.xml')).find('s:workbookProtection',NS) is None
    check(role+' serialized styles including default unlocked',True,{'protected_ranges':[]})
    p=Probe(w)
    check(role+' blank template formula probe',p.cell('Checks','B3')==0 and p.cell('Checks','B4')==91 and p.cell('Checks','B5')==0 and p.cell('Checks','B6')=='ACTION NEEDED')
    for r in range(2,93):assert p.cell('Decisions',f'H{r}')==''
    for tab in w:
        for row in tab:
            for cell in row:
                if cell.data_type=='f':p.cell(tab.title,cell.coordinate)
    check(role+' all template formulas traversed by bounded probe',True)
    return w

def scenarios(template):
    w=openpyxl.load_workbook(template);s=next(s for s in w if s.title.startswith('C_'));name=s.title
    w['Setup']['B2']='TEST_REVIEWER'
    # Exercise untouched formulas through every domain/disposition branch.
    for domain in OPTIONS['E']:
        for disposition in [None,'retain','exclude','discuss']:
            s['D2']=disposition;s['E2']=domain
            p=Probe(w)
            applicable=set('HIJ' if domain=='statistical_results' else 'K' if domain in ['plan_material','report_internal_inconsistency'] else 'L')
            for col in 'HIJKL':
                expect='' if disposition in [None,'retain'] and col in applicable else 'not_applicable'
                assert p.cell(name,col+'2')==expect
    check('probe untouched applicability defaults across all domains/dispositions',True)
    s['D2']=None;s['E2']='statistical_results'
    original={c:s[c].value for c in ['A2','D2','E2','F2','H2','I2','J2','K2','L2','R2']}
    def run(label,changes,expected,outcome=None,identity='TEST_REVIEWER'):
        for c,v in original.items():s[c]=v
        w['Setup']['B2']=identity
        for c,v in {'D2':'retain','E2':'statistical_results','F2':'primary_done','H2':'confirmed_issue','I2':'substantive','J2':None,'K2':None,'L2':None,**changes}.items():s[c]=v
        p=Probe(w);actual=p.cell(name,'G2')
        check('probe '+label,actual==expected,{'expected':expected,'actual':actual})
        if outcome is not None:check('outcome '+label,p.cell('Decisions','H2')==outcome,{'expected':outcome,'actual':p.cell('Decisions','H2')})
    run('statistical confirmed',{},'PRIMARY READY','primary_temporal')
    run('statistical probable',{'H2':'probable_issue'},'PRIMARY READY','primary_temporal')
    run('minor',{'I2':'minor_presentation'},'PRIMARY READY','not_eligible')
    run('indeterminate issue',{'H2':'indeterminate'},'PRIMARY READY','not_eligible')
    run('indeterminate substance',{'I2':'indeterminate'},'PRIMARY READY','not_eligible')
    run('not issue',{'H2':'not_issue'},'PRIMARY READY','not_eligible')
    for col in ['H','I']:
        for value in [None,'not_applicable','invalid',' confirmed_issue ']:run('required '+col+' '+str(value),{col+'2':value},'CHECK FIELDS','')
    for col,domain in [('J','statistical_results'),('L','protocol_report')]:
        base={} if col=='J' else {'E2':domain,'H2':None,'I2':None}
        for value in [None]+OPTIONS[col]:run('optional '+col+' '+str(value),{**base,col+'2':value},'PRIMARY READY')
        for value in ['not_applicable','invalid']:run('optional invalid '+col+' '+value,{**base,col+'2':value},'CHECK FIELDS','')
    run('report internal',{'E2':'report_internal_inconsistency','H2':None,'I2':None,'K2':'substantive_inconsistency'},'PRIMARY READY','secondary_temporal')
    run('plan material',{'E2':'plan_material','H2':None,'I2':None,'K2':'substantive_inconsistency'},'PRIMARY READY','not_eligible')
    for domain in ['report_internal_inconsistency','plan_material']:
        for v in [None,'not_applicable','invalid']:run('required K '+domain+' '+str(v),{'E2':domain,'H2':None,'I2':None,'K2':v},'CHECK FIELDS','')
    run('protocol optional absent',{'E2':'protocol_report','H2':None,'I2':None},'PRIMARY READY','not_eligible')
    for disposition in ['exclude','discuss']:
        run(disposition,{'D2':disposition,'H2':None,'I2':None},'PRIMARY READY','not_eligible' if disposition=='exclude' else '')
        run(disposition+' requires E',{'D2':disposition,'E2':None,'H2':None,'I2':None},'CHECK FIELDS','')
        run(disposition+' stale',{'D2':disposition},'CLEAR INACTIVE VALUES','')
    for identity in [None,'','   ','\t\n','\u00a0']:run('identity '+repr(identity),{},'ENTER REVIEWER ID','',identity)
    for col in ['D','E']:run('invalid pasted '+col,{col+'2':'invalid','H2':None,'I2':None},'CHECK FIELDS','')
    run('missing D',{'D2':None,'H2':None,'I2':None},'CHECK FIELDS','')
    run('inactive K stale',{'K2':'indeterminate'},'CLEAR INACTIVE VALUES','')
    run('domain change stale',{'E2':'protocol_report'},'CLEAR INACTIVE VALUES','')
    run('domain changed cleared',{'E2':'protocol_report','H2':None,'I2':None},'PRIMARY READY')
    run('manual reset',{'F2':'no'},'PRIMARY DRAFT','')
    run('blank complete',{'F2':None},'PRIMARY DRAFT','')
    for flag in ['checker_done','invalid','yes']:run('invalid complete '+flag,{'F2':flag},'CHECK COMPLETE','')
    run('missing key',{'A2':None},'CHECK KEY','')
    run('unknown key',{'A2':'2024::unknown::C001'},'CHECK KEY','')
    run('duplicate key',{'A2':s['A3'].value},'CHECK KEY','')
    # Cross-paper key must not become a valid row in the wrong paper.
    other=next(t for t in w if t.title.startswith('C_') and t.title!=name)
    run('key from different paper',{'A2':other['A2'].value},'CHECK KEY','')
    for rel,expect in [('P_R','protocol_descriptive'),('P_P','plan_descriptive'),('S_C','primary_temporal'),('S_W','primary_temporal')]:
        w['Decisions']['D2']=rel;run('relationship precedence '+rel,{},'PRIMARY READY',expect)
    # Fresh copy: every row complete, using each frozen suggested domain.
    w=openpyxl.load_workbook(template);w['Setup']['B2']='TEST_ALL_PRIMARY';counts={}
    for s in w:
        if not s.title.startswith('C_'):continue
        for r in range(2,s.max_row+1):
            s[f'D{r}']='retain';s[f'F{r}']='primary_done';domain=s[f'E{r}'].value
            if domain=='statistical_results':s[f'H{r}']='confirmed_issue';s[f'I{r}']='substantive'
            elif domain in ['plan_material','report_internal_inconsistency']:s[f'K{r}']='substantive_inconsistency'
    p=Probe(w)
    check('probe all 91 primary ready',p.cell('Checks','B3')==91 and p.cell('Checks','B4')==0 and p.cell('Checks','B6')=='PRIMARY READY')
    for r in range(2,12):
        assert p.cell('Paper_counts',f'F{r}')=='yes' and p.cell('Paper_counts',f'G{r}')=='yes'
        assert sum(p.cell('Paper_counts',f'{c}{r}') for c in 'HIJKL')==w['Paper_counts'][f'C{r}'].value
    for r in range(2,93):
        outcome=p.cell('Decisions',f'H{r}');counts[outcome]=counts.get(outcome,0)+1
    for tab in w:
        for row in tab:
            for cell in row:
                if cell.data_type=='f':p.cell(tab.title,cell.coordinate)
    check('probe all completed formulas including descriptive table',True)
    for r in range(2,w['Protocol_plan_table'].max_row+1):
        assert p.cell('Protocol_plan_table',f'G{r}')=='PRIMARY READY'
        assert p.cell('Protocol_plan_table',f'L{r}') in ['protocol_descriptive','plan_descriptive']
    check('probe protocol/plan status and outcomes',True)
    check('probe all-paper outcome totals reconcile',sum(counts.values())==91,counts)
    w['Start']['B2']='TEST ONLY — SYNTHETIC ANSWERS — NEVER DELIVER'
    fixture=HERE/'tests/all_primary_complete.xlsx';w.save(fixture)
    reopened=openpyxl.load_workbook(fixture)
    check('probe saved / reopened synthetic values and formulas',Probe(reopened).cell('Checks','B3')==91)
    # Count gate remains withheld with discuss, while independent entry is complete.
    s=next(s for s in w if s.title.startswith('C_'));s['D2']='discuss'
    for c in 'HIJKL':s[c+'2']=None
    p=Probe(w)
    check('probe discuss completion vs paper gate',p.cell('Checks','B6')=='PRIMARY READY' and p.cell('Checks','B5')==1 and p.cell('Checks','B7')==1 and p.cell('Paper_counts','F2')=='yes' and p.cell('Paper_counts','G2')=='no' and all(p.cell('Paper_counts',c+'2')=='' for c in 'HIJKL') and p.cell('Decisions','H2')=='')
    w.save(HERE/'tests/discuss_complete.xlsx')
    s['F2']='no';p=Probe(w)
    check('probe unfinished paper gate',p.cell('Paper_counts','F2')=='no' and all(p.cell('Paper_counts',c+'2')=='' for c in 'HIJKL'))
    # Reorder whole rows and translate relative formulas as Excel does during a row sort.
    w=reopened;s=next(s for s in w if s.title.startswith('C_'));before=Probe(w)
    expected={w['Decisions'][f'C{r}'].value:tuple(before.cell('Decisions',f'{c}{r}') for c in 'GHKLMNOPQR') for r in range(2,93)}
    s['R2']='TEST_KEY_NOTE'
    expected[s['A2'].value]=expected[s['A2'].value][:-1]+('TEST_KEY_NOTE',)
    rows=[[(c.value,c.coordinate) for c in row] for row in s.iter_rows(min_row=2)]
    for dest,old in enumerate(reversed(rows),2):
        for col,(v,origin) in enumerate(old,1):
            c=s.cell(dest,col);c.value=Translator(v,origin=origin).translate_formula(c.coordinate) if isinstance(v,str) and v.startswith('=') else v
    p=Probe(w)
    actual={w['Decisions'][f'C{r}'].value:tuple(p.cell('Decisions',f'{c}{r}') for c in 'GHKLMNOPQR') for r in range(2,93)}
    check('probe whole-row reordering key matching',actual==expected)
    w.save(HERE/'tests/whole_row_reordered.xlsx')
    # Isolated one-candidate 2024 fixture tests the n=1 range construction.
    tiny=openpyxl.Workbook();tiny.remove(tiny.active)
    for sn in ['Setup','C_fixture','Decisions']:tiny.create_sheet(sn)
    tiny['Setup']['B2']='TEST_SINGLE';ts=tiny['C_fixture'];ts.append(['key']);ts.append(['2024::fixture::C001',None,None,'retain','statistical_results','primary_done',None,'probable_issue','substantive'])
    tiny['Decisions'].append(['year','paper','key']);tiny['Decisions'].append([2024,'fixture','2024::fixture::C001'])
    ts['G2']=hint(2,1,'C_fixture');tiny['Decisions']['G2']='='+lookup('C_fixture','G','C2',1)
    check('probe single-candidate lookup',Probe(tiny).cell('Decisions','G2')=='PRIMARY READY')
    tiny.save(HERE/'tests/single_candidate.xlsx')
    check('probe optional cleared blank effective result',Probe(openpyxl.load_workbook(fixture)).cell('Decisions','O2')=='')
    return counts

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--year',type=int,required=True);ap.add_argument('--input-dir',default='staging');args=ap.parse_args()
    if args.year!=2024:ap.error('Only 2024 is authorized')
    directory=(HERE/args.input_dir).resolve()
    if not directory.is_relative_to(HERE):ap.error('Input must be inside this year folder')
    baseline=openpyxl.load_workbook(ANNUAL);catalog=openpyxl.load_workbook(BASE/'resources/source_catalog_v_2_1_5_3.xlsx')
    paths=[directory/f'human_adjudication_v_2_3_1_2024_{role}.xlsx' for role in ['R_A','R_B']]
    workbooks=[static(path,role,baseline,catalog) for path,role in zip(paths,['R_A','R_B'])]
    a,b=workbooks
    # Identical templates except role-bearing labels; file mutation cannot affect the other file.
    differences=[]
    for sa in a:
        sb=b[sa.title]
        for row in sa:
            for c in row:
                v=sb[c.coordinate].value
                if c.value!=v:differences.append((sa.title,c.coordinate))
    check('independence template differences restricted to role labels',differences==[('Start','B2'),('Setup','B4'),('Manifest','B5')],differences)
    bhash=sha(paths[1]);a['Setup']['B2']='TEST_A_ONLY';a.save(HERE/'tests/independence_A.xlsx')
    check('independence edit A leaves B untouched',sha(paths[1])==bhash and openpyxl.load_workbook(paths[1])['Setup']['B2'].value is None)
    counts=scenarios(paths[0])
    for item in json.loads((HERE/'source_hash_manifest.json').read_text()):assert sha(Path(item['path']))==item['sha256']
    check('source hashes unchanged after validation',True)
    report={'scope':'2024 only; construction validation','input_dir':str(directory.relative_to(HERE)),'checks':RESULTS,'passed':len(RESULTS),'failed':0,'formula_probe':'Local bounded Python interpreter of generated formulas; not Excel or a compatible spreadsheet engine. No cache or UI validation claimed.','synthetic_outcome_totals':counts,'not_run':['Excel or compatible engine full recalculation and saved formula-cache inspection','Desktop Excel reviewer smoke test, including actual sort, save/close/reopen','Spreadsheet rendering / screenshot comparison, clipping and dropdown usability'],'protected_ranges':[],'outputs':{p.name:sha(p) for p in paths}}
    (directory/'validation_report.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['checks','outputs']},indent=2))
if __name__=='__main__':main()
