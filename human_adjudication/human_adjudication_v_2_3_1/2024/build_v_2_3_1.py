#!/usr/bin/env python3
"""Construction only. Never run against completed reviews or released templates."""
import argparse, copy, hashlib, json, os, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.workbook.properties import CalcProperties
from openpyxl.comments import Comment

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
BASE=ROOT/'human_adjudication/previous/human_adjudication_v_2_1_5_3'
ANNUAL=BASE/'human_adjudication_v_2_1_5_3_2024.xlsx'
SOURCES=['Candidates','Evidence','Links','Relations']
SCHEMA='xlsx-adjudication-independent-primary-2.3.1'
OPTIONS={'D':['retain','exclude','discuss'],'E':['statistical_results','report_internal_inconsistency','plan_material','protocol_report'],'F':['no','primary_done'],'H':['confirmed_issue','probable_issue','not_issue','indeterminate'],'I':['substantive','minor_presentation','indeterminate'],'J':['direct','possible','unlikely','unknown'],'K':['substantive_inconsistency','minor_presentation','non_inconsistency','indeterminate'],'L':['substantive_difference','minor_presentation','version_timing','non_inconsistency','indeterminate']}
NAVY='173F57'; YELLOW='FFF2CC'; BLUE='EAF3FC'; GRAY='EDF2F5'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def quote(s): return '"'+str(s).replace('"','""')+'"'
def sheet(s): return "'"+s.replace("'","''")+"'"
def valid(e,opts): return 'OR('+','.join(f'{e}={quote(o)}' for o in opts)+')'
def active(col,d,e,potential=False):
    disp=f'OR({d}="retain",{d}="")' if potential else f'{d}="retain"'
    domain=f'{e}="statistical_results"' if col in 'HIJ' else f'OR({e}="report_internal_inconsistency",{e}="plan_material")' if col=='K' else f'{e}="protocol_report"'
    return f'AND({disp},{domain})'
def preserve(e): return f'IF({e}="","",{e})'
def lookup(name,col,key,n,keycol='A'):
    return f'INDEX({sheet(name)}!${col}$2:${col}${n+1},MATCH({key},{sheet(name)}!${keycol}$2:${keycol}${n+1},0))'
def hint(r,n,name):
    d,e,f,a=(f'{x}{r}' for x in 'DEFA')
    # Membership is checked against the original paper as well as the frozen key set.
    key=f'OR({a}="",COUNTIF($A$2:$A${n+1},{a})<>1,COUNTIFS(Decisions!$C$2:$C$92,{a},Decisions!$B$2:$B$92,{quote(name[2:])})<>1)'
    errors=[f'NOT({valid(d,OPTIONS["D"])})',f'NOT({valid(e,OPTIONS["E"])})']
    stale=[]
    for col in 'HIJKL':
        v=f'{col}{r}'; act=active(col,d,e)
        stale.append(f'AND(NOT({act}),{v}<>"",{v}<>"not_applicable")')
        optional=f'{v}<>"",' if col in 'JL' else ''
        errors.append(f'AND({act},{optional}NOT({valid(v,OPTIONS[col])}))')
    # CLEAN removes ASCII tabs/newlines; SUBSTITUTE handles nonbreaking spaces.
    identity='LEN(TRIM(CLEAN(SUBSTITUTE(Setup!$B$2,CHAR(160)," "))))=0'
    branches=[(identity,'ENTER REVIEWER ID'),(key,'CHECK KEY'),('OR('+','.join(stale)+')','CLEAR INACTIVE VALUES'),('OR('+','.join(errors)+')','CHECK FIELDS'),(f'AND({f}<>"",NOT({valid(f,OPTIONS["F"])}))','CHECK COMPLETE'),(f'{f}<>"primary_done"','PRIMARY DRAFT')]
    result='"PRIMARY READY"'
    for condition,msg in reversed(branches): result=f'IF({condition},{quote(msg)},{result})'
    return f'=IFERROR({result},"CHECK KEY / INPUT")'
def table(w,name,rows,widths=None,height=54):
    if name in w: del w[name]
    s=w.create_sheet(name)
    for row in rows: s.append(row)
    for row in s:
        for c in row:
            c.font=Font(name='Calibri',size=11,color=NAVY)
            c.alignment=Alignment(wrap_text=True,vertical='center')
            c.protection=Protection(locked=False,hidden=False)
            if c.row==1:
                c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
            elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F4F8FA')
    for r in range(1,s.max_row+1):s.row_dimensions[r].height=42 if r==1 else height
    for c in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(c)].width=(widths or [])[c-1] if widths and c<=len(widths) else 28
    s.freeze_panes='B2';s.sheet_view.showGridLines=False
    return s

def start_rows(role):return [
 ['Step','Instructions'],['Release',f'v.2.3.1 | 2024 | {role} | 10 papers / 91 candidates | Independent primary review in Excel'],
 ['Coordinator','Assign two different people to the two role files. Distribute only the corresponding file to each person. Role labels are not person IDs. A standalone workbook cannot verify the other person’s identity.'],
 ['1. Begin','Save As a working copy in this release folder, or preserve the same relative source-folder layout. In Setup B2 enter your own reviewer ID/name. B3 optionally records your full name.'],
 ['2. Read evidence','Use Papers to open a C_ paper sheet. Click a Candidate key to reach its embedded Evidence. Candidates, Links and Relations retain the frozen source context. Source-file links require access to the original repository folders; all structured evidence is embedded.'],
 ['3. Assess','Choose D: retain, exclude or discuss. Confirm or change the AI-suggested domain in E. Read Codebook definitions and answer applicable H–L judgments. Add optional English notes in R. M–Q are blank spacers.'],
 ['4. Applicability','For retain: statistical_results requires H and I, with J optional; report_internal_inconsistency and plan_material require K; protocol_report has optional L and no extra required judgment. Other H–L fields are inactive. D and valid E are required even for exclude/discuss.'],
 ['5. Defaults and changes','Untouched H–L formulas show blank when potentially applicable, otherwise not_applicable. Selecting an answer replaces that input formula. Blank is no answer. Changing D/E does not erase literal answers. Clear stale inactive values or select not_applicable; never use not_applicable in an active field.'],
 ['6. Complete a row','After manual review set F=primary_done. Inspect G for PRIMARY READY. A dropdown does not catch every pasted value; G also validates entries. Optional active answers may be cleared to blank.'],
 ['7. Revisit a row','Before changing a completed row, manually set F=no, make changes and review again, then set primary_done. Excel formulas cannot automatically erase an entered completion flag.'],
 ['8. Check and deliver','Inspect Checks and filter Decisions readiness. When every row is PRIMARY READY, save, close and deliver this completed workbook. Reopen to confirm saved entries. Built-in results update in Excel; no scripts, export, refresh or external calculation service are needed.'],
 ['Unresolved assessments','discuss is a completed independent primary response when otherwise valid. It is unresolved scientific uncertainty, not confirmed absence of error. Its outcome stays blank. Checks shows discuss separately; completing data entry does not resolve uncertainty.'],
 ['Results and paper gate','Decisions shows this reviewer’s effective assessment and ready outcomes. Inactive results normalize to not_applicable only after G flags stale inputs. Paper_counts withholds outcome counts until all paper rows are ready AND none is discuss. Completion and unresolved counts are separate.'],
 ['Outcome rules','Retained P_R/P_P are protocol_descriptive/plan_descriptive and never temporal. Retained S_C/S_W confirmed/probable substantive statistical findings are primary_temporal; substantive report-internal inconsistencies are secondary_temporal. Other ready determinate records are not_eligible. Unfinished outcomes stay blank.'],
 ['Editing and backups','All cells, including unused cells, are unlocked. No passwords or unlock step. Edit Setup and D/E/F/H–L/R; preserve source keys, evidence and result formulas. Undo accidental edits or recover from a normal Excel backup. Do not rebuild completed reviews.'],
 ['Limits','Checks validates entries, not scientific accuracy. Fixed rules are visible in Rule_defaults and are unprotected. Formula cells remain editable, so preserve them. Construction validation and pending desktop Excel checks are recorded in VALIDATION.md alongside the release.']]

def build(role,outdir):
    w=openpyxl.load_workbook(ANNUAL)
    rows=list(w['Candidates'].values); candidates=[dict(zip(rows[0],r)) for r in rows[1:]]
    papers=[s.title for s in w if s.title.startswith('C_')]
    assert len(papers)==10 and len(candidates)==91
    # Preserve source and primary layout; replace all legacy workflow/calculation objects.
    for name in list(w.sheetnames):
        if name not in SOURCES and name not in papers: del w[name]
    w.defined_names.clear();w.security=None
    for s in w:
        s.tables.clear(); s.auto_filter.ref=f'A1:{openpyxl.utils.get_column_letter(s.max_column)}{s.max_row}'
        s.data_validations.dataValidation=[]
        s.conditional_formatting._cf_rules.clear()
    table(w,'Start',start_rows(role),[31,120],84)
    setup=table(w,'Setup',[['Field','Value'],['Reviewer ID / name',None],['Full name (optional)',None],['Reviewer role',role],['Year',2024]],[34,70],42)
    for c in ['B2','B3']:setup[c].fill=PatternFill('solid',fgColor=YELLOW)
    table(w,'Checks',[
      ['Check','Result'],['Candidates',91],['Primary ready','=COUNTIF(Decisions!G2:G92,"PRIMARY READY")'],['Unfinished','=B2-B3'],['Discuss / unresolved','=COUNTIF(Decisions!K2:K92,"discuss")'],['Workflow status','=IF(B3=B2,"PRIMARY READY","ACTION NEEDED")'],['Ready with discuss','=COUNTIFS(Decisions!G2:G92,"PRIMARY READY",Decisions!K2:K92,"discuss")'],['Meaning','PRIMARY READY means independent data entry is complete; discuss remains unresolved and prevents that paper’s outcome counts.'],['Next action','Inspect G on C_ sheets or filter Decisions readiness. Save the completed file and deliver it; preserve normal Excel backups.']],[34,105],60)
    ps=table(w,'Papers',[['Paper','Review worksheet','Candidates']]+[[name[2:],name,w[name].max_row-1] for name in papers],[32,34,18],36)
    for r in range(2,ps.max_row+1):ps.cell(r,2).hyperlink=f'#{sheet(ps.cell(r,2).value)}!A1'
    erows={}
    for r in range(2,w['Evidence'].max_row+1):erows.setdefault(w['Evidence'].cell(r,1).value,r)
    for name in papers:
        s=w[name];n=s.max_row-1
        s['E1'].comment=Comment('AI-suggested starting domain; confirm or edit after reviewing the evidence.','Construction')
        # Blank headers preclude an Excel Table here; use an AutoFilter on complete A:R rows.
        for col in 'MNOPQ':
            s.column_dimensions[col].width=3;s.column_dimensions[col].hidden=True
            for r in range(1,n+2):
                s[f'{col}{r}']=None
                s[f'{col}{r}'].comment=None
        s.column_dimensions['E'].width=36
        s.column_dimensions['H'].width=29
        s.column_dimensions['K'].width=36
        for col,opts in OPTIONS.items():
            values=opts+(['not_applicable'] if col in 'HIJKL' else [])
            dv=DataValidation(type='list',formula1='"'+','.join(values)+'"',allow_blank=True)
            dv.showErrorMessage=True;dv.errorStyle='stop';dv.errorTitle='Choose a listed option';dv.error='Use an option from the dropdown; G checks pasted values too.'
            dv.showInputMessage=True;dv.promptTitle='Independent primary input';dv.prompt='Confirm suggested E. Clear inactive answers after changes. Reset F=no before revising.'
            dv.showDropDown=False;s.add_data_validation(dv);dv.add(f'{col}2:{col}{n+1}')
        for r in range(2,n+2):
            s[f'D{r}']=None;s[f'F{r}']='no';s[f'R{r}']=None;s[f'G{r}']=hint(r,n,name)
            for col in 'HIJKL':s[f'{col}{r}']=f'=IF({active(col,f"D{r}",f"E{r}",True)},"","not_applicable")'
            s[f'A{r}'].hyperlink=f"#'Evidence'!A{erows[s[f'A{r}'].value]}"
            for col in 'DEFHIK':s[f'{col}{r}'].fill=PatternFill('solid',fgColor=YELLOW)
            for col in 'JLR':s[f'{col}{r}'].fill=PatternFill('solid',fgColor=BLUE)
        for col in 'HIJKL':
            s.conditional_formatting.add(f'{col}2:{col}{n+1}',FormulaRule(formula=[f'NOT({active(col,"$D2","$E2",True)})'],fill=PatternFill('solid',fgColor=GRAY),font=Font(color='687783')))
        s.conditional_formatting.add(f'G2:G{n+1}',FormulaRule(formula=['G2<>"PRIMARY READY"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color='9C0006')))
    headers=['year','paper','candidate_key','relationship','reviewer_role','reviewer_id','readiness','outcome','ai_disposition','ai_domain','disposition','domain','statistical_issue','substantiveness','downstream_impact','internal_inconsistency','protocol_interpretation','notes']
    ds=table(w,'Decisions',[headers]+[[2024,c['package'],c['key'],c['primary']]+[None]*14 for c in candidates],[12,26,44,18,20,28,28,28,22,36,20,36,28,24,24,32,30,50],75)
    for r,c in enumerate(candidates,2):
        name='C_'+c['package']; n=w[name].max_row-1
        key=f'C{r}'
        def get(col):return lookup(name,col,key,n)
        ds[f'E{r}']='=Setup!$B$4';ds[f'F{r}']='='+preserve('Setup!$B$2')
        ds[f'G{r}']=f'=IFERROR({get("G")},"CHECK KEY")'
        ds[f'I{r}']=c['ai_disposition'];ds[f'J{r}']=c['ai_domain']
        for dest,col in zip('KLMNOPQR','DEHIJKLR'):
            expr=preserve(get(col))
            if col in 'HIJKL':expr=f'IF({active(col,f"K{r}",f"L{r}")},{expr},"not_applicable")'
            ds[f'{dest}{r}']=f'=IFERROR({expr},"")'
        ds[f'H{r}']=f'=IF(OR(G{r}<>"PRIMARY READY",K{r}="discuss"),"",IF(K{r}<>"retain","not_eligible",IF(D{r}="P_R","protocol_descriptive",IF(D{r}="P_P","plan_descriptive",IF(AND(OR(D{r}="S_C",D{r}="S_W"),L{r}="statistical_results",OR(M{r}="confirmed_issue",M{r}="probable_issue"),N{r}="substantive"),"primary_temporal",IF(AND(OR(D{r}="S_C",D{r}="S_W"),L{r}="report_internal_inconsistency",P{r}="substantive_inconsistency"),"secondary_temporal","not_eligible"))))))'
    ds.auto_filter.ref=f'A1:R{ds.max_row}'
    counts=table(w,'Paper_counts',[['year','paper','candidates','unfinished','discuss_unresolved','primary_complete','outcomes_available','primary_temporal','secondary_temporal','protocol_descriptive','plan_descriptive','not_eligible']]+[[2024,name[2:],w[name].max_row-1]+[None]*9 for name in papers],[12,30,18,18,24,24,24,24,26,26,26,22],50)
    for r in range(2,12):
        match=f'Decisions!$B$2:$B$92,B{r}'
        counts[f'D{r}']=f'=C{r}-COUNTIFS({match},Decisions!$G$2:$G$92,"PRIMARY READY")'
        counts[f'E{r}']=f'=COUNTIFS({match},Decisions!$K$2:$K$92,"discuss")'
        counts[f'F{r}']=f'=IF(D{r}=0,"yes","no")'
        counts[f'G{r}']=f'=IF(AND(D{r}=0,E{r}=0),"yes","no")'
        for col,o in zip('HIJKL',['primary_temporal','secondary_temporal','protocol_descriptive','plan_descriptive','not_eligible']):counts[f'{col}{r}']=f'=IF(G{r}<>"yes","",COUNTIFS({match},Decisions!$H$2:$H$92,"{o}"))'
    pc=[c for c in candidates if c['primary'] in ['P_R','P_P']]
    plans=table(w,'Protocol_plan_table',[['year','paper','candidate_key','relationship','title','statement','readiness','disposition','protocol_interpretation','internal_inconsistency','notes','outcome']]+[[2024,c['package'],c['key'],c['primary'],c['title'],c['statement']]+[None]*6 for c in pc],[12,27,44,18,60,90,28,20,30,32,60,28],100)
    for r in range(2,plans.max_row+1):
        for dest,col in zip('GHIJKL','GKQPRH'):plans[f'{dest}{r}']='='+preserve(lookup('Decisions',col,f'C{r}',91,'C'))
    # Preserve original scientific definitions; remove obsolete workflow and person-2 fields.
    old=json.loads((BASE/'maintenance/codebook.json').read_text()); code=[old[0]];skip=False
    overrides={'disposition':'Keep, exclude, or defer as discuss. Blank is untouched. No AI disposition is preselected. discuss is an independent unresolved response, not a shared workflow.', 'reviewComplete':'Use no while drafting and primary_done after manual review. Before revising a completed row reset to no; review again, then mark primary_done. Excel cannot automatically reset an entered flag.', 'readiness':'Calculated entry checks for this reviewer. PRIMARY READY means independent entry complete. discuss may remain unresolved. Checks does not establish scientific accuracy or protect source/formula cells.', 'locations':'Where to check the source claim. Use Candidates, Evidence and Links for context and the original paper files.', 'notes':'Optional English context. Clearing removes the text; blank means no answer.'}
    for row in old[1:]:
        row=list(row)
        if row[0] and str(row[0]).startswith('5  |'):skip=True;continue
        if row[0] and str(row[0]).startswith('6  |'):skip=False;row[0]='5  |  Candidate context'
        if skip or row[4]=='primarySummary' or row[1]=='GUIDE' or row[1] in ['primary_done','checker_done']:continue
        if row[4] in overrides and row[1]=='FIELD':row[3]=overrides[row[4]];row[2]=row[2].replace('each stage','each row').replace('both roles','this reviewer')
        if row[1]=='no':row[3]='Draft. The reviewer has not yet confirmed manual review complete.'
        if row[1]=='discuss':row[3]+=' In this version, it can be PRIMARY READY but remains unresolved; no shared resolution step is required.'
        code.append(row)
    code.append(['Complete','primary_done','Primary review','Human reviewed; independent entry complete. Scientific uncertainty may remain as discuss.','reviewComplete'])
    table(w,'Codebook',code,[36,30,40,100,28],100)
    table(w,'Rule_defaults',[
        ['Domain when D=retain','Required','Optional','Inactive','Default'],
        ['statistical_results','H, I','J','K, L','not_applicable'],
        ['report_internal_inconsistency','K','None','H, I, J, L','not_applicable'],
        ['plan_material','K','None','H, I, J, L','not_applicable'],
        ['protocol_report','None','L','H, I, J, K','not_applicable'],
        ['exclude / discuss','D and E remain required','R notes','H, I, J, K, L','not_applicable'],
        ['Source','WORKFLOW.md sections 5–6; v.2.1.5.3 Codebook',None,None,None],
        ['Protection','None. Fixed rules are a visible reference; all cells remain unlocked.',None,None,None]],[38,54,25,35,26],60)
    manifest=[['field','value'],['version','v.2.3.1'],['schema',SCHEMA],['year',2024],['role',role],['papers / candidates','10 / 91'],['workflow','Independent primary review; all review-time checks and results use internal Excel formulas.'],['protection_ranges','None. All cells and styles, including unused cells, explicitly unlocked; no passwords or workbook structure protection.'],['baseline_protection','Inspected v.2.1.5.3 was entirely unprotected.'],['source','Frozen annual baseline and matching 2024 source-catalog subset; no human answers imported.'],['paper_gate_adaptation','Outcome counts withheld until every paper row is PRIMARY READY and no disposition is discuss. Primary completion is independent of scientific resolution.'],['key_matching','Result lookups use candidate key within the original paper. Whole-row sorting is supported; preserve complete A:R rows and all keys.'],['source_links','Source cell values preserved verbatim. Hyperlink targets separately rebased from repository root for this output directory; keep folder layout when making working copies.'],['columns A-C','Frozen candidate key, title and source location.'],['columns D-G','Disposition; editable AI-suggested domain; manual completion; calculated primary validation hint.'],['columns H-L','Statistical issue; substantiveness; optional downstream impact; internal/planning inconsistency; optional protocol interpretation.'],['columns M-Q','Blank hidden spacers, no inputs or calculations.'],['column R','Optional English notes.'],['validation_limit','No desktop Excel or compatible calculation/rendering engine available during construction. See VALIDATION.md; interactive smoke test pending.']]
    for source in json.loads((HERE/'baseline_inventory.current.json').read_text())['source_files']:manifest.append(['source_sha256: '+Path(source['path']).name,source['sha256']])
    table(w,'Manifest',manifest,[43,112],68)
    # Original href text remains untouched. Links are new OOXML relationships.
    for name,col in [('Links',5),('Candidates',15)]:
        for r in range(2,w[name].max_row+1):
            c=w[name].cell(r,col);href=c.value
            if not href:continue
            path,sep,fragment=href.partition('#')
            assert path.startswith('../../'),href
            target=(ROOT/path[6:]).resolve()
            c.hyperlink=os.path.relpath(target,outdir)+(sep+fragment if sep else '')
    # Candidate table provides direct navigation into embedded evidence.
    for r in range(2,93):w['Candidates'].cell(r,7).hyperlink=f"#'Evidence'!A{erows[w['Candidates'].cell(r,7).value]}"
    order=['Start','Setup','Checks','Papers']+papers+['Decisions','Paper_counts','Protocol_plan_table']+SOURCES+['Codebook','Rule_defaults','Manifest']
    w._sheets=[w[name] for name in order];w.active=0
    w.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    # The default style is critical: unlock even unused cells, not just instantiated inputs.
    unlocked=Protection(locked=False,hidden=False);pid=w._protections.add(unlocked)
    for st in w._cell_styles:st.protectionId=pid
    for ns in w._named_styles:ns.protection=copy.copy(unlocked)
    for s in w:
        s.protection.sheet=False
        for dim in list(s.column_dimensions.values())+list(s.row_dimensions.values()):
            dim.protection=copy.copy(unlocked)
        for row in s:
            for c in row:c.protection=copy.copy(unlocked)
    return w

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--year',type=int,required=True);parser.add_argument('--output-dir',default='staging');args=parser.parse_args()
    if args.year!=2024:parser.error('Only the 2024 cohort is authorized.')
    out=(HERE/args.output_dir).resolve()
    if not out.is_relative_to(HERE) or out==HERE/'release' or (HERE/'release') in out.parents:parser.error('Use a construction directory inside this year; release is never a build target.')
    inventory=json.loads((HERE/'baseline_inventory.json').read_text());current=json.loads((HERE/'baseline_inventory.current.json').read_text())
    if inventory!=current:raise RuntimeError('Preflight inventory changed; investigate before build.')
    for f in current['source_files']:
        if sha(Path(f['path']))!=f['sha256']:raise RuntimeError('Source hash changed: '+f['path'])
    out.mkdir(parents=True,exist_ok=True)
    paths=[out/f'human_adjudication_v_2_3_1_2024_{role}.xlsx' for role in ['R_A','R_B']]
    if any(p.exists() for p in paths):raise FileExistsError('Never overwrite an existing workbook; choose a fresh staging directory.')
    for role,path in zip(['R_A','R_B'],paths):build(role,out).save(path);print(path.relative_to(HERE))
    (HERE/'source_hash_manifest.json').write_text(json.dumps(current['source_files'],indent=2)+'\n')
    (out/'output_hashes.json').write_text(json.dumps({p.name:sha(p) for p in paths},indent=2)+'\n')
if __name__=='__main__':main()
