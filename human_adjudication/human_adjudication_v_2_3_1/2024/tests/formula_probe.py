"""Small construction-only formula probe, NOT an Excel-compatible calculation engine.
Evaluates the exact generated formula subset for regression testing. It does not
validate Excel serialization, UI behavior, recalculation, formatting, or caches.
"""
import re
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries

class Probe:
    def __init__(self,w): self.w=w; self.cache={};self.asts={};self.busy=set()
    def cell(self,s,a):
        a=a.replace('$','');key=(s,a)
        if key in self.cache:return self.cache[key]
        if key in self.busy:raise ValueError('Circular reference '+str(key))
        self.busy.add(key)
        v=self.w[s][a].value
        if isinstance(v,str) and v.startswith('='):
            if v not in self.asts:self.asts[v]=parse(v)
            v=self.eval(self.asts[v],s)
        if v is None:v=''
        self.busy.remove(key);self.cache[key]=v;return v
    def ref(self,v,s):
        if '!' in v:s,v=v.rsplit('!',1);s=s.strip("'").replace("''", "'")
        v=v.replace('$','')
        if ':' not in v:return self.cell(s,v)
        x1,y1,x2,y2=range_boundaries(v)
        return [self.cell(s,self.w[s].cell(y,x).coordinate) for y in range(y1,y2+1) for x in range(x1,x2+1)]
    def eval(self,n,s):
        tag=n[0]
        if tag=='literal':return n[1]
        if tag=='ref':return self.ref(n[1],s)
        if tag=='binary':
            a=self.eval(n[2],s);b=self.eval(n[3],s)
            if n[1]=='=':return same(a,b)
            if n[1]=='<>':return not same(a,b)
            if n[1]=='+':return a+b
            if n[1]=='-':return a-b
            if n[1]=='>':return a>b
            if n[1]=='<':return a<b
            raise ValueError(n)
        fn,args=n[1],n[2]
        if fn=='IF':return self.eval(args[1] if self.eval(args[0],s) else args[2],s)
        if fn=='IFERROR':
            try:return self.eval(args[0],s)
            except (ValueError,KeyError,IndexError):return self.eval(args[1],s)
        a=[self.eval(x,s) for x in args]
        if fn=='OR':return any(a)
        if fn=='AND':return all(a)
        if fn=='NOT':return not a[0]
        if fn=='LEN':return len(str(a[0]))
        if fn=='TRIM':return re.sub(' +',' ',str(a[0]).strip(' '))
        if fn=='CLEAN':return ''.join(c for c in str(a[0]) if ord(c)>=32)
        if fn=='SUBSTITUTE':return str(a[0]).replace(a[1],a[2])
        if fn=='CHAR':return chr(a[0])
        if fn=='COUNTIF':return sum(same(v,a[1]) for v in a[0])
        if fn=='COUNTIFS':return sum(all(same(a[j][i],a[j+1]) for j in range(0,len(a),2)) for i in range(len(a[0])))
        if fn=='MATCH':return match(a[0],a[1])
        if fn=='INDEX':return a[0][a[1]-1]
        raise ValueError('Unsupported '+fn)

def same(a,b):return a.casefold()==b.casefold() if isinstance(a,str) and isinstance(b,str) else a==b

def match(x,arr):
    for i,v in enumerate(arr):
        if same(x,v):return i+1
    raise ValueError('Missing MATCH')

def parse(formula):
    ts=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE'];pos=0
    def expr():
        nonlocal pos
        node=atom()
        while pos<len(ts) and ts[pos].type=='OPERATOR-INFIX':
            op=ts[pos].value;pos+=1;node=('binary',op,node,atom())
        return node
    def atom():
        nonlocal pos
        t=ts[pos];pos+=1
        if t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            while ts[pos].subtype!='CLOSE':
                args.append(expr())
                if ts[pos].type=='SEP':pos+=1
                else:break
            assert ts[pos].subtype=='CLOSE';pos+=1
            return ('func',t.value[:-1].upper(),args)
        if t.type=='OPERAND':
            if t.subtype=='TEXT':return ('literal',t.value[1:-1].replace('""','"'))
            if t.subtype=='NUMBER':return ('literal',int(t.value) if '.' not in t.value else float(t.value))
            if t.subtype=='LOGICAL':return ('literal',t.value=='TRUE')
            if t.subtype=='RANGE':return ('ref',t.value)
        raise ValueError((formula,t))
    result=expr()
    assert pos==len(ts),(formula,pos,len(ts))
    return result
