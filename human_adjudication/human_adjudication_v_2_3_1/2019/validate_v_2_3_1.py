#!/usr/bin/env python3
"""Construction-only static validation of the user-authorized 2019 subset."""
import argparse
from collections import Counter
from copy import deepcopy
import json
from pathlib import Path
import re
import subprocess
import sys
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries
from build_v_2_3_1 import HERE, SOURCE, BASE, EXCLUDED, SOURCE_SHEETS, OPTIONS, SCHEMA, sha, verify_sources

NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REPORT=[]


def check(name, condition, detail=None):
    REPORT.append({'check':name,'passed':bool(condition),'detail':detail})
    if not condition:raise AssertionError(f'{name}: {detail}')


def included(rows):
    rows=list(rows);ki=rows[0].index('key')
    return [rows[0]]+[r for r in rows[1:] if str(r[ki]).startswith('2019::') and ('::'+EXCLUDED+'::') not in str(r[ki])]


def baseline_check():
    verify_sources(); check('All seven source hashes and original archive hash unchanged',True)
    a=json.loads((HERE/'baseline_inventory.json').read_text());b=json.loads((HERE/'baseline_inventory.current.json').read_text())
    for v in (a,b):
        for f in v['source_files']:f['path']=f['path'].split('/human_adjudication_v_2_1_5_3/')[-1]
    check('Preflight matches supplied historical inventory after source-path relocation',a==b)
    return load_workbook(BASE),load_workbook(SOURCE/'resources/source_catalog_v_2_1_5_3.xlsx',read_only=True)


def inspect(path,baseline,catalog):
    role='R_A' if path.stem.endswith('R_A') else 'R_B'
    wb=load_workbook(path);cnames=[n for n in wb.sheetnames if n.startswith('C_')]
    expected=[n for n in baseline.sheetnames if n.startswith('C_') and n!='C_'+EXCLUDED]
    check(role+' paper order / total',cnames==expected and len(cnames)==9)
    keys=[row[0].value for n in cnames for row in list(wb[n])[1:]]
    source_keys=[r[6] for r in included(baseline['Candidates'].values)[1:]]
    check(role+' 62 unique keys in source order',keys==source_keys and len(set(keys))==62)
    check(role+' blank reviewer identity',wb['Setup']['B2'].value is None and wb['Setup']['B3'].value is None)
    check(role+' role/year/schema',wb['Setup']['B4'].value==role and wb['Setup']['B5'].value==2019 and wb['Setup']['B6'].value==SCHEMA)
    check(role+' expected sheets',set(wb.sheetnames)==set(['Start','Setup','Checks','Papers',*expected,'Decisions','Paper_counts','Protocol_plan_table',*SOURCE_SHEETS,'Codebook','Rule_defaults','Manifest']))
    check(role+' no names',not wb.defined_names and not any(w.defined_names for w in wb))
    check(role+' auto/full recalculation',wb.calculation.calcMode=='auto' and wb.calculation.fullCalcOnLoad and wb.calculation.forceFullCalc)
    for name in SOURCE_SHEETS:
        values=list(wb[name].values)
        check(role+' '+name+' baseline cell equality',values==included(baseline[name].values))
        check(role+' '+name+' catalog subset equality',values==included(catalog[name].values))
    blank=True;defaults=True;layout=True
    for name in cnames:
        s=wb[name];b=baseline[name]
        for r in range(2,s.max_row+1):
            blank &= s[f'D{r}'].value is None and s[f'R{r}'].value is None and s[f'F{r}'].value=='no'
            defaults &= all(s.cell(r,col).value==b.cell(r,col).value for col in (1,2,3,5,8,9,10,11,12))
            check(role+' '+name+f' row {r} all editable',all(not c.protection.locked and not c.protection.hidden for c in s[r]))
        check(role+' '+name+' empty spacers',all(s.cell(r,col).value is None and s.cell(r,col).hyperlink is None for r in range(1,s.max_row+1) for col in range(13,18)))
        check(role+' '+name+' spacer widths',all(s.column_dimensions[c].hidden and s.column_dimensions[c].width<=3 for c in 'MNOPQ'))
        layout &= s.max_column==18 and s.freeze_panes==b.freeze_panes and s.auto_filter.ref==f'A1:R{s.max_row}'
        dvs=s.data_validations.dataValidation
        check(role+' '+name+' eight editable dropdowns',len(dvs)==8)
        for col,options in OPTIONS.items():
            dv=next((v for v in dvs if f'{col}2' in v),None)
            expected_formula='"'+','.join(options+(['not_applicable'] if col in 'HIJKL' else []))+'"'
            check(role+' '+name+' dropdown '+col,dv is not None and dv.formula1==expected_formula and not dv.showDropDown and dv.errorStyle=='stop' and dv.showErrorMessage)
        check(role+' '+name+' conditional applicability',len(s.conditional_formatting)==7)
        for k,rules in s.conditional_formatting._cf_rules.items():
            check(role+' CF no spacer rules '+str(k),not any(c in k.sqref for c in ['M2','N2','O2','P2','Q2']))
    check(role+' blank human inputs',blank)
    check(role+' A-C, E suggestions and H-L formulas exact baseline preservation',defaults)
    check(role+' 18-column layout and frozen panes',layout)
    check(role+' primary field colors',all(wb[n]['D2'].fill.fgColor.rgb[-6:]=='FFF2CC' and wb[n]['J2'].fill.fgColor.rgb[-6:]=='EAF3FC' and wb[n]['A1'].fill.fgColor.rgb[-6:]=='173F57' for n in cnames))
    check(role+' headers wrapped and readable dropdown width',all(wb[n]['E1'].alignment.wrap_text and wb[n].column_dimensions['E'].width>=40 for n in cnames))
    rawcodes=json.loads((SOURCE/'maintenance/codebook.json').read_text()); code=list(wb['Codebook'].values)
    primary_options=[]
    for row in rawcodes:
        if row[0]=='5  |  Second-person checking and discussion':break
        if row[2]=='Option' and row[1]!='no':primary_options.append(tuple(row))
    check(role+' all baseline primary option definitions preserved verbatim',all(r in code for r in primary_options))
    check(role+' plan descriptions and relationship precedence source',[(r[2],r[3],r[4],r[5]) for r in list(wb['Protocol_plan_table'].values)[1:]]==[(r[6],r[12],r[22],r[18]) for r in included(baseline['Candidates'].values)[1:] if r[12] in ('P_R','P_P')])
    formulas=[];links=[];functions=set();missing=[]
    forbidden=re.compile(r'Primary_snapshot|checker_done|FINAL READY|checkStatus|checkerId|checkerName|primarySummary|resolutionStatus|resolutionNote|agree_clarification',re.I)
    for s in wb:
        check(role+' unprotected '+s.title,not s.protection.sheet)
        for row in s:
            for c in row:
                value=str(c.value or '')
                checkname=f'{s.title}!{c.coordinate}'
                if EXCLUDED in value:missing.append('Excluded identifier '+checkname)
                if forbidden.search(value):missing.append('Obsolete workflow '+checkname)
                if any(x in value for x in ('#REF!','#NAME?','#DIV/0!')) or c.data_type=='e':missing.append('Error cell '+checkname)
                if len(value)>32767:missing.append('Text overflow '+checkname)
                if c.protection.locked or c.protection.hidden:missing.append('Locked cell '+checkname)
                if c.data_type=='f':
                    formulas.append((s.title,c.coordinate,c.value))
                    if len(c.value)>8192:missing.append('Formula length '+checkname)
                    if '[' in c.value or '_xlfn' in c.value or '\\' in c.value:missing.append('External/unsupported formula '+checkname)
                    for t in Tokenizer(c.value).items:
                        if t.type=='FUNC' and t.subtype=='OPEN':functions.add(t.value[:-1])
                        if t.type=='OPERAND' and t.subtype=='RANGE':
                            token=t.value
                            if '!' in token:
                                sheet,ref=token.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
                            else:sheet,ref=s.title,token
                            try:
                                assert sheet in wb
                                x1,y1,x2,y2=range_boundaries(ref);assert all(v is not None for v in (x1,y1,x2,y2))
                                assert 1<=x1<=x2<=16384 and 1<=y1<=y2<=1048576
                            except Exception:missing.append('Broken reference '+checkname+' '+token)
                if c.hyperlink:
                    h=c.hyperlink
                    if h.location:
                        target_sheet,cell=h.location.rsplit('!',1);target_sheet=target_sheet.strip("'")
                        if target_sheet not in wb:missing.append('Internal link '+checkname)
                        else:
                            try:assert wb[target_sheet][cell].value is not None
                            except Exception:missing.append('Empty link destination '+checkname)
                    elif h.target:
                        target=urlsplit(h.target)
                        if target.scheme not in ('http','https') and not (path.parent/unquote(target.path)).resolve().is_file():missing.append('Source link '+checkname)
                        if EXCLUDED in h.target:missing.append('Excluded link '+checkname)
                    links.append(checkname)
    check(role+' all content/references/links and editable cells valid',not missing,missing)
    with ZipFile(path) as z:
        entries=z.namelist()
        check(role+' no external workbook links, macros, queries or connections',not any(re.search(r'externalLinks/|vbaProject|connections.xml|queryTables/',n) for n in entries))
        styles=ET.fromstring(z.read('xl/styles.xml'))
        xfs=styles.findall('m:cellXfs/m:xf',NS)+styles.findall('m:cellStyleXfs/m:xf',NS)
        check(role+' serialized ALL styles incl unused default unlocked',all(x.get('applyProtection')=='1' and x.find('m:protection',NS) is not None and x.find('m:protection',NS).get('locked')=='0' and x.find('m:protection',NS).get('hidden')=='0' for x in xfs))
        check(role+' no structure lock or password',ET.fromstring(z.read('xl/workbook.xml')).find('m:workbookProtection',NS) is None)
        check(role+' no sheetProtection or stored error values',all(b'sheetProtection' not in z.read(n) and b't="e"' not in z.read(n) for n in entries if re.match(r'xl/worksheets/sheet\d+\.xml$',n)))
        check(role+' exclusion absent from all workbook XML',all(EXCLUDED.encode() not in z.read(n) for n in entries if n.endswith(('.xml','.rels'))))
    check(role+' documented no protection',any(r[0]=='protected_ranges' and r[1].startswith('None.') for r in wb['Manifest'].values))
    return wb,{'file':path.name,'sha256':sha(path),'papers':len(cnames),'candidates':len(keys),'formulas':len(formulas),'max_formula_length':max(len(f[2]) for f in formulas),'hyperlinks':len(links),'excel_functions':sorted(functions),'protected_ranges':[]}


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--year',type=int,required=True);p.add_argument('--input-dir',default='staging');args=p.parse_args()
    if args.year!=2019:p.error('Only 2019 is supported')
    directory=(HERE/args.input_dir).resolve()
    if not directory.is_relative_to(HERE):p.error('Input/report directory must be inside this 2019 folder')
    reportfile=directory/'validation_report.json'
    try:
        baseline,catalog=baseline_check();books=[];details=[]
        for role in ('R_A','R_B'):
            wb,info=inspect(directory/f'human_adjudication_v_2_3_1_2019_{role}.xlsx',baseline,catalog);books.append(wb);details.append(info)
        differences=[]
        for s in books[0]:
            b=books[1][s.title]
            for row in s:
                for c in row:
                    a=c.value;v=b[c.coordinate].value
                    if isinstance(a,str):a=a.replace('R_A','ROLE')
                    if isinstance(v,str):v=v.replace('R_B','ROLE')
                    if a!=v:differences.append(s.title+'!'+c.coordinate)
        check('Two templates identical except their own role labels',not differences,differences)
        original_b=sha(directory/'human_adjudication_v_2_3_1_2019_R_B.xlsx')
        a=books[0];a['Setup']['B2']='TEST_ONLY_A';a[next(n for n in a.sheetnames if n.startswith('C_'))]['D2']='exclude'
        (HERE/'tests').mkdir(exist_ok=True);a.save(HERE/'tests/independence_A.xlsx')
        b=load_workbook(directory/'human_adjudication_v_2_3_1_2019_R_B.xlsx')
        check('Editing saved synthetic A cannot change B',sha(directory/'human_adjudication_v_2_3_1_2019_R_B.xlsx')==original_b and b['Setup']['B2'].value is None)
        verify_sources();check('Source hashes unchanged after validation',True)
        report={'status':'STATIC PASS','scope':'2019 user-authorized subset: 9 papers / 62 candidates','workbooks':details,'passed':len(REPORT),'failed':0,'checks':REPORT,'calculation':'See tests/calculation_report.json; openpyxl itself does not evaluate formulas.','desktop_excel_smoke_test':'NOT RUN','rendering':'NOT RUN','protected_ranges':[]}
        reportfile.write_text(json.dumps(report,indent=2)+'\n')
        print(json.dumps({k:v for k,v in report.items() if k!='checks'},indent=2))
    except Exception as exc:
        reportfile.write_text(json.dumps({'status':'FAIL','checks':REPORT,'error':str(exc)},indent=2)+'\n');raise

if __name__=='__main__':main()
