#!/usr/bin/env python3
"""Construction only. Build the two blank 2019 primary templates, never a release overwrite."""
from __future__ import annotations
import argparse
from copy import copy
import hashlib
import json
import math
import os
from pathlib import Path
import re
import textwrap
from urllib.parse import urlsplit, unquote
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE = HERE / 'sources/human_adjudication_v_2_1_5_3'
BASE = SOURCE / 'human_adjudication_v_2_1_5_3_2019.xlsx'
SCHEMA = 'xlsx-adjudication-independent-primary-2.3.1'
SOURCE_SHEETS = ('Candidates', 'Evidence', 'Links', 'Relations')
EXCLUDED = 'jama.2019.0556'
PAPER_TOTAL, CANDIDATE_TOTAL = 9, 62
NAVY, YELLOW, BLUE, GRAY = '173F57', 'FFF2CC', 'EAF3FC', 'EDF2F5'
OPTIONS = {
    'D': ['retain', 'exclude', 'discuss'],
    'E': ['statistical_results', 'report_internal_inconsistency', 'plan_material', 'protocol_report'],
    'F': ['no', 'primary_done'],
    'H': ['confirmed_issue', 'probable_issue', 'not_issue', 'indeterminate'],
    'I': ['substantive', 'minor_presentation', 'indeterminate'],
    'J': ['direct', 'possible', 'unlikely', 'unknown'],
    'K': ['substantive_inconsistency', 'minor_presentation', 'non_inconsistency', 'indeterminate'],
    'L': ['substantive_difference', 'minor_presentation', 'version_timing', 'non_inconsistency', 'indeterminate'],
}
FIELDS = ['disposition', 'domain', 'statisticalIssueStatus', 'substantiveStatus',
          'downstreamImpact', 'internalStatus', 'protocolInterpretation', 'notes']
OUTCOMES = ['primary_temporal', 'secondary_temporal', 'protocol_descriptive', 'plan_descriptive', 'not_eligible']


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify_sources():
    manifest = json.loads((HERE / 'source_hash_manifest.json').read_text())
    for item in manifest['files']:
        if sha(HERE / item['local_copy']) != item['sha256']:
            raise ValueError('Baseline hash mismatch: ' + item['local_copy'])
    if sha(HERE / manifest['archive']) != manifest['archive_sha256']:
        raise ValueError('Source archive changed; investigate before construction')
    return manifest


def q(value):
    return '"' + str(value).replace('"', '""') + '"'


def sq(sheet):
    return "'" + sheet.replace("'", "''") + "'"


def lookup(sheet, col, key, end, keycol='A'):
    # INDEX/MATCH also works for a one-cell range and for whole-row reordering.
    return f'INDEX({sq(sheet)}!${col}$2:${col}${end},MATCH({key},{sq(sheet)}!${keycol}$2:${keycol}${end},0))'


def preserve(expr):
    return f'IF({expr}="","",{expr})'


def valid(cell, codes):
    # EXACT rejects uppercase or whitespace-padded codes pasted past dropdown validation.
    return 'OR(' + ','.join(f'EXACT({cell},{q(code)})' for code in codes) + ')'


def active(col, d, e, untouched=False):
    disp = f'OR({d}="retain",{d}="")' if untouched else f'{d}="retain"'
    domain = (f'{e}="statistical_results"' if col in 'HIJ' else
              f'OR({e}="report_internal_inconsistency",{e}="plan_material")' if col == 'K' else
              f'{e}="protocol_report"')
    return f'AND({disp},{domain})'


def hint(row, end, paper, role, total=62):
    d, e, f, a = (f'{c}{row}' for c in 'DEFA')
    cleaned = 'TRIM(CLEAN(SUBSTITUTE(Setup!$B$2,CHAR(160)," ")))'
    identity = f'OR(LEN({cleaned})=0,Setup!$B$2<>{cleaned})'
    keybad = f'OR({a}="",COUNTIF($A$2:$A${end},{a})<>1,COUNTIFS(Candidates!$G$2:$G${total+1},{a},Candidates!$K$2:$K${total+1},{q(paper)})<>1)'
    stale, invalid, missing = [], [], []
    for col in 'HIJKL':
        c = f'{col}{row}'; act = active(col, d, e)
        stale.append(f'AND(NOT({act}),{c}<>"",NOT(EXACT({c},"not_applicable")))')
        invalid.append(f'AND({act},{c}<>"",NOT({valid(c,OPTIONS[col])}))')
        if col in 'HIK':
            missing.append(f'AND({act},OR({c}="",EXACT({c},"not_applicable")))')
    checks = [
        (f'OR(NOT(EXACT(Setup!$B$4,{q(role)})),Setup!$B$5<>2019)', 'CHECK SETUP'),
        (identity, 'CHECK REVIEWER ID'),
        (keybad, 'CHECK KEY'),
        (f'NOT({valid(f,OPTIONS["F"]+[""])})', 'CHECK COMPLETE FLAG'),
        (f'OR({d}="",{e}="")', 'MISSING D / E'),
        (f'OR(NOT({valid(d,OPTIONS["D"])}),NOT({valid(e,OPTIONS["E"])}))', 'INVALID D / E'),
        ('OR('+','.join(stale)+')', 'CLEAR INACTIVE VALUES'),
        ('OR('+','.join(missing)+')', 'MISSING REQUIRED JUDGMENT'),
        ('OR('+','.join(invalid)+')', 'INVALID ACTIVE JUDGMENT'),
        (f'{f}<>"primary_done"', 'PRIMARY DRAFT'),
    ]
    expr = '"PRIMARY READY"'
    for cond, message in reversed(checks):
        expr = f'IF({cond},{q(message)},{expr})'
    return f'=IFERROR({expr},"CHECK INPUT ERROR")'


def outcome(row):
    r = row
    return (f'=IF(OR(G{r}<>"PRIMARY READY",K{r}="discuss"),"",'
            f'IF(K{r}<>"retain","not_eligible",IF(D{r}="P_R","protocol_descriptive",'
            f'IF(D{r}="P_P","plan_descriptive",IF(AND(OR(D{r}="S_C",D{r}="S_W"),'
            f'L{r}="statistical_results",OR(M{r}="confirmed_issue",M{r}="probable_issue"),'
            f'N{r}="substantive"),"primary_temporal",IF(AND(OR(D{r}="S_C",D{r}="S_W"),'
            f'L{r}="report_internal_inconsistency",P{r}="substantive_inconsistency"),'
            f'"secondary_temporal","not_eligible"))))))')


def table(wb, name, rows, widths, height=42):
    if name in wb:
        index = wb.sheetnames.index(name); del wb[name]
        ws = wb.create_sheet(name, index)
    else:
        ws = wb.create_sheet(name)
    for row in rows:
        ws.append(row)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws:
        for c in row:
            c.font = Font(name='Calibri', size=11, color=NAVY, bold=c.row==1)
            c.alignment = Alignment(vertical='center', wrap_text=True)
            c.protection = Protection(locked=False, hidden=False)
            if c.row == 1:
                c.fill = PatternFill('solid', fgColor=NAVY); c.font = Font(name='Calibri',size=11,bold=True,color='FFFFFF')
            elif c.row % 2 == 0:
                c.fill = PatternFill('solid', fgColor='F4F8FA')
        ws.row_dimensions[row[0].row].height = 44 if row[0].row==1 else height
    ws.freeze_panes = 'B2'
    ws.sheet_view.showGridLines = False
    return ws


def link(cell, location=None, target=None):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location, target=target)
    f = copy(cell.font); f.color='0563C1'; f.underline='single'; cell.font=f


def source_target(href, output_dir):
    p = urlsplit(href)
    if p.scheme in ('https', 'http'):
        return href, None
    rel = unquote(p.path)
    while rel.startswith('../'):
        rel = rel[3:]
    path = ROOT / rel
    if not path.is_file():
        path = ROOT / 'archived/local_only' / rel
    if not path.is_file():
        raise FileNotFoundError(f'Unresolved source target: {href}')
    target = Path(os.path.relpath(path, output_dir)).as_posix()
    if p.fragment:
        target += '#' + p.fragment
    return target, path


def codebook_rows():
    raw = json.loads((SOURCE / 'maintenance/codebook.json').read_text())
    out=[]; obsolete=False
    for row in raw:
        row=list(row)
        if row[0] == '5  |  Second-person checking and discussion':
            obsolete=True
        if row[0] == '6  |  Candidate context':
            obsolete=False; row[0]='5  |  Candidate context'
        if obsolete or row[4] == 'primarySummary' or row[1] == 'checker_done':
            continue
        if row[0] == 'Working method':
            row[2]='Independent primary review'
        if row[4] == 'disposition' and row[1]=='FIELD':
            row[3]='Keep, exclude, or defer for discussion. Blank means untouched, not discuss. No AI disposition is preselected. Enter your own independent assessment.'
        if row[4] == 'reviewComplete':
            row[2]='Independent primary review'
            row[3]=('Human reviewed; independent primary assessment complete. Scientific uncertainty can remain as discuss.' if row[1]=='primary_done' else
                    'Use no while drafting, then primary_done after manual review. Before changing a completed row, reset F to no, review again, then mark primary_done. Formulas cannot erase a completion flag.')
        if row[1]=='no':
            row[3]='Draft. Independent primary review is not yet confirmed complete; AI suggestions and partially filled fields do not count as completed human review.'
        if row[4] == 'readiness':
            row[3]='Calculated entry checks for this reviewer only. PRIMARY READY requires valid identity, key, fields and primary_done. Discuss can be ready and remains unresolved. Checks do not establish scientific accuracy.'
        if row[4] == 'locations':
            row[2]='Source reference'
            row[3]='Where to check the source claim. Use candidate-key links, Candidates, Evidence and Links for context and original documents.'
        out.append(row)
    out.extend([
        ['Suggested review domain','GUIDE','Column E','The initial domain is the unchanged baseline AI suggestion; confirm or edit it. D remains blank. AI suggestions are not human judgments.',None],
        ['Untouched defaults','GUIDE','H-L','An untouched formula displays blank for a potentially applicable field while D is blank or retain; otherwise not_applicable. Choosing a judgment replaces that formula. Literal answers never clear automatically; clear stale inactive answers.',None],
        ['Discuss and counts','GUIDE','Results','A valid discuss row can be PRIMARY READY but its outcome stays blank. Paper outcome counts stay blank until every candidate is primary-ready and none is discuss. Completion does not mean uncertainty is resolved.',None],
        ['Result eligibility','GUIDE','Results','Ready retained P_R/P_P records are protocol_descriptive/plan_descriptive before domain rules. S_C/S_W confirmed/probable substantive statistical issues are primary_temporal; substantive report-internal inconsistency is secondary_temporal. Other ready non-discuss records are not_eligible. This code is not proof of no error.',None],
        ['Protection','GUIDE','All cells','Protected ranges: none. Every style and cell, including unused cells, is explicitly unlocked. Preserve reference and calculation cells; undo accidental edits or use backups.',None],
    ])
    return out


def start_rows(role):
    return [
        ['Step','Instructions'],
        ['Release',f'v.2.3.1 | 2019 | {role} | 9 papers / 62 candidates | Independent primary review in Excel'],
        ['Coordinator','Assign two different people to the two role files. Distribute only the corresponding file to each person. The role label is not a person ID. This standalone workbook cannot verify the identity assigned to the other file.'],
        ['1. Save a working copy','Use Excel Save As before entering answers. Keep ordinary Excel backups. Setup B2: enter your own reviewer ID or name (required, no leading/trailing whitespace). B3: optional full name. B4/B5 identify this role and year.'],
        ['2. Read candidate and evidence','Start with Papers. Open a C_ paper sheet; click its key for Evidence, or the Candidates rationale if no Evidence rows exist. Filter Candidates, Evidence and Links by key; source links open original files. Structured evidence is embedded. Source PDFs require the repository folders to remain accessible.'],
        ['3. Enter your assessment','Choose D (retain / exclude / discuss), confirm or edit the AI-suggested E domain, answer applicable H-L and add optional English notes in R. Yellow = common/required; blue = optional; gray = inactive. M-Q are empty hidden spacers. Read Codebook for meanings.'],
        ['4. Finish a row','After manual review, set F=primary_done. G must show PRIMARY READY. Use F=no while drafting. Dropdowns help entry; G also checks pasted values. The complete flag alone never makes invalid fields ready.'],
        ['5. Inspect progress','Checks shows total, ready, unfinished and discuss/unresolved rows. Open Decisions and filter readiness for remaining work, then return to the matching C_ sheet. All calculations run in Excel; no refresh, terminal, scripts or export commands.'],
        ['6. Save and deliver','When Checks says PRIMARY READY, save the completed workbook and deliver this file. Built-in Decisions, Paper_counts and Protocol_plan_table update in Excel. Keep your backups; do not regenerate or overwrite completed review files.'],
        ['Changing completed answers','Manually reset F to no BEFORE editing an already completed row. Review the revised row again, then set primary_done. Excel formulas cannot automatically erase an entered completion flag.'],
        ['Defaults / changing branches','Untouched H-L formulas show blank for potentially applicable fields while D is blank or retain, and not_applicable otherwise. Choosing a value overwrites its formula. Changing D/E never clears a literal answer: clear inactive entries or use not_applicable, then answer newly required fields. A blank is blank.'],
        ['Required and optional','Retained statistical_results: H and I required; J optional. Retained report_internal_inconsistency or plan_material: K required. Retained protocol_report: L optional, no extra required judgment. Exclude/discuss: all H-L inactive. Active optional fields may be blank, but may not contain not_applicable.'],
        ['Discuss remains unresolved','Discuss is a completed independent response when the fields are valid and F=primary_done. Its outcome stays blank. Checks highlights the unresolved total. No shared discussion or resolution is required to deliver your independent review.'],
        ['Meaning of paper counts','Paper outcome counts appear only when every candidate in the paper is primary-ready AND no candidate is discuss. Entry completion and unresolved count are shown separately. Blank counts are not zero. Retained P_R/P_P results are descriptive and never temporal comparisons.'],
        ['Editable cells and recovery','All cells are editable; no passwords or unlock steps. Protected ranges: none. Use Setup B2/B3 and C_ D/E/F/H-L/R for review. Preserve keys, source data and formulas. If damaged, Undo or recover your own saved copy. These checks do not validate scientific accuracy.'],
        ['Navigation and sorting','Use Papers links to open review sheets and candidate-key links to read evidence. When sorting, select the entire A:R candidate range including hidden M-Q; never sort one column alone. Results look up candidate keys and follow whole-row reordering.'],
        ['Source access','Original href text is preserved in Candidates and Links; clickable targets are adjusted for this file location. Embedded source and evidence text remains available if original documents are not distributed.'],
    ]


def serialize_unlocked(path):
    """Explicitly unlock ALL style XFs, including default/unused cells; no sheet locks."""
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; tag=lambda s:f'{{{ns}}}{s}'
    with ZipFile(path) as z:
        data={name:z.read(name) for name in z.namelist()}
    root=ET.fromstring(data['xl/styles.xml'])
    for container in ('cellStyleXfs','cellXfs'):
        for xf in root.find(tag(container)):
            for p in list(xf.findall(tag('protection'))): xf.remove(p)
            xf.set('applyProtection','1'); ET.SubElement(xf,tag('protection'),{'locked':'0','hidden':'0'})
    data['xl/styles.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    w=ET.fromstring(data['xl/workbook.xml'])
    for p in list(w.findall(tag('workbookProtection'))):w.remove(p)
    data['xl/workbook.xml']=ET.tostring(w,encoding='utf-8',xml_declaration=True)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for name,content in data.items():z.writestr(name,content)


def build(role, output_dir):
    wb=load_workbook(BASE)
    # User scope correction: omitted package has no review/source/result rows.
    del wb['C_'+EXCLUDED]
    for name in SOURCE_SHEETS:
        ws=wb[name]; headers=list(next(ws.values)); keycol=headers.index('key')+1
        for r in range(ws.max_row,1,-1):
            if ('::'+EXCLUDED+'::') in str(ws.cell(r,keycol).value):ws.delete_rows(r)
    for r in range(wb['Papers'].max_row,1,-1):
        if wb['Papers'].cell(r,1).value==EXCLUDED:wb['Papers'].delete_rows(r)
    candidates=[dict(zip(next(wb['Candidates'].values), row)) for row in list(wb['Candidates'].values)[1:]]
    papers=list(dict.fromkeys(c['package'] for c in candidates)); total=len(candidates)
    assert len(papers)==PAPER_TOTAL and total==CANDIDATE_TOTAL and len({c['key'] for c in candidates})==CANDIDATE_TOTAL
    del wb['Primary_snapshot']; wb.defined_names.clear()
    for ws in wb:
        ws.tables.clear(); ws.defined_names.clear()
    st=table(wb,'Start',start_rows(role),[31,112],height=82)
    st.row_dimensions[2].height=42
    setup=table(wb,'Setup',[
        ['Field','Value'],['Reviewer ID or name (required)',None],['Reviewer name (optional)',None],
        ['Reviewer role',role],['Year',2019],['Schema',SCHEMA],
        ['Identity guidance','Enter your own ID or name in B2. Do not enter the role label as your identity. B3 is optional.'],
    ],[37,90],height=45)
    for c in ('B2','B3'):setup[c].fill=PatternFill('solid',fgColor=YELLOW)
    setup['B2'].comment=Comment('Required identity. Whitespace-only or padded input is invalid.','Review instructions')
    dec=table(wb,'Decisions',[
        ['year','paper','key','relationship','reviewerRole','reviewerId','readiness','outcome',
         'aiDisposition','aiDomain',*FIELDS,'reviewerName','complete','unresolved'],
        *[[2019,c['package'],c['key'],c['primary'],None,None,None,None,c['ai_disposition'],c['ai_domain']]+[None]*11 for c in candidates]
    ],[12,27,44,18,18,26,31,27,20,40,18,40,26,24,24,30,28,60,27,20,18],height=80)
    evidence_row={}
    for r in range(2,wb['Evidence'].max_row+1):evidence_row.setdefault(wb['Evidence'].cell(r,1).value,r)
    destinations={c['key']: (f"'Evidence'!A{evidence_row[c['key']]}" if c['key'] in evidence_row else f"'Candidates'!G{r}") for r,c in enumerate(candidates,2)}
    paper_end={p:1+sum(c['package']==p for c in candidates) for p in papers}
    for paper in papers:
        ws=wb['C_'+paper]; end=ws.max_row
        ws.data_validations=DataValidationList(); ws.conditional_formatting._cf_rules.clear()
        ws.auto_filter.ref=f'A1:R{end}'
        # Preserve all primary source/default values and baseline row order.
        for row in ws:
            for cell in row:
                cell.protection=Protection(locked=False,hidden=False)
        for r in range(1,end+1):
            for col in range(13,18):
                cell=ws.cell(r,col); cell.value=None; cell.comment=None; cell.hyperlink=None
                cell.fill=PatternFill(fill_type=None)
            if r==1:continue
            for col in ('D','R'):ws[f'{col}{r}']=None
            ws[f'F{r}']='no'; ws[f'G{r}']=hint(r,end,paper,role,total)
            link(ws[f'A{r}'],location=destinations[ws[f'A{r}'].value])
            if ws[f'A{r}'].value not in evidence_row:
                ws[f'A{r}'].comment=Comment('The baseline has no separate Evidence rows for this candidate. This link opens its embedded Candidates statement and rationale. Consult Links for original documents.','Source navigation')
            for col in 'DEFHIK':ws[f'{col}{r}'].fill=PatternFill('solid',fgColor=YELLOW)
            for col in 'JLR':ws[f'{col}{r}'].fill=PatternFill('solid',fgColor=BLUE)
            ws[f'E{r}'].comment=Comment('Baseline AI-suggested domain. Confirm or edit; this is not a human decision.','Review instructions')
        ws['E1']='Review domain (AI suggestion; editable)'
        ws['H1']='Statistical/numerical issue'
        for col in 'MNOPQ':ws.column_dimensions[col].width=2;ws.column_dimensions[col].hidden=True
        ws.column_dimensions['E'].width=41;ws.column_dimensions['G'].width=32
        ws.column_dimensions['L'].width=29
        ws.row_dimensions[1].height=58
        # Increase baseline row heights only where long wrapped source text needs it.
        # This conservative character estimate is not a desktop Excel rendering test.
        for r in range(2,end+1):
            lines=max(sum(max(1,len(textwrap.wrap(part,max(5,int(ws.column_dimensions[col].width)-2))))
                          for part in str(ws[f'{col}{r}'].value or '').split('\n')) for col in 'ABC')
            ws.row_dimensions[r].height=max(92,min(409,lines*14+8))
        for col,opts in OPTIONS.items():
            dv=DataValidation(type='list',formula1=q(','.join(opts+(['not_applicable'] if col in 'HIJKL' else []))),allow_blank=True)
            dv.showErrorMessage=True;dv.errorStyle='stop';dv.errorTitle='Choose a listed option'
            dv.error='Use an exact option code. G checks required fields, applicability and pasted values.'
            dv.showInputMessage=True;dv.promptTitle='Independent primary review'
            dv.prompt=('AI suggestion: confirm or edit.' if col=='E' else 'Use Codebook for definitions. Reset F=no before revising a completed row.')
            dv.showDropDown=False;ws.add_data_validation(dv);dv.add(f'{col}2:{col}{end}')
        for col in 'HIJKL':
            ws.conditional_formatting.add(f'{col}2:{col}{end}',FormulaRule(formula=[f'NOT({active(col,"$D2","$E2",True)})'],fill=PatternFill('solid',fgColor=GRAY),font=Font(color='687783')))
        ws.conditional_formatting.add(f'G2:G{end}',FormulaRule(formula=['$G2<>"PRIMARY READY"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color='9C0006')))
        ws.conditional_formatting.add(f'G2:G{end}',FormulaRule(formula=['$G2="PRIMARY READY"'],fill=PatternFill('solid',fgColor='E2F0D9'),font=Font(color='375623')))
        ws.conditional_formatting.add(f'D2:D{end}',FormulaRule(formula=['$D2="discuss"'],fill=PatternFill('solid',fgColor='FCE4D6')))
    for r,c in enumerate(candidates,2):
        name='C_'+c['package'];end=paper_end[c['package']]
        get=lambda col:lookup(name,col,f'C{r}',end)
        dec[f'E{r}']='=Setup!$B$4';dec[f'F{r}']='='+preserve('Setup!$B$2')
        dec[f'G{r}']=f'=IFERROR(IF(COUNTIF({sq(name)}!$A$2:$A${end},C{r})<>1,"CHECK KEY",{get("G")}),"CHECK KEY")'
        for dest,source in zip('KLMNOPQR','DEHIJKLR'):
            expr=preserve(get(source))
            if source in 'HIJKL':expr=f'IF({active(source,f"K{r}",f"L{r}")},{expr},"not_applicable")'
            dec[f'{dest}{r}']=f'=IFERROR({expr},"")'
        dec[f'H{r}']=outcome(r)
        dec[f'S{r}']='='+preserve('Setup!$B$3');dec[f'T{r}']=f'=IFERROR({preserve(get("F"))},"")'
        dec[f'U{r}']=f'=IF(K{r}="discuss","yes","no")'
        link(dec[f'C{r}'],location=f"'{name}'!A1")
    dec.auto_filter.ref=f'A1:U{total+1}'
    counts=table(wb,'Paper_counts',[
        ['year','paper','candidates','primary_ready','unfinished','discuss_unresolved','review_complete','counts_available',*OUTCOMES],
        *[[2019,p,paper_end[p]-1]+[None]*10 for p in papers]
    ],[12,29,18,22,20,25,24,24,25,26,28,26,23],height=54)
    for r,p in enumerate(papers,2):
        match=f'Decisions!$B$2:$B${total+1},B{r}'
        counts[f'D{r}']=f'=COUNTIFS({match},Decisions!$G$2:$G${total+1},"PRIMARY READY")'
        counts[f'E{r}']=f'=C{r}-D{r}'
        counts[f'F{r}']=f'=COUNTIFS({match},Decisions!$K$2:$K${total+1},"discuss")'
        counts[f'G{r}']=f'=IF(D{r}=C{r},"yes","no")'
        counts[f'H{r}']=f'=IF(AND(G{r}="yes",F{r}=0),"yes","no")'
        for col,out in zip('IJKLM',OUTCOMES):
            counts[f'{col}{r}']=f'=IF(H{r}<>"yes","",COUNTIFS({match},Decisions!$H$2:$H${total+1},{q(out)}))'
    plans=[c for c in candidates if c['primary'] in ('P_R','P_P')]
    plan=table(wb,'Protocol_plan_table',[
        ['year','paper','key','relationship','title','statement','readiness','disposition','domain','interpretation','internalStatus','notes','outcome','reviewerRole','reviewerId'],
        *[[2019,c['package'],c['key'],c['primary'],c['title'],c['statement']]+[None]*9 for c in plans]
    ],[12,27,44,18,60,90,31,20,40,29,30,60,28,18,26],height=100)
    for r,c in enumerate(plans,2):
        for dest,source in zip('GHIJKLMNO','GKLQPRHEF'):
            plan[f'{dest}{r}']=f'=IFERROR({preserve(lookup("Decisions",source,f"C{r}",total+1,"C"))},"")'
        link(plan[f'C{r}'],location=destinations[c['key']])
    plan.auto_filter.ref=f'A1:O{plan.max_row}'
    checks=table(wb,'Checks',[
        ['Check','Result'],['Reviewer role','=Setup!$B$4'],['Total candidates',total],
        ['Primary ready',f'=COUNTIF(Decisions!$G$2:$G${total+1},"PRIMARY READY")'],['Unfinished','=B3-B4'],
        ['Discuss / unresolved',f'=COUNTIF(Decisions!$K$2:$K${total+1},"discuss")'],
        ['Workflow status','=IF(B4=B3,"PRIMARY READY","ACTION NEEDED")'],
        ['Ready discuss rows',f'=COUNTIFS(Decisions!$K$2:$K${total+1},"discuss",Decisions!$G$2:$G${total+1},"PRIMARY READY")'],
        ['Papers with counts',f'=COUNTIF(Paper_counts!$H$2:$H${len(papers)+1},"yes")'],
        ['Next action','=IF(B7="PRIMARY READY",IF(B6>0,"Save and deliver this independent review. Discuss rows remain unresolved; affected paper counts are blank.","Save and deliver this independent review with its built-in results."),"Filter Decisions by readiness, return to the C_ sheet, correct entries and manually confirm F.")'],
        ['Scope','Entry checks and independent completion only. Discuss is not a confirmed absence of error. Review the evidence; preserve ordinary Excel backups.'],
    ],[36,108],height=53)
    checks['B6'].fill=PatternFill('solid',fgColor='FCE4D6');checks.row_dimensions[10].height=76
    for r,p in enumerate(papers,2):link(wb['Papers'][f'B{r}'],location=f"'C_{p}'!A1")
    table(wb,'Codebook',codebook_rows(),[35,29,43,99,28],height=113)
    table(wb,'Rule_defaults',[
        ['Retained domain','Required','Optional','Inactive','Inactive effective value'],
        ['statistical_results','H, I','J','K, L','not_applicable'],
        ['report_internal_inconsistency','K','None','H, I, J, L','not_applicable'],
        ['plan_material','K','None','H, I, J, L','not_applicable'],
        ['protocol_report','None','L','H, I, J, K','not_applicable'],
        ['exclude / discuss','D and valid E','R notes','H, I, J, K, L','not_applicable'],
        ['Reference only','Protected ranges: none','All inputs remain editable','WORKFLOW.md sections 5-6','not_applicable'],
    ],[42,29,27,33,28],height=54)
    manifest=table(wb,'Manifest',[
        ['field','value'],['version','v.2.3.1'],['schema',SCHEMA],['year',2019],['reviewer_role',role],
        ['cohort','User-authorized subset of the annual v.2.1.5.3 baseline: 9 papers / 62 candidates.'],
        ['source_baseline','Frozen v.2.1.5.3 annual workbook and catalog, recovered from archived_not_touch.zip. All seven source hashes match supplied inventory.'],
        ['source_sheets','Included Candidates, Evidence, Links and Relations rows preserved exactly, in baseline order. Clickable source targets added separately.'],
        ['protected_ranges','None. All cell styles, cells and unused cells explicitly unlocked; no sheet or workbook protection, passwords or hidden formulas.'],
        ['C_columns','A=key; B=title; C=source location; D=disposition; E=editable AI-suggested domain; F=manual completion; G=readiness; H=issue; I=substantiveness; J=optional impact; K=inconsistency; L=optional protocol interpretation; M-Q=empty hidden spacers; R=optional English notes.'],
        ['Setup_cells','B2=required reviewer ID or name; B3=optional reviewer name; B4=this role; B5=2019; B6=schema.'],
        ['Decisions_columns','A-D=year/paper/key/relationship; E-F=role/ID; G=readiness; H=outcome; I-J=baseline AI suggestions; K-R=effective disposition/domain/H-L/notes; S=name; T=completion; U=unresolved flag.'],
        ['Paper_counts_columns','A-C=year/paper/total; D=ready; E=unfinished; F=discuss; G=entry complete; H=counts available; I-M=primary/secondary temporal, protocol/plan descriptive, not eligible.'],
        ['Protocol_plan_table_columns','A-F=source year/paper/key/relationship/title/statement; G=readiness; H=disposition; I=domain; J=interpretation; K=internalStatus; L=notes; M=outcome; N-O=role/ID.'],
        ['counts_gate_adaptation','v.2.3.1: all paper rows primary-ready AND no discuss before outcome counts appear. Independent completion may include discuss.'],
        ['default_behavior','Original H-L formulas preserved. Selecting an answer replaces the formula. Inactive effective values normalize to not_applicable; stale inactive inputs block readiness first.'],
        ['outcome_precedence','Ready retained P_R/P_P descriptive before S_C/S_W domain eligibility. Discuss and unfinished outcomes blank. Other ready records not_eligible, including indeterminate judgments; not_eligible is not proof of no issue.'],
        ['navigation','Papers to review sheets; review keys to embedded Evidence. Source targets rebased for this output directory.'],
        ['source_evidence_gap','The six jama.2019.14901 candidates have no separate Evidence rows in the frozen baseline. Their keys open embedded Candidates rationale/ambiguity; clickable reportHref opens the original report. Source blanks are preserved.'],
        ['runtime','Excel formulas, dropdowns and conditional formatting only. Automatic full recalculation on opening. No external workbook links or runtime automation.'],
        ['validation_limit','Construction checks are recorded in VALIDATION.md. Desktop Excel calculation, rendering and reviewer smoke test pending; openpyxl does not calculate formulas.'],
    ],[36,116],height=69)
    for item in verify_sources()['files']:
        manifest.append(['sha256 '+item['local_copy'].split('/human_adjudication_v_2_1_5_3/')[-1],item['sha256']])
        r=manifest.max_row;manifest.row_dimensions[r].height=48
        for cell in manifest[r]:cell.alignment=Alignment(wrap_text=True,vertical='center');cell.font=Font(name='Calibri',size=11,color=NAVY)
    source_links=[]
    for name,col in [('Candidates',15),('Links',5)]:
        ws=wb[name]
        for r in range(2,ws.max_row+1):
            cell=ws.cell(r,col)
            if cell.value:
                target,path=source_target(cell.value,output_dir);link(cell,target=target)
                source_links.append({'cell':f'{name}!{cell.coordinate}','href':cell.value,'target':target,'resolved':str(path)})
    for r,c in enumerate(candidates,2):
        if c['key'] in evidence_row:link(wb['Candidates'].cell(r,7),location=destinations[c['key']])
    for name in SOURCE_SHEETS:
        ws=wb[name];ws.auto_filter.ref=ws.dimensions
    for ws in wb:
        ws.protection.sheet=False
        for row in ws:
            for cell in row:cell.protection=Protection(locked=False,hidden=False)
    wb._sheets=[wb[n] for n in ['Start','Setup','Checks','Papers',*['C_'+p for p in papers],'Decisions','Paper_counts','Protocol_plan_table',*SOURCE_SHEETS,'Codebook','Rule_defaults','Manifest']]
    wb.active=0
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    wb.properties.title=f'2019 independent primary review {role} v.2.3.1'
    wb.properties.subject='Blank independent primary assessment template'
    wb.properties.creator='Human adjudication construction';wb.properties.description=SCHEMA
    out=output_dir/f'human_adjudication_v_2_3_1_2019_{role}.xlsx'
    if out.exists():raise FileExistsError('Never overwrite a constructed/review workbook: '+str(out))
    wb.save(out);serialize_unlocked(out)
    return {'file':out.name,'sha256':sha(out),'papers':PAPER_TOTAL,'candidates':CANDIDATE_TOTAL,'role':role,'source_links':source_links,'protected_ranges':[]}


def safe_directory(value):
    out=(HERE/value).resolve()
    if not out.is_relative_to(HERE):raise ValueError('All writes must remain inside the 2019 folder')
    if out==HERE/'release' or (HERE/'release') in out.parents:
        raise ValueError('Builder cannot write release/. Validate staging and promote clean templates separately.')
    for role in ('R_A','R_B'):
        if (HERE/'release'/f'human_adjudication_v_2_3_1_2019_{role}.xlsx').exists():
            raise FileExistsError('A released workbook exists. Corrections require a separately named new version.')
    return out


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--year',type=int,required=True);ap.add_argument('--output-dir',default='staging');args=ap.parse_args()
    if args.year!=2019:ap.error('This builder supports only 2019')
    out=safe_directory(args.output_dir);verify_sources()
    if any(out.glob('*.xlsx')):raise FileExistsError('Output already contains workbooks; never overwrite review files')
    out.mkdir(parents=True,exist_ok=True)
    results=[build(role,out) for role in ('R_A','R_B')]
    (out/'build_manifest.json').write_text(json.dumps(results,indent=2)+'\n')
    print(json.dumps([{k:v for k,v in r.items() if k!='source_links'} for r in results],indent=2))

if __name__=='__main__':main()
