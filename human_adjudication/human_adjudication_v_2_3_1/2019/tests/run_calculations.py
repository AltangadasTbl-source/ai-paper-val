#!/usr/bin/env python3
"""Exercise real workbook formulas with formulas 1.3.4; never modifies release templates.
This interpreter is not desktop Excel and provides no rendering/UI verification.
"""
from __future__ import annotations
from copy import copy
import json
from pathlib import Path
import re
import sys
import time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import formulas
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from build_v_2_3_1 import HERE, OPTIONS, SOURCE_SHEETS, hint, sha

TESTS=HERE/'tests'; CA=[]; started=time.time()
INPUT=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else HERE/'staging'
FILE=INPUT/'human_adjudication_v_2_3_1_2019_R_A.xlsx'


def assert_case(name,actual,expected):
    if hasattr(actual,'item'):actual=actual.item()
    CA.append({'case':name,'actual':actual,'expected':expected,'passed':actual==expected})
    if actual!=expected:raise AssertionError(f'{name}: {actual!r} != {expected!r}')


def node(path,s,c):return f"'[{path.name}]{s.upper()}'!{c}"

def scalar(solution,path,s,c):return solution[node(path,s,c)].value[0,0]


def all_calculate(path,label):
    m=formulas.ExcelModel().loads(str(path)).finish();sol=m.calculate()
    wb=load_workbook(path);errs=[];count=0
    for ws in wb:
        for row in ws:
            for c in row:
                if c.data_type=='f':
                    count+=1;name=node(path,ws.title,c.coordinate)
                    if name not in sol:errs.append(name+' not evaluated');continue
                    v=sol[name].value[0,0]
                    if str(v) in ('#REF!','#VALUE!','#NAME?','#N/A','#DIV/0!','#NUM!','#NULL!'):errs.append(name+' '+str(v))
    assert_case(label+' all formulas evaluated without Excel error values',errs,[])
    return m,sol,wb,count


def save_values(wb,path,solution,model_path):
    # A disposable, explicitly named values-only engine export for save/reopen checks.
    for ws in wb:
        for row in ws:
            for c in row:
                if c.data_type=='f':
                    v=scalar(solution,model_path,ws.title,c.coordinate)
                    if hasattr(v,'item'):v=v.item()
                    c.value=v
    wb.save(path)
    return load_workbook(path,data_only=True)


def fill_all(wb):
    wb['Setup']['B2']='TEST_ONLY_REVIEWER';wb['Setup']['B3']='Synthetic calculation fixture'
    for ws in wb:
        if not ws.title.startswith('C_'):continue
        for r in range(2,ws.max_row+1):
            ws[f'D{r}']='retain';ws[f'F{r}']='primary_done';ws[f'R{r}']='TEST note for '+ws[f'A{r}'].value
            domain=ws[f'E{r}'].value
            if domain=='statistical_results':ws[f'H{r}']='confirmed_issue';ws[f'I{r}']='substantive'
            elif domain in ('report_internal_inconsistency','plan_material'):ws[f'K{r}']='substantive_inconsistency'
    wb['Start']['B2']='TEST ONLY - synthetic answers; never deliver as a review template'


def test_scenarios(model):
    wb=load_workbook(FILE);sheet=next(n for n in wb.sheetnames if n.startswith('C_'));key=wb[sheet]['A2'].value
    fields=['id','A','D','E','F','H','I','J','K','L','relationship']
    inputs=[node(FILE,'Setup','B2')]+[node(FILE,sheet,c+'2') for c in fields[1:-1]]+[node(FILE,'Decisions','D2')]
    outputs=[node(FILE,sheet,'G2'),node(FILE,'Decisions','G2'),node(FILE,'Decisions','H2')]+[node(FILE,'Decisions',c+'2') for c in 'MNOPQ']
    fn=model.compile(inputs,outputs)
    base=dict(zip(fields,['TEST_REVIEWER',key,'retain','statistical_results','primary_done','confirmed_issue','substantive','','','','S_W']))
    def case(name,changes,status='PRIMARY READY',result=None,effective=None):
        vals=dict(base,**changes);ans=[r.value[0,0] for r in fn(*[vals[f] for f in fields])]
        assert_case(name+' readiness',ans[0],status)
        if result is None:result='primary_temporal' if status=='PRIMARY READY' else ''
        assert_case(name+' outcome',ans[2],result)
        if effective is not None:assert_case(name+' effective H-L',ans[3:],effective)
        return ans
    case('confirmed substantive statistical',{})
    case('probable substantive statistical',{'H':'probable_issue'})
    case('cross-document statistical',{'relationship':'S_C'})
    case('minor statistical',{'I':'minor_presentation'},result='not_eligible')
    case('indeterminate issue',{'H':'indeterminate'},result='not_eligible')
    case('indeterminate importance',{'I':'indeterminate'},result='not_eligible')
    case('not an issue',{'H':'not_issue'},result='not_eligible')
    for relationship,result in [('P_R','protocol_descriptive'),('P_P','plan_descriptive')]:
        case(relationship+' precedence over statistical',{'relationship':relationship},result=result)
        case(relationship+' precedence over internal',{'relationship':relationship,'E':'report_internal_inconsistency','H':'','I':'','K':'substantive_inconsistency'},result=result)
    case('retained report inconsistency',{'E':'report_internal_inconsistency','H':'','I':'','K':'substantive_inconsistency'},result='secondary_temporal')
    case('cross-document report inconsistency',{'relationship':'S_C','E':'report_internal_inconsistency','H':'','I':'','K':'substantive_inconsistency'},result='secondary_temporal')
    case('minor report inconsistency',{'E':'report_internal_inconsistency','H':'','I':'','K':'minor_presentation'},result='not_eligible')
    case('retained planning',{'relationship':'P_P','E':'plan_material','H':'','I':'','K':'substantive_inconsistency'},result='plan_descriptive')
    protocol={'relationship':'P_R','E':'protocol_report','H':'','I':''}
    case('protocol optional blank',protocol,result='protocol_descriptive',effective=['not_applicable']*4+[''])
    for v in OPTIONS['J']:
        case('valid optional impact '+v,{'J':v})
    for v in OPTIONS['L']:
        case('valid protocol interpretation '+v,dict(protocol,L=v),result='protocol_descriptive')
    for col in 'HI':
        for v in ('','not_applicable'):
            case('required '+col+' '+repr(v),{col:v},status='MISSING REQUIRED JUDGMENT')
        for v in ('BAD',' confirmed_issue','CONFIRMED_ISSUE'):
            case('invalid required '+col+' '+repr(v),{col:v},status='INVALID ACTIVE JUDGMENT')
    for domain in ('report_internal_inconsistency','plan_material'):
        for v in ('','not_applicable'):
            case(domain+' required blank/NA '+repr(v),{'E':domain,'H':'','I':'','K':v},status='MISSING REQUIRED JUDGMENT')
        case(domain+' invalid K',{'E':domain,'H':'','I':'','K':'BAD'},status='INVALID ACTIVE JUDGMENT')
    for v in ('not_applicable','BAD',' '):
        case('invalid optional J '+repr(v),{'J':v},status='INVALID ACTIVE JUDGMENT')
        case('invalid optional L '+repr(v),dict(protocol,L=v),status='INVALID ACTIVE JUDGMENT')
    for d,out in [('exclude','not_eligible'),('discuss','')]:
        case(d,{'D':d,'H':'','I':''},result=out,effective=['not_applicable']*5)
        for col in 'HIJKL':case(d+' stale '+col,{'D':d,'H':'','I':'',col:OPTIONS[col][0]},status='CLEAR INACTIVE VALUES')
    for ident in ('',' ','   ','\t','\n','\r','\u00a0',' \t\u00a0\n ',' reviewer','reviewer '):
        case('missing or whitespace identity '+repr(ident),{'id':ident},status='CHECK REVIEWER ID')
    for col in 'DE':
        case('missing '+col,{col:''},status='MISSING D / E')
        for v in ('BAD',' ',base[col].upper(),base[col]+' '):case('pasted '+col+' '+repr(v),{col:v},status='INVALID D / E')
    for flag in ('checker_done','yes','PRIMARY_DONE',' '):case('forbidden completion '+repr(flag),{'F':flag},status='CHECK COMPLETE FLAG')
    for flag in ('','no'):case('draft or reset completion '+repr(flag),{'F':flag},status='PRIMARY DRAFT')
    case('domain change leaves stale statistical answers',{'E':'protocol_report'},status='CLEAR INACTIVE VALUES')
    case('domain change cleared, no extra protocol judgment',dict(protocol),result='protocol_descriptive')
    case('optional clearing is blank',{'J':''},effective=['confirmed_issue','substantive','','not_applicable','not_applicable'])
    case('inactive not_applicable literals',{'K':'not_applicable','L':'not_applicable'},effective=['confirmed_issue','substantive','','not_applicable','not_applicable'])
    case('inactive uppercase code rejected',{'K':'NOT_APPLICABLE'},status='CLEAR INACTIVE VALUES')
    case('missing key',{'A':''},status='CHECK KEY')
    case('unknown key',{'A':'2019::unknown::C000'},status='CHECK KEY')
    case('duplicate key',{'A':wb[sheet]['A3'].value},status='CHECK KEY')
    foreign=next(n for n in wb.sheetnames if n.startswith('C_') and n!=sheet)
    case('foreign paper key',{'A':wb[foreign]['A2'].value},status='CHECK KEY')
    # Explicitly exercise untouched formulas with actual workbook dependencies.
    for domain in OPTIONS['E']:
        for disp in ('','retain','exclude','discuss'):
            sol=model.calculate(inputs={node(FILE,sheet,'D2'):disp,node(FILE,sheet,'E2'):domain},outputs=[node(FILE,sheet,c+'2') for c in 'HIJKL'])
            actual=[scalar(sol,FILE,sheet,c+'2') for c in 'HIJKL']
            expected=[]
            for col in 'HIJKL':
                applies=disp in ('','retain') and ((col in 'HIJ' and domain=='statistical_results') or (col=='K' and domain in ('plan_material','report_internal_inconsistency')) or (col=='L' and domain=='protocol_report'))
                expected.append('' if applies else 'not_applicable')
            assert_case('untouched defaults '+repr(disp)+' / '+domain,actual,expected)


def run():
    initial_hashes={p.name:sha(p) for p in INPUT.glob('*.xlsx')}
    model,sol,wb,count=all_calculate(FILE,'blank R_A')
    for cell,want in [('B3',62),('B4',0),('B5',62),('B6',0),('B7','ACTION NEEDED')]:assert_case('blank Checks '+cell,scalar(sol,FILE,'Checks',cell),want)
    print('Blank workbook calculation passed.',flush=True)
    test_scenarios(model);print(f'Row scenarios passed: {len(CA)} assertions.',flush=True)
    bfile=INPUT/'human_adjudication_v_2_3_1_2019_R_B.xlsx'
    bm,bs,bw,bc=all_calculate(bfile,'blank R_B');assert_case('blank R_B ready count',scalar(bs,bfile,'Checks','B4'),0)
    full=load_workbook(FILE);fill_all(full);fullpath=TESTS/'all_complete.xlsx';full.save(fullpath)
    fm,fs,fw,fc=all_calculate(fullpath,'all complete')
    assert_case('all complete ready',scalar(fs,fullpath,'Checks','B4'),62)
    assert_case('all complete unfinished',scalar(fs,fullpath,'Checks','B5'),0)
    assert_case('all complete workflow',scalar(fs,fullpath,'Checks','B7'),'PRIMARY READY')
    assert_case('all nine paper count gates open',scalar(fs,fullpath,'Checks','B9'),9)
    candidate_rows=list(full['Candidates'].values)[1:];outcome_counts={};expected_total={}
    for r,c in enumerate(candidate_rows,2):
        expected=('protocol_descriptive' if c[12]=='P_R' else 'plan_descriptive' if c[12]=='P_P' else 'primary_temporal' if c[26]=='statistical_results' else 'secondary_temporal' if c[26]=='report_internal_inconsistency' else 'not_eligible')
        assert_case('completed outcome '+c[6],scalar(fs,fullpath,'Decisions','H'+str(r)),expected)
        expected_total[expected]=expected_total.get(expected,0)+1
        assert_case('key-based notes '+c[6],scalar(fs,fullpath,'Decisions','R'+str(r)),'TEST note for '+c[6])
    for r in range(2,full['Paper_counts'].max_row+1):
        vals=[scalar(fs,fullpath,'Paper_counts',c+str(r)) for c in 'IJKLM']
        assert_case('paper '+str(r)+' counts reconcile',sum(vals),full['Paper_counts']['C'+str(r)].value)
    for r in range(2,full['Protocol_plan_table'].max_row+1):
        key=full['Protocol_plan_table']['C'+str(r)].value
        assert_case('plan table notes '+key,scalar(fs,fullpath,'Protocol_plan_table','L'+str(r)),'TEST note for '+key)
        assert_case('plan table readiness '+key,scalar(fs,fullpath,'Protocol_plan_table','G'+str(r)),'PRIMARY READY')
    reopened=save_values(fw,TESTS/'all_complete_engine_values.xlsx',fs,fullpath)
    assert_case('engine values export saved/reopened workflow',reopened['Checks']['B7'].value,'PRIMARY READY')
    assert_case('engine values export saved/reopened ready count',reopened['Checks']['B4'].value,62)
    # Reload the formula-preserving physical copy too; no cache claim is made.
    assert_case('saved/reopened formula copy retains Excel formula',load_workbook(fullpath)['Decisions']['G2'].data_type,'f')
    print('All-paper completion, counts, protocol table and save/reopen passed.',flush=True)
    paper=next(n for n in full.sheetnames if n.startswith('C_'))
    key=full[paper]['A2'].value
    changes={node(fullpath,paper,'D2'):'discuss',**{node(fullpath,paper,c+'2'):'' for c in 'HIJKL'}}
    ss=fm.calculate(inputs=changes)
    assert_case('ready discuss completes entry',scalar(ss,fullpath,'Checks','B7'),'PRIMARY READY')
    assert_case('ready discuss unresolved count',scalar(ss,fullpath,'Checks','B6'),1)
    assert_case('ready discuss outcome blank',scalar(ss,fullpath,'Decisions','H2'),'')
    assert_case('ready discuss paper complete',scalar(ss,fullpath,'Paper_counts','G2'),'yes')
    assert_case('ready discuss paper counts withheld',[scalar(ss,fullpath,'Paper_counts',c+'2') for c in 'HIJKLM'],['no','','','','',''])
    changes[node(fullpath,paper,'F2')]='no';ss=fm.calculate(inputs=changes)
    assert_case('draft discuss unfinished',scalar(ss,fullpath,'Checks','B5'),1)
    assert_case('draft discuss still counted unresolved',scalar(ss,fullpath,'Checks','B6'),1)
    excluded=fm.calculate(inputs={node(fullpath,paper,'D2'):'exclude',**{node(fullpath,paper,c+'2'):'' for c in 'HIJKL'}})
    assert_case('completed exclusion keeps paper count gate open',scalar(excluded,fullpath,'Paper_counts','H2'),'yes')
    assert_case('completed exclusion counted as not eligible',scalar(excluded,fullpath,'Paper_counts','M2'),1)
    assert_case('completed exclusion not unresolved',scalar(excluded,fullpath,'Checks','B6'),0)
    for edit,label in [('', 'missing'),(full[paper]['A3'].value,'duplicate')]:
        ss=fm.calculate(inputs={node(fullpath,paper,'A2'):edit})
        assert_case(label+' original candidate Decisions blocks',scalar(ss,fullpath,'Decisions','G2'),'CHECK KEY')
        assert_case(label+' original outcome blank',scalar(ss,fullpath,'Decisions','H2'),'')
        assert_case(label+' prevents primary completion',scalar(ss,fullpath,'Checks','B7'),'ACTION NEEDED')
    reset=fm.calculate(inputs={node(fullpath,paper,'F2'):'no'})
    assert_case('manual completion reset clears outcome',scalar(reset,fullpath,'Decisions','H2'),'')
    assert_case('manual completion reset withholds paper counts',scalar(reset,fullpath,'Paper_counts','I2'),'')
    # Whole A:R rows moved together, formulas translated as relative Excel row formulas.
    sortedbook=load_workbook(fullpath);ws=sortedbook[paper];old=[[copy(c) for c in row] for row in list(ws)[1:]]
    for dest,oldrow in enumerate(reversed(old),2):
        for col,cell in enumerate(oldrow,1):
            new=ws.cell(dest,col);new.value=Translator(cell.value,origin=cell.coordinate).translate_formula(new.coordinate) if cell.data_type=='f' else cell.value
            new._style=copy(cell._style)
            new.comment=copy(cell.comment)
            new._hyperlink=copy(cell.hyperlink)
            if new.hyperlink:new.hyperlink.ref=new.coordinate
    sortpath=TESTS/'whole_row_reordered.xlsx';sortedbook.save(sortpath)
    sm,ss,sw,sc=all_calculate(sortpath,'whole rows reordered')
    assert_case('reordered workbook all ready',scalar(ss,sortpath,'Checks','B4'),62)
    for r,c in enumerate(candidate_rows,2):
        assert_case('reordered key-to-note '+c[6],scalar(ss,sortpath,'Decisions','R'+str(r)),'TEST note for '+c[6])
        assert_case('reordered outcome unchanged '+c[6],scalar(ss,sortpath,'Decisions','H'+str(r)),scalar(fs,fullpath,'Decisions','H'+str(r)))
    # Smallest possible fixture uses one included candidate; all range endpoints reduced.
    one=load_workbook(fullpath)
    for name in list(one.sheetnames):
        if name.startswith('C_') and name!=paper:del one[name]
    one[paper].delete_rows(3,one[paper].max_row)
    one[paper]['G2']=hint(2,2,paper[2:],'R_A',total=1)
    for name in ('Decisions','Papers','Paper_counts'):one[name].delete_rows(3,one[name].max_row)
    one['Protocol_plan_table'].delete_rows(2,one['Protocol_plan_table'].max_row)
    for name in SOURCE_SHEETS:
        ws=one[name];headers=list(next(ws.values));ki=headers.index('key')+1
        for r in range(ws.max_row,1,-1):
            if ws.cell(r,ki).value!=key:ws.delete_rows(r)
    for name in ('Decisions','Paper_counts','Checks'):
        for row in one[name]:
            for c in row:
                if c.data_type=='f':c.value=re.sub(r'(\$[A-Z]+\$)(?:63|10|4)\b',r'\g<1>2',c.value)
    one['Checks']['B3']=1;one['Papers']['C2']=1;one['Paper_counts']['C2']=1
    onepath=TESTS/'single_candidate.xlsx';one.save(onepath)
    om,os,ow,oc=all_calculate(onepath,'single candidate')
    assert_case('single candidate workflow',scalar(os,onepath,'Checks','B7'),'PRIMARY READY')
    assert_case('single candidate lookup outcome',scalar(os,onepath,'Decisions','H2'),'primary_temporal')
    assert_case('single candidate counts',scalar(os,onepath,'Paper_counts','I2'),1)
    assert_case('input templates byte-identical after all tests',{p.name:sha(p) for p in INPUT.glob('*.xlsx')},initial_hashes)
    return {'status':'PASS','engine':'formulas '+formulas.__version__+' (Python Excel formula interpreter)','assertions':len(CA),'failed':0,'cases':CA,'all_complete_outcome_totals':expected_total,'duration_seconds':round(time.time()-started,2),'formula_cells_per_template':count,'desktop_excel_smoke_test':'NOT RUN','rendering':'NOT RUN','limitations':['This interpreter is not desktop Excel or LibreOffice. Dropdown UI, Excel row-sort behavior, visual clipping and native Excel save/reopen remain pending.','The saved/reopened engine output is a disposable values-only test export. Delivered workbooks retain normal Excel formulas and automatic full recalculation; no fabricated cached values are inserted.','The interpreter can retain a preassembled entirely blank lookup range when individual blank cells are overridden. Notes and sorting are therefore tested using physically saved filled copies, not blank-cell override assumptions.']}

if __name__=='__main__':
    try:report=run()
    except Exception as exc:
        (TESTS/'calculation_report.json').write_text(json.dumps({'status':'FAIL','error':str(exc),'cases':CA},indent=2,default=str)+'\n');raise
    (TESTS/'calculation_report.json').write_text(json.dumps(report,indent=2,default=str)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='cases'},indent=2))
