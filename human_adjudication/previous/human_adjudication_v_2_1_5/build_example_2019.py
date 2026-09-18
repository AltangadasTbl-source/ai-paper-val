#!/usr/bin/env python3
"""Create explicitly synthetic, completed-primary and ready-checker 2019 examples."""
from pathlib import Path
import tempfile

from openpyxl import load_workbook

import build_v_2_1_5 as app


def build_example(output_dir=None):
    output_dir = Path(output_dir) if output_dir is not None else app.HERE / 'examples'
    primary_path = output_dir / 'human_adjudication_v_2_1_5_2019_TEST_primary_complete.xlsx'
    checker_path = output_dir / 'human_adjudication_v_2_1_5_2019_TEST_checker_ready.xlsx'
    if primary_path.exists() or checker_path.exists():
        raise ValueError('Example output already exists; existing files are never overwritten')
    candidates, signature = app.load_catalog()
    annual = [c for c in candidates if str(c['year']) == '2019']
    wb = app.make_review(annual, signature, primary_path, primary_id='TEST_primary_2019')
    app.put(wb['Setup']['B3'], 'SYNTHETIC TEST REVIEWER - not a real adjudicator')
    app.put(wb['Start']['B2'], app.VERSION + ' | 2019 | SYNTHETIC TEST DATA ONLY | primary complete')
    for name in app.review_names(annual, 'primary'):
        ws = wb[name]
        for row in range(2, ws.max_row + 1):
            values = {
                'disposition': 'exclude',
                'reviewComplete': 'yes',
                'notes': 'SYNTHETIC TEST DATA ONLY. Exclude is a placeholder for testing the checker workflow, not a scientific assessment. Do not use this sample in research results.',
            }
            for field, value in values.items():
                app.put(ws.cell(row, app.REVIEW_FIELDS.index(field) + 1), value)
    app.save_new(wb, primary_path)
    primary_report = app.validate(primary_path, require_complete=True)
    app.fail_on_errors(primary_report)
    with tempfile.TemporaryDirectory(prefix='adjudication-v215-example-') as directory:
        prepared = Path(directory) / 'prepared_checker.xlsx'
        app.prepare_checker(primary_path, prepared, 'TEST_checker_2019')
        checker = load_workbook(prepared)
        app.put(checker['Setup']['B5'], 'SYNTHETIC TEST REVIEWER - not a real adjudicator')
        app.put(checker['Start']['B2'], app.VERSION + ' | 2019 | SYNTHETIC TEST DATA ONLY | ready for checking')
        app.save_new(checker, checker_path)
    checker_report = app.validate(checker_path)
    app.fail_on_errors(checker_report)
    assert len(checker_report['primary']) == 71
    assert checker_report['primary'] == primary_report['primary']
    assert all(r['reviewComplete'] == 'yes' for r in checker_report['primary'].values())
    assert all(r['reviewComplete'] == 'no' and r['checkStatus'] == '' for r in checker_report['overrides'].values())
    print('2019 test sample ready: 71 completed synthetic primary records, 71 pending checker records.')


if __name__ == '__main__':
    build_example()
