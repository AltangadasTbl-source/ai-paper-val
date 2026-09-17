"""Behavioral checks using real XLSX save/reopen cycles; no human data modified."""
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator

import build_v_2_1_2 as app

WORKBOOK_DIR = Path(os.environ.get("ADJUDICATION_TEST_WORKBOOK_DIR", str(app.HERE)))


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.all_candidates, cls.signature = app.load_catalog()
        cls.candidates = [c for c in cls.all_candidates if str(c["year"]) == "2019"]
        cls.template = WORKBOOK_DIR / "human_adjudication_v_2_1_2_2019.xlsx"

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="adjudication-v212-test-")
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "review.xlsx"
        self.wb = load_workbook(self.template)
        self.catalog_mock = patch.object(app, "load_catalog", return_value=(self.all_candidates, self.signature))
        self.catalog_mock.start()
        self.addCleanup(self.catalog_mock.stop)

    def set(self, sheet, row, field, value, literal=True):
        if sheet == "Primary":
            cell = self.primary_cell(row, field)
        else:
            ws = self.wb[sheet]
            headers = [c.value for c in ws[1]]
            cell = ws.cell(row, headers.index(app.LABELS[field]) + 1)
        if literal:
            app.put(cell, value)
        else:
            cell.value = value

    def primary_cell(self, global_row, field):
        candidate = self.candidates[global_row - 2]
        ws = self.wb[app.primary_sheet_name(candidate["package"])]
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == candidate["key"])
        return ws.cell(row, app.REVIEW_FIELDS.index(field) + 1)

    def validate(self, **kwargs):
        self.wb.save(self.path)
        return app.validate(self.path, **kwargs)

    def fill_primary(self, disposition="exclude"):
        app.put(self.wb["Setup"]["B2"], "primary_reviewer")
        for row in range(2, len(self.candidates) + 2):
            self.set("Primary", row, "disposition", disposition)
            self.set("Primary", row, "reviewComplete", "yes")

    def fill_checker(self, status="agree"):
        app.put(self.wb["Setup"]["B4"], "checker_reviewer")
        for row in range(2, len(self.candidates) + 2):
            self.set("Checker", row, "checkStatus", status)
            self.set("Checker", row, "reviewComplete", "yes")

    def assert_valid(self, report):
        self.assertFalse([i for i in report["issues"] if i[1] == "ERROR"], report["issues"][:5])

    def assert_error(self, report, fragment):
        self.assertTrue(any(i[1] == "ERROR" and fragment in i[2] for i in report["issues"]), report["issues"][:5])

    def test_source_conversion_is_lossless_for_all_412_candidates(self):
        original = []
        for year in app.YEARS:
            p = app.HERE.parent / "human_adjudication_v_2_0_4" / f"candidate_flag_review_v_2_0_4_{year}.html"
            data = json.loads(re.search(r'<script id="candidate-data" type="application/json">(.*?)</script>', p.read_text(), re.S)[1])
            original.extend(data["candidates"])
        self.assertEqual(original, self.all_candidates)
        self.assertEqual(Counter(c["primary"] for c in original), app.RELATIONSHIPS)

    def test_all_annual_templates_have_no_human_answers(self):
        for year, count in app.YEARS.items():
            path = WORKBOOK_DIR / f"human_adjudication_v_2_1_2_{year}.xlsx"
            report = app.validate(path)
            self.assert_valid(report)
            self.assertEqual(len(report["results"]), count)
            self.assertEqual({r["status"] for r in report["results"]}, {"pending"})
            self.assertEqual(set(report["setup"].values()), {""})
            self.assertTrue(all(r["reviewComplete"] == "no" for r in report["primary"].values()))
            self.assertTrue(all(r["disposition"] == "" for r in report["primary"].values()))
            self.assertTrue(all(r["reviewComplete"] == "no" and r["checkStatus"] == "" for r in report["overrides"].values()))

    def test_dropdowns_and_protection_survive_roundtrip_without_hyperlinks(self):
        for sheet in app.primary_sheet_names(self.candidates) + ["Checker"]:
            ws = self.wb[sheet]
            self.assertTrue(ws.protection.sheet)
            self.assertTrue(ws.auto_filter.ref)
            self.assertTrue(ws.freeze_panes)
            for dv in ws.data_validations.dataValidation:
                self.assertTrue(dv.showErrorMessage)
                self.assertEqual(dv.errorStyle, "stop")
                self.assertIn(dv.formula1, self.wb.defined_names)
                self.assertIn(str(ws.max_row), str(dv.sqref))
            fields = app.CHECK_FIELDS if sheet == "Checker" else app.REVIEW_FIELDS
            for f in app.ASSESSMENT:
                self.assertFalse(ws.cell(2, fields.index(f)+1).protection.locked)
            self.assertTrue(ws["A2"].protection.locked)
        self.assertFalse(any(cell.hyperlink for ws in self.wb for row in ws for cell in row))
        self.assertFalse(self.wb.vba_archive)

    def test_primary_handoff_does_not_require_checker(self):
        self.fill_primary()
        self.assert_valid(self.validate(require_complete=True, primary_only=True))

    def test_ai_prefills_cannot_be_exported_as_human_results(self):
        self.validate()
        with self.assertRaisesRegex(ValueError, "Unfinished"):
            app.export_analysis([self.path], Path(self.tmp.name) / "analysis.xlsx")

    def test_simple_agreement_inherits_all_primary_fields(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "notes", "Reviewed manually")
        report = self.validate(require_complete=True)
        self.assert_valid(report)
        self.assertEqual({r["status"] for r in report["results"]}, {"resolved"})
        self.assertEqual(report["reviews"][self.candidates[0]["key"]]["notes"], "Reviewed manually")

    def test_checker_cannot_finish_before_primary(self):
        self.fill_checker()
        self.assert_error(self.validate(), "before primary")

    def test_distinct_reviewer_ids_required(self):
        self.fill_primary()
        self.fill_checker()
        self.wb["Setup"]["B4"] = "PRIMARY_REVIEWER"
        self.assert_error(self.validate(), "different person")

    def test_missing_primary_id_detected(self):
        self.fill_primary()
        self.wb["Setup"]["B2"] = ""
        self.assert_error(self.validate(), "stable primary")

    def test_pasted_enum_and_boolean_completion_rejected(self):
        self.set("Primary", 2, "disposition", "retian")
        self.set("Primary", 3, "reviewComplete", True)
        report = self.validate()
        self.assert_error(report, "invalid option")
        self.assert_error(report, "expected text")

    def test_required_fields_for_all_four_modules(self):
        for domain, fields in app.MODULES.items():
            record = app.default_record(self.candidates[0])
            record.update(disposition="retain", domain=domain, reviewComplete="yes")
            for f in fields:
                self.assertIn(f"{f}: required", app.assessment_errors(record, True))
                record[f] = next(iter(app.OPTIONS[f]))
            self.assertEqual(app.assessment_errors(record, True), [])

    def test_nonretain_dispositions_do_not_need_modules(self):
        for disposition in ("exclude", "discuss", "out_of_scope"):
            for domain in app.MODULES:
                r = app.default_record(self.candidates[0])
                r.update(disposition=disposition, domain=domain, reviewComplete="yes")
                self.assertEqual(app.assessment_errors(r, True), [])

    def test_agree_cannot_hide_modified_assessment(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Checker", 2, "notes", "Changed")
        self.assert_error(self.validate(), "unchanged primary")

    def test_nonagreement_comment_required(self):
        self.fill_primary()
        self.fill_checker()
        for status in ("agree_clarification", "disagree", "unable"):
            self.set("Checker", 2, "checkStatus", status)
            self.assert_error(self.validate(), "checkerComment")

    def test_unresolved_disagreement_cannot_enter_final_counts(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Checker", 2, "checkStatus", "disagree")
        self.set("Checker", 2, "checkerComment", "Discuss scope")
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report["results"][0]["status"], "unresolved")
        with self.assertRaisesRegex(ValueError, "Unfinished/unresolved"):
            app.export_analysis([self.path], Path(self.tmp.name) / "final.xlsx")

    def test_resolution_requires_note_and_non_discuss_disposition(self):
        self.fill_primary("discuss")
        self.fill_checker("unable")
        for row in range(2, len(self.candidates)+2):
            self.set("Checker", row, "checkerComment", "Needs discussion")
        self.set("Checker", 2, "resolutionStatus", "resolved")
        report = self.validate()
        self.assert_error(report, "resolutionNote")
        self.assert_error(report, "cannot retain disposition=discuss")

    def test_primary_discussion_needs_documented_resolution(self):
        self.fill_primary("discuss")
        self.fill_checker()
        self.set("Checker", 2, "disposition", "exclude")
        self.set("Checker", 2, "checkStatus", "agree_clarification")
        self.set("Checker", 2, "checkerComment", "Reviewed together")
        self.assertEqual(self.validate()["results"][0]["status"], "unresolved")
        self.set("Checker", 2, "resolutionStatus", "resolved")
        self.set("Checker", 2, "resolutionNote", "Primary and checker agreed to exclude")
        self.assertEqual(self.validate()["results"][0]["status"], "resolved")

    def test_empty_override_and_explicit_clear_are_distinct(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "notes", "Old note")
        self.set("Checker", 2, "notes", "<clear>")
        self.set("Checker", 2, "checkStatus", "agree_clarification")
        self.set("Checker", 2, "checkerComment", "Remove obsolete note")
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report["primary"][self.candidates[0]["key"]]["notes"], "Old note")
        self.assertEqual(report["reviews"][self.candidates[0]["key"]]["notes"], "")

    def test_domain_override_requires_its_new_module(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Checker", 2, "checkStatus", "disagree")
        self.set("Checker", 2, "checkerComment", "Different assessment domain")
        self.set("Checker", 2, "disposition", "retain")
        self.set("Checker", 2, "domain", "internal_nonstatistical")
        self.assert_error(self.validate(), "internalStatus: required")

    def test_missing_and_duplicate_candidate_rows_rejected(self):
        self.wb["Checker"].delete_rows(2)
        with self.assertRaisesRegex(ValueError, "key.*mismatch"):
            self.validate()
        self.wb = load_workbook(self.template)
        self.set("Primary", 3, "key", self.candidates[0]["key"])
        with self.assertRaisesRegex(ValueError, "duplicate candidate"):
            self.validate()

    def test_source_edit_and_wrong_version_rejected(self):
        self.wb["Evidence"]["D2"] = "Altered evidence"
        with self.assertRaisesRegex(ValueError, "Frozen candidate"):
            self.validate()
        self.wb = load_workbook(self.template)
        self.wb["Manifest"]["B2"] = "v.2.0.4"
        with self.assertRaisesRegex(ValueError, "Unsupported"):
            self.validate()

    def test_formula_in_answer_rejected_but_literal_math_preserved(self):
        self.set("Primary", 2, "notes", "=1+1", literal=False)
        self.assert_error(self.validate(), "formula/error cell")
        self.set("Primary", 2, "notes", "=1+1", literal=True)
        self.assert_valid(self.validate())
        cell = self.primary_cell(2, "notes")
        self.assertEqual(cell.data_type, "s")

    def test_row_reordering_preserves_key_based_inheritance(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "notes", "Key-specific note")
        ws = self.wb["Checker"]
        row2, row3 = [[c.value for c in ws[r]] for r in (2, 3)]
        for col, value in enumerate(row3, 1):
            if isinstance(value, str) and value.startswith("="):
                ws.cell(2, col).value = Translator(value, origin=ws.cell(3, col).coordinate).translate_formula(ws.cell(2, col).coordinate)
            else:
                app.put(ws.cell(2, col), value)
        for col, value in enumerate(row2, 1):
            if isinstance(value, str) and value.startswith("="):
                ws.cell(3, col).value = Translator(value, origin=ws.cell(2, col).coordinate).translate_formula(ws.cell(3, col).coordinate)
            else:
                app.put(ws.cell(3, col), value)
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report["reviews"][self.candidates[0]["key"]]["notes"], "Key-specific note")

    def test_zero_outcomes_include_all_papers(self):
        self.fill_primary()
        self.fill_checker()
        self.validate()
        output = Path(self.tmp.name) / "analysis.xlsx"
        app.export_analysis([self.path], output)
        wb = load_workbook(output)
        papers = app.records(wb["Paper_counts"])
        self.assertEqual(len(papers), 10)
        self.assertEqual(sum(p["candidates"] for p in papers), 71)
        self.assertTrue(all(p["paper_complete"] == "yes" and p["primary_temporal"] == 0 for p in papers))

    def test_provisional_export_uses_blank_not_zero_for_incomplete_papers(self):
        self.validate()
        output = Path(self.tmp.name) / "provisional.xlsx"
        app.export_analysis([self.path], output, allow_incomplete=True)
        wb = load_workbook(output)
        papers = app.records(wb["Paper_counts"])
        self.assertTrue(all(p["primary_temporal"] == "" and p["paper_complete"] == "no" for p in papers))
        self.assertIn("PROVISIONAL", str(wb["Readme"]["B5"].value))

    def test_duplicate_years_and_existing_outputs_rejected(self):
        self.fill_primary()
        self.fill_checker()
        self.validate()
        with self.assertRaisesRegex(ValueError, "Duplicate year"):
            app.export_analysis([self.path, self.path], Path(self.tmp.name) / "duplicate.xlsx")
        original = hashlib.sha256(self.path.read_bytes()).hexdigest()
        with self.assertRaises(FileExistsError):
            app.save_new(self.wb, self.path)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), original)

    def test_optional_frozen_handoff_detects_primary_mutation(self):
        self.fill_primary()
        self.validate()
        output = Path(self.tmp.name) / "frozen.xlsx"
        app.prepare_checker(self.path, output, "checker_reviewer")
        self.assert_valid(app.validate(output))
        frozen = load_workbook(output)
        disposition_col = app.REVIEW_FIELDS.index("disposition") + 1
        for sheet in app.primary_sheet_names(self.candidates):
            self.assertTrue(frozen[sheet].cell(2, disposition_col).protection.locked)
        frozen[app.primary_sheet_name(self.candidates[0]["package"])].cell(2, disposition_col).value = "discuss"
        frozen.save(output)
        with self.assertRaisesRegex(ValueError, "Locked primary"):
            app.validate(output)

    def test_optional_handoff_does_not_discard_existing_checker_work(self):
        self.fill_primary()
        self.set("Checker", 2, "checkerComment", "Started reading")
        self.validate()
        with self.assertRaisesRegex(ValueError, "already started"):
            app.prepare_checker(self.path, Path(self.tmp.name) / "fresh.xlsx")

    def test_outcome_boundaries(self):
        for relation in ("S_C", "S_W", "P_R", "P_P"):
            c = {"primary": relation}
            r = {"disposition": "retain", "domain": "statistical_results", "statisticalErrorStatus": "confirmed_error", "substantiveStatus": "substantive"}
            expected = {"S_C": "primary_temporal", "S_W": "primary_temporal", "P_R": "protocol_descriptive", "P_P": "plan_descriptive"}[relation]
            self.assertEqual(app.outcome(c, r), expected)
            if relation in ("S_C", "S_W"):
                r["statisticalErrorStatus"] = "probable_error"
                self.assertEqual(app.outcome(c, r), "primary_temporal")
                r["substantiveStatus"] = "minor_presentation"
                self.assertEqual(app.outcome(c, r), "not_eligible")
                r.update(domain="internal_nonstatistical", internalStatus="substantive_inconsistency")
                self.assertEqual(app.outcome(c, r), "secondary_temporal")
            r["disposition"] = "exclude"
            self.assertEqual(app.outcome(c, r), "not_eligible")

    def test_exactly_one_primary_sheet_per_paper_all_years(self):
        total = 0
        for year in app.YEARS:
            wb = load_workbook(WORKBOOK_DIR / f"human_adjudication_v_2_1_2_{year}.xlsx")
            annual = [c for c in self.all_candidates if str(c["year"]) == year]
            expected = app.primary_sheet_names(annual)
            self.assertNotIn("Primary", wb.sheetnames)
            self.assertEqual({s for s in wb.sheetnames if s.startswith("P_")}, set(expected))
            self.assertEqual(len(app.read_primary(wb, annual)), app.YEARS[year])
            for paper, candidates in app.paper_groups(annual).items():
                ws = wb[app.primary_sheet_name(paper)]
                self.assertLessEqual(len(ws.title), 31)
                self.assertEqual(ws.max_row - 1, len(candidates))
                self.assertNotIn("Paper", [c.value for c in ws[1]])
                self.assertEqual({ws.cell(r, 1).value.split("::")[1] for r in range(2, ws.max_row+1)}, {paper})
            total += len(expected)
        self.assertEqual(total, 45)

    def test_missing_primary_sheet_rejected(self):
        self.wb.remove(self.wb[app.primary_sheet_names(self.candidates)[0]])
        with self.assertRaisesRegex(ValueError, "worksheet set mismatch"):
            self.validate()

    def test_candidate_moved_to_another_paper_rejected(self):
        sheets = app.primary_sheet_names(self.candidates)
        a, b = self.wb[sheets[0]], self.wb[sheets[1]]
        a_row, b_row = [[cell.value for cell in ws[2]] for ws in (a, b)]
        for ws, row in ((a, b_row), (b, a_row)):
            for i, value in enumerate(row, 1):
                app.put(ws.cell(2, i), value)
        with self.assertRaisesRegex(ValueError, "key set mismatch for paper"):
            self.validate()

    def test_papers_index_tampering_rejected(self):
        self.wb["Papers"]["C2"] = 999
        with self.assertRaisesRegex(ValueError, "Papers index mismatch"):
            self.validate()

    def test_all_release_xlsx_packages_are_english_only(self):
        for path in list(WORKBOOK_DIR.glob("*.xlsx")) + [app.CATALOG]:
            app.require_english_xlsx(path)

    def test_chinese_answers_and_hidden_content_rejected(self):
        for place in ("answer", "comment", "hidden", "prompt", "property"):
            self.wb = load_workbook(self.template)
            if place == "answer":
                self.primary_cell(2, "notes").value = "\u4e2d\u6587"
            elif place == "comment":
                self.wb["Setup"]["B2"].comment = Comment("\u4e2d\u6587", "Test")
            elif place == "hidden":
                self.wb["Links"]["A2"] = "\u4e2d\u6587"
            elif place == "prompt":
                self.wb["Checker"].data_validations.dataValidation[0].prompt = "\u4e2d\u6587"
            else:
                self.wb.properties.description = "\u4e2d\u6587"
            with self.assertRaisesRegex(ValueError, "English-only"):
                self.validate()

    def test_chinese_output_rejected_before_creating_file(self):
        self.primary_cell(2, "notes").value = "\u4e2d\u6587"
        target = Path(self.tmp.name) / "must-not-exist.xlsx"
        with self.assertRaisesRegex(ValueError, "English-only"):
            app.save_new(self.wb, target)
        self.assertFalse(target.exists())

    def test_english_theme_retains_scientific_unicode(self):
        self.set("Primary", 2, "notes", "The estimated \u03b2 differs by 0.30 percentage points; 95% CI \u2260 0.")
        self.assert_valid(self.validate())
        target = Path(self.tmp.name) / "scientific-unicode.xlsx"
        app.save_new(self.wb, target)
        app.require_english_xlsx(target)

    def test_blank_disposition_requires_an_explicit_human_choice(self):
        self.wb["Setup"]["B2"] = "primary_reviewer"
        self.set("Primary", 2, "reviewComplete", "yes")
        self.assert_error(self.validate(), "disposition: required")
        self.set("Primary", 2, "disposition", "discuss")
        self.assert_valid(self.validate())

    def test_ai_disposition_is_reference_only_for_all_candidates(self):
        self.assertEqual(Counter(c["aiInitialAssessment"]["disposition"] for c in self.all_candidates), {"retain": 402, "exclude": 9, "discuss": 1})
        for c in self.all_candidates:
            self.assertEqual(app.default_record(c)["disposition"], "")

    def test_live_defaults_cover_every_domain_and_disposition(self):
        key = self.candidates[0]["key"]
        for domain in app.MODULES:
            for disposition in ("", "retain", "exclude", "out_of_scope", "discuss"):
                self.set("Primary", 2, "domain", domain)
                self.set("Primary", 2, "disposition", disposition)
                report = self.validate()
                self.assert_valid(report)
                record = report["primary"][key]
                active = app.APPLICABLE[domain] if disposition in ("", "retain") else ()
                for field in app.MODULE_FIELDS:
                    self.assertEqual(record[field], "" if field in active else app.NA)
                self.assertEqual(record["reviewComplete"], "no")

    def test_optional_assessments_do_not_block_completed_primary(self):
        self.fill_primary()
        for domain in ("statistical_results", "protocol_report"):
            self.wb = load_workbook(self.template)
            self.fill_primary()
            self.set("Primary", 2, "domain", domain)
            self.set("Primary", 2, "disposition", "retain")
            for field in app.MODULES[domain]:
                self.set("Primary", 2, field, next(iter(app.OPTIONS[field])))
            report = self.validate(require_complete=True, primary_only=True)
            self.assert_valid(report)
            for field in app.OPTIONAL_MODULES.get(domain, ()):
                self.assertEqual(report["primary"][self.candidates[0]["key"]][field], "")

    def test_inactive_old_answers_are_warned_and_excluded_from_final_values(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "statisticalErrorStatus", "confirmed_error")
        report = self.validate()
        self.assert_valid(report)
        self.assertTrue(any(i[1] == "WARNING" and "statisticalErrorStatus" in i[2] for i in report["issues"]))
        key = self.candidates[0]["key"]
        self.assertEqual(report["primary"][key]["statisticalErrorStatus"], "confirmed_error")
        self.assertEqual(report["reviews"][key]["statisticalErrorStatus"], app.NA)

    def test_manual_na_cannot_complete_applicable_statistical_questions(self):
        self.fill_primary()
        self.set("Primary", 2, "domain", "statistical_results")
        self.set("Primary", 2, "disposition", "retain")
        self.set("Primary", 2, "statisticalErrorStatus", app.NA)
        self.set("Primary", 2, "substantiveStatus", app.NA)
        self.assert_error(self.validate(primary_only=True), "not_applicable is invalid")

    def test_protocol_wording_only_is_a_human_answer_but_na_is_not(self):
        self.fill_primary()
        self.set("Primary", 2, "domain", "protocol_report")
        self.set("Primary", 2, "disposition", "retain")
        self.set("Primary", 2, "protocolInterpretation", app.NA)
        self.assert_error(self.validate(primary_only=True), "not_applicable is invalid")
        self.set("Primary", 2, "protocolInterpretation", "wording_only")
        self.assert_valid(self.validate(primary_only=True, require_complete=True))

    def test_checker_reactivation_cannot_inherit_nonretain_protocol_defaults(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "domain", "protocol_report")
        self.set("Checker", 2, "disposition", "retain")
        self.set("Checker", 2, "checkStatus", "agree_clarification")
        self.set("Checker", 2, "checkerComment", "Retain after discussion")
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report["reviews"][self.candidates[0]["key"]]["protocolInterpretation"], "")

    def test_domain_change_does_not_inherit_stale_answer_in_new_module(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "statisticalErrorStatus", "confirmed_error")
        self.set("Checker", 2, "domain", "statistical_results")
        self.set("Checker", 2, "disposition", "retain")
        self.set("Checker", 2, "checkStatus", "agree_clarification")
        self.set("Checker", 2, "checkerComment", "New branch requires fresh review")
        self.assert_error(self.validate(), "statisticalErrorStatus: required")

    def test_default_formula_tampering_is_rejected_without_evaluating_it(self):
        cell = self.primary_cell(2, "statisticalErrorStatus")
        cell.value = '=IF(TRUE,"confirmed_error","")'
        self.assert_error(self.validate(), "formula/error cell")

    def test_codebook_covers_all_review_headers_and_exact_dropdown_options(self):
        ws = self.wb["Codebook"]
        field_rows = {ws.cell(r, 5).value: r for r in range(2, ws.max_row+1) if ws.cell(r, 2).value == "FIELD"}
        self.assertEqual(set(field_rows), set(app.REVIEW_FIELDS) | set(app.CHECK_FIELDS))
        for field, row in field_rows.items():
            self.assertEqual(ws.cell(row, 1).value, app.LABELS[field])
            self.assertTrue(ws.cell(row, 1).font.bold)
            self.assertEqual(ws.cell(row, 1).alignment.horizontal, "center")
            self.assertEqual(ws.cell(row, 1).border.top.style, "medium")
            self.assertTrue(ws.cell(row, 3).value)
            self.assertTrue(ws.cell(row, 4).value)
        for field, options in app.OPTIONS.items():
            name = self.wb.defined_names["opt_" + field]
            values = []
            for sheet, reference in name.destinations:
                values.extend(cell.value for row in self.wb[sheet][reference] for cell in row)
            self.assertEqual(values, list(options))
        self.assertGreater(len(ws.merged_cells.ranges), 5)

    def test_optional_checker_judgment_can_be_cleared(self):
        self.fill_primary()
        self.fill_checker()
        self.set("Primary", 2, "domain", "statistical_results")
        self.set("Primary", 2, "disposition", "retain")
        self.set("Primary", 2, "statisticalErrorStatus", "confirmed_error")
        self.set("Primary", 2, "substantiveStatus", "substantive")
        self.set("Primary", 2, "downstreamImpact", "possible")
        self.set("Checker", 2, "downstreamImpact", "<clear>")
        self.set("Checker", 2, "checkStatus", "agree_clarification")
        self.set("Checker", 2, "checkerComment", "Do not rate optional impact")
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report["reviews"][self.candidates[0]["key"]]["downstreamImpact"], "")

    def test_restore_defaults_keeps_human_answers_and_does_not_overwrite_input(self):
        self.set("Primary", 2, "domain", "statistical_results")
        self.set("Primary", 2, "disposition", "retain")
        self.set("Primary", 2, "statisticalErrorStatus", "probable_error")
        self.set("Primary", 2, "substantiveStatus", "")
        self.validate()
        before = self.path.read_bytes()
        output = Path(self.tmp.name) / "restored.xlsx"
        app.restore_defaults(self.path, output)
        self.assertEqual(before, self.path.read_bytes())
        wb = load_workbook(output)
        ws = wb[app.primary_sheet_name(self.candidates[0]["package"])]
        row = next(r for r in range(2, ws.max_row+1) if ws.cell(r, 1).value == self.candidates[0]["key"])
        self.assertEqual(ws.cell(row, app.REVIEW_FIELDS.index("statisticalErrorStatus")+1).value, "probable_error")
        self.assertEqual(ws.cell(row, app.REVIEW_FIELDS.index("substantiveStatus")+1).data_type, "f")
        self.assert_valid(app.validate(output))

    def test_checker_lookup_ranges_cannot_be_redirected(self):
        name = app.primary_range_name(self.candidates[0]["package"], "domain")
        self.wb.defined_names[name].attr_text = "'Setup'!$B$2:$B$3"
        with self.assertRaisesRegex(ValueError, "lookup range changed"):
            self.validate()

    def test_removed_protocol_investigations_absent_from_review_and_exports(self):
        removed = {"documentationStatus", "justificationStatus", "Documentation", "Justification", "documentation", "justification"}
        self.assertFalse(set(app.LABELS) & removed)
        self.assertFalse(set(app.OPTIONS) & removed)
        for name in app.primary_sheet_names(self.candidates) + ["Checker", "Codebook"]:
            self.assertFalse(any(cell.value in removed for row in self.wb[name] for cell in row if isinstance(cell.value, str)))
        self.fill_primary()
        self.fill_checker()
        self.validate()
        output = Path(self.tmp.name) / "minimal-protocol.xlsx"
        app.export_analysis([self.path], output)
        wb = load_workbook(output)
        for name in ("Decisions", "Protocol_plan_table"):
            self.assertFalse(any(any(word in str(cell.value) for word in ("documentation", "justification")) for cell in wb[name][1]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
