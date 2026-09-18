"""Exercise four annual Checker-layout templates and both review stages."""
from collections import Counter
from copy import copy
import io
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

import build_v_2_1_5_1 as app


class AnnualWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.all_candidates, cls.signature = app.load_catalog()
        cls.candidates = [c for c in cls.all_candidates if str(c['year']) == '2019']
        cls.primary_template = app.HERE / 'human_adjudication_v_2_1_5_1_2019.xlsx'
        cls.checker_template = cls.primary_template

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='adjudication-v2151-test-')
        self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name)
        self.path = self.directory / 'working.xlsx'
        self.wb = load_workbook(self.primary_template)

    def cell(self, index, field):
        c = self.candidates[index]
        stage = app.meta(self.wb)['stage']
        name = app.primary_sheet_name(c['package']) if stage == 'primary' else app.checker_sheet_name(c['package'])
        ws = self.wb[name]
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == c['key'])
        fields = app.REVIEW_FIELDS if stage == 'primary' else app.CHECK_FIELDS
        return ws.cell(row, fields.index(field) + 1)

    def set(self, index, field, value):
        app.put(self.cell(index, field), value)

    def validate(self, **kwargs):
        self.wb.save(self.path)
        return app.validate(self.path, **kwargs)

    def assert_valid(self, report):
        self.assertFalse([i for i in report['issues'] if i[1] == 'ERROR'], report['issues'][:8])

    def assert_error(self, report, text):
        self.assertTrue(any(i[1] == 'ERROR' and text in i[2] for i in report['issues']), report['issues'][:8])

    def fill_primary(self, disposition='exclude'):
        self.wb['Setup']['B2'] = 'primary_reviewer'
        for i in range(len(self.candidates)):
            self.set(i, 'disposition', disposition)
            self.set(i, 'reviewComplete', 'yes')

    def handoff(self):
        self.wb.save(self.path)
        original = self.path.read_bytes()
        output = self.directory / 'ready_checker.xlsx'
        app.prepare_checker(self.path, output, 'checker_reviewer')
        self.assertEqual(self.path.read_bytes(), original)
        self.wb = load_workbook(output)
        return output

    def agree(self):
        for i in range(len(self.candidates)):
            self.set(i, 'checkStatus', 'agree')
            self.set(i, 'reviewComplete', 'yes')

    def test_four_annual_workbooks_use_revised_layout(self):
        self.assertEqual(len(list(app.HERE.glob('*.xlsx'))), 4)
        total = 0
        for year, count in app.YEARS.items():
            annual = [c for c in self.all_candidates if str(c['year']) == year]
            path = app.HERE / f'human_adjudication_v_2_1_5_1_{year}.xlsx'
            report = app.validate(path)
            self.assert_valid(report)
            self.assertEqual(len(report['results']), count)
            self.assertEqual({r['status'] for r in report['results']}, {'pending'})
            wb = report['wb']
            # v2.1.4 supplies the unchanged baseline styles; the v2.1.5
            # working copy may have user edits or Excel-normalized styles.
            old = load_workbook(app.HERE.parent / 'human_adjudication_v_2_1_4' / f'human_adjudication_v_2_1_4_{year}.xlsx')
            expected = set(app.checker_sheet_names(annual))
            self.assertEqual({name for name in wb.sheetnames if name.startswith(('P_', 'C_'))}, expected)
            for name in ('Start', 'Setup', 'Papers', 'Candidates', 'Evidence', 'Codebook'):
                self.assertEqual(wb[name].sheet_state, 'visible')
            for name in expected:
                sheet, reference = wb[name], old[name]
                self.assertEqual([c.value for c in sheet[1]], [app.LABELS[f] for f in app.CHECK_FIELDS])
                self.assertEqual(sheet.max_column, 18)
                self.assertEqual(sheet.max_row, reference.max_row)
                self.assertEqual(sheet.freeze_panes, reference.freeze_panes)
                self.assertEqual(sheet.auto_filter.ref, f'A1:R{sheet.max_row}')
                old_headers = [c.value for c in reference[1]]
                for col, field in enumerate(app.CHECK_FIELDS, 1):
                    label = app.LABELS[field]
                    previous = {'statisticalIssueStatus': 'Statistical error', 'internalStatus': 'Internal or plan assessment', 'primarySummary': 'Primary assessment (read only)'}.get(field, label)
                    old_col = old_headers.index(previous) + 1
                    if field not in ('internalStatus', 'statisticalIssueStatus'):
                        self.assertEqual(sheet.column_dimensions[get_column_letter(col)].width, reference.column_dimensions[get_column_letter(old_col)].width)
                    self.assertEqual(copy(sheet.cell(1, col).fill), copy(reference.cell(1, old_col).fill))
                    self.assertEqual(copy(sheet.cell(2, col).alignment), copy(reference.cell(2, old_col).alignment))
                self.assertEqual(sheet['M1'].value, 'Primary assessment')
                self.assertEqual(sheet['R1'].value, 'Notes (optional)')
                self.assertTrue(sheet['M2'].protection.locked)
                self.assertFalse(sheet['R2'].protection.locked)
                self.assertEqual(sheet.row_dimensions[2].height, reference.row_dimensions[2].height)
                for row in range(2, sheet.max_row + 1):
                    for field in app.CHECK_ONLY_FIELDS:
                        cell = sheet.cell(row, app.REVIEW_FIELDS.index(field)+1)
                        self.assertIsNone(cell.value)
                        self.assertTrue(cell.protection.locked)
            self.assertTrue(all(r['disposition'] == '' and r['reviewComplete'] == 'no' for r in report['primary'].values()))
            self.assertTrue(all(r['reviewComplete'] == 'no' and r['checkStatus'] == '' for r in report['overrides'].values()))
            self.assertFalse(any(c.hyperlink for ws in wb for row in ws for c in row))
            self.assertFalse(wb._external_links)
            self.assertFalse(wb.vba_archive)
            total += len(expected)
        self.assertEqual(total, 45)

    def test_catalog_preserves_all_v215_sources(self):
        for p in list(app.HERE.glob('*.xlsx')) + [app.CATALOG]:
            app.require_english_xlsx(p)
        old = load_workbook(app.HERE.parent / 'human_adjudication_v_2_1_5/resources/source_catalog_v_2_1_5.xlsx')
        self.assertEqual(app.read_candidates(old), self.all_candidates)
        current = load_workbook(app.CATALOG)
        self.assertEqual(app.source_signature(current), app.source_signature(old))
        self.assertEqual(app.meta(current)['parentCatalogSignature'], app.source_signature(old))
        self.assertEqual(app.meta(current)['parentCatalogVersion'], 'v.2.1.5')
        self.assertEqual(app.meta(current)['suggestedDomainsRenamed'], '0')
        self.assertEqual(Counter(c['primary'] for c in self.all_candidates), app.RELATIONSHIPS)

    def test_primary_completes_without_checker_and_cannot_export(self):
        self.fill_primary()
        report = self.validate(require_complete=True)
        self.assert_valid(report)
        self.assertEqual({r['status'] for r in report['results']}, {'primary_complete'})
        with self.assertRaisesRegex(ValueError, 'prepared Checker'):
            app.export_analysis([self.path], self.directory / 'bad.xlsx', allow_incomplete=True)

    def test_removed_codes_rejected_in_both_review_stages(self):
        for field, value in [('disposition', 'out_of_scope'), ('domain', 'internal_nonstatistical')]:
            original = self.cell(0, field).value
            self.set(0, field, value)
            self.assert_error(self.validate(), field + ': invalid option')
            self.set(0, field, original or '')
        self.fill_primary()
        self.handoff()
        self.agree()
        for field, value in [('disposition', 'out_of_scope'), ('domain', 'internal_nonstatistical')]:
            self.set(0, field, value)
            self.assert_error(self.validate(), field + ': invalid option')
            self.set(0, field, '')

    def test_release_and_example_dropdowns_use_new_terminology(self):
        paths = list(app.HERE.glob('*.xlsx')) + list((app.HERE / 'examples').glob('*.xlsx'))
        self.assertEqual(len(paths), 6)
        for path in paths:
            wb = load_workbook(path)
            for field, expected in [('disposition', ['retain', 'exclude', 'discuss']), ('domain', list(app.OPTIONS['domain']))]:
                actual = [c.value for sheet, ref in wb.defined_names['opt_' + field].destinations for row in wb[sheet][ref] for c in row]
                self.assertEqual(actual, expected)
            for sheet in wb:
                for row in sheet:
                    for cell in row:
                        for retired in ('out_of_scope', 'internal_nonstatistical', 'Outside review scope', 'Internal nonstatistical'):
                            self.assertNotIn(retired, str(cell.value))
                        if cell.comment:
                            self.assertNotIn('out_of_scope', cell.comment.text)
                            self.assertNotIn('internal_nonstatistical', cell.comment.text)
        description = app.OPTIONS['domain']['report_internal_inconsistency']
        self.assertIn('main paper', description)
        self.assertIn('updated supplement', description)

    def test_report_internal_inconsistency_handoff_and_export(self):
        self.fill_primary()
        index = next(i for i, c in enumerate(self.candidates) if c['primary'] in ('S_C', 'S_W'))
        self.set(index, 'domain', 'report_internal_inconsistency')
        self.set(index, 'disposition', 'retain')
        self.assert_error(self.validate(require_complete=True), 'internalStatus: required')
        self.set(index, 'internalStatus', 'substantive_inconsistency')
        self.assert_valid(self.validate(require_complete=True))
        self.handoff()
        self.agree()
        report = self.validate(require_complete=True)
        self.assert_valid(report)
        self.assertEqual(report['results'][index]['outcome'], 'secondary_temporal')
        self.assertIn('report_internal_inconsistency', self.cell(index, 'primarySummary').value)
        output = self.directory / 'report_inconsistency_results.xlsx'
        app.export_analysis([self.path], output)
        rows = app.records(load_workbook(output)['Decisions'])
        decision = next(r for r in rows if r['key'] == self.candidates[index]['key'])
        self.assertEqual(decision['primary_domain'], 'report_internal_inconsistency')
        self.assertEqual(decision['final_domain'], 'report_internal_inconsistency')
        self.assertEqual(decision['outcome'], 'secondary_temporal')

    def test_issue_vocabulary_and_removed_cluster_in_every_release_file(self):
        self.assertEqual(app.LABELS['statisticalIssueStatus'], 'Statistcal_or_numerical issue')
        self.assertNotIn('relatedClusterId', app.ASSESSMENT)
        for path in list(app.HERE.glob('*.xlsx')) + list((app.HERE / 'examples').glob('*.xlsx')):
            wb = load_workbook(path)
            for row in wb['Codebook']:
                for cell in row:
                    self.assertNotIn('error', str(cell.value).lower())
                    self.assertNotIn('wording_only', str(cell.value))
                    self.assertNotIn('relatedClusterId', str(cell.value))
            for field in ('internalStatus', 'protocolInterpretation'):
                values = [c.value for name, ref in wb.defined_names['opt_' + field].destinations for row in wb[name][ref] for c in row]
                self.assertIn('minor_presentation', values)
                self.assertIn('non_inconsistency', values)
                self.assertNotIn('wording_only', values)
            for name in [n for n in wb.sheetnames if n.startswith('C_')]:
                self.assertEqual(wb[name]['M1'].value, 'Primary assessment')
                self.assertEqual(wb[name]['R1'].value, 'Notes (optional)')
                for cell in wb[name][1]:
                    self.assertNotIn('error', cell.value.lower())
                    self.assertNotIn('error', cell.comment.text.lower())

    def test_retired_verdicts_and_wording_rejected_in_both_stages(self):
        self.fill_primary()
        self.assert_valid(self.validate(require_complete=True))
        for stage in ('primary', 'checker'):
            if stage == 'checker':
                self.handoff()
                self.agree()
            for field, value in [('statisticalIssueStatus', 'confirmed_error'), ('statisticalIssueStatus', 'probable_error'), ('statisticalIssueStatus', 'not_error'), ('internalStatus', 'wording_only'), ('protocolInterpretation', 'wording_only'), ('internalStatus', 'not_inconsistency'), ('protocolInterpretation', 'not_discrepancy')]:
                with self.subTest(stage=stage, field=field, value=value):
                    original = self.cell(0, field).value
                    self.set(0, field, value)
                    self.assert_error(self.validate(), field + ': invalid option')
                    self.cell(0, field).value = original

    def test_new_verdicts_preserve_outcomes_through_handoff_and_export(self):
        self.fill_primary()
        statistical = [i for i, c in enumerate(self.candidates) if c['primary'] in ('S_C', 'S_W')][:3]
        internal = next(i for i, c in enumerate(self.candidates) if c['primary'] in ('S_C', 'S_W') and i not in statistical)
        protocol = [i for i, c in enumerate(self.candidates) if c['primary'] == 'P_R'][:2]
        plan = next(i for i, c in enumerate(self.candidates) if c['primary'] == 'P_P')
        cases = [(i, 'statistical_results', 'statisticalIssueStatus', value) for i, value in zip(statistical, ('confirmed_issue', 'probable_issue', 'not_issue'))]
        cases += [(internal, 'report_internal_inconsistency', 'internalStatus', 'minor_presentation'), (plan, 'plan_material', 'internalStatus', 'non_inconsistency')]
        cases += [(i, 'protocol_report', 'protocolInterpretation', value) for i, value in zip(protocol, ('minor_presentation', 'non_inconsistency'))]
        for i, domain, field, value in cases:
            self.set(i, 'domain', domain)
            self.set(i, 'disposition', 'retain')
            self.set(i, field, value)
            if domain == 'statistical_results':
                self.set(i, 'substantiveStatus', 'substantive')
        self.set(statistical[0], 'notes', 'Context retained across reordered columns')
        self.assert_valid(self.validate(require_complete=True))
        self.handoff()
        self.agree()
        self.set(statistical[1], 'notes', 'Checker clarification in final column')
        self.set(statistical[1], 'checkStatus', 'agree_clarification')
        self.set(statistical[1], 'checkerComment', 'Added context')
        report = self.validate(require_complete=True)
        self.assert_valid(report)
        for i, domain, field, value in cases:
            self.assertEqual(report['reviews'][self.candidates[i]['key']][field], value)
        self.assertEqual(report['results'][statistical[0]]['outcome'], 'primary_temporal')
        self.assertEqual(report['results'][statistical[1]]['outcome'], 'primary_temporal')
        self.assertNotEqual(report['results'][statistical[2]]['outcome'], 'primary_temporal')
        for i in protocol:
            self.assertEqual(report['results'][i]['outcome'], 'protocol_descriptive')
        output = self.directory / 'new_vocab.xlsx'
        app.export_analysis([self.path], output)
        decisions = {r['key']: r for r in app.records(load_workbook(output)['Decisions'])}
        for i, domain, field, value in cases:
            row = decisions[self.candidates[i]['key']]
            self.assertEqual(row['primary_' + field], value)
            self.assertEqual(row['final_' + field], value)
            self.assertFalse(any('relatedCluster' in h or 'statisticalError' in h for h in row))
        self.assertEqual(decisions[self.candidates[statistical[0]]['key']]['final_notes'], 'Context retained across reordered columns')
        self.assertEqual(decisions[self.candidates[statistical[1]]['key']]['final_notes'], 'Checker clarification in final column')

    def test_unfinished_primary_handoff_rejected(self):
        self.validate()
        with self.assertRaisesRegex(ValueError, 'Validation failed'):
            app.prepare_checker(self.path, self.directory / 'bad.xlsx')

    def test_annual_primary_cannot_accept_checker_work_or_export(self):
        self.wb = load_workbook(self.checker_template)
        self.assert_valid(self.validate())
        with self.assertRaisesRegex(ValueError, 'prepared Checker'):
            app.export_analysis([self.path], self.directory / 'bad.xlsx', allow_incomplete=True)
        self.set(0, 'checkerComment', 'Started before handoff')
        self.assert_error(self.validate(), 'Checker fields must remain blank')

    def test_handoff_embeds_snapshot_and_visible_summary(self):
        self.fill_primary()
        self.set(0, 'notes', 'Paper-specific primary context')
        self.handoff()
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report['metadata']['stage'], 'checker')
        self.assertEqual(len(report['metadata']['primaryFileSHA256']), 64)
        self.assertEqual(self.wb['_Primary'].sheet_state, 'veryHidden')
        self.assertTrue(all(cell.protection.locked for row in self.wb['_Primary'] for cell in row))
        self.assertIn('Paper-specific primary context', self.cell(0, 'primarySummary').value)
        self.assertTrue(self.cell(0, 'primarySummary').protection.locked)
        self.assertEqual(self.wb['Setup']['B2'].value, 'primary_reviewer')
        self.assertTrue(self.wb['Setup']['B2'].protection.locked)
        self.assertFalse(self.wb['Setup']['B4'].protection.locked)

    def test_checker_cannot_be_reimported_as_primary(self):
        self.fill_primary()
        self.handoff()
        self.validate()
        with self.assertRaisesRegex(ValueError, 'Primary file'):
            app.prepare_checker(self.path, self.directory / 'bad.xlsx')

    def test_complete_agreement_export_contains_all_papers_and_zeroes(self):
        self.fill_primary()
        self.handoff()
        self.agree()
        report = self.validate(require_complete=True)
        self.assert_valid(report)
        self.assertEqual({r['status'] for r in report['results']}, {'resolved'})
        output = self.directory / 'analysis.xlsx'
        app.export_analysis([self.path], output)
        wb = load_workbook(output)
        papers = app.records(wb['Paper_counts'])
        self.assertEqual(len(papers), 10)
        self.assertEqual(sum(p['candidates'] for p in papers), 71)
        self.assertTrue(all(p['primary_temporal'] == 0 and p['paper_complete'] == 'yes' for p in papers))
        for name in ('Decisions', 'Protocol_plan_table'):
            self.assertFalse(any(any(word in str(c.value) for word in ('documentation', 'justification')) for c in wb[name][1]))

    def test_conditional_defaults_and_required_fields(self):
        for domain in app.MODULES:
            for disposition in ('', 'retain', 'exclude', 'discuss'):
                self.set(0, 'domain', domain)
                self.set(0, 'disposition', disposition)
                r = self.validate()['primary'][self.candidates[0]['key']]
                active = app.APPLICABLE[domain] if disposition in ('', 'retain') else ()
                for field in app.MODULE_FIELDS:
                    self.assertEqual(r[field], '' if field in active else app.NA)
        self.fill_primary()
        self.set(0, 'domain', 'statistical_results')
        self.set(0, 'disposition', 'retain')
        self.assert_error(self.validate(require_complete=True), 'statisticalIssueStatus: required')
        self.set(0, 'statisticalIssueStatus', 'confirmed_issue')
        self.set(0, 'substantiveStatus', 'substantive')
        self.assert_valid(self.validate(require_complete=True))

    def test_protocol_description_optional_and_excluded_from_temporal_analysis(self):
        self.fill_primary()
        index = next(i for i, c in enumerate(self.candidates) if c['primary'] == 'P_R')
        self.set(index, 'domain', 'protocol_report')
        self.set(index, 'disposition', 'retain')
        self.assert_valid(self.validate(require_complete=True))
        self.handoff()
        self.agree()
        report = self.validate(require_complete=True)
        self.assert_valid(report)
        self.assertEqual(report['results'][index]['outcome'], 'protocol_descriptive')

    def test_new_checker_domain_requires_fresh_judgment(self):
        self.fill_primary()
        self.handoff()
        self.agree()
        for field, value in [('disposition', 'retain'), ('domain', 'report_internal_inconsistency'), ('checkStatus', 'agree_clarification'), ('checkerComment', 'Change review domain')]:
            self.set(0, field, value)
        self.assert_error(self.validate(), 'internalStatus: required')
        self.set(0, 'internalStatus', 'substantive_inconsistency')
        self.assert_valid(self.validate())

    def test_agree_cannot_hide_modified_answers(self):
        self.fill_primary()
        self.handoff()
        self.agree()
        self.set(0, 'notes', 'Changed judgment')
        self.assert_error(self.validate(), 'unchanged primary')

    def test_disagreement_requires_comment_and_joint_resolution(self):
        self.fill_primary()
        self.handoff()
        self.agree()
        self.set(0, 'checkStatus', 'disagree')
        self.assert_error(self.validate(), 'checkerComment')
        self.set(0, 'checkerComment', 'Discuss the scope')
        self.assertEqual(self.validate()['results'][0]['status'], 'unresolved')
        self.set(0, 'resolutionStatus', 'resolved')
        self.assert_error(self.validate(), 'resolutionNote')
        self.set(0, 'resolutionNote', 'Primary and checker agreed to exclude')
        self.assertEqual(self.validate()['results'][0]['status'], 'resolved')

    def test_primary_discuss_still_needs_joint_resolution(self):
        self.fill_primary('discuss')
        self.handoff()
        self.agree()
        self.set(0, 'disposition', 'exclude')
        self.set(0, 'checkStatus', 'agree_clarification')
        self.set(0, 'checkerComment', 'Discussed together')
        self.assertEqual(self.validate()['results'][0]['status'], 'unresolved')

    def test_snapshot_summary_and_lookup_tampering_rejected(self):
        self.fill_primary()
        self.handoff()
        original = io.BytesIO()
        self.wb.save(original)
        self.wb['_Primary']['B2'] = 'discuss'
        with self.assertRaisesRegex(ValueError, 'snapshot/identity changed'):
            self.validate()
        self.wb = load_workbook(io.BytesIO(original.getvalue()))
        self.set(0, 'primarySummary', 'Wrong reference answer')
        self.assert_error(self.validate(), 'primary summary changed')
        self.wb = load_workbook(io.BytesIO(original.getvalue()))
        name = app.primary_range_name(self.candidates[0]['package'], 'domain')
        self.wb.defined_names[name].attr_text = "'Setup'!$B$2:$B$3"
        with self.assertRaisesRegex(ValueError, 'lookup range changed'):
            self.validate()

    def test_checker_missing_duplicate_or_misplaced_paper_rejected(self):
        self.wb = load_workbook(self.checker_template)
        a, b = [self.wb[name] for name in app.checker_sheet_names(self.candidates)[:2]]
        a.cell(2, 1).value = b.cell(2, 1).value
        with self.assertRaisesRegex(ValueError, 'key set mismatch'):
            self.validate()
        self.wb = load_workbook(self.checker_template)
        self.wb.remove(self.wb[app.checker_sheet_names(self.candidates)[0]])
        with self.assertRaisesRegex(ValueError, 'worksheet set mismatch'):
            self.validate()

    def test_complete_row_reordering_preserves_inheritance(self):
        self.fill_primary()
        self.set(0, 'notes', 'Unique context')
        self.handoff()
        self.agree()
        ws = self.wb[app.checker_sheet_names(self.candidates)[0]]
        rows = [[c.value for c in ws[i]] for i in (2, 3)]
        for target, source, values in ((3, 2, rows[0]), (2, 3, rows[1])):
            for col, value in enumerate(values, 1):
                cell = ws.cell(target, col)
                if isinstance(value, str) and value.startswith('='):
                    cell.value = Translator(value, origin=ws.cell(source, col).coordinate).translate_formula(cell.coordinate)
                else:
                    app.put(cell, value)
        report = self.validate()
        self.assert_valid(report)
        self.assertEqual(report['reviews'][self.candidates[0]['key']]['notes'], 'Unique context')

    def test_codebook_covers_exact_headers_and_dropdowns(self):
        self.fill_primary()
        ready = self.handoff()
        for path in (self.primary_template, ready):
            wb = load_workbook(path)
            ws = wb['Codebook']
            mapping = {ws.cell(r, 5).value: ws.cell(r, 1).value for r in range(2, ws.max_row+1) if ws.cell(r, 2).value == 'FIELD'}
            self.assertEqual(mapping, app.LABELS)
            for name in app.review_names(self.candidates, app.meta(wb)['stage']):
                for cell in wb[name][1]:
                    self.assertIn(cell.value, mapping.values())
            self.assertEqual(len(ws.merged_cells.ranges), 7)
            for field, options in app.OPTIONS.items():
                values = [c.value for sheet, reference in wb.defined_names['opt_' + field].destinations for row in wb[sheet][reference] for c in row]
                self.assertEqual(values, list(options))
            self.assertNotIn('Documentation', mapping.values())
            self.assertNotIn('Justification', mapping.values())

    def test_formula_and_nonenglish_input_rejected(self):
        self.cell(0, 'statisticalIssueStatus').value = '=IF(TRUE,"confirmed_issue","")'
        self.assert_error(self.validate(), 'formula/error cell')
        self.wb = load_workbook(self.primary_template)
        self.set(0, 'notes', '中文')
        with self.assertRaisesRegex(ValueError, 'English-only'):
            self.validate()

    def test_source_mutation_and_papers_index_rejected(self):
        self.wb['Evidence']['D2'] = 'Changed evidence'
        with self.assertRaisesRegex(ValueError, 'Frozen candidate'):
            self.validate()
        self.wb = load_workbook(self.primary_template)
        self.wb['Papers']['C2'] = 999
        with self.assertRaisesRegex(ValueError, 'Papers index mismatch'):
            self.validate()

    def test_distinct_ids_and_no_overwrites(self):
        self.fill_primary()
        self.validate()
        with self.assertRaisesRegex(ValueError, 'distinct'):
            app.prepare_checker(self.path, self.directory / 'bad.xlsx', 'PRIMARY_REVIEWER')
        before = self.path.read_bytes()
        with self.assertRaises(FileExistsError):
            app.prepare_checker(self.path, self.path)
        self.assertEqual(before, self.path.read_bytes())

    def test_provisional_export_keeps_unfinished_counts_blank(self):
        self.fill_primary()
        self.handoff()
        self.validate()
        output = self.directory / 'provisional.xlsx'
        app.export_analysis([self.path], output, allow_incomplete=True)
        rows = app.records(load_workbook(output)['Paper_counts'])
        self.assertTrue(all(r['primary_temporal'] == '' and r['paper_complete'] == 'no' for r in rows))
        with self.assertRaisesRegex(ValueError, 'Duplicate year'):
            app.export_analysis([self.path, self.path], self.directory / 'bad.xlsx', allow_incomplete=True)

    def test_restore_defaults_preserves_frozen_snapshot(self):
        self.fill_primary()
        self.handoff()
        self.set(0, 'internalStatus', '')
        before = self.validate()
        output = self.directory / 'restored.xlsx'
        app.restore_defaults(self.path, output)
        after = app.validate(output)
        self.assert_valid(after)
        self.assertEqual(before['primary'], after['primary'])
        self.assertEqual(before['metadata']['primarySignature'], after['metadata']['primarySignature'])

    def test_all_years_handoff_preserves_layout_and_both_review_stages(self):
        paths = []
        for year in app.YEARS:
            template = app.HERE / f'human_adjudication_v_2_1_5_1_{year}.xlsx'
            wb = load_workbook(template)
            app.put(wb['Setup']['B2'], 'primary_' + year)
            app.put(wb['Setup']['B3'], 'Primary Reviewer')
            names = [n for n in wb.sheetnames if n.startswith('C_')]
            for name in names:
                for row in range(2, wb[name].max_row + 1):
                    for field, value in [('disposition', 'exclude'), ('reviewComplete', 'yes'), ('notes', year + ' primary context')]:
                        app.put(wb[name].cell(row, app.REVIEW_FIELDS.index(field)+1), value)
            primary_path = self.directory / f'{year}_primary.xlsx'
            wb.save(primary_path)
            primary_report = app.validate(primary_path, require_complete=True)
            self.assert_valid(primary_report)
            ready = self.directory / f'{year}_checker.xlsx'
            app.prepare_checker(primary_path, ready, 'checker_' + year)
            report = app.validate(ready)
            self.assert_valid(report)
            self.assertEqual(report['primary'], primary_report['primary'])
            self.assertEqual(report['wb']['Setup']['B3'].value, 'Primary Reviewer')
            self.assertEqual([n for n in report['wb'].sheetnames if n.startswith('C_')], names)
            for name in names:
                ws = report['wb'][name]
                self.assertEqual([c.value for c in ws[1]], [c.value for c in wb[name][1]])
                self.assertEqual(ws.freeze_panes, wb[name].freeze_panes)
                for row in range(2, ws.max_row + 1):
                    self.assertEqual(ws.cell(row, app.CHECK_FIELDS.index('reviewComplete')+1).value, 'no')
                    self.assertIsNone(ws.cell(row, app.CHECK_FIELDS.index('disposition')+1).value)
                    for field, value in [('checkStatus', 'agree'), ('reviewComplete', 'yes')]:
                        cell = ws.cell(row, app.CHECK_FIELDS.index(field)+1)
                        self.assertFalse(cell.protection.locked)
                        app.put(cell, value)
            report['wb'].save(ready)
            completed = app.validate(ready, require_complete=True)
            self.assert_valid(completed)
            self.assertEqual({r['status'] for r in completed['results']}, {'resolved'})
            paths.append(ready)
        output = self.directory / 'four_year_analysis.xlsx'
        app.export_analysis(paths, output)
        exported = load_workbook(output)
        self.assertEqual(len(app.records(exported['Decisions'])), 412)
        self.assertEqual(len(app.records(exported['Paper_counts'])), 45)


if __name__ == '__main__':
    unittest.main(verbosity=2)
