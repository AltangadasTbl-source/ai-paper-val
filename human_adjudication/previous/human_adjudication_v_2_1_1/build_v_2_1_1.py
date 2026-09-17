#!/usr/bin/env python3
"""XLSX-only adjudication: freeze legacy candidates once, build, check and export.

Human judgments are never inferred from prefilled AI values. Formula caches are
never used by validation or analysis. Existing output files are never overwritten.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import re
import sys
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter as letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.views import Selection
from openpyxl.writer.theme import theme_xml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
VERSION = "v.2.1.1"
REVISION = "v.2.1.1"
SCHEMA = "xlsx-adjudication-3"
CATALOG_SCHEMA = "xlsx-adjudication-1"
HAN = re.compile("[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\U00020000-\U000323af]")
YEARS = {"2018": 119, "2019": 71, "2024": 91, "2025": 131}
RELATIONSHIPS = {"P_R": 45, "P_P": 50, "S_C": 65, "S_W": 252}
CATALOG = HERE / "source_catalog_v_2_1_1.xlsx"
SOURCE_SHEETS = ("Candidates", "Evidence", "Links", "Relations")
MODULES = {
    "statistical_results": ("statisticalErrorStatus", "substantiveStatus", "downstreamImpact"),
    "internal_nonstatistical": ("internalStatus",),
    "plan_material": ("internalStatus",),
    "protocol_report": ("documentationStatus", "justificationStatus", "protocolInterpretation"),
}
OPTIONS = {
    "disposition": {
        "retain": "Retain candidate. After human review, keep this finding in the applicable review results; retaining a protocol discrepancy does not automatically confirm an error.",
        "exclude": "Do not retain. The candidate has been assessed and should not remain as a finding, for example because the apparent mismatch has a satisfactory explanation.",
        "out_of_scope": "Outside review scope. The candidate falls outside the agreed review question or eligible material, regardless of whether the observation is correct.",
        "discuss": "Needs discussion. The reviewer has considered the candidate but cannot settle the disposition without further discussion. This is a decision to defer, not an untouched blank.",
    },
    "domain": {
        "statistical_results": "Statistical results (S_C / S_W). Assess reported numbers, estimates, intervals, denominators or statistical claims in the current report materials.",
        "internal_nonstatistical": "Internal nonstatistical (S_C / S_W). Assess inconsistent labels, descriptions or other nonstatistical content within or between current report materials.",
        "protocol_report": "Protocol-report (P_R). Assess a difference between planned and reported conduct or analyses, including whether the change was documented and justified.",
        "plan_material": "Plan material (P_P). Assess an inconsistency involving planning materials themselves rather than a planned-versus-reported result. The original relationship code remains unchanged.",
    },
    "statisticalErrorStatus": {
        "confirmed_error": "Confirmed statistical reporting error. The available evidence establishes a reporting error or contradiction, even if the correct replacement value is not known.",
        "probable_error": "Probable statistical reporting error. The evidence favors a reporting error, but an unresolved limitation or plausible alternative prevents confirmation.",
        "not_error": "Not a statistical error. The apparent discrepancy is consistent with a valid calculation, definition, analysis method or reporting context.",
        "indeterminate": "Indeterminate. The available evidence does not support deciding whether a statistical reporting error exists; neither confirmation nor dismissal is justified.",
    },
    "substantiveStatus": {
        "substantive": "Substantive. The issue could meaningfully change interpretation, data extraction or downstream analysis, rather than only the presentation of the result.",
        "minor_presentation": "Minor or presentation only. The intended result remains clear and the issue is a small formatting, labeling or typographic problem without a meaningful change in interpretation.",
        "indeterminate": "Indeterminate. There is not enough information to decide whether the issue is substantive or limited to presentation.",
    },
    "downstreamImpact": {
        "direct": "Direct extraction or analysis impact. The affected value or definition is itself an input to data extraction or subsequent analysis; this describes potential impact, not proven downstream harm.",
        "possible": "Possible downstream impact. A plausible route to affecting later interpretation, extraction or analysis exists, but the connection is indirect or uncertain.",
        "unlikely": "Unlikely downstream impact. The issue is unlikely to change extracted information or later analysis because the relevant result remains clear and usable.",
        "unknown": "Unable to assess. The materials do not establish enough about the affected information or its likely use to judge downstream impact.",
    },
    "internalStatus": {
        "substantive_inconsistency": "Substantive inconsistency. The conflicting descriptions or planning statements could materially change how the study, population, procedures or results are understood.",
        "minor_presentation": "Minor presentation inconsistency. A real discrepancy exists, but it mainly affects presentation and the intended meaning remains clear.",
        "wording_only": "Wording only. Different phrases express the same substantive meaning; there is no meaningful conflict between the statements.",
        "not_inconsistency": "Not an inconsistency. The statements are compatible once their definitions, populations, time points or other context are taken into account.",
        "indeterminate": "Indeterminate. The available material is insufficient to establish whether the statements conflict or how important any conflict is.",
    },
    "documentationStatus": {
        "explicit": "Explicitly documented. The reviewed material directly identifies the change from the plan, rather than merely presenting a different method or result.",
        "partial": "Partial or implicit documentation. Some acknowledgment or description of the change exists, but its nature or extent is not fully stated.",
        "none_located": "No documentation located. No acknowledgment of the change was found in the reviewed materials; this does not prove that documentation does not exist elsewhere.",
        "cannot_assess": "Cannot assess. Missing material or unclear version history prevents judging whether the change was documented.",
        "not_applicable": "Not applicable. Documentation of a change is not relevant to this candidate, for example when no substantive change is being assessed.",
    },
    "justificationStatus": {
        "explicit": "Explicit justification located. The reviewed material states a reason for the change. Recording that reason exists does not by itself establish that the reason is adequate.",
        "none_located": "No justification located. A reason for the change was not found in the reviewed materials; this is distinct from proving that the change had no valid reason.",
        "cannot_assess": "Cannot assess. The materials are insufficient to determine whether a reason was provided or how it relates to the change.",
        "not_applicable": "Not applicable. A justification for change is not relevant to the candidate being assessed, for example when the difference is wording only.",
    },
    "protocolInterpretation": {
        "documented_justified": "Documented and justified. The change is acknowledged and its stated rationale adequately explains it in the reviewer's judgment; the difference is not automatically a reporting error.",
        "documented_no_justification": "Documented; no justification located. The change is acknowledged, but the reviewed materials do not provide a reason explaining it.",
        "undocumented_change_candidate": "Possible undocumented change. A substantive planned-versus-reported difference appears present, but acknowledgment of the change was not located in the reviewed materials.",
        "wording_only": "Wording only. The apparent planned-versus-reported difference does not represent a substantive change in meaning or conduct.",
        "version_timing": "Version or timing ambiguity. The difference may reflect which protocol version or date is being compared; the sequence cannot be resolved from the available material.",
        "indeterminate": "Indeterminate. The available evidence cannot support a more specific interpretation of the planned-versus-reported difference.",
    },
    "reviewComplete": {
        "no": "Draft. This review stage is not yet confirmed complete; AI suggestions or partially filled fields do not count as completed human review.",
        "yes": "Human reviewed and complete. The reviewer has checked this row and its required fields. A discussion item may still be unresolved even though this review stage is complete.",
    },
    "checkStatus": {
        "agree": "Agree, assessment unchanged. Accept the primary assessment as recorded. Leave override cells blank; no checker comment is required.",
        "agree_clarification": "Agree with clarification. Accept the main assessment but add context or a clarifying amendment. Explain it in a checker comment and fill only any changed fields.",
        "disagree": "Disagree, proposal pending discussion. Recommend a different assessment, explain why in a comment, and enter proposed changes. The proposal is not final consensus until discussion is resolved.",
        "unable": "Unable to resolve. The checker cannot settle the issue from the available information. State what remains uncertain in a comment and refer it for discussion.",
    },
    "resolutionStatus": {
        "pending": "Discussion not resolved. A required discussion is not yet settled. For an unchanged agreement this default can be left as-is because no separate resolution is required.",
        "resolved": "Joint discussion resolved. Record the participants and agreed conclusion in Resolution note, update the proposed assessment, and do not leave disposition as discuss.",
    },
}
LABELS = {
    "key": 'Candidate key', "package": 'Paper', "title": 'Candidate title', "locations": 'Source location',
    "disposition": 'Disposition', "domain": 'Review domain', "reviewComplete": 'Complete', "readiness": 'Validation hint',
    "statisticalErrorStatus": 'Statistical error', "substantiveStatus": 'Substantiveness', "downstreamImpact": 'Downstream impact',
    "internalStatus": 'Internal or plan assessment', "documentationStatus": 'Documentation', "justificationStatus": 'Justification', "protocolInterpretation": 'Protocol interpretation',
    "relatedClusterId": 'Related cluster (optional)', "notes": 'Notes (optional)', "checkStatus": 'Checker result', "checkerComment": 'Checker comment', "resolutionStatus": 'Discussion resolution', "resolutionNote": 'Resolution note',
}
ASSESSMENT = ("disposition", "domain", "statisticalErrorStatus", "substantiveStatus", "downstreamImpact", "internalStatus", "documentationStatus", "justificationStatus", "protocolInterpretation", "relatedClusterId", "notes")
REVIEW_FIELDS = ("key", "title", "locations", "disposition", "domain", "reviewComplete", "readiness") + ASSESSMENT[2:]
CHECK_FIELDS = ("key", "package", "title", "checkStatus", "checkerComment", "reviewComplete", "readiness", "resolutionStatus", "resolutionNote") + ASSESSMENT
BLUE = "173F57"
YELLOW = "FFF2CC"
GRAY = "EDF2F5"


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def now():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def require_english_xlsx(path_or_stream):
    """Check decoded XML, including hidden cells, themes, comments and properties."""
    with zipfile.ZipFile(path_or_stream) as archive:
        for name in archive.namelist():
            if HAN.search(name):
                raise ValueError("English-only XLSX required: non-English package member name")
            if not name.endswith((".xml", ".rels", ".vml")):
                continue
            root = ET.fromstring(archive.read(name))
            for node in root.iter():
                values = [node.text or "", node.tail or "", *node.attrib.values()]
                if any(HAN.search(value) for value in values):
                    raise ValueError(f"English-only XLSX required: CJK text in {name}")


def primary_sheet_name(package):
    name = "P_" + package
    if len(name) > 31 or re.search(r"[\\/*?:\[\]]", name):
        raise ValueError(f"Invalid paper worksheet name: {name}")
    return name


def paper_groups(candidates):
    groups = defaultdict(list)
    for c in candidates:
        groups[c["package"]].append(c)
    return groups


def primary_sheet_names(candidates):
    return [primary_sheet_name(paper) for paper in paper_groups(candidates)]


def put(cell, value):
    """Force all source/free text to literal strings, including leading '='."""
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError(f"Text exceeds Excel cell limit: {cell.coordinate}")
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value


def table(wb, name, headers, rows, widths=None, protect=True):
    ws = wb.create_sheet(name)
    for i, header in enumerate(headers, 1):
        put(ws.cell(1, i), header)
        ws.cell(1, i).font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        ws.cell(1, i).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(1, i).alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[letter(i)].width = (widths or {}).get(header, 25)
    for row in rows:
        for col, value in enumerate(row, 1):
            cell = ws.cell(ws.max_row + 1, col) if col == 1 else ws.cell(target_row, col)
            if col == 1:
                target_row = cell.row
            put(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=11)
            if target_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F8FA")
        lines = max((sum(max(1, math.ceil(len(part) / max(8, ws.column_dimensions[letter(i)].width - 3))) for part in str(value).splitlines()) for i, value in enumerate(row, 1)), default=1)
        ws.row_dimensions[target_row].height = min(300, max(30, lines * 16 + 8))
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    if protect:
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
    return ws


def empty_book():
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def metadata(wb, values):
    table(wb, "Manifest", ["field", "value"], list(values.items()), {"field": 30, "value": 110})


def meta(wb):
    rows = records(wb["Manifest"], ["field", "value"])
    if len({r["field"] for r in rows}) != len(rows):
        raise ValueError("Duplicate Manifest fields")
    return {r["field"]: r["value"] for r in rows}


def records(ws, expected=None):
    headers = [c.value for c in ws[1]]
    if expected is not None and headers != list(expected):
        raise ValueError(f"{ws.title}: headers changed; use the original template")
    if len(set(headers)) != len(headers) or None in headers:
        raise ValueError(f"{ws.title}: duplicate or blank headers")
    result = []
    for cells in ws.iter_rows(min_row=2):
        values = [c.value for c in cells]
        if all(v is None for v in values):
            continue
        result.append(dict(zip(headers, ["" if v is None else v for v in values])))
    return result


def source_tables(candidates):
    containers = {"aiInitialAssessment", "evidence", "sourceLinks", "evidencePaths", "relations", "legacyRelations"}
    scalar = sorted(set().union(*(c.keys() for c in candidates)) - containers)
    fields = scalar + ["ai_" + f for f in ("source", "relationship", "domain", "disposition")]
    rows = [[c.get(f, "") if not f.startswith("ai_") else c["aiInitialAssessment"][f[3:]] for f in fields] for c in candidates]
    evidence, links, relations = [], [], []
    for c in candidates:
        evidence.extend([c["key"], i, k, v] for i, (k, v) in enumerate(c["evidence"].items()))
        for kind in ("sourceLinks", "evidencePaths"):
            links.extend([c["key"], kind, i, v["label"], v["href"]] for i, v in enumerate(c[kind]))
        for kind in ("relations", "legacyRelations"):
            relations.extend([c["key"], kind, i, v] for i, v in enumerate(c[kind]))
    return {"Candidates": (fields, rows), "Evidence": (["key", "order", "section", "text"], evidence), "Links": (["key", "kind", "order", "label", "href"], links), "Relations": (["key", "kind", "order", "relationship"], relations)}


def source_signature(wb):
    # Sorting whole rows permits user sorting without hiding edits or duplicate rows.
    return digest({name: sorted(records(wb[name]), key=lambda r: json.dumps(r, sort_keys=True, ensure_ascii=False)) for name in SOURCE_SHEETS})


def read_candidates(wb):
    candidates = []
    for row in records(wb["Candidates"]):
        c = {k: v for k, v in row.items() if not k.startswith("ai_")}
        c["aiInitialAssessment"] = {k[3:]: v for k, v in row.items() if k.startswith("ai_")}
        c.update(evidence={}, sourceLinks=[], evidencePaths=[], relations=[], legacyRelations=[])
        candidates.append(c)
    by_key = {c["key"]: c for c in candidates}
    if len(by_key) != len(candidates):
        raise ValueError("Duplicate source candidate keys")
    for r in sorted(records(wb["Evidence"]), key=lambda r: (r["key"], r["order"])):
        by_key[r["key"]]["evidence"][r["section"]] = r["text"]
    for r in sorted(records(wb["Links"]), key=lambda r: (r["key"], r["kind"], r["order"])):
        by_key[r["key"]][r["kind"]].append({"label": r["label"], "href": r["href"]})
    for r in sorted(records(wb["Relations"]), key=lambda r: (r["key"], r["kind"], r["order"])):
        by_key[r["key"]][r["kind"]].append(r["relationship"])
    return candidates


def local_target(href):
    parsed = urlsplit(href)
    if parsed.scheme in ("http", "https"):
        return None
    if parsed.scheme or not href.startswith(tuple(f"../../{y}/" for y in YEARS)):
        raise ValueError(f"Nonportable source link: {href}")
    target = (HERE / unquote(parsed.path)).resolve()
    if not target.is_relative_to(ROOT) or not target.is_file():
        raise ValueError(f"Missing or invalid source link: {href}")
    return target


def validate_sources(candidates, all_years=True, check_files=False):
    if len({c["key"] for c in candidates}) != len(candidates):
        raise ValueError("Duplicate candidate keys")
    if all_years and (Counter(str(c["year"]) for c in candidates) != YEARS or Counter(c["primary"] for c in candidates) != RELATIONSHIPS):
        raise ValueError("Candidate counts differ from the frozen 412-candidate universe")
    count = 0
    for c in candidates:
        if c["key"] != f'{c["year"]}::{c["package"]}::{c["id"]}':
            raise ValueError(f"Candidate identity mismatch: {c['key']}")
        if check_files:
            for href in [c["reportHref"], "../../" + c["reportPath"]] + [r["href"] for kind in ("sourceLinks", "evidencePaths") for r in c[kind]]:
                local_target(href)
                count += 1
    return count


def add_sources(wb, candidates, output):
    for name, (headers, rows) in source_tables(candidates).items():
        widths = {"key": 45, "title": 65, "statement": 90, "text": 115, "section": 38, "label": 60, "href": 95, "rationale": 90, "locations": 80}
        ws = table(wb, name, headers, rows, widths)
        if name == "Evidence":
            ws.column_dimensions["B"].hidden = True
            for r in range(2, ws.max_row + 1):
                text = str(ws.cell(r, 4).value or "")
                lines = sum(max(1, math.ceil(len(line) / 105)) for line in text.splitlines())
                ws.row_dimensions[r].height = min(400, max(48, lines * 16 + 10))
        elif name == "Links":
            ws.sheet_state = "hidden"
        elif name == "Relations":
            ws.sheet_state = "hidden"


def save_new(wb, path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Output must be an .xlsx workbook")
    # Use standard English font names, without changing evidence or answers.
    theme = wb.loaded_theme or theme_xml
    if isinstance(theme, bytes):
        theme = theme.decode("utf-8")
    fonts = {"\u5b8b\u4f53": "SimSun", "\u65b0\u7d30\u660e\u9ad4": "PMingLiU"}
    theme_root = ET.fromstring(theme)
    for node in theme_root.iter():
        if node.get("typeface") in fonts:
            node.set("typeface", fonts[node.get("typeface")])
    wb.loaded_theme = ET.tostring(theme_root, encoding="utf-8", xml_declaration=True)
    content = io.BytesIO()
    wb.save(content)
    content.seek(0)
    require_english_xlsx(content)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation protects completed human records and template files.
    with path.open("xb") as handle:
        handle.write(content.getvalue())
    # Reopen the actual zip to catch serialization problems immediately.
    with path.open("rb") as handle:
        check = load_workbook(handle)
        check.close()
    print(f"Created {path}")


def freeze_source(output):
    candidates, inputs = [], []
    for year in YEARS:
        path = HERE.parent / "human_adjudication_v_2_0_4" / f"candidate_flag_review_v_2_0_4_{year}.html"
        content = path.read_text(encoding="utf-8")
        match = re.search(r'<script id="candidate-data" type="application/json">(.*?)</script>', content, re.S)
        if not match:
            raise ValueError(f"Missing legacy candidate payload: {path}")
        payload = json.loads(match[1])
        rows = payload["candidates"]
        legacy_digest = hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        if payload["reviewVersion"] != "v.2.0.4" or payload["schemaVersion"] != 7 or payload["years"] != [year] or payload["keySet"] != [c["key"] for c in rows] or legacy_digest != payload["sourceDigest"]:
            raise ValueError(f"Legacy payload integrity error: {year}")
        candidates.extend(rows)
        inputs.append((f"legacy_{year}_sha256", hashlib.sha256(path.read_bytes()).hexdigest()))
    validate_sources(candidates, check_files=True)
    if Counter(c["aiInitialAssessment"]["disposition"] for c in candidates) != {"retain": 402, "exclude": 9, "discuss": 1}:
        raise ValueError("Unexpected AI starting values")
    wb = empty_book()
    add_sources(wb, candidates, output)
    # Verify a lossless normalized conversion, including all evidence and metadata.
    if read_candidates(wb) != candidates:
        raise ValueError("Lossy legacy-to-XLSX source conversion")
    metadata(wb, {"version": VERSION, "schema": CATALOG_SCHEMA, "stage": "catalog", "createdAt": now(), "sourceSignature": source_signature(wb), **dict(inputs), "provenance": "v.2.0.1 -> v.2.0.2 -> v.2.0.3 -> v.2.0.4; no v_1_0_* inputs", "humanRecordsImported": "none; bundled tester snapshot explicitly excluded"})
    save_new(wb, output)


def load_catalog(path=CATALOG):
    require_english_xlsx(path)
    wb = load_workbook(path)
    m = meta(wb)
    if m.get("version") != VERSION or m.get("stage") != "catalog" or m.get("schema") != CATALOG_SCHEMA or m.get("sourceSignature") != source_signature(wb):
        raise ValueError("Source catalog integrity/version mismatch")
    candidates = read_candidates(wb)
    validate_sources(candidates)
    return candidates, m["sourceSignature"]


def add_codebook(wb):
    rows = []
    for field, options in OPTIONS.items():
        start = len(rows) + 2
        rows.extend([field, code, label] for code, label in options.items())
        wb.defined_names.add(DefinedName("opt_" + field, attr_text=f"'Codebook'!$B${start}:$B${len(rows)+1}"))
    rows.extend([
        ["disposition", "", "Blank means no human disposition has been selected. It is not retain, exclude or discuss. Choose an option before marking primary review complete; previous AI disposition is reference-only on Candidates."],
        ["relatedClusterId", "", 'Optional grouping of distinct candidates; never deduplicates'],
        ["domain", "", 'All four choices are editable for every candidate; source relationship remains fixed'],
        ["reviewComplete", "", 'Choose yes only after checking this row. After any edit, set no and recheck'],
        ["notes", "", 'Optional. No mandatory reproduction log or primary rationale'],
        ["resolutionNote", "", 'Required for resolved disagreement/unable or a primary discuss changed to a decision. Record participants and conclusion'],
        ["Checker overrides", "", 'Blank assessment cells inherit the matching candidate on its P_<paper> worksheet. Fill only changes; <clear> clears inherited notes/cluster.'],
    ])
    table(wb, "Codebook", ["field", "code", 'meaning'], rows, {"field": 30, "code": 40, 'meaning': 115})


def instructions(wb, year, stage):
    rows = [
        ['Release', f"{VERSION} | revision {REVISION} | {year} | {stage} | XLSX-only human adjudication"],
        ['1. Identity', 'Fill primary reviewer ID on Setup; the second person fills checker ID.'],
        ['2. Paper worksheets', 'Each P_<paper> worksheet contains primary review for one paper only. Papers lists worksheet names and candidate counts. Checker remains the annual second-person review table.'],
        ['Evidence', 'Candidates contains descriptions and source locations; Evidence contains full evidence text. Filter by key and consult originals in the shared 20XX folders. No document links or page navigation are provided.'],
        ['3. Enter', 'Only yellow cells are inputs. Disposition starts blank: choose it yourself after review. The suggested domain remains editable; previous AI disposition is reference-only on Candidates. Use English for names, notes and comments; Codebook explains when to use each option.'],
        ['4. Retain module', 'retain + statistical_results: error, substantiveness, impact; protocol_report: documentation, justification, interpretation; other two domains: internalStatus. Other dispositions only need a valid domain.'],
        ['5. Save', 'Set Complete=yes only after review; Ctrl+S saves the XLSX. After editing a completed row, reset to no, recheck, then yes. No macros and no automatic per-edit timestamp/reset.'],
        ['6. Handoff', 'Save the primary-completed workbook as a dated copy and pass it to the checker. The checker fills Checker in the same XLSX, leaving all P_<paper> worksheets unchanged. No script is needed for handoff.'],
        ['7. Checker', 'Non-agree requires a comment. agree requires identical assessment. Disagree/unable stays unresolved until resolutionStatus=resolved plus resolutionNote. Primary discuss changed to a decision also needs documented resolution.'],
        ['8. Validation', 'Dropdowns restrict choices. Validation hint is a convenience; optional Python validation checks saved answers, including pasted values.'],
        ['9. Minimal checker entry', 'For agreement: select agree and Complete=yes. Blank assessment cells inherit the same candidate key from its P_<paper> worksheet. Fill only changes; <clear> clears inherited notes/cluster.'],
        ['10. Research', 'Exploratory audit aid: human judgment remains responsible. Protocol/plan discrepancies remain descriptive and are not automatically errors. Temporal outcomes require substantive S_C/S_W judgments. Visible-answer checking is not blinded inter-rater reliability.'],
        ['11. Current status', '412 frozen AI candidates in 45 papers; human review is pending. Do not report AI retained counts as confirmed errors. See USAGE.md for workflow and project rationale.'],
    ]
    ws = table(wb, "Start", ['Step', 'Instructions'], rows, {'Step': 32, 'Instructions': 130})
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 68


def default_record(c):
    r = dict.fromkeys(ASSESSMENT, "")
    r.update(disposition="", domain=c["aiInitialAssessment"]["domain"], reviewComplete="no")
    return r


def row_hint(row, fields, stage):
    col = {f: f"{letter(fields.index(f)+1)}{row}" for f in fields}
    d, domain, done = col["disposition"], col["domain"], col["reviewComplete"]
    if stage == "checker":
        status, comment = col["checkStatus"], col["checkerComment"]
        return f'=IF({status}="","PENDING",IF(AND({status}<>"agree",LEN(TRIM({comment}))=0),"COMMENT NEEDED",IF({done}="yes","CHECKED - see resolution","DRAFT")))'
    missing = [f'COUNTIF(opt_disposition,{d})<>1', f'COUNTIF(opt_domain,{domain})<>1']
    for name, required in MODULES.items():
        invalid = ",".join(f'COUNTIF(opt_{field},{col[field]})<>1' for field in required)
        missing.append(f'AND({d}="retain",{domain}="{name}",OR({invalid}))')
    if stage == "checker":
        missing.extend([f'COUNTIF(opt_checkStatus,{col["checkStatus"]})<>1', f'AND({col["checkStatus"]}<>"agree",LEN(TRIM({col["checkerComment"]}))=0)'])
    return f'=IF(OR({",".join(missing)}),"CHECK FIELDS",IF({done}="yes","READY - validate saved file","DRAFT"))'


def apply_review_layout(ws, stage):
    """Keep only the candidate key frozen, leaving room for several answer columns."""
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    widths = {
        "key": 22, "package": 20, "title": 28, "locations": 22,
        "disposition": 18, "domain": 24, "reviewComplete": 12,
        "readiness": 22, "internalStatus": 22,
        "notes": 32, "checkerComment": 32, "resolutionNote": 32,
    }
    for index, field in enumerate(fields, 1):
        ws.column_dimensions[letter(index)].width = widths.get(field, 18)
    ws.freeze_panes = "B2"
    # Reset the pane selections: changing an existing freeze can otherwise append
    # duplicate selections pointing to the old frozen region.
    ws.sheet_view.selection = [
        Selection(pane="topRight", activeCell="B1", sqref="B1"),
        Selection(pane="bottomLeft", activeCell="A2", sqref="A2"),
        Selection(pane="bottomRight", activeCell="B2", sqref="B2"),
    ]


def add_review(wb, candidates, reviews, stage, name="Primary", locked=False):
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    rows = []
    for c in candidates:
        values = {"key": c["key"], "package": c["package"], "title": c["title"], "locations": c["locations"], **reviews[c["key"]]}
        rows.append([values.get(f, "") for f in fields])
    ws = table(wb, name, [LABELS[f] for f in fields], rows)
    apply_review_layout(ws, stage)
    ws.sheet_properties.tabColor = "D7B95C" if not locked else "8B9CA6"
    for index, c in enumerate(candidates, 2):
        ws.row_dimensions[index].height = 76
        for field in fields:
            cell = ws.cell(index, fields.index(field) + 1)
            if field == "readiness":
                cell.value = row_hint(index, fields, stage)
            elif field in ASSESSMENT + ("reviewComplete", "checkStatus", "checkerComment", "resolutionStatus", "resolutionNote"):
                cell.protection = Protection(locked=locked)
                cell.fill = PatternFill("solid", fgColor=GRAY if locked else YELLOW)
    for field in fields:
        col = letter(fields.index(field) + 1)
        ws[f"{col}1"].comment = Comment(LABELS[field] + "\n" + ("\n".join(f"{k}: {v}" for k, v in OPTIONS[field].items()) if field in OPTIONS else "See Start and Codebook for instructions."), "Adjudication v.2.1.1")
        if field in OPTIONS and not locked:
            validation = DataValidation(type="list", formula1="opt_" + field, allow_blank=True)
            validation.errorTitle = 'Invalid option'
            validation.error = 'Choose a listed value.'
            validation.errorStyle = "stop"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            validation.promptTitle = field[:32]
            validation.prompt = 'See Codebook for meanings.'
            validation.showDropDown = False
            ws.add_data_validation(validation)
            validation.add(f"{col}2:{col}{ws.max_row}")
    d = letter(fields.index("disposition") + 1)
    domain = letter(fields.index("domain") + 1)
    for field in (() if stage == "checker" else ASSESSMENT[2:9]):
        col = letter(fields.index(field) + 1)
        active = [k for k, v in MODULES.items() if field in v]
        condition = ",".join(f'${domain}2="{k}"' for k in active)
        formula = f'OR(${d}2<>"retain",NOT(OR({condition})))'
        ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=GRAY)))
    return ws


def make_review(candidates, catalog_signature, output, stage="workbook", primary=None, primary_id="", checker_id="", primary_hash=""):
    wb = empty_book()
    year = str(candidates[0]["year"])
    instructions(wb, year, stage)
    setup = table(wb, "Setup", ["field", "value"], [["primaryReviewerId", primary_id], ["primaryReviewerName", ""], ["checkerId", checker_id], ["checkerName", ""]], {"field": 32, "value": 65})
    for row in ((2, 3, 4, 5) if stage == "workbook" else (4, 5)):
        setup.cell(row, 2).protection = Protection(locked=False)
        setup.cell(row, 2).fill = PatternFill("solid", fgColor=YELLOW)
    add_sources(wb, candidates, output)
    add_codebook(wb)
    primary = primary or {c["key"]: default_record(c) for c in candidates}
    papers = paper_groups(candidates)
    table(wb, "Papers", ["paper", "primary_worksheet", "candidates"], [[paper, primary_sheet_name(paper), len(rows)] for paper, rows in papers.items()], {"paper": 35, "primary_worksheet": 38})
    for paper, rows in papers.items():
        add_review(wb, rows, primary, "primary", name=primary_sheet_name(paper), locked=stage == "checker")
    reviews = {key: {**dict.fromkeys(ASSESSMENT, ""), "reviewComplete": "no", "checkStatus": "", "checkerComment": "", "resolutionStatus": "pending", "resolutionNote": ""} for key in primary}
    add_review(wb, candidates, reviews, "checker", name="Checker")
    metadata(wb, {"version": VERSION, "schema": SCHEMA, "stage": stage, "year": year, "createdAt": now(), "catalogSignature": catalog_signature, "sourceSignature": source_signature(wb), "primarySignature": digest({"id": primary_id, "reviews": primary}) if stage == "checker" else "", "primaryFileSHA256": primary_hash, "revision": REVISION})
    # The normalized source sheets preserve provenance without adding daily forms.
    visible_context = {"key", "package", "title", "statement", "locations", "primary", "ai_domain", "ai_disposition"}
    for cell in wb["Candidates"][1]:
        wb["Candidates"].column_dimensions[letter(cell.column)].hidden = cell.value not in visible_context
    for row in range(2, wb["Candidates"].max_row + 1):
        wb["Candidates"].row_dimensions[row].height = 95
    wb["Manifest"].sheet_state = "hidden"
    order = ["Start", "Setup", "Papers"] + primary_sheet_names(candidates) + ["Checker", "Candidates", "Evidence", "Codebook", "Manifest", "Links", "Relations"]
    wb._sheets = [wb[s] for s in order]
    wb.active = 0
    wb.security.lockStructure = True
    return wb


def build(output_dir, catalog=CATALOG):
    candidates, signature = load_catalog(catalog)
    targets = [output_dir / f"human_adjudication_v_2_1_1_{year}.xlsx" for year in YEARS]
    if any(p.exists() for p in targets):
        raise ValueError("Output already exists; use a new --output-dir. Human files are never overwritten.")
    for year, output in zip(YEARS, targets):
        annual = [c for c in candidates if str(c["year"]) == year]
        save_new(make_review(annual, signature, output), output)


def review_records(wb, sheet="Primary", stage="primary"):
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    rows = records(wb[sheet], [LABELS[f] for f in fields])
    result = {}
    for row in rows:
        r = {f: row[LABELS[f]] for f in fields}
        key = r.pop("key")
        if key in result:
            raise ValueError(f"{sheet}: duplicate candidate key {key}")
        for field in ("package", "title", "locations", "readiness"):
            r.pop(field, None)
        result[key] = r
    return result


def read_primary(wb, candidates):
    expected = set(primary_sheet_names(candidates))
    actual = {s for s in wb.sheetnames if s.startswith("P_")}
    if expected != actual or "Primary" in wb.sheetnames:
        raise ValueError("Primary paper worksheet set mismatch")
    result = {}
    for paper, rows in paper_groups(candidates).items():
        records_for_paper = review_records(wb, primary_sheet_name(paper))
        if set(records_for_paper) != {c["key"] for c in rows}:
            raise ValueError(f"Primary candidate-key set mismatch for paper {paper}")
        result.update(records_for_paper)
    expected_index = {paper: (primary_sheet_name(paper), len(rows)) for paper, rows in paper_groups(candidates).items()}
    index_rows = records(wb["Papers"], ["paper", "primary_worksheet", "candidates"])
    actual_index = {r["paper"]: (r["primary_worksheet"], r["candidates"]) for r in index_rows}
    if actual_index != expected_index or len(index_rows) != len(expected_index):
        raise ValueError("Papers index mismatch")
    return result


def assessment_errors(r, require_fields):
    errors = []
    for field in ASSESSMENT + ("reviewComplete",):
        value = r.get(field, "")
        if not isinstance(value, str):
            errors.append(f"{field}: expected text")
        elif field in OPTIONS and value and value not in OPTIONS[field]:
            errors.append(f"{field}: invalid option {value!r}")
    required = ["disposition", "domain", "reviewComplete"]
    if r.get("disposition") == "retain":
        required += list(MODULES.get(r.get("domain"), ()))
    if require_fields:
        errors.extend(f"{field}: required" for field in required if not r.get(field))
    return errors


def unresolved(r, primary):
    if r["disposition"] == "discuss":
        return True
    needs_resolution = r.get("checkStatus") in ("disagree", "unable") or primary["disposition"] == "discuss"
    return needs_resolution and r.get("resolutionStatus") != "resolved"


def outcome(c, r):
    if r["disposition"] != "retain":
        return "not_eligible"
    if c["primary"] == "P_R":
        return "protocol_descriptive"
    if c["primary"] == "P_P":
        return "plan_descriptive"
    if r["domain"] == "statistical_results" and r["statisticalErrorStatus"] in ("confirmed_error", "probable_error") and r["substantiveStatus"] == "substantive":
        return "primary_temporal"
    if r["domain"] == "internal_nonstatistical" and r["internalStatus"] == "substantive_inconsistency":
        return "secondary_temporal"
    return "not_eligible"


def validate(path, catalog=CATALOG, require_complete=False, primary_only=False):
    require_english_xlsx(path)
    wb = load_workbook(path, data_only=False)
    m = meta(wb)
    if m.get("version") != VERSION or m.get("revision") != REVISION or m.get("schema") != SCHEMA or m.get("stage") not in ("workbook", "checker"):
        raise ValueError("Unsupported workbook version/schema/stage")
    candidates, signature = load_catalog(catalog)
    candidates = [c for c in candidates if str(c["year"]) == m["year"]]
    if not candidates or m["catalogSignature"] != signature:
        raise ValueError("Wrong year or catalog provenance")
    expected = empty_book()
    add_sources(expected, candidates, Path(path))
    if source_signature(wb) != source_signature(expected) or m["sourceSignature"] != source_signature(expected):
        raise ValueError("Frozen candidate/evidence content changed or candidate set incomplete")
    stage = m["stage"]
    setup_rows = records(wb["Setup"], ["field", "value"])
    if [r["field"] for r in setup_rows] != ["primaryReviewerId", "primaryReviewerName", "checkerId", "checkerName"]:
        raise ValueError("Setup structure changed")
    setup = {r["field"]: r["value"] for r in setup_rows}
    primary_id, checker_id = setup["primaryReviewerId"], setup["checkerId"]
    issues = []
    def issue(key, severity, message):
        issues.append((key, severity, message))
    def valid_id(value):
        return isinstance(value, str) and bool(value.strip()) and value == value.strip()
    identity_needed = require_complete
    overrides = review_records(wb, "Checker", stage="checker")
    keys = {c["key"] for c in candidates}
    if set(overrides) != keys:
        raise ValueError("Checker candidate-key set mismatch (missing, extra, or wrong-year rows)")
    primary = read_primary(wb, candidates)
    if set(primary) != keys:
        raise ValueError("Primary candidate-key set mismatch")
    if stage == "checker" and m["primarySignature"] != digest({"id": primary_id, "reviews": primary}):
        raise ValueError("Locked primary snapshot/identity changed")
    # Reject formulas in inputs/source cells, including pasted formulas with stale caches.
    review_sheets = tuple(primary_sheet_names(candidates)) + ("Checker",)
    for name in ("Setup", "Manifest", "Papers") + SOURCE_SHEETS + review_sheets:
        ws = wb[name]
        allowed = LABELS["readiness"] if name in review_sheets else None
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type in ("f", "e") and ws.cell(1, cell.column).value != allowed:
                    issue(str(ws.cell(cell.row, 1).value), "ERROR", f"{name}!{cell.coordinate}: formula/error cell is not an answer")
    # Read-only context in review rows must still match its key after sorting/pasting.
    by_key = {c["key"]: c for c in candidates}
    for name in review_sheets:
        for r in records(wb[name]):
            c = by_key[r[LABELS["key"]]]
            if (name == "Checker" and r[LABELS["package"]] != c["package"]) or r[LABELS["title"]] != c["title"]:
                issue(c["key"], "ERROR", f"{name}: paper/title no longer matches candidate key")
    results, reviews = [], {}
    for c in candidates:
        key = c["key"]
        p, changes = primary[key], overrides[key]
        r = {**changes, **{f: changes[f] if changes[f] != "" else p[f] for f in ASSESSMENT}}
        for f in ("notes", "relatedClusterId"):
            if changes[f] == "<clear>":
                r[f] = ""
        reviews[key] = r
        complete = r["reviewComplete"] == "yes"
        primary_complete = p["reviewComplete"] == "yes"
        identity_needed |= complete or primary_complete
        errors = ["Primary: " + e for e in assessment_errors(p, primary_complete or require_complete)]
        if not primary_complete:
            issue(key, "ERROR" if require_complete else "PENDING", "Primary: human completion not confirmed")
        if not complete and not primary_only:
            issue(key, "ERROR" if require_complete else "PENDING", "Human completion not confirmed")
        if not primary_only:
            errors += assessment_errors(r, complete or require_complete)
            if complete and not primary_complete:
                errors.append("Cannot complete checker before primary")
            for field in ("checkerComment", "resolutionNote"):
                if not isinstance(r[field], str):
                    errors.append(f"{field}: expected text")
            for field in ("checkStatus", "resolutionStatus"):
                if r[field] and r[field] not in OPTIONS[field]:
                    errors.append(f"{field}: invalid option")
            if complete or require_complete:
                if not r["checkStatus"]:
                    errors.append("checkStatus: required")
                if r["checkStatus"] != "agree" and not str(r["checkerComment"]).strip():
                    errors.append("checkerComment: required for non-agree")
                if r["checkStatus"] == "agree" and any(r[f] != p[f] for f in ASSESSMENT):
                    errors.append("agree requires an unchanged primary assessment")
            if r["resolutionStatus"] == "resolved":
                if not str(r["resolutionNote"]).strip():
                    errors.append("resolutionNote: required when resolved")
                if r["disposition"] == "discuss":
                    errors.append("resolved cannot retain disposition=discuss")
        for error in errors:
            issue(key, "ERROR", error)
        active = MODULES.get(r["domain"], ()) if r["disposition"] == "retain" else ()
        stale = [f for f in ASSESSMENT[2:9] if f not in active and r[f]]
        if stale:
            issue(key, "WARNING", "Inactive fields ignored in analysis; clear if obsolete: " + ", ".join(stale))
        status = "invalid" if errors else "pending" if not (primary_complete and complete) else "unresolved" if unresolved(r, p) else "resolved"
        if status == "unresolved":
            issue(key, "UNRESOLVED", "Discussion remains unresolved; not eligible for final analysis")
        results.append({"key": key, "status": status, "outcome": outcome(c, r) if status == "resolved" else "undetermined"})
    if identity_needed and not valid_id(primary_id):
        issue("Setup", "ERROR", "A stable primary reviewer ID is required (no surrounding whitespace)")
    if not primary_only and (require_complete or any(r["reviewComplete"] == "yes" for r in reviews.values())) and not valid_id(checker_id):
        issue("Setup", "ERROR", "Distinct primary and checker IDs are required")
    if checker_id and str(primary_id).casefold() == str(checker_id).casefold():
        issue("Setup", "ERROR", "Checker must be a different person from primary")
    return {"wb": wb, "metadata": m, "setup": setup, "candidates": candidates, "reviews": reviews, "primary": primary, "overrides": overrides, "issues": issues, "results": results}


def fail_on_errors(report):
    errors = [i for i in report["issues"] if i[1] == "ERROR"]
    if errors:
        raise ValueError(f"Validation failed ({len(errors)} errors): " + "; ".join(f"{k}: {msg}" for k, _, msg in errors[:8]))


def validation_report(report, output):
    wb = empty_book()
    table(wb, "Summary", ["field", "value"], [["version", VERSION], ["stage", report["metadata"]["stage"]], ["validatedAt", now()], *sorted(Counter(i[1] for i in report["issues"]).items()), ["note", "PASS means structural/entry validity, not scientific confirmation; unresolved items remain excluded."]], {"value": 105})
    table(wb, "Issues", ["key", "severity", "message"], report["issues"], {"key": 45, "message": 110})
    table(wb, "Status", ["key", "status", "outcome"], [[r[f] for f in ("key", "status", "outcome")] for r in report["results"]], {"key": 45})
    save_new(wb, output)


def prepare_checker(path, output, checker_id="", catalog=CATALOG):
    report = validate(path, catalog, require_complete=True, primary_only=True)
    fail_on_errors(report)
    if report["metadata"]["stage"] != "workbook":
        raise ValueError("prepare-checker requires a completed primary XLSX")
    if any(r["reviewComplete"] != "no" or r["checkStatus"] or r["checkerComment"] or r["resolutionNote"] or r["resolutionStatus"] != "pending" or any(r[f] for f in ASSESSMENT) for r in report["overrides"].values()):
        raise ValueError("Checker work has already started; continue that XLSX instead of creating a fresh checker copy")
    primary_id = report["setup"]["primaryReviewerId"]
    if checker_id and (checker_id != checker_id.strip() or checker_id.casefold() == primary_id.casefold()):
        raise ValueError("Checker must have a distinct, trimmed ID")
    wb = make_review(report["candidates"], report["metadata"]["catalogSignature"], output, "checker", report["primary"], primary_id, checker_id, hashlib.sha256(Path(path).read_bytes()).hexdigest())
    put(wb["Setup"]["B3"], report["setup"]["primaryReviewerName"])
    save_new(wb, output)


def export_analysis(paths, output, catalog=CATALOG, allow_incomplete=False):
    reports = [validate(path, catalog) for path in paths]
    for r in reports:
        fail_on_errors(r)
    if len({r["metadata"]["year"] for r in reports}) != len(reports):
        raise ValueError("Duplicate year: supply exactly one authoritative checker workbook per year")
    if not allow_incomplete and any(s["status"] != "resolved" for r in reports for s in r["results"]):
        raise ValueError("Unfinished/unresolved candidates remain; use --allow-incomplete for a labeled provisional workbook")
    wb = empty_book()
    table(wb, "Readme", ["field", "value"], [["version", VERSION], ["createdAt", now()], ["scope", "Supplied years only: " + ", ".join(r["metadata"]["year"] for r in reports)], ["status", "PROVISIONAL; unresolved/pending are not zero" if any(s["status"] != "resolved" for r in reports for s in r["results"]) else "Completed visible-answer checker records for supplied years"], ["interpretation", "Exploratory candidate counts, not sensitivity or causal AI effects. Protocol/plan retained counts are not automatically confirmed errors. Confirmed and probable statistical findings remain separable in Decisions."], ["counts", "Paper complete counts are blank until all its candidates are resolved. Resolved-so-far counts are explicitly partial; all sampled papers in supplied annual candidate sets are included, including zero eligible outcomes."], ["agreement", "Visible primary answers: checker agreement describes process consistency, not blinded inter-rater reliability."]], {"field": 30, "value": 130})
    decision_fields = ["year", "package", "key", "relationship", "primaryReviewerId", "checkerId", "status", "outcome", "aiDisposition", "aiDomain"] + ["primary_" + f for f in ASSESSMENT] + ["final_" + f for f in ASSESSMENT] + ["checkStatus", "checkerComment", "resolutionStatus", "resolutionNote"]
    decisions, paper_rows, protocol_rows, lineage = [], [], [], []
    for path, report in zip(paths, reports):
        lineage.append([Path(path).name, hashlib.sha256(Path(path).read_bytes()).hexdigest(), report["metadata"]["primaryFileSHA256"], report["metadata"]["catalogSignature"]])
        states = {r["key"]: r for r in report["results"]}
        papers = defaultdict(list)
        for c in report["candidates"]:
            key = c["key"]
            r, p, s = report["reviews"][key], report["primary"][key], states[key]
            decisions.append([str(c["year"]), c["package"], key, c["primary"], report["setup"]["primaryReviewerId"], report["setup"]["checkerId"], s["status"], s["outcome"], c["aiInitialAssessment"]["disposition"], c["aiInitialAssessment"]["domain"]] + [p[f] for f in ASSESSMENT] + [r[f] for f in ASSESSMENT] + [r[f] for f in ("checkStatus", "checkerComment", "resolutionStatus", "resolutionNote")])
            papers[c["package"]].append(s)
            if c["primary"] in ("P_R", "P_P"):
                protocol_rows.append([str(c["year"]), c["package"], key, c["primary"], c["title"], c["statement"], s["status"], r["disposition"], r["documentationStatus"], r["justificationStatus"], r["protocolInterpretation"], r["internalStatus"], r["notes"]])
        for paper, states_for_paper in papers.items():
            counts = Counter(s["outcome"] for s in states_for_paper)
            pending = sum(s["status"] == "pending" for s in states_for_paper)
            uncertain = sum(s["status"] == "unresolved" for s in states_for_paper)
            complete = not (pending or uncertain)
            outcomes = [counts[k] for k in ("primary_temporal", "secondary_temporal", "protocol_descriptive", "plan_descriptive")]
            paper_rows.append([report["metadata"]["year"], paper, len(states_for_paper), pending, uncertain, "yes" if complete else "no"] + (outcomes if complete else [""] * 4) + outcomes)
    table(wb, "Decisions", decision_fields, decisions, {"key": 45, "final_notes": 75, "checkerComment": 75, "resolutionNote": 75})
    table(wb, "Paper_counts", ["year", "paper", "candidates", "pending", "unresolved", "paper_complete", "primary_temporal", "secondary_temporal", "protocol_descriptive", "plan_descriptive", "primary_resolved_so_far", "secondary_resolved_so_far", "protocol_resolved_so_far", "plan_resolved_so_far"], paper_rows, {"paper": 32})
    table(wb, "Protocol_plan_table", ["year", "paper", "key", "relationship", "title", "statement", "status", "disposition", "documentation", "justification", "interpretation", "internalStatus", "notes"], protocol_rows, {"key": 45, "title": 70, "statement": 110, "notes": 80})
    table(wb, "Provenance", ["input_filename", "input_SHA256", "primary_file_SHA256", "catalog_signature"], lineage, {"input_filename": 65, "input_SHA256": 70, "primary_file_SHA256": 70, "catalog_signature": 70})
    save_new(wb, output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("freeze-source", help="One-time lossless migration of v.2.0.4 candidate data only")
    p.add_argument("--output", type=Path, default=CATALOG)
    p = sub.add_parser("build", help="Build annual primary/checker XLSX workbooks from the XLSX catalog")
    p.add_argument("--output-dir", type=Path, default=HERE)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p = sub.add_parser("validate", help="Validate actual saved answers; never trust cached formulas")
    p.add_argument("input", type=Path)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p.add_argument("--report", type=Path)
    p.add_argument("--require-complete", action="store_true")
    p.add_argument("--primary-only", action="store_true", help="Validate primary handoff before checker work begins")
    p = sub.add_parser("prepare-checker", help="Freeze completed primary records into a new checker XLSX")
    p.add_argument("input", type=Path)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--checker-id", default="")
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p = sub.add_parser("export", help="Produce XLSX decisions, paper counts, and protocol/plan table")
    p.add_argument("inputs", nargs="+", type=Path)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p.add_argument("--allow-incomplete", action="store_true")
    args = parser.parse_args()
    try:
        if args.command == "freeze-source":
            freeze_source(args.output)
        elif args.command == "build":
            build(args.output_dir, args.catalog)
        elif args.command == "validate":
            report = validate(args.input, args.catalog, args.require_complete, args.primary_only)
            if args.report:
                validation_report(report, args.report)
            print(dict(Counter(i[1] for i in report["issues"])))
            fail_on_errors(report)
        elif args.command == "prepare-checker":
            prepare_checker(args.input, args.output, args.checker_id, args.catalog)
        else:
            export_analysis(args.inputs, args.output, args.catalog, args.allow_incomplete)
    except (ValueError, KeyError, FileExistsError, FileNotFoundError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
