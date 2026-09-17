#!/usr/bin/env python3
"""XLSX-only adjudication: freeze legacy candidates once, build, check and export.

Human judgments are never inferred from prefilled AI values. Formula caches are
never used by validation or analysis. Existing output files are never overwritten.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import re
import sys
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.views import Selection
from openpyxl.writer.theme import theme_xml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
VERSION = "v.2.1.5.1"
REVISION = "v.2.1.5.1"
SCHEMA = "xlsx-adjudication-8"
CATALOG_SCHEMA = "xlsx-adjudication-1"
HAN = re.compile("[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\U00020000-\U000323af]")
YEARS = {"2018": 119, "2019": 71, "2024": 91, "2025": 131}
RELATIONSHIPS = {"P_R": 45, "P_P": 50, "S_C": 65, "S_W": 252}
CATALOG = HERE / "resources" / "source_catalog_v_2_1_5_1.xlsx"
SOURCE_SHEETS = ("Candidates", "Evidence", "Links", "Relations")
MODULES = {
    "statistical_results": ("statisticalIssueStatus", "substantiveStatus"),
    "report_internal_inconsistency": ("internalStatus",),
    "plan_material": ("internalStatus",),
    "protocol_report": (),
}
OPTIONAL_MODULES = {"statistical_results": ("downstreamImpact",), "protocol_report": ("protocolInterpretation",)}
APPLICABLE = {domain: fields + OPTIONAL_MODULES.get(domain, ()) for domain, fields in MODULES.items()}
MODULE_FIELDS = tuple(dict.fromkeys(field for fields in APPLICABLE.values() for field in fields))
NA = "not_applicable"
OPTIONS = {
    "disposition": {
        "retain": "Retain candidate. After human review, keep this finding in the applicable review results; retaining a protocol discrepancy does not automatically confirm an issue.",
        "exclude": "Do not retain. The candidate has been assessed and should not remain as a finding, for example because the apparent mismatch has a satisfactory explanation.",
        "discuss": "Needs discussion. The reviewer has considered the candidate but cannot settle the disposition without further discussion. This is a decision to defer, not an untouched blank.",
    },
    "domain": {
        "statistical_results": "Statistical results (S_C / S_W). Assess reported numbers, estimates, intervals, denominators or statistical claims in the current report materials.",
        "report_internal_inconsistency": "Report internal inconsistency (S_C / S_W). Assess inconsistencies within the main paper, within the updated supplement, or between the main paper and updated supplement. The name describes the report materials being compared; it is not restricted to nonstatistical content. Use statistical_results when specifically judging a statistical or numerical reporting issue.",
        "protocol_report": "Protocol-report (P_R). Assess whether planned and reported conduct or analyses meaningfully differ, using the compared materials. No search for acknowledgment or reasons for changes is required.",
        "plan_material": "Plan material (P_P). Assess an inconsistency involving planning materials themselves rather than a planned-versus-reported result. The original relationship code remains unchanged.",
    },
    "statisticalIssueStatus": {
        "confirmed_issue": "Confirmed statistical or numerical reporting issue. The available evidence establishes a reporting issue or contradiction, even if the correct replacement value is not known.",
        "probable_issue": "Probable statistical or numerical reporting issue. The evidence favors a reporting issue, but an unresolved limitation or plausible alternative prevents confirmation.",
        "not_issue": "Not a statistical or numerical issue. The apparent discrepancy is consistent with a valid calculation, definition, analysis method or reporting context.",
        "indeterminate": "Indeterminate. The available evidence does not support deciding whether a statistical or numerical reporting issue exists; neither confirmation nor dismissal is justified.",
    },
    "substantiveStatus": {
        "substantive": "Substantive. The issue could meaningfully change interpretation, data extraction or downstream analysis, rather than only the presentation of the result.",
        "minor_presentation": "Minor or presentation only. The intended result remains clear and the issue is a small formatting, labeling or typographic problem without a meaningful change in interpretation.",
        "indeterminate": "Indeterminate. There is not enough information to decide whether the issue is substantive or limited to presentation.",
    },
    "downstreamImpact": {
        "direct": "Direct extraction or analysis impact. The affected value or definition is itself an input to data extraction or subsequent analysis; this describes potential impact, not proven downstream harm.",
        "possible": "Possible downstream impact. A plausible route to affecting later interpretation, extraction or analysis exists, but the connection is indirect or uncertain.",
        "unlikely": "Unlikely downstream impact. The issue is unlikely to change extracted information or later analysis because the relevant result remains clear and usable.",
        "unknown": "Unable to assess. The materials do not establish enough about the affected information or its likely use to judge downstream impact.",
    },
    "internalStatus": {
        "substantive_inconsistency": "Substantive inconsistency. Conflicting statements, values, labels, or descriptions in the compared report or planning materials could materially change how the study, population, procedures or results are understood.",
        "minor_presentation": "Minor presentation inconsistency. A real discrepancy exists, but it mainly affects presentation, including minor wording or labeling problems, and the intended meaning remains clear. Example: an imprecise label whose intended meaning is clear from context.",
        "non_inconsistency": "Not an inconsistency. The statements are compatible once their definitions, populations, time points or other context are taken into account. This includes different wording with the same substantive meaning. Example: participants and study participants refer to the same group.",
        "indeterminate": "Indeterminate. The available material is insufficient to establish whether the statements conflict or how important any conflict is.",
    },
    "protocolInterpretation": {
        "substantive_difference": "Substantive difference. The compared plan and report meaningfully differ in an outcome, population, procedure, or analysis. This describes the difference; it does not establish an issue or an unjustified change. Example: the planned primary time point is 12 months but the report uses 6 months.",
        "minor_presentation": "Minor presentation difference. A minor wording, labeling or presentation problem exists, but the intended planned and reported meaning or conduct remains clear. Example: an imprecise outcome label clarified by the same definition in both documents.",
        "version_timing": "Version or timing ambiguity. The difference may reflect which protocol version or date is being compared; the sequence cannot be resolved from the available material.",
        "non_inconsistency": "No inconsistency. The compared statements are compatible after checking the population, definitions, time points, or context, including equivalent wording with no change in meaning or conduct. Example: week 12 and 12 weeks after baseline describe the same time point. Normally consider Disposition=exclude for this candidate.",
        "indeterminate": "Indeterminate. The available evidence cannot support a more specific interpretation of the planned-versus-reported difference.",
    },
    "reviewComplete": {
        "no": "Draft. This review stage is not yet confirmed complete; AI suggestions or partially filled fields do not count as completed human review.",
        "yes": "Human reviewed and complete. The reviewer has checked this row and its required fields. A discussion item may still be unresolved even though this review stage is complete.",
    },
    "checkStatus": {
        "agree": "Agree, assessment unchanged. Accept the primary assessment as recorded. Leave override cells blank; no checker comment is required.",
        "agree_clarification": "Agree with clarification. Accept the main assessment but add context or a clarifying amendment. Explain it in a checker comment and fill only any changed fields.",
        "disagree": "Disagree, proposal pending discussion. Recommend a different assessment, explain why in a comment, and enter proposed changes. The proposal is not final consensus until discussion is resolved.",
        "unable": "Unable to resolve. The checker cannot settle the issue from the available information. State what remains uncertain in a comment and refer it for discussion.",
    },
    "resolutionStatus": {
        "pending": "Discussion not resolved. A required discussion is not yet settled. For an unchanged agreement this default can be left as-is because no separate resolution is required.",
        "resolved": "Joint discussion resolved. Record the participants and agreed conclusion in Resolution note, update the proposed assessment, and do not leave disposition as discuss.",
    },
}
LABELS = {
    "key": 'Candidate key', "title": 'Candidate title', "locations": 'Source location', "primarySummary": 'Primary assessment',
    "disposition": 'Disposition', "domain": 'Review domain', "reviewComplete": 'Complete', "readiness": 'Validation hint',
    "statisticalIssueStatus": 'Statistcal_or_numerical issue', "substantiveStatus": 'Substantiveness', "downstreamImpact": 'Downstream impact',
    "internalStatus": 'Report-internal or planning-material inconsistency', "protocolInterpretation": 'Protocol interpretation',
    "notes": 'Notes (optional)', "checkStatus": 'Checker result', "checkerComment": 'Checker comment', "resolutionStatus": 'Discussion resolution', "resolutionNote": 'Resolution note',
}
ASSESSMENT = ("disposition", "domain", "statisticalIssueStatus", "substantiveStatus", "downstreamImpact", "internalStatus", "protocolInterpretation", "notes")
REVIEW_FIELDS = ("key", "title", "locations", "disposition", "domain", "reviewComplete", "readiness") + ASSESSMENT[2:]
CHECK_FIELDS = REVIEW_FIELDS[:-1] + ("primarySummary", "checkStatus", "checkerComment", "resolutionStatus", "resolutionNote", "notes")
REVIEW_FIELDS = CHECK_FIELDS
CHECK_ONLY_FIELDS = ("checkStatus", "checkerComment", "resolutionStatus", "resolutionNote")
PRIMARY_STAGE_SUMMARY = "PRIMARY REVIEW - complete this row before handoff. The primary snapshot appears here in the prepared checker copy."
SNAPSHOT_FIELDS = ("key",) + ASSESSMENT + ("reviewComplete",)
BLUE = "173F57"
YELLOW = "FFF2CC"
GRAY = "EDF2F5"
OPTIONAL_BLUE = "EAF3FC"
for _field in MODULE_FIELDS:
    if NA not in OPTIONS[_field]:
        OPTIONS[_field][NA] = "Automatic default outside this field's review domain or for a non-retained candidate. It is not an issue verdict, missing evidence, or an unanswered applicable question. Do not choose it for an applicable field."
for _field in ("downstreamImpact", "protocolInterpretation"):
    OPTIONS[_field]["<clear>"] = "Checker only: remove an inherited optional judgment. This is a control value, not an assessment option for primary review. The final optional answer becomes blank."


def applicable_fields(record, include_pending=False):
    if record.get("disposition") == "retain" or (include_pending and not record.get("disposition")):
        return APPLICABLE.get(record.get("domain"), ())
    return ()


def effective_record(record):
    """Preserve raw answers elsewhere; represent inactive fields consistently."""
    active = applicable_fields(record, include_pending=True)
    return {**record, **{field: NA for field in MODULE_FIELDS if field not in active}}


def excel_context(wb, row, stage, key=None):
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    refs = {f: f"${letter(fields.index(f)+1)}{row}" for f in ("disposition", "domain")}
    if stage == "checker":
        package = key.split("::")[1]
        for field in refs:
            lookup = f'INDEX({primary_range_name(package, field)},MATCH($A{row},{primary_range_name(package, "key")},0))'
            refs[field] = f'IF({refs[field]}="",{lookup},{refs[field]})'
    return refs


def primary_range_name(package, field):
    return "primary_" + package.replace(".", "_") + "_" + field


def primary_ranges(wb, candidates):
    for package in paper_groups(candidates):
        sheet = "_Primary"
        for field in ("key", "disposition", "domain"):
            col = letter(SNAPSHOT_FIELDS.index(field) + 1)
            yield primary_range_name(package, field), f"'{sheet}'!${col}$2:${col}${wb[sheet].max_row}"


def checker_sheet_name(package):
    return "C_" + primary_sheet_name(package)[2:]


def checker_sheet_names(candidates):
    return [checker_sheet_name(paper) for paper in paper_groups(candidates)]


def review_names(candidates, stage):
    return checker_sheet_names(candidates)


def empty_checker():
    return {**dict.fromkeys(ASSESSMENT, ""), "reviewComplete": "no", "checkStatus": "", "checkerComment": "", "resolutionStatus": "pending", "resolutionNote": ""}


def primary_summary(record):
    if record["reviewComplete"] != "yes":
        return PRIMARY_STAGE_SUMMARY
    fields = ("disposition", "domain") + tuple(f for f in ASSESSMENT[2:] if f in applicable_fields(record) or f == "notes")
    return "\n".join(f"{LABELS[f]}: {record[f]}" for f in fields if record[f] not in ("", NA))


def inactive_condition(wb, row, field, stage, key=None):
    refs = excel_context(wb, row, stage, key)
    domains = ",".join(f'{refs["domain"]}="{d}"' for d, fields in APPLICABLE.items() if field in fields)
    d = refs["disposition"]
    return f'OR(AND({d}<>"",{d}<>"retain"),NOT(OR({domains})))'


def default_formula(wb, row, field, stage, key=None):
    return f'=IF({inactive_condition(wb, row, field, stage, key)},"{NA}","")'


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def now():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def require_english_xlsx(path_or_stream):
    """Check decoded XML, including hidden cells, themes, comments and properties."""
    with zipfile.ZipFile(path_or_stream) as archive:
        for name in archive.namelist():
            if HAN.search(name):
                raise ValueError("English-only XLSX required: non-English package member name")
            if not name.endswith((".xml", ".rels", ".vml")):
                continue
            root = ET.fromstring(archive.read(name))
            for node in root.iter():
                values = [node.text or "", node.tail or "", *node.attrib.values()]
                if any(HAN.search(value) for value in values):
                    raise ValueError(f"English-only XLSX required: CJK text in {name}")


def primary_sheet_name(package):
    name = "C_" + package
    if len(name) > 31 or re.search(r"[\\/*?:\[\]]", name):
        raise ValueError(f"Invalid paper worksheet name: {name}")
    return name


def paper_groups(candidates):
    groups = defaultdict(list)
    for c in candidates:
        groups[c["package"]].append(c)
    return groups


def primary_sheet_names(candidates):
    return [primary_sheet_name(paper) for paper in paper_groups(candidates)]


def put(cell, value):
    """Force all source/free text to literal strings, including leading '='."""
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError(f"Text exceeds Excel cell limit: {cell.coordinate}")
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value


def table(wb, name, headers, rows, widths=None, protect=True):
    ws = wb.create_sheet(name)
    for i, header in enumerate(headers, 1):
        put(ws.cell(1, i), header)
        ws.cell(1, i).font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        ws.cell(1, i).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(1, i).alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[letter(i)].width = (widths or {}).get(header, 25)
    for row in rows:
        for col, value in enumerate(row, 1):
            cell = ws.cell(ws.max_row + 1, col) if col == 1 else ws.cell(target_row, col)
            if col == 1:
                target_row = cell.row
            put(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=11)
            if target_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F8FA")
        lines = max((sum(max(1, math.ceil(len(part) / max(8, ws.column_dimensions[letter(i)].width - 3))) for part in str(value).splitlines()) for i, value in enumerate(row, 1)), default=1)
        ws.row_dimensions[target_row].height = min(300, max(30, lines * 16 + 8))
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    if protect:
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
    return ws


def empty_book():
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def metadata(wb, values):
    table(wb, "Manifest", ["field", "value"], list(values.items()), {"field": 30, "value": 110})


def meta(wb):
    rows = records(wb["Manifest"], ["field", "value"])
    if len({r["field"] for r in rows}) != len(rows):
        raise ValueError("Duplicate Manifest fields")
    return {r["field"]: r["value"] for r in rows}


def records(ws, expected=None):
    headers = [c.value for c in ws[1]]
    if expected is not None and headers != list(expected):
        raise ValueError(f"{ws.title}: headers changed; use the original template")
    if len(set(headers)) != len(headers) or None in headers:
        raise ValueError(f"{ws.title}: duplicate or blank headers")
    result = []
    for cells in ws.iter_rows(min_row=2):
        values = [c.value for c in cells]
        if all(v is None for v in values):
            continue
        result.append(dict(zip(headers, ["" if v is None else v for v in values])))
    return result


def source_tables(candidates):
    containers = {"aiInitialAssessment", "evidence", "sourceLinks", "evidencePaths", "relations", "legacyRelations"}
    scalar = sorted(set().union(*(c.keys() for c in candidates)) - containers)
    fields = scalar + ["ai_" + f for f in ("source", "relationship", "domain", "disposition")]
    rows = [[c.get(f, "") if not f.startswith("ai_") else c["aiInitialAssessment"][f[3:]] for f in fields] for c in candidates]
    evidence, links, relations = [], [], []
    for c in candidates:
        evidence.extend([c["key"], i, k, v] for i, (k, v) in enumerate(c["evidence"].items()))
        for kind in ("sourceLinks", "evidencePaths"):
            links.extend([c["key"], kind, i, v["label"], v["href"]] for i, v in enumerate(c[kind]))
        for kind in ("relations", "legacyRelations"):
            relations.extend([c["key"], kind, i, v] for i, v in enumerate(c[kind]))
    return {"Candidates": (fields, rows), "Evidence": (["key", "order", "section", "text"], evidence), "Links": (["key", "kind", "order", "label", "href"], links), "Relations": (["key", "kind", "order", "relationship"], relations)}


def source_signature(wb):
    # Sorting whole rows permits user sorting without hiding edits or duplicate rows.
    return digest({name: sorted(records(wb[name]), key=lambda r: json.dumps(r, sort_keys=True, ensure_ascii=False)) for name in SOURCE_SHEETS})


def read_candidates(wb):
    candidates = []
    for row in records(wb["Candidates"]):
        c = {k: v for k, v in row.items() if not k.startswith("ai_")}
        c["aiInitialAssessment"] = {k[3:]: v for k, v in row.items() if k.startswith("ai_")}
        c.update(evidence={}, sourceLinks=[], evidencePaths=[], relations=[], legacyRelations=[])
        candidates.append(c)
    by_key = {c["key"]: c for c in candidates}
    if len(by_key) != len(candidates):
        raise ValueError("Duplicate source candidate keys")
    for r in sorted(records(wb["Evidence"]), key=lambda r: (r["key"], r["order"])):
        by_key[r["key"]]["evidence"][r["section"]] = r["text"]
    for r in sorted(records(wb["Links"]), key=lambda r: (r["key"], r["kind"], r["order"])):
        by_key[r["key"]][r["kind"]].append({"label": r["label"], "href": r["href"]})
    for r in sorted(records(wb["Relations"]), key=lambda r: (r["key"], r["kind"], r["order"])):
        by_key[r["key"]][r["kind"]].append(r["relationship"])
    return candidates


def local_target(href):
    parsed = urlsplit(href)
    if parsed.scheme in ("http", "https"):
        return None
    if parsed.scheme or not href.startswith(tuple(f"../../{y}/" for y in YEARS)):
        raise ValueError(f"Nonportable source link: {href}")
    target = (HERE / unquote(parsed.path)).resolve()
    if not target.is_relative_to(ROOT) or not target.is_file():
        raise ValueError(f"Missing or invalid source link: {href}")
    return target


def validate_sources(candidates, all_years=True, check_files=False):
    if len({c["key"] for c in candidates}) != len(candidates):
        raise ValueError("Duplicate candidate keys")
    if all_years and (Counter(str(c["year"]) for c in candidates) != YEARS or Counter(c["primary"] for c in candidates) != RELATIONSHIPS):
        raise ValueError("Candidate counts differ from the frozen 412-candidate universe")
    count = 0
    for c in candidates:
        if c["key"] != f'{c["year"]}::{c["package"]}::{c["id"]}':
            raise ValueError(f"Candidate identity mismatch: {c['key']}")
        if check_files:
            for href in [c["reportHref"], "../../" + c["reportPath"]] + [r["href"] for kind in ("sourceLinks", "evidencePaths") for r in c[kind]]:
                local_target(href)
                count += 1
    return count


def add_sources(wb, candidates, output):
    for name, (headers, rows) in source_tables(candidates).items():
        widths = {"key": 45, "title": 65, "statement": 90, "text": 115, "section": 38, "label": 60, "href": 95, "rationale": 90, "locations": 80}
        ws = table(wb, name, headers, rows, widths)
        if name == "Evidence":
            ws.column_dimensions["B"].hidden = True
            for r in range(2, ws.max_row + 1):
                text = str(ws.cell(r, 4).value or "")
                lines = sum(max(1, math.ceil(len(line) / 105)) for line in text.splitlines())
                ws.row_dimensions[r].height = min(400, max(48, lines * 16 + 10))
        elif name == "Links":
            ws.sheet_state = "hidden"
        elif name == "Relations":
            ws.sheet_state = "hidden"


def save_new(wb, path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Output must be an .xlsx workbook")
    # Use standard English font names, without changing evidence or answers.
    theme = wb.loaded_theme or theme_xml
    if isinstance(theme, bytes):
        theme = theme.decode("utf-8")
    fonts = {"\u5b8b\u4f53": "SimSun", "\u65b0\u7d30\u660e\u9ad4": "PMingLiU"}
    theme_root = ET.fromstring(theme)
    for node in theme_root.iter():
        if node.get("typeface") in fonts:
            node.set("typeface", fonts[node.get("typeface")])
    wb.loaded_theme = ET.tostring(theme_root, encoding="utf-8", xml_declaration=True)
    content = io.BytesIO()
    wb.save(content)
    content.seek(0)
    require_english_xlsx(content)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation protects completed human records and template files.
    with path.open("xb") as handle:
        handle.write(content.getvalue())
    # Reopen the actual zip to catch serialization problems immediately.
    with path.open("rb") as handle:
        check = load_workbook(handle)
        check.close()
    print(f"Created {path}")


def freeze_source(output):
    candidates, inputs = [], []
    for year in YEARS:
        path = HERE.parent / "human_adjudication_v_2_0_4" / f"candidate_flag_review_v_2_0_4_{year}.html"
        content = path.read_text(encoding="utf-8")
        match = re.search(r'<script id="candidate-data" type="application/json">(.*?)</script>', content, re.S)
        if not match:
            raise ValueError(f"Missing legacy candidate payload: {path}")
        payload = json.loads(match[1])
        rows = payload["candidates"]
        legacy_digest = hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        if payload["reviewVersion"] != "v.2.0.4" or payload["schemaVersion"] != 7 or payload["years"] != [year] or payload["keySet"] != [c["key"] for c in rows] or legacy_digest != payload["sourceDigest"]:
            raise ValueError(f"Legacy payload integrity error: {year}")
        candidates.extend(rows)
        inputs.append((f"legacy_{year}_sha256", hashlib.sha256(path.read_bytes()).hexdigest()))
    validate_sources(candidates, check_files=True)
    if Counter(c["aiInitialAssessment"]["disposition"] for c in candidates) != {"retain": 402, "exclude": 9, "discuss": 1}:
        raise ValueError("Unexpected AI starting values")
    wb = empty_book()
    add_sources(wb, candidates, output)
    # Verify a lossless normalized conversion, including all evidence and metadata.
    if read_candidates(wb) != candidates:
        raise ValueError("Lossy legacy-to-XLSX source conversion")
    metadata(wb, {"version": VERSION, "schema": CATALOG_SCHEMA, "stage": "catalog", "createdAt": now(), "sourceSignature": source_signature(wb), **dict(inputs), "provenance": "v.2.0.1 -> v.2.0.2 -> v.2.0.3 -> v.2.0.4; no v_1_0_* inputs", "humanRecordsImported": "none; bundled tester snapshot explicitly excluded"})
    save_new(wb, output)


def import_catalog(source, output):
    """Version the v.2.1.5 frozen catalog without changing source records."""
    require_english_xlsx(source)
    wb = load_workbook(source)
    manifest = meta(wb)
    signature = source_signature(wb)
    if manifest.get("stage") != "catalog" or manifest.get("schema") != CATALOG_SCHEMA or manifest.get("sourceSignature") != signature:
        raise ValueError("Invalid source catalog")
    if manifest.get("version") not in ("v.2.1.5", VERSION):
        raise ValueError("Import requires the v.2.1.5 source catalog or this version")
    candidates = read_candidates(wb)
    validate_sources(candidates)
    if any(c["aiInitialAssessment"]["domain"] not in MODULES or c.get("suggestedDomain") not in MODULES for c in candidates):
        raise ValueError("Unknown suggested domain in source catalog")
    wb.remove(wb["Manifest"])
    metadata(wb, {**manifest, "version": VERSION, "sourceSignature": signature,
                  "parentCatalogVersion": manifest["version"], "parentCatalogSignature": signature,
                  "suggestedDomainsRenamed": "0",
                  "catalogChange": "No source record changes; human-review terminology and layout updated"})
    save_new(wb, output)


def load_catalog(path=CATALOG):
    require_english_xlsx(path)
    wb = load_workbook(path)
    m = meta(wb)
    if m.get("version") != VERSION or m.get("stage") != "catalog" or m.get("schema") != CATALOG_SCHEMA or m.get("sourceSignature") != source_signature(wb):
        raise ValueError("Source catalog integrity/version mismatch")
    candidates = read_candidates(wb)
    validate_sources(candidates)
    if any(c["aiInitialAssessment"]["domain"] not in MODULES or c.get("suggestedDomain") not in MODULES for c in candidates):
        raise ValueError("Source catalog contains an unsupported suggested domain")
    return candidates, m["sourceSignature"]


def add_codebook(wb):
    guides = {
        "key": ("Read only", "Unique candidate identity used for matching, even after complete rows are reordered. Never edit this value."),
        "primarySummary": ("Read only", "During primary review this column explains the handoff. In the prepared checker copy it shows the completed primary judgments for this candidate. Review the summary before choosing Checker result. Blank applicable judgment cells then inherit the matching snapshot."),
        "title": ("Read only", "Short candidate title. A candidate is an AI flag for human assessment, not a confirmed issue."),
        "locations": ("Read only; both roles", "Where to check the source claim. Both Primary and Checker show this column in the same position. Use Candidates and Evidence for context and the shared paper folders for originals."),
        "disposition": ("Required; all domains", "Keep, exclude, or defer for discussion. Blank means untouched, not discuss. No AI disposition is preselected. Checker blanks inherit the primary answer."),
        "domain": ("Required; all domains", "Choose the assessment branch first. It controls automatic not_applicable defaults and which judgments are required. All four choices are available; the frozen document relationship code does not change."),
        "reviewComplete": ("Required; each stage", "Mark yes only after human review. Automatic defaults never complete a row. Reset to no after changing a completed assessment, then review again."),
        "readiness": ("Calculated; read only", "CHECK FIELDS means an applicable required judgment is missing or invalid. CLEAR INACTIVE VALUES means an old answer remains in an irrelevant field. This hint is not the full scientific or handoff validation."),
        "statisticalIssueStatus": ("Required; retained statistical_results", "Does a statistical or numerical reporting issue exist? Example: a reported percentage conflicts with its stated numerator and denominator after checking definitions. Certainty is separate from importance."),
        "substantiveStatus": ("Required; retained statistical_results", "Could the issue meaningfully affect interpretation or data use? A harmless spacing typo is minor_presentation; an ambiguous extractable effect estimate may be substantive. An issue can be confirmed yet minor."),
        "downstreamImpact": ("Optional; retained statistical_results", "Add only useful information about later extraction or analysis. direct means the affected value itself may be extracted, not that harm has occurred. Blank is allowed and means not assessed; it is not unlikely or not_applicable. This field does not determine temporal eligibility."),
        "internalStatus": ("Required; retained report_internal_inconsistency or plan_material", "Judge whether the compared materials are inconsistent and how much that matters. For report_internal_inconsistency, compare content within the main paper, within the updated supplement, or between those two report materials. For plan_material, compare planning materials such as protocol versus protocol or SAP. Plan versus final report belongs to protocol_report. Use statistical_results when specifically judging a statistical or numerical reporting issue."),
        "protocolInterpretation": ("Optional; retained protocol_report", "Describe the compared plan and report only when useful: meaningful difference, minor presentation issue, version/timing ambiguity, no inconsistency, or indeterminate. Use minor_presentation for a small wording problem and non_inconsistency for equivalent wording with no actual conflict. Use the available comparison evidence; no additional search for change records or reasons is required. Blank is allowed and never blocks completion. Retained protocol differences are descriptive and excluded from temporal analysis."),
        "notes": ("Optional", "Add relevant context or uncertainty in English. No mandatory reproduction log or primary rationale. In Checker, blank inherits; <clear> removes the inherited text."),
        "checkStatus": ("Required; completed Checker", "The checker sees primary answers. For an unchanged agreement, select agree and Complete=yes only. Other choices need a Checker comment. This is not blinded independent rating."),
        "checkerComment": ("Required unless Checker result is agree", "Explain the clarification, disagreement, or uncertainty in English. An unchanged agreement does not require a comment."),
        "resolutionStatus": ("Conditional; discussion", "Use resolved only after joint discussion settles a disagreement, an unable result, or a primary discuss changed to a decision. Record who participated and the conclusion in Resolution note."),
        "resolutionNote": ("Required when resolved", "Record participants and the agreed conclusion in English. Update proposed judgments and do not leave Disposition=discuss. Checker completion alone does not settle disagreement."),
    }
    groups = [
        ("1  |  Common decisions and optional notes", "173F57", ["disposition", "domain", "reviewComplete", "readiness", "notes"]),
        ("2  |  Statistical results: two required judgments", "176B75", ["statisticalIssueStatus", "substantiveStatus", "downstreamImpact"]),
        ("3  |  Report internal or planning-material inconsistency", "526947", ["internalStatus"]),
        ("4  |  Protocol versus report: optional description only", "70578A", ["protocolInterpretation"]),
        ("5  |  Second-person checking and discussion", "8B6038", ["checkStatus", "checkerComment", "resolutionStatus", "resolutionNote"]),
        ("6  |  Candidate context", "506575", ["key", "title", "locations", "primarySummary"]),
    ]
    rows, kinds = [], []
    def add(values, kind, color="173F57"):
        rows.append(values)
        kinds.append((kind, color))
    add(["How to read this Codebook", "", "", "", ""], "section")
    add(["Hierarchy and names", "GUIDE", "Field names match review headers exactly", "Dark bands identify modules. Pale bold rows identify fields. Indented option rows explain each dropdown value. Internal keys are for software only. Colors indicate structure, not scientific severity.", ""], "guide")
    add(["Automatic defaults", "GUIDE", "Gray = inactive; yellow = required input; pale blue = optional input", "Default formulas show not_applicable outside the selected domain or after exclude/discuss. Applicable questions start blank. Blank Disposition remains an unfinished decision. Changing the suggested domain changes untouched defaults.", ""], "guide")
    add(["Blank, unknown, and not_applicable", "GUIDE", "Different meanings", "Blank = no human answer (or no optional answer); indeterminate/unknown = applicable but uncertain; not_applicable = outside the selected branch or a non-retained candidate. Never use not_applicable to answer an applicable question or treat a default as a scientific judgment.", ""], "guide")
    add(["Changing a domain after answering", "GUIDE", "Recheck the whole row", "A dropdown choice replaces that cell's default formula. Later domain changes update untouched formulas, but cannot erase old literal answers. Red cells and CLEAR INACTIVE VALUES flag those answers: clear them or choose not_applicable. Newly applicable questions still need valid answers. A cleared cell does not restore its formula; use restore-defaults on a new copy if needed.", ""], "guide")
    add(["Checker defaults and inheritance", "GUIDE", "Only enter changes", "Automatic not_applicable in Checker is not an override. Applicable blank cells inherit the primary answer by Candidate key. When the domain/disposition makes a previously inactive field applicable, the old default is not inherited as an answer. Add the newly required judgment. Optional judgments may use <clear> to remove an inherited answer.", ""], "guide")
    for title, color, fields in groups:
        add([title, "", "", "", ""], "section", color)
        for field in fields:
            requirement, explanation = guides[field]
            add([LABELS[field], "FIELD", requirement, explanation, field], "field", color)
            start = len(rows) + 2
            for code, meaning in OPTIONS.get(field, {}).items():
                add(["", code, "Option", meaning, ""], "option", color)
            if field in OPTIONS:
                wb.defined_names.add(DefinedName("opt_" + field, attr_text=f"'Codebook'!$B${start}:$B${len(rows)+1}"))
    headers = ["Field name", "Option / value", "Requirement / scope", "Explanation and example", "Internal key"]
    ws = table(wb, "Codebook", headers, rows, dict(zip(headers, (30, 35, 38, 90, 29))))
    ws.auto_filter.ref = None
    ws.freeze_panes = "C2"
    ws.sheet_properties.tabColor = "176B75"
    edge = Side(style="thin", color="CFDCE4")
    for row, (kind, color) in enumerate(kinds, 2):
        for cell in ws[row]:
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.alignment = Alignment(horizontal="center" if cell.column in (1, 2, 3, 5) else "left", vertical="center", wrap_text=True)
            cell.font = Font(name="Calibri", size=11, color="253746")
        if kind == "section":
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            cell = ws.cell(row, 1)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[row].height = 32
        elif kind == "field":
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor="E5EDF3")
                cell.border = Border(left=edge, right=edge, top=Side(style="medium", color=color), bottom=edge)
                cell.font = Font(name="Calibri", size=11, bold=cell.column != 4, color=color)
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height, 90)
        elif kind == "option":
            ws.row_dimensions[row].outlineLevel = 1
            ws.cell(row, 2).font = Font(name="Consolas", size=10, color=color)
            ws.cell(row, 4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height, 52)
        else:
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height, 100)
    ws.sheet_properties.outlinePr.summaryBelow = False
    for cell in ws[1]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    assert {f for _, _, fields in groups for f in fields} == set(LABELS)


def instructions(wb, year, stage):
    rows = [
        ['Release', f"{VERSION} | {year} | {stage} | XLSX-only human adjudication"],
        ['1. Identity', 'Fill primary reviewer ID on Setup; the second person fills checker ID.'],
        ['2. Paper worksheets', 'Four annual templates, one file per year. Each paper has a C_<paper> worksheet with the same 18-column Checker layout in both stages. Every file contains Papers, Candidates, Evidence and Codebook. The C_ prefix identifies the shared layout; see File status for the current stage.'],
        ['Evidence', 'Candidates contains descriptions and source locations; Evidence contains full evidence text. Filter by key and consult originals in the shared 20XX folders. No document links or page navigation are provided.'],
        ['3. Enter', 'Yellow cells are required or common inputs; pale blue cells are optional inputs. Gray module cells are inactive and default to not_applicable. Disposition starts blank. Use English throughout. Codebook field names match review headers exactly.'],
        ['4. Retain module', 'statistical_results: Statistcal_or_numerical issue + Substantiveness required; Downstream impact optional. protocol_report: Protocol interpretation optional; no extra required module fields. Other two domains: Report-internal or planning-material inconsistency required. Other dispositions require no module judgments.'],
        ['Conditional defaults', 'Choose Review domain before answering. Untouched formulas switch between blank (applicable) and not_applicable (inactive). These defaults are not human answers. Optional blank means not assessed. No search for protocol change records or reasons is required.'],
        ['Changing a branch', 'A dropdown choice replaces its formula. Later branch changes cannot erase that answer: red cells / CLEAR INACTIVE VALUES identify obsolete answers. Clear them or select not_applicable, then answer newly required fields. Clearing does not restore formulas; the optional restore-defaults command can repair eligible empty/default cells in a new copy.'],
        ['5. Save', 'Set Complete=yes only after review; Ctrl+S saves the XLSX. After editing a completed row, reset to no, recheck, then yes. No macros and no automatic per-edit timestamp/reset.'],
        ['6. Handoff', 'Start primary review in the annual template. Checker-only columns are locked until handoff. Save a completed primary backup; the manager runs prepare-checker to create the next working copy with the same paper sheets and layout. It embeds the primary snapshot and resets Complete for checking. Both review stages are retained in that file; no external workbook links or manual answer copying are needed.'],
        ['7. Checker', 'Non-agree requires a comment. agree requires identical assessment. Disagree/unable stays unresolved until resolutionStatus=resolved plus resolutionNote. Primary discuss changed to a decision also needs documented resolution.'],
        ['8. Validation', 'Dropdowns restrict choices. Validation hint is a convenience; optional Python validation checks saved answers, including pasted values.'],
        ['9. Minimal checker entry', 'For agreement: select agree and Complete=yes. Blank applicable fields inherit primary answers by Candidate key. Formula-generated not_applicable is not an override. Newly applicable fields cannot inherit old inactive defaults. Fill only changes; <clear> clears inherited notes or optional judgments.'],
        ['10. Research', 'Exploratory audit aid: human judgment remains responsible. Protocol/plan discrepancies remain descriptive and are not automatically confirmed issues. Temporal outcomes require substantive S_C/S_W judgments. Visible-answer checking is not blinded inter-rater reliability.'],
        ['11. File status', 'PRIMARY REVIEW: fill the assessment columns. Checker result, Checker comment, Discussion resolution and Resolution note are locked until prepare-checker creates the next working copy.' if stage == 'primary' else 'READY FOR CHECKING: primary answers are preserved in this file. Assessment cells now record changes; Complete refers to checker completion.'],
    ]
    ws = table(wb, "Start", ['Step', 'Instructions'], rows, {'Step': 32, 'Instructions': 130})
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 68


def default_record(c):
    r = dict.fromkeys(ASSESSMENT, "")
    r.update(disposition="", domain=c["aiInitialAssessment"]["domain"], reviewComplete="no")
    return r


def row_hint(row, fields, stage, wb=None, key=None):
    col = {f: f"{letter(fields.index(f)+1)}{row}" for f in fields}
    d, domain, done = col["disposition"], col["domain"], col["reviewComplete"]
    stale = []
    if wb is not None:
        stale = [f'AND({inactive_condition(wb, row, f, stage, key)},{col[f]}<>"",{col[f]}<>"{NA}")' for f in MODULE_FIELDS]
    if stage == "checker":
        status, comment = col["checkStatus"], col["checkerComment"]
        formula = f'IF({status}="","PENDING",IF(AND({status}<>"agree",LEN(TRIM({comment}))=0),"COMMENT NEEDED",IF({done}="yes","CHECKED - see resolution","DRAFT")))'
        return f'=IF(OR({",".join(stale)}),"CLEAR INACTIVE VALUES",{formula})' if stale else "=" + formula
    missing = [f'COUNTIF(opt_disposition,{d})<>1', f'COUNTIF(opt_domain,{domain})<>1']
    for name, required in MODULES.items():
        if not required:
            continue
        invalid = ",".join(f'OR(COUNTIF(opt_{field},{col[field]})<>1,{col[field]}="{NA}")' for field in required)
        missing.append(f'AND({d}="retain",{domain}="{name}",OR({invalid}))')
    if stage == "checker":
        missing.extend([f'COUNTIF(opt_checkStatus,{col["checkStatus"]})<>1', f'AND({col["checkStatus"]}<>"agree",LEN(TRIM({col["checkerComment"]}))=0)'])
    formula = f'IF(OR({",".join(missing)}),"CHECK FIELDS",IF({done}="yes","READY - validate saved file","DRAFT"))'
    return f'=IF(OR({",".join(stale)}),"CLEAR INACTIVE VALUES",{formula})' if stale else "=" + formula


def apply_review_layout(ws, stage):
    """Keep only the candidate key frozen, leaving room for several answer columns."""
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    widths = {
        "key": 22, "package": 20, "title": 28, "locations": 22,
        "disposition": 18, "domain": 24, "reviewComplete": 12,
        "readiness": 22, "internalStatus": 34, "statisticalIssueStatus": 28,
        "notes": 32, "checkerComment": 32, "resolutionNote": 32, "primarySummary": 58,
    }
    for index, field in enumerate(fields, 1):
        ws.column_dimensions[letter(index)].width = widths.get(field, 18)
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = "B2"
    # Reset the pane selections: changing an existing freeze can otherwise append
    # duplicate selections pointing to the old frozen region.
    ws.sheet_view.selection = [
        Selection(pane="topRight", activeCell="B1", sqref="B1"),
        Selection(pane="bottomLeft", activeCell="A2", sqref="A2"),
        Selection(pane="bottomRight", activeCell="B2", sqref="B2"),
    ]


def add_review(wb, candidates, reviews, stage, name="Primary", locked=False):
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    rows = []
    for c in candidates:
        values = {"key": c["key"], "package": c["package"], "title": c["title"], "locations": c["locations"], **reviews[c["key"]]}
        rows.append([values.get(f, "") for f in fields])
    ws = table(wb, name, [LABELS[f] for f in fields], rows)
    apply_review_layout(ws, stage)
    ws.sheet_properties.tabColor = "D7B95C" if not locked else "8B9CA6"
    for index, c in enumerate(candidates, 2):
        ws.row_dimensions[index].height = 76
        for field in fields:
            cell = ws.cell(index, fields.index(field) + 1)
            if field == "readiness":
                cell.value = row_hint(index, fields, stage, wb, c["key"])
            elif field in ASSESSMENT + ("reviewComplete", "checkStatus", "checkerComment", "resolutionStatus", "resolutionNote"):
                field_locked = locked or (stage == "primary" and field in CHECK_ONLY_FIELDS)
                cell.protection = Protection(locked=field_locked)
                optional = field in ("downstreamImpact", "protocolInterpretation", "notes")
                cell.fill = PatternFill("solid", fgColor=GRAY if field_locked else OPTIONAL_BLUE if optional else YELLOW)
                cell.alignment = Alignment(horizontal="center" if field in OPTIONS else "left", vertical="center", wrap_text=True)
                if field in MODULE_FIELDS and not locked and values_for_default(reviews[c["key"]], field, stage):
                    cell.value = default_formula(wb, index, field, stage, c["key"])
    for field in fields:
        col = letter(fields.index(field) + 1)
        ws[f"{col}1"].comment = Comment(LABELS[field] + "\n" + ("\n".join(f"{k}: {v}" for k, v in OPTIONS[field].items()) if field in OPTIONS else "See Start and Codebook for instructions."), "Adjudication v.2.1.5.1")
        if field in OPTIONS and not locked and not (stage == "primary" and field in CHECK_ONLY_FIELDS):
            validation = DataValidation(type="list", formula1="opt_" + field, allow_blank=True)
            validation.errorTitle = 'Invalid option'
            validation.error = 'Choose a listed value.'
            validation.errorStyle = "stop"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            validation.promptTitle = LABELS[field][:32]
            validation.prompt = 'See Codebook for meanings.'
            validation.showDropDown = False
            ws.add_data_validation(validation)
            validation.add(f"{col}2:{col}{ws.max_row}")
    for field in MODULE_FIELDS:
        col = letter(fields.index(field) + 1)
        rows = enumerate(candidates, 2) if stage == "checker" else [(2, candidates[0])]
        for row, candidate in rows:
            condition = inactive_condition(wb, row, field, stage, candidate["key"])
            target = f"{col}{row}" if stage == "checker" else f"{col}2:{col}{ws.max_row}"
            stale = f'AND({condition},{col}{row}<>"",{col}{row}<>"{NA}")'
            ws.conditional_formatting.add(target, FormulaRule(formula=[stale], fill=PatternFill("solid", fgColor="FADBD8"), font=Font(color="9C0006"), stopIfTrue=True))
            ws.conditional_formatting.add(target, FormulaRule(formula=[condition], fill=PatternFill("solid", fgColor=GRAY), font=Font(color="687783")))
    for field in MODULE_FIELDS:
        header = ws.cell(1, fields.index(field) + 1)
        color = "176B75" if field in APPLICABLE["statistical_results"] else "526947" if field == "internalStatus" else "70578A"
        header.fill = PatternFill("solid", fgColor=color)
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.border = Border(left=Side(style="thin", color="FFFFFF"), right=Side(style="thin", color="FFFFFF"))
    return ws


def values_for_default(record, field, stage):
    # Never replace a human answer when building a locked primary handoff.
    return record.get(field, "") == "" or (record.get(field) == NA and field not in applicable_fields(record, include_pending=True))


def make_review(candidates, catalog_signature, output, stage="primary", primary=None, primary_id="", checker_id="", primary_hash=""):
    if stage not in ("primary", "checker"):
        raise ValueError("Unknown workbook role")
    wb = empty_book()
    year = str(candidates[0]["year"])
    instructions(wb, year, stage)
    setup = table(wb, "Setup", ["field", "value"], [["primaryReviewerId", primary_id], ["primaryReviewerName", ""], ["checkerId", checker_id], ["checkerName", ""]], {"field": 32, "value": 65})
    for row in ((2, 3) if stage == "primary" else (4, 5)):
        setup.cell(row, 2).protection = Protection(locked=False)
        setup.cell(row, 2).fill = PatternFill("solid", fgColor=YELLOW)
    add_sources(wb, candidates, output)
    add_codebook(wb)
    primary = primary if primary is not None else {c["key"]: default_record(c) for c in candidates}
    papers = paper_groups(candidates)
    names = review_names(candidates, stage)
    table(wb, "Papers", ["paper", "review_worksheet", "candidates"], [[paper, name, len(rows)] for (paper, rows), name in zip(papers.items(), names)], {"paper": 35, "review_worksheet": 38})
    if stage != "primary":
        snapshot = table(wb, "_Primary", SNAPSHOT_FIELDS, [[c["key"]] + [primary[c["key"]][f] for f in SNAPSHOT_FIELDS[1:]] for c in candidates])
        snapshot.sheet_state = "veryHidden"
        for name, reference in primary_ranges(wb, candidates):
            wb.defined_names.add(DefinedName(name, attr_text=reference))
    reviews = {key: {**record, **dict.fromkeys(CHECK_ONLY_FIELDS, ""), "primarySummary": PRIMARY_STAGE_SUMMARY} for key, record in primary.items()} if stage == "primary" else {key: {**empty_checker(), "primarySummary": primary_summary(record)} for key, record in primary.items()}
    for (paper, rows), name in zip(papers.items(), names):
        add_review(wb, rows, reviews, "primary" if stage == "primary" else "checker", name=name)
    metadata(wb, {"version": VERSION, "schema": SCHEMA, "stage": stage, "year": year, "createdAt": now(), "catalogSignature": catalog_signature, "sourceSignature": source_signature(wb), "primarySignature": digest({"id": primary_id, "reviews": primary}) if stage != "primary" else "", "primaryFileSHA256": primary_hash, "revision": REVISION})
    visible_context = {"key", "package", "title", "statement", "locations", "primary", "ai_domain", "ai_disposition"}
    for cell in wb["Candidates"][1]:
        wb["Candidates"].column_dimensions[letter(cell.column)].hidden = cell.value not in visible_context
    for row in range(2, wb["Candidates"].max_row + 1):
        wb["Candidates"].row_dimensions[row].height = 95
    wb["Manifest"].sheet_state = "hidden"
    order = ["Start", "Setup", "Papers"] + names + ["Candidates", "Evidence", "Codebook", "Manifest", "Links", "Relations"]
    if stage != "primary":
        order.append("_Primary")
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    wb.security.lockStructure = True
    return wb


def build(output_dir, catalog=CATALOG):
    candidates, signature = load_catalog(catalog)
    targets = [(year, output_dir / f"human_adjudication_v_2_1_5_1_{year}.xlsx") for year in YEARS]
    if any(path.exists() for _, path in targets):
        raise ValueError("Output already exists; use a new --output-dir. Human files are never overwritten.")
    for year, output in targets:
        annual = [c for c in candidates if str(c["year"]) == year]
        save_new(make_review(annual, signature, output, "primary"), output)


def review_records(wb, sheet="Primary", stage="primary"):
    fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
    rows = records(wb[sheet], [LABELS[f] for f in fields])
    result = {}
    for index, row in enumerate(rows, 2):
        r = {f: row[LABELS[f]] for f in fields}
        key = r.pop("key")
        if key in result:
            raise ValueError(f"{sheet}: duplicate candidate key {key}")
        for field in ("package", "title", "locations", "readiness", "primarySummary"):
            r.pop(field, None)
        if stage == "primary":
            for field in CHECK_ONLY_FIELDS:
                r.pop(field)
        for field in MODULE_FIELDS:
            cell = wb[sheet].cell(index, fields.index(field) + 1)
            if cell.data_type == "f" and cell.value == default_formula(wb, index, field, stage, key):
                # Evaluate known defaults from saved inputs, never from Excel caches.
                # In Checker these formulas are UI defaults, never override values.
                r[field] = "" if stage == "checker" or field in applicable_fields(r, include_pending=True) else NA
        result[key] = r
    return result


def read_paper_reviews(wb, candidates, stage):
    names = review_names(candidates, stage)
    expected = set(names)
    actual = {name for name in wb.sheetnames if name.startswith(("P_", "C_"))}
    if expected != actual or "Primary" in wb.sheetnames or "Checker" in wb.sheetnames:
        raise ValueError("Paper review worksheet set mismatch")
    result = {}
    for (paper, rows), name in zip(paper_groups(candidates).items(), names):
        part = review_records(wb, name, "primary" if stage == "primary" else "checker")
        if set(part) != {c["key"] for c in rows}:
            raise ValueError(f"Candidate-key set mismatch for paper {paper}")
        result.update(part)
    expected_index = {paper: (name, len(rows)) for (paper, rows), name in zip(paper_groups(candidates).items(), names)}
    index_rows = records(wb["Papers"], ["paper", "review_worksheet", "candidates"])
    actual_index = {r["paper"]: (r["review_worksheet"], r["candidates"]) for r in index_rows}
    if actual_index != expected_index or len(index_rows) != len(expected_index):
        raise ValueError("Papers index mismatch")
    return result


def read_primary(wb, candidates):
    if meta(wb)["stage"] == "primary":
        return read_paper_reviews(wb, candidates, "primary")
    rows = records(wb["_Primary"], SNAPSHOT_FIELDS)
    if len(rows) != len(candidates) or {r["key"] for r in rows} != {c["key"] for c in candidates}:
        raise ValueError("Primary snapshot candidate-key set mismatch")
    return {r["key"]: {f: r[f] for f in SNAPSHOT_FIELDS[1:]} for r in rows}


def assessment_errors(r, require_fields):
    errors = []
    for field in ASSESSMENT + ("reviewComplete",):
        value = r.get(field, "")
        if not isinstance(value, str):
            errors.append(f"{field}: expected text")
        elif field in OPTIONS and value and value not in OPTIONS[field]:
            errors.append(f"{field}: invalid option {value!r}")
    required = ["disposition", "domain", "reviewComplete"]
    if r.get("disposition") == "retain":
        required += list(MODULES.get(r.get("domain"), ()))
    for field in applicable_fields(r, include_pending=True):
        if r.get(field) == NA:
            errors.append(f"{field}: not_applicable is invalid for the selected domain; answer or clear the applicable field")
    for field in ("downstreamImpact", "protocolInterpretation"):
        if r.get(field) == "<clear>":
            errors.append(f"{field}: <clear> is a Checker-only control")
    if require_fields:
        errors.extend(f"{field}: required" for field in required if not r.get(field))
    return errors


def unresolved(r, primary):
    if r["disposition"] == "discuss":
        return True
    needs_resolution = r.get("checkStatus") in ("disagree", "unable") or primary["disposition"] == "discuss"
    return needs_resolution and r.get("resolutionStatus") != "resolved"


def outcome(c, r):
    if r["disposition"] != "retain":
        return "not_eligible"
    if c["primary"] == "P_R":
        return "protocol_descriptive"
    if c["primary"] == "P_P":
        return "plan_descriptive"
    if r["domain"] == "statistical_results" and r["statisticalIssueStatus"] in ("confirmed_issue", "probable_issue") and r["substantiveStatus"] == "substantive":
        return "primary_temporal"
    if r["domain"] == "report_internal_inconsistency" and r["internalStatus"] == "substantive_inconsistency":
        return "secondary_temporal"
    return "not_eligible"


def validate(path, catalog=CATALOG, require_complete=False, primary_only=False):
    require_english_xlsx(path)
    wb = load_workbook(path, data_only=False)
    m = meta(wb)
    if m.get("version") != VERSION or m.get("revision") != REVISION or m.get("schema") != SCHEMA or m.get("stage") not in ("primary", "checker"):
        raise ValueError("Unsupported workbook version/schema/stage")
    candidates, signature = load_catalog(catalog)
    candidates = [c for c in candidates if str(c["year"]) == m["year"]]
    if not candidates or m["catalogSignature"] != signature:
        raise ValueError("Wrong year or catalog provenance")
    expected = empty_book()
    add_sources(expected, candidates, Path(path))
    if source_signature(wb) != source_signature(expected) or m["sourceSignature"] != source_signature(expected):
        raise ValueError("Frozen candidate/evidence content changed or candidate set incomplete")
    stage = m["stage"]
    if primary_only and stage != "primary":
        raise ValueError("--primary-only requires a Primary file")
    primary_only = stage == "primary"
    if stage == "primary" and "_Primary" in wb.sheetnames:
        raise ValueError("A Primary file must not contain an imported primary snapshot")
    setup_rows = records(wb["Setup"], ["field", "value"])
    if [r["field"] for r in setup_rows] != ["primaryReviewerId", "primaryReviewerName", "checkerId", "checkerName"]:
        raise ValueError("Setup structure changed")
    setup = {r["field"]: r["value"] for r in setup_rows}
    primary_id, checker_id = setup["primaryReviewerId"], setup["checkerId"]
    issues = []
    def issue(key, severity, message):
        issues.append((key, severity, message))
    def valid_id(value):
        return isinstance(value, str) and bool(value.strip()) and value == value.strip()
    identity_needed = require_complete
    overrides = {c["key"]: empty_checker() for c in candidates} if primary_only else read_paper_reviews(wb, candidates, "checker")
    keys = {c["key"] for c in candidates}
    if set(overrides) != keys:
        raise ValueError("Checker candidate-key set mismatch (missing, extra, or wrong-year rows)")
    primary = read_primary(wb, candidates)
    if set(primary) != keys:
        raise ValueError("Primary candidate-key set mismatch")
    if stage != "primary" and m["primarySignature"] != digest({"id": primary_id, "reviews": primary}):
        raise ValueError("Locked primary snapshot/identity changed")
    if stage == "checker" and not re.fullmatch(r"[0-9a-f]{64}", m.get("primaryFileSHA256", "")):
        raise ValueError("Ready Checker file requires primary file provenance")
    for name, reference in (() if primary_only else primary_ranges(wb, candidates)):
        if name not in wb.defined_names or wb.defined_names[name].attr_text != reference:
            raise ValueError("Primary lookup range changed: " + name)
    # Reject formulas in inputs/source cells, including pasted formulas with stale caches.
    review_sheets = tuple(review_names(candidates, stage))
    for name in ("Setup", "Manifest", "Papers") + SOURCE_SHEETS + review_sheets + (() if primary_only else ("_Primary",)):
        ws = wb[name]
        allowed = LABELS["readiness"] if name in review_sheets else None
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                known_default = False
                if name in review_sheets and cell.data_type == "f":
                    fields = REVIEW_FIELDS if primary_only else CHECK_FIELDS
                    field = fields[cell.column - 1]
                    known_default = field in MODULE_FIELDS and cell.value == default_formula(wb, cell.row, field, "primary" if primary_only else "checker", ws.cell(cell.row, 1).value)
                if cell.data_type in ("f", "e") and ws.cell(1, cell.column).value != allowed and not known_default:
                    issue(str(ws.cell(cell.row, 1).value), "ERROR", f"{name}!{cell.coordinate}: formula/error cell is not an answer")
    # Read-only context in review rows must still match its key after sorting/pasting.
    by_key = {c["key"]: c for c in candidates}
    for name in review_sheets:
        for r in records(wb[name]):
            c = by_key[r[LABELS["key"]]]
            if r[LABELS["title"]] != c["title"] or r[LABELS["locations"]] != c["locations"]:
                issue(c["key"], "ERROR", f"{name}: paper/title no longer matches candidate key")
            if not primary_only and r[LABELS["primarySummary"]] != primary_summary(primary[c["key"]]):
                issue(c["key"], "ERROR", f"{name}: read-only primary summary changed")
            if primary_only:
                if any(r[LABELS[f]] != "" for f in CHECK_ONLY_FIELDS):
                    issue(c["key"], "ERROR", "Checker fields must remain blank during primary review; use prepare-checker after completion")
                if r[LABELS["primarySummary"]] != PRIMARY_STAGE_SUMMARY:
                    issue(c["key"], "ERROR", f"{name}: primary-stage instruction changed")
    results, reviews = [], {}
    for c in candidates:
        key = c["key"]
        p, changes = primary[key], overrides[key]
        r = {**changes, **{f: changes[f] if changes[f] != "" else p[f] for f in ASSESSMENT}}
        for f in MODULE_FIELDS:
            if not changes[f] and f in applicable_fields(r, include_pending=True) and f not in applicable_fields(p, include_pending=True):
                r[f] = ""
        for f in ("notes", "downstreamImpact", "protocolInterpretation"):
            if changes[f] == "<clear>":
                r[f] = ""
        complete = r["reviewComplete"] == "yes"
        primary_complete = p["reviewComplete"] == "yes"
        if stage == "checker" and not primary_complete:
            errors_ready = "Ready Checker snapshot contains unfinished primary review"
            issue(key, "ERROR", errors_ready)
        identity_needed |= complete or primary_complete
        errors = ["Primary: " + e for e in assessment_errors(p, primary_complete or require_complete)]
        if not primary_complete:
            issue(key, "ERROR" if require_complete else "PENDING", "Primary: human completion not confirmed")
        if not complete and not primary_only:
            issue(key, "ERROR" if require_complete else "PENDING", "Human completion not confirmed")
        if not primary_only:
            errors += assessment_errors(r, complete or require_complete)
            if complete and not primary_complete:
                errors.append("Cannot complete checker before primary")
            for field in ("checkerComment", "resolutionNote"):
                if not isinstance(r[field], str):
                    errors.append(f"{field}: expected text")
            for field in ("checkStatus", "resolutionStatus"):
                if r[field] and r[field] not in OPTIONS[field]:
                    errors.append(f"{field}: invalid option")
            if complete or require_complete:
                if not r["checkStatus"]:
                    errors.append("checkStatus: required")
                if r["checkStatus"] != "agree" and not str(r["checkerComment"]).strip():
                    errors.append("checkerComment: required for non-agree")
                if r["checkStatus"] == "agree" and any(effective_record(r)[f] != effective_record(p)[f] for f in ASSESSMENT):
                    errors.append("agree requires an unchanged primary assessment")
            if r["resolutionStatus"] == "resolved":
                if not str(r["resolutionNote"]).strip():
                    errors.append("resolutionNote: required when resolved")
                if r["disposition"] == "discuss":
                    errors.append("resolved cannot retain disposition=discuss")
        for error in errors:
            issue(key, "ERROR", error)
        for label, record in (("Primary", p), ("Final", r)):
            stale = [f for f in MODULE_FIELDS if f not in applicable_fields(record, include_pending=True) and record[f] not in ("", NA)]
            if stale:
                issue(key, "WARNING", label + ": inactive answers preserved in the input but ignored in analysis; clear obsolete values: " + ", ".join(stale))
        reviews[key] = effective_record(r)
        status = "invalid" if errors else "primary_complete" if primary_only and primary_complete else "pending" if not (primary_complete and complete) else "unresolved" if unresolved(r, p) else "resolved"
        if status == "unresolved":
            issue(key, "UNRESOLVED", "Discussion remains unresolved; not eligible for final analysis")
        results.append({"key": key, "status": status, "outcome": outcome(c, r) if status == "resolved" else "undetermined"})
    if identity_needed and not valid_id(primary_id):
        issue("Setup", "ERROR", "A stable primary reviewer ID is required (no surrounding whitespace)")
    if not primary_only and (require_complete or any(r["reviewComplete"] == "yes" for r in reviews.values())) and not valid_id(checker_id):
        issue("Setup", "ERROR", "Distinct primary and checker IDs are required")
    if checker_id and str(primary_id).casefold() == str(checker_id).casefold():
        issue("Setup", "ERROR", "Checker must be a different person from primary")
    return {"wb": wb, "metadata": m, "setup": setup, "candidates": candidates, "reviews": reviews, "primary": primary, "overrides": overrides, "issues": issues, "results": results}


def fail_on_errors(report):
    errors = [i for i in report["issues"] if i[1] == "ERROR"]
    if errors:
        raise ValueError(f"Validation failed ({len(errors)} errors): " + "; ".join(f"{k}: {msg}" for k, _, msg in errors[:8]))


def validation_report(report, output):
    wb = empty_book()
    table(wb, "Summary", ["field", "value"], [["version", VERSION], ["stage", report["metadata"]["stage"]], ["validatedAt", now()], *sorted(Counter(i[1] for i in report["issues"]).items()), ["note", "PASS means structural/entry validity, not scientific confirmation; unresolved items remain excluded."]], {"value": 105})
    table(wb, "Issues", ["key", "severity", "message"], report["issues"], {"key": 45, "message": 110})
    table(wb, "Status", ["key", "status", "outcome"], [[r[f] for f in ("key", "status", "outcome")] for r in report["results"]], {"key": 45})
    save_new(wb, output)


def prepare_checker(path, output, checker_id="", catalog=CATALOG):
    report = validate(path, catalog, require_complete=True, primary_only=True)
    fail_on_errors(report)
    if report["metadata"]["stage"] != "primary":
        raise ValueError("prepare-checker requires a completed primary XLSX")
    primary_id = report["setup"]["primaryReviewerId"]
    if checker_id and (checker_id != checker_id.strip() or checker_id.casefold() == primary_id.casefold()):
        raise ValueError("Checker must have a distinct, trimmed ID")
    wb = make_review(report["candidates"], report["metadata"]["catalogSignature"], output, "checker", report["primary"], primary_id, checker_id, hashlib.sha256(Path(path).read_bytes()).hexdigest())
    put(wb["Setup"]["B3"], report["setup"]["primaryReviewerName"])
    save_new(wb, output)


def export_analysis(paths, output, catalog=CATALOG, allow_incomplete=False):
    reports = [validate(path, catalog) for path in paths]
    for r in reports:
        fail_on_errors(r)
        if r["metadata"]["stage"] != "checker":
            raise ValueError("Export requires prepared Checker files; annual files still in primary review are not final review inputs")
    if len({r["metadata"]["year"] for r in reports}) != len(reports):
        raise ValueError("Duplicate year: supply exactly one authoritative checker workbook per year")
    if not allow_incomplete and any(s["status"] != "resolved" for r in reports for s in r["results"]):
        raise ValueError("Unfinished/unresolved candidates remain; use --allow-incomplete for a labeled provisional workbook")
    wb = empty_book()
    table(wb, "Readme", ["field", "value"], [["version", VERSION], ["createdAt", now()], ["scope", "Supplied years only: " + ", ".join(r["metadata"]["year"] for r in reports)], ["status", "PROVISIONAL; unresolved/pending are not zero" if any(s["status"] != "resolved" for r in reports for s in r["results"]) else "Completed visible-answer checker records for supplied years"], ["interpretation", "Exploratory candidate counts, not sensitivity or causal AI effects. Protocol/plan retained counts are not automatically confirmed issues. Confirmed and probable statistical findings remain separable in Decisions."], ["counts", "Paper complete counts are blank until all its candidates are resolved. Resolved-so-far counts are explicitly partial; all sampled papers in supplied annual candidate sets are included, including zero eligible outcomes."], ["agreement", "Visible primary answers: checker agreement describes process consistency, not blinded inter-rater reliability."]], {"field": 30, "value": 130})
    decision_fields = ["year", "package", "key", "relationship", "primaryReviewerId", "checkerId", "status", "outcome", "aiDisposition", "aiDomain"] + ["primary_" + f for f in ASSESSMENT] + ["final_" + f for f in ASSESSMENT] + ["checkStatus", "checkerComment", "resolutionStatus", "resolutionNote"]
    decisions, paper_rows, protocol_rows, lineage = [], [], [], []
    for path, report in zip(paths, reports):
        lineage.append([Path(path).name, hashlib.sha256(Path(path).read_bytes()).hexdigest(), report["metadata"]["primaryFileSHA256"], report["metadata"]["catalogSignature"]])
        states = {r["key"]: r for r in report["results"]}
        papers = defaultdict(list)
        for c in report["candidates"]:
            key = c["key"]
            r, p, s = report["reviews"][key], report["primary"][key], states[key]
            decisions.append([str(c["year"]), c["package"], key, c["primary"], report["setup"]["primaryReviewerId"], report["setup"]["checkerId"], s["status"], s["outcome"], c["aiInitialAssessment"]["disposition"], c["aiInitialAssessment"]["domain"]] + [p[f] for f in ASSESSMENT] + [r[f] for f in ASSESSMENT] + [r[f] for f in ("checkStatus", "checkerComment", "resolutionStatus", "resolutionNote")])
            papers[c["package"]].append(s)
            if c["primary"] in ("P_R", "P_P"):
                protocol_rows.append([str(c["year"]), c["package"], key, c["primary"], c["title"], c["statement"], s["status"], r["disposition"], r["protocolInterpretation"], r["internalStatus"], r["notes"]])
        for paper, states_for_paper in papers.items():
            counts = Counter(s["outcome"] for s in states_for_paper)
            pending = sum(s["status"] == "pending" for s in states_for_paper)
            uncertain = sum(s["status"] == "unresolved" for s in states_for_paper)
            complete = not (pending or uncertain)
            outcomes = [counts[k] for k in ("primary_temporal", "secondary_temporal", "protocol_descriptive", "plan_descriptive")]
            paper_rows.append([report["metadata"]["year"], paper, len(states_for_paper), pending, uncertain, "yes" if complete else "no"] + (outcomes if complete else [""] * 4) + outcomes)
    table(wb, "Decisions", decision_fields, decisions, {"key": 45, "final_notes": 75, "checkerComment": 75, "resolutionNote": 75})
    table(wb, "Paper_counts", ["year", "paper", "candidates", "pending", "unresolved", "paper_complete", "primary_temporal", "secondary_temporal", "protocol_descriptive", "plan_descriptive", "primary_resolved_so_far", "secondary_resolved_so_far", "protocol_resolved_so_far", "plan_resolved_so_far"], paper_rows, {"paper": 32})
    table(wb, "Protocol_plan_table", ["year", "paper", "key", "relationship", "title", "statement", "status", "disposition", "interpretation", "internalStatus", "notes"], protocol_rows, {"key": 45, "title": 70, "statement": 110, "notes": 80})
    table(wb, "Provenance", ["input_filename", "input_SHA256", "primary_file_SHA256", "catalog_signature"], lineage, {"input_filename": 65, "input_SHA256": 70, "primary_file_SHA256": 70, "catalog_signature": 70})
    save_new(wb, output)


def restore_defaults(path, output, catalog=CATALOG):
    """Restore only empty or inactive N/A cells, preserving all human answers."""
    report = validate(path, catalog)
    fail_on_errors(report)
    wb = report["wb"]
    count = 0
    role = "primary" if report["metadata"]["stage"] == "primary" else "checker"
    for stage, names in ((role, review_names(report["candidates"], role)),):
        fields = CHECK_FIELDS if stage == "checker" else REVIEW_FIELDS
        records_by_key = report["primary"] if stage == "primary" else report["reviews"]
        for name in names:
            ws = wb[name]
            for row in range(2, ws.max_row + 1):
                key = ws.cell(row, 1).value
                active = applicable_fields(records_by_key[key], include_pending=True)
                for field in MODULE_FIELDS:
                    cell = ws.cell(row, fields.index(field) + 1)
                    if cell.value in (None, "") or (cell.value == NA and field not in active):
                        cell.value = default_formula(wb, row, field, stage, key)
                        count += 1
    save_new(wb, output)
    print(f"Restored {count} default formulas; existing judgments and completion markers preserved")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("freeze-source", help="One-time lossless migration of v.2.0.4 candidate data only")
    p.add_argument("--output", type=Path, default=CATALOG)
    p = sub.add_parser("import-catalog", help="Copy and version a verified frozen XLSX catalog")
    p.add_argument("input", type=Path)
    p.add_argument("--output", type=Path, default=CATALOG)
    p = sub.add_parser("build", help="Build four annual XLSX workbooks using the Checker layout for both stages")
    p.add_argument("--output-dir", type=Path, default=HERE)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p = sub.add_parser("validate", help="Validate actual saved answers; never trust cached formulas")
    p.add_argument("input", type=Path)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p.add_argument("--report", type=Path)
    p.add_argument("--require-complete", action="store_true")
    p.add_argument("--primary-only", action="store_true", help="Validate primary handoff before checker work begins")
    p = sub.add_parser("prepare-checker", help="Freeze completed primary records into a new checker XLSX")
    p.add_argument("input", type=Path)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--checker-id", default="")
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p = sub.add_parser("export", help="Produce XLSX decisions, paper counts, and protocol/plan table")
    p.add_argument("inputs", nargs="+", type=Path)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    p.add_argument("--allow-incomplete", action="store_true")
    p = sub.add_parser("restore-defaults", help="Restore empty/inactive default formulas in a new workbook copy")
    p.add_argument("input", type=Path)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--catalog", type=Path, default=CATALOG)
    args = parser.parse_args()
    try:
        if args.command == "freeze-source":
            freeze_source(args.output)
        elif args.command == "import-catalog":
            import_catalog(args.input, args.output)
        elif args.command == "build":
            build(args.output_dir, args.catalog)
        elif args.command == "validate":
            report = validate(args.input, args.catalog, args.require_complete, args.primary_only)
            if args.report:
                validation_report(report, args.report)
            print(dict(Counter(i[1] for i in report["issues"])))
            fail_on_errors(report)
        elif args.command == "prepare-checker":
            prepare_checker(args.input, args.output, args.checker_id, args.catalog)
        elif args.command == "restore-defaults":
            restore_defaults(args.input, args.output, args.catalog)
        else:
            export_analysis(args.inputs, args.output, args.catalog, args.allow_incomplete)
    except (ValueError, KeyError, FileExistsError, FileNotFoundError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
