#!/usr/bin/env python3
"""Build the v.2.4.1 primary-review workbooks for one year (BUILD_SPEC.md).

Reads every <repo>/<year>/<paper>/.ai_paper_validation/final_report_1_6_2.md
(read-only), writes candidates_<year>.json, the R_A and R_B workbooks and
output_hashes.json. The 2018 v.2.3.1 R_A workbook is the structural template.
Any unexpected report structure raises an error; nothing is guessed.

Usage: build_v_2_4_1.py [--out-dir DIR]   (default: this folder)
"""
from __future__ import annotations

import argparse
from copy import copy, deepcopy
import hashlib
import io
import json
from pathlib import Path, PurePosixPath
import re
from urllib.parse import unquote
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import MultiCellRange

YEAR = 2019
VERSION = 'v_2_4_1'
HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
YEAR_DIR = REPO / str(YEAR)
REPORT_REL = '.ai_paper_validation/final_report_1_6_2.md'
TEMPLATE = REPO / 'human_adjudication/human_adjudication_v_2_3_1/2018/release/human_adjudication_v_2_3_1_2018_R_A.xlsx'
TEMPLATE_YEAR = 2018
ROLES = ('R_A', 'R_B')
DOMAINS = ('statistical_results', 'report_internal_inconsistency', 'plan_material', 'protocol_report')
RELATIONSHIPS = ('P_R', 'P_P', 'S_C', 'S_W')
H1_OLD, H1_NEW = 'Statistcal_or_numerical issue', 'Statistical_or_numerical issue'
CB_SECTION_OLD, CB_SECTION_NEW = '6  |  Candidate context', '5  |  Candidate context'
FIXED_ZIP_TIME = (2026, 10, 11, 0, 0, 0)
FIXED_MODIFIED = '2026-10-11T00:00:00Z'
FORMULA_COLS = 'GHIJKL'

# Card field labels (normalized lower case) -> canonical name. Only exact
# matches are accepted; a missing or duplicated required field is an error.
FIELD_ALIASES = {
    'candidate statement': 'statement',
    'review domain': 'domain',
    'relationship type': 'relationship',
    'primary relationship type': 'relationship',
    'exact source locations': 'locations',
    'exact source location': 'locations',
}
REQUIRED_FIELDS = ('statement', 'domain', 'relationship', 'locations')

PENDING_RE = re.compile(r'\s*\bPending Human Adjudication\b\.?\s*', re.I)
HEADING_RE = re.compile(r'^(#{1,6})\s+(.*?)\s*#*\s*$')
CARD_RE = re.compile(r'^C(\d{3})\s+—\s+(.+?)(?:\s*\{#[^}]*\})?\s*$')
CARD_LIKE_RE = re.compile(r'^C\d{3}\b')
LABEL_RE = re.compile(r'^\*\*([^*]+?):\*\*\s*(.*)$')
ANCHOR_LINE_RE = re.compile(r'^\s*<a\s+id="[^"]*"\s*>\s*</a>\s*$', re.I)
LINK_RE = re.compile(r'(!?)\[([^\]]*)\]\((<[^>]*>|[^)\s]*)\)')
URL_RE = re.compile(r'<?https?://[^\s>)]+>?')
TABLE_SEP_RE = re.compile(r'^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$')
ODD_CHARS_RE = re.compile('[\u0000-\u0008\u000b\u000c\u000e-\u001f\u007f ­​-‏  ⁠﻿�]')


class BuildError(RuntimeError):
    pass


def fail(msg):
    raise BuildError(msg)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# --------------------------------------------------------------------------
# Markdown helpers
# --------------------------------------------------------------------------

def link_target_file(target: str) -> str:
    t = target.strip()
    if t.startswith('<') and t.endswith('>'):
        t = t[1:-1]
    t = t.split('#', 1)[0].split('?', 1)[0]
    return PurePosixPath(unquote(t)).name if t else ''


def replace_links(text: str, prefix_files: bool, notes: list | None = None) -> str:
    def repl(m):
        if m.group(1):
            fail(f'image markdown not expected: {m.group(0)!r}')
        label, target = m.group(2).strip(), m.group(3)
        fname = link_target_file(target)
        if URL_RE.fullmatch(target.strip('<>')):
            return label
        if prefix_files and fname and fname.lower() not in label.lower():
            if notes is not None:
                notes.append(f'prefixed {fname!r} to link text {label!r}')
            return f'{fname} — {label}' if label else fname
        if prefix_files and not fname and notes is not None:
            notes.append(f'link without file target kept as text: {m.group(0)!r}')
        return label
    return LINK_RE.sub(repl, text)


def strip_inline(text: str) -> str:
    text = text.replace('**', '').replace('`', '')
    text = URL_RE.sub('', text)
    text = PENDING_RE.sub(' ', text)
    return re.sub(r'\s+', ' ', text).strip()


def split_row(line: str) -> list[str]:
    s = line.strip()
    if not (s.startswith('|') and s.endswith('|')):
        fail(f'table row without outer pipes: {line!r}')
    s = s[1:-1].replace('\\|', '\x00')
    return [c.replace('\x00', '|').strip() for c in s.split('|')]


def parse_tables(lines: list[str]) -> list[list[list[str]]]:
    tables, cur = [], []
    for ln in lines + ['']:
        if ln.strip().startswith('|'):
            cur.append(ln)
        elif cur:
            tables.append(cur)
            cur = []
    out = []
    for t in tables:
        if len(t) < 2 or not TABLE_SEP_RE.match(t[1].strip()):
            fail(f'malformed table starting {t[0]!r}')
        header = split_row(t[0])
        rows = [split_row(r) for r in t[2:]]
        for r in rows:
            if len(r) != len(header):
                fail(f'table row has {len(r)} cells, header has {len(header)}: {r}')
        out.append([header] + rows)
    return out


def cell_text(c: str) -> str:
    return strip_inline(replace_links(c, prefix_files=False))


def block_to_plain(lines: list[str], notes: list) -> str:
    """Field block (label remainder + continuation lines) to plain text."""
    parts, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith('|'):
            j = i
            while j < len(lines) and lines[j].strip().startswith('|'):
                j += 1
            table = parse_tables(lines[i:j])[0]
            notes.append(f'table with {len(table) - 1} rows rendered as text')
            rows = []
            for r in table[1:]:
                cells = [strip_inline(replace_links(c, True, notes)) for c in r]
                rows.append(': '.join(c for c in cells if c))
            parts.append('; '.join(rows) + '.')
            i = j
            continue
        m = re.match(r'^\s*(?:[-*+]|\d+[.)])\s+(.*)$', ln)
        if m:
            item = strip_inline(replace_links(m.group(1), True, notes))
            if item and item[-1] not in '.;:':
                item += ';'
            parts.append(item)
        else:
            parts.append(strip_inline(replace_links(ln, True, notes)))
        i += 1
    text = re.sub(r'\s+', ' ', ' '.join(p for p in parts if p)).strip()
    return re.sub(r';$', '.', text)


# --------------------------------------------------------------------------
# Report parsing
# --------------------------------------------------------------------------

def headings(lines):
    out = []
    for i, ln in enumerate(lines):
        m = HEADING_RE.match(ln)
        if m:
            out.append((i, len(m.group(1)), m.group(2).strip()))
    return out


def section(lines, heads, title):
    hits = [h for h in heads if h[2] == title]
    if len(hits) != 1:
        fail(f'expected exactly one heading {title!r}, found {len(hits)}')
    i, level, _ = hits[0]
    end = next((h[0] for h in heads if h[0] > i and h[1] <= level), len(lines))
    return i, level, end


def parse_domain_summary(lines, heads):
    i, _, end = section(lines, heads, 'Review-Domain Summary')
    tables = [t for t in parse_tables(lines[i + 1:end]) if cell_text(t[0][0]) == 'Review domain']
    if len(tables) != 1:
        fail(f'expected one Review-domain table in Review-Domain Summary, found {len(tables)}')
    counts, total = {}, None
    for r in tables[0][1:]:
        name, val = cell_text(r[0]), cell_text(r[-1])
        if not re.fullmatch(r'\d+', val):
            fail(f'non-integer domain count {val!r}')
        if name == 'Total':
            total = int(val)
        elif name in DOMAINS and name not in counts:
            counts[name] = int(val)
        else:
            fail(f'unexpected domain-summary row {name!r}')
    if set(counts) != set(DOMAINS) or total is None:
        fail(f'domain summary incomplete: {counts}, total={total}')
    if sum(counts.values()) != total:
        fail(f'domain summary rows sum {sum(counts.values())} != Total {total}')
    return counts, total


def parse_index(lines, heads):
    i, _, end = section(lines, heads, 'Candidate Index')
    tables = parse_tables(lines[i + 1:end])
    if len(tables) != 1:
        fail(f'expected one table under Candidate Index, found {len(tables)}')
    hdr = [cell_text(c) for c in tables[0][0]]
    id_col = [k for k, h in enumerate(hdr) if h in ('ID', 'Stable ID')]
    if len(id_col) != 1 or 'Review domain' not in hdr or 'Relationship type' not in hdr:
        fail(f'unexpected Candidate Index header {hdr}')
    ic, dc, rc = id_col[0], hdr.index('Review domain'), hdr.index('Relationship type')
    tc = next((k for k, h in enumerate(hdr) if h.startswith('Candidate')), None)
    rows = []
    for r in tables[0][1:]:
        cid = cell_text(r[ic])
        if not re.fullmatch(r'C\d{3}', cid):
            fail(f'bad index ID {r[ic]!r}')
        rows.append({'id': cid, 'domain': cell_text(r[dc]), 'relationship': cell_text(r[rc]),
                     'title': cell_text(r[tc]) if tc is not None else None})
    return rows


def parse_cards(lines, heads, paper):
    i, level, _ = section(lines, heads, 'Candidate Evidence Cards')
    all_card_heads = [h for h in heads if CARD_LIKE_RE.match(h[2])]
    stop_level = max(level, 2)
    end = next((h[0] for h in heads if h[0] > i and h[1] <= stop_level
                and not (h[1] == 2 and CARD_LIKE_RE.match(h[2]))), len(lines))
    card_heads = [h for h in heads if i < h[0] < end and h[1] == 2 and CARD_LIKE_RE.match(h[2])]
    if len(card_heads) != len(all_card_heads) or any(h[1] != 2 for h in all_card_heads):
        fail(f'{paper}: card-like headings outside the cards section or at another level')
    cards = []
    for k, (li, _, title) in enumerate(card_heads):
        m = CARD_RE.match(title)
        if not m:
            fail(f'{paper}: unparseable card heading {title!r}')
        cid = 'C' + m.group(1)
        body_end = card_heads[k + 1][0] if k + 1 < len(card_heads) else end
        body = [ln for ln in lines[li + 1:body_end] if not ANCHOR_LINE_RE.match(ln)]
        if any('<a id' in ln.lower() for ln in body):
            fail(f'{paper} {cid}: inline <a id> anchor in card body')
        fields, raw = {}, {}
        j = 0
        while j < len(body):
            lm = LABEL_RE.match(body[j])
            if lm:
                canon = FIELD_ALIASES.get(lm.group(1).strip().lower())
                blk = [lm.group(2)]
                j += 1
                while j < len(body) and not body[j].lstrip().startswith('**') and not HEADING_RE.match(body[j]) \
                        and not (body[j].strip() == '' and j + 1 < len(body) and LABEL_RE.match(body[j + 1])):
                    blk.append(body[j])
                    j += 1
                if canon:
                    if canon in fields:
                        fail(f'{paper} {cid}: duplicate field {canon}')
                    while blk and not blk[-1].strip():
                        blk.pop()
                    fields[canon] = blk
                    raw[canon] = lm.group(1).strip()
                continue
            j += 1
        missing = [f for f in REQUIRED_FIELDS if f not in fields]
        if missing:
            fail(f'{paper} {cid}: missing fields {missing}')
        notes = []
        clean_title = strip_inline(replace_links(m.group(2), False))
        statement = strip_inline(replace_links(' '.join(fields['statement']), False))
        domain = strip_inline(' '.join(fields['domain'])).rstrip('.')
        rel = strip_inline(' '.join(fields['relationship'])).rstrip('.')
        locs_md = '\n'.join(fields['locations']).strip()
        locs = block_to_plain([ln for ln in fields['locations'] if ln.strip()], notes)
        if domain not in DOMAINS:
            fail(f'{paper} {cid}: invalid review domain {domain!r}')
        if rel not in RELATIONSHIPS:
            fail(f'{paper} {cid}: invalid relationship type {rel!r}')
        for name, val in (('title', clean_title), ('statement', statement), ('locations', locs)):
            if not val:
                fail(f'{paper} {cid}: empty {name}')
            for bad in ('](', '**', '<a id', '`'):
                if bad in val:
                    fail(f'{paper} {cid}: markdown residue {bad!r} in {name}: {val!r}')
            if ODD_CHARS_RE.search(val):
                notes.append(f'odd character(s) {sorted(set(ODD_CHARS_RE.findall(val)))!r} in {name}')
        cards.append({'id': cid, 'title': clean_title, 'statement': statement, 'domain': domain,
                      'relationship_type': rel, 'source_locations': locs,
                      'source_locations_markdown': locs_md,
                      'source_link_targets': [t.strip('<>') for _, _, t in LINK_RE.findall(locs_md)],
                      'field_labels_used': raw, 'parse_notes': notes})
    return cards


def parse_report(paper: str, path: Path) -> dict:
    text = path.read_text(encoding='utf-8')
    if '```' in text:
        fail(f'{paper}: code fences present; parser does not handle them')
    lines = text.split('\n')
    heads = headings(lines)
    counts, total = parse_domain_summary(lines, heads)
    index = parse_index(lines, heads)
    cards = parse_cards(lines, heads, paper)
    n = len(cards)
    if not (len(index) == n == total):
        fail(f'{paper}: index rows {len(index)}, cards {n}, domain-summary total {total} disagree')
    expected_ids = [f'C{k:03d}' for k in range(1, n + 1)]
    if [c['id'] for c in cards] != expected_ids or [r['id'] for r in index] != expected_ids:
        fail(f'{paper}: IDs not sequential C001..C{n:03d} in both index and cards')
    for c, r in zip(cards, index):
        if c['domain'] != r['domain'] or c['relationship_type'] != r['relationship']:
            fail(f'{paper} {c["id"]}: card domain/relationship {c["domain"]}/{c["relationship_type"]} '
                 f'!= index {r["domain"]}/{r["relationship"]}')
        c['index_title'] = r['title']
    card_counts = {d: sum(c['domain'] == d for c in cards) for d in DOMAINS}
    if card_counts != counts:
        fail(f'{paper}: card domain counts {card_counts} != summary {counts}')
    return {'paper': paper, 'cards': cards, 'index_rows': len(index), 'summary_total': total,
            'domain_counts': counts}


def discover_papers() -> list[str]:
    papers = sorted(p.name for p in YEAR_DIR.iterdir() if p.is_dir() and not p.name.startswith('.'))
    for p in papers:
        if not (YEAR_DIR / p / REPORT_REL).is_file():
            fail(f'{p}: no {REPORT_REL}')
    return papers


def source_hashes(papers):
    return {p: sha256(YEAR_DIR / p / REPORT_REL) for p in papers}


# --------------------------------------------------------------------------
# Template formula patterns
# --------------------------------------------------------------------------

REL_REF = re.compile(r'(?<![$A-Za-z0-9_])([A-R])2(?![0-9])')


def template_patterns(tpl_ws, tpl_paper):
    """Turn template row-2 formulas into placeholder patterns."""
    last = tpl_ws.max_row
    prefix = f'{TEMPLATE_YEAR}::{tpl_paper}::'
    rng = f'$A$2:$A${last}'
    pats = {}
    for col in FORMULA_COLS:
        f = tpl_ws[f'{col}2'].value
        if not isinstance(f, str) or not f.startswith('='):
            fail(f'template {col}2 is not a formula')
        p = f.replace(prefix, '{PREFIX}').replace(rng, '{RANGE}')
        p = REL_REF.sub(lambda m: m.group(1) + '{ROW}', p)
        if '{ROW}' not in p:
            fail(f'no row reference found in template {col}2')
        pats[col] = p
    return pats


def render(pattern, row, last, prefix):
    return pattern.replace('{ROW}', str(row)).replace('{RANGE}', f'$A$2:$A${last}').replace('{PREFIX}', prefix)


def checks_formula(sheets_last, col, word):
    return '=' + '+'.join(f"COUNTIF('{s}'!${col}$2:${col}${last},\"{word}\")" for s, last in sheets_last)


def verify_template(tpl):
    """Self-check: patterns regenerate every template formula exactly."""
    cnames = [n for n in tpl.sheetnames if n.startswith('C_')]
    first = tpl[cnames[0]]
    pats = template_patterns(first, cnames[0][2:])
    for n in cnames:
        ws = tpl[n]
        last = ws.max_row
        for r in range(2, last + 1):
            for col in FORMULA_COLS:
                exp = render(pats[col], r, last, f'{TEMPLATE_YEAR}::{n[2:]}::')
                if ws[f'{col}{r}'].value != exp:
                    fail(f'template pattern does not reproduce {n}!{col}{r}')
    sl = [(n, tpl[n].max_row) for n in cnames]
    if tpl['Checks']['B4'].value != checks_formula(sl, 'G', 'PRIMARY READY') or \
            tpl['Checks']['B5'].value != checks_formula(sl, 'D', 'discuss'):
        fail('template Checks B4/B5 pattern not reproduced')
    return pats


# --------------------------------------------------------------------------
# Workbook construction
# --------------------------------------------------------------------------

def build_c_sheet(wb, proto, title, year, paper, cards, pats):
    ws = wb.create_sheet(title)
    last = len(cards) + 1
    prefix = f'{year}::{paper}::'
    ws.sheet_properties = deepcopy(proto.sheet_properties)
    ws.sheet_format = deepcopy(proto.sheet_format)
    ws.views = deepcopy(proto.views)
    ws.page_margins = deepcopy(proto.page_margins)
    ws.page_setup = deepcopy(proto.page_setup)
    ws.print_options = deepcopy(proto.print_options)
    for key, dim in proto.column_dimensions.items():
        ws.column_dimensions[key] = copy(dim)
        ws.column_dimensions[key].worksheet = ws
    for r in range(1, last + 1):
        src_r = 1 if r == 1 else (2 if r % 2 == 0 else 3)
        rd = copy(proto.row_dimensions[src_r])
        rd.worksheet = ws
        rd.index = r
        ws.row_dimensions[r] = rd
        for c in range(1, 19):
            src = proto.cell(src_r, c)
            cell = ws.cell(r, c)
            cell._style = copy(src._style)
            if r == 1:
                cell.value = src.value
    if ws['H1'].value != H1_OLD:
        fail(f'template H1 unexpected: {ws["H1"].value!r}')
    ws['H1'].value = H1_NEW
    for r, cd in enumerate(cards, start=2):
        ws[f'A{r}'] = f'{prefix}{cd["id"]}'
        ws[f'B{r}'] = cd['title']
        ws[f'C{r}'] = cd['source_locations']
        ws[f'E{r}'] = cd['domain']
        ws[f'F{r}'] = 'no'
        for col in FORMULA_COLS:
            ws[f'{col}{r}'] = render(pats[col], r, last, prefix)
    for dv in proto.data_validations.dataValidation:
        new = deepcopy(dv)
        cols = {rng.coord.split(':')[0].rstrip('0123456789') for rng in dv.sqref.ranges}
        if len(cols) != 1:
            fail('unexpected template DV sqref')
        col = cols.pop()
        new.sqref = MultiCellRange(f'{col}2:{col}{last}')
        ws.data_validations.append(new)
    for cf in proto.conditional_formatting:
        cols = {rng.coord.split(':')[0].rstrip('0123456789') for rng in cf.sqref.ranges}
        if len(cols) != 1:
            fail('unexpected template CF sqref')
        col = cols.pop()
        for rule in cf.rules:
            ws.conditional_formatting.add(f'{col}2:{col}{last}', deepcopy(rule))
    ws.auto_filter.ref = f'A1:R{last}'
    return ws, last


def build_workbook(catalog, role, year, out_path):
    tpl = load_workbook(TEMPLATE)
    pats = verify_template(tpl)
    wb = load_workbook(TEMPLATE)
    tpl_c = [n for n in tpl.sheetnames if n.startswith('C_')]
    proto = tpl[tpl_c[0]]
    for n in tpl_c:
        wb.remove(wb[n])
    sheets_last = []
    new_sheets = []
    for entry in catalog:
        title = f'C_{entry["paper"]}'
        if len(title) > 31:
            fail(f'sheet name too long: {title}')
        ws, last = build_c_sheet(wb, proto, title, year, entry['paper'], entry['cards'], pats)
        sheets_last.append((title, last))
        new_sheets.append(ws)
    order = [wb['Setup'], wb['Checks'], wb['Papers'], *new_sheets, wb['Codebook']]
    if set(wb.sheetnames) != {s.title for s in order}:
        fail(f'unexpected sheets {wb.sheetnames}')
    wb._sheets = order
    wb.active = 0

    st = wb['Setup']
    if st['B4'].value != 'R_A' or st['B5'].value != TEMPLATE_YEAR or st['B2'].value is not None or st['B3'].value is not None:
        fail('template Setup not as expected')
    st['B4'] = role
    st['B5'] = int(year)

    total = sum(len(e['cards']) for e in catalog)
    ck = wb['Checks']
    ck['B3'] = total
    ck['B4'] = checks_formula(sheets_last, 'G', 'PRIMARY READY')
    ck['B5'] = checks_formula(sheets_last, 'D', 'discuss')

    pp = wb['Papers']
    tpl_rows = pp.max_row
    styles = {par: [copy(pp.cell(2 + par, c)._style) for c in range(1, 4)] for par in (0, 1)}
    rdim = copy(pp.row_dimensions[2])
    if pp.max_row > len(catalog) + 1:
        pp.delete_rows(len(catalog) + 2, pp.max_row - len(catalog) - 1)
    for k, entry in enumerate(catalog):
        r = k + 2
        vals = (entry['paper'], f'C_{entry["paper"]}', len(entry['cards']))
        for c, v in enumerate(vals, start=1):
            cell = pp.cell(r, c)
            cell.value = v
            cell._style = copy(styles[k % 2][c - 1])
        if r > tpl_rows:
            d = copy(rdim)
            d.worksheet = pp
            d.index = r
            pp.row_dimensions[r] = d

    cb = wb['Codebook']
    hits = {H1_OLD: [], CB_SECTION_OLD: []}
    for row in cb.iter_rows():
        for c in row:
            if c.value in hits:
                hits[c.value].append(c)
    for old, new in ((H1_OLD, H1_NEW), (CB_SECTION_OLD, CB_SECTION_NEW)):
        if len(hits[old]) != 1:
            fail(f'Codebook: expected one {old!r}, found {len(hits[old])}')
        hits[old][0].value = new

    buf = io.BytesIO()
    wb.save(buf)
    normalize_zip(buf.getvalue(), out_path)


def normalize_zip(data: bytes, out_path: Path):
    """Fixed core 'modified' stamp and zip entry times for reproducible bytes."""
    src = ZipFile(io.BytesIO(data))
    with ZipFile(out_path, 'w', ZIP_DEFLATED) as z:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == 'docProps/core.xml':
                content, n = re.subn(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',
                                     rb'\g<1>' + FIXED_MODIFIED.encode() + rb'\g<2>', content)
                if n != 1:
                    fail('core.xml modified stamp not found')
            zi = ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = ZIP_DEFLATED
            zi.external_attr = info.external_attr
            z.writestr(zi, content)


# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--out-dir', default=str(HERE))
    args = ap.parse_args()
    out = Path(args.out_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)

    papers = discover_papers()
    before = source_hashes(papers)
    tpl_sha_before = sha256(TEMPLATE)
    catalog = [parse_report(p, YEAR_DIR / p / REPORT_REL) for p in papers]

    records = []
    for e in catalog:
        for c in e['cards']:
            records.append({
                'paper': e['paper'], 'id': c['id'], 'key': f'{YEAR}::{e["paper"]}::{c["id"]}',
                'title': c['title'], 'index_title': c['index_title'], 'statement': c['statement'],
                'domain': c['domain'], 'relationship_type': c['relationship_type'],
                'source_locations': c['source_locations'],
                'source_locations_markdown': c['source_locations_markdown'],
                'source_link_targets': c['source_link_targets'],
                'field_labels_used': c['field_labels_used'], 'parse_notes': c['parse_notes'],
                'report_path': f'{YEAR}/{e["paper"]}/{REPORT_REL}', 'report_sha256': before[e['paper']],
            })
    keys = [r['key'] for r in records]
    if len(keys) != len(set(keys)):
        fail('duplicate candidate keys')
    cand_doc = {
        'version': 'v.2.4.1', 'year': YEAR, 'source': 'workflow 1.6.2 final_report_1_6_2.md',
        'papers': [{'paper': e['paper'], 'candidates': len(e['cards']), 'index_rows': e['index_rows'],
                    'domain_summary_total': e['summary_total'], 'domain_counts': e['domain_counts'],
                    'report_path': f'{YEAR}/{e["paper"]}/{REPORT_REL}', 'report_sha256': before[e['paper']]}
                   for e in catalog],
        'total_candidates': len(records), 'candidates': records,
    }
    (out / f'candidates_{YEAR}.json').write_text(json.dumps(cand_doc, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')

    outputs = {}
    for role in ROLES:
        path = out / f'human_adjudication_{VERSION}_{YEAR}_{role}.xlsx'
        build_workbook(catalog, role, YEAR, path)
        outputs[path.name] = sha256(path)

    after = source_hashes(papers)
    if after != before or sha256(TEMPLATE) != tpl_sha_before:
        fail('source report or template hash changed during build')
    (out / 'output_hashes.json').write_text(json.dumps({
        'workbooks': outputs,
        'source_reports': {f'{YEAR}/{p}/{REPORT_REL}': h for p, h in before.items()},
        'template': {str(TEMPLATE.relative_to(REPO)): tpl_sha_before},
        'source_hashes_unchanged_during_build': True,
    }, indent=2) + '\n')
    for e in catalog:
        print(e['paper'], len(e['cards']), e['domain_counts'])
    print('total', len(records))


if __name__ == '__main__':
    main()
