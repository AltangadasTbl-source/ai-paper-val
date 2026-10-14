#!/usr/bin/env python3
"""Structural and recalculation checks for the v.2.4.1 workbooks (BUILD_SPEC.md,
"Required checks" 1-6, plus a template-fidelity check). Writes
validation_report.json next to this script. Test copies are written only to
--scratch-dir, which must lie outside this folder.

Usage: validate_v_2_4_1.py --scratch-dir DIR
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

from openpyxl import load_workbook

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_v_2_4_1 as B  # noqa: E402

HERE = B.HERE
YEAR = B.YEAR
RESULTS = []


def check(num, name, ok, detail=None):
    RESULTS.append({'check': num, 'name': name, 'passed': bool(ok), 'detail': detail})
    print(('PASS' if ok else 'FAIL'), num, name, '' if ok or detail is None else detail)
    return ok


def wb_path(role):
    return HERE / f'human_adjudication_{B.VERSION}_{YEAR}_{role}.xlsx'


def sheet_xml_map(z):
    """sheet title -> worksheet part name, via workbook.xml + rels."""
    wbx = z.read('xl/workbook.xml').decode()
    rels = z.read('xl/_rels/workbook.xml.rels').decode()
    rid_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rid_target.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet name="([^"]+)" sheetId="\d+" state="visible" r:id="(rId\d+)"', wbx):
        t = rid_target[rid]
        out[name.replace('&amp;', '&')] = t.lstrip('/') if t.startswith('/') else 'xl/' + t
    return out


# --------------------------------------------------------------------------

def check1_counts(catalog_json, papers, wbs):
    ok_all = True
    for p in papers:
        text = (B.YEAR_DIR / p / B.REPORT_REL).read_text(encoding='utf-8')
        lines = text.split('\n')
        n_cards = len(re.findall(r'(?m)^## C\d{3}\b', text))
        heads = B.headings(lines)
        n_index = len(B.parse_index(lines, heads))
        _, total = B.parse_domain_summary(lines, heads)
        rows = {role: wb[f'C_{p}'].max_row - 1 for role, wb in wbs.items()}
        papers_c = {role: next(r[2] for r in wb['Papers'].iter_rows(min_row=2, values_only=True) if r[0] == p)
                    for role, wb in wbs.items()}
        js = next(x['candidates'] for x in catalog_json['papers'] if x['paper'] == p)
        ok = n_cards == n_index == total == js and all(v == n_cards for v in rows.values()) \
            and all(v == n_cards for v in papers_c.values())
        ok_all &= check(1, f'{p}: rows = index rows = cards = domain-summary total', ok,
                        {'cards': n_cards, 'index': n_index, 'summary_total': total, 'json': js,
                         'sheet_rows': rows, 'papers_C': papers_c})
    return ok_all


def check2_content(catalog_json, wb, role):
    cands = catalog_json['candidates']
    problems = []
    seen = []
    for p in sorted({c['paper'] for c in cands}):
        ws = wb[f'C_{p}']
        exp = [c for c in cands if c['paper'] == p]
        for k, c in enumerate(exp, start=2):
            a = ws[f'A{k}'].value
            seen.append(a)
            want_key = f'{YEAR}::{p}::C{k - 1:03d}'
            if a != want_key or c['key'] != want_key:
                problems.append(f'{ws.title}!A{k} key {a!r} != {want_key!r}')
            if ws[f'E{k}'].value not in B.DOMAINS or ws[f'E{k}'].value != c['domain']:
                problems.append(f'{ws.title}!E{k} domain {ws[f"E{k}"].value!r}')
            if ws[f'D{k}'].value is not None:
                problems.append(f'{ws.title}!D{k} not blank')
            if ws[f'F{k}'].value != 'no':
                problems.append(f'{ws.title}!F{k} != no')
            if ws[f'R{k}'].value is not None or any(ws[f'{col}{k}'].value is not None for col in 'MNOPQ'):
                problems.append(f'{ws.title} row {k}: M-R not blank')
            for col, field in (('B', 'title'), ('C', 'source_locations')):
                v = ws[f'{col}{k}'].value
                if not isinstance(v, str) or not v.strip():
                    problems.append(f'{ws.title}!{col}{k} empty')
                    continue
                if v != c[field]:
                    problems.append(f'{ws.title}!{col}{k} differs from catalog')
                for bad in ('](', '**', '<a id', '`'):
                    if bad in v:
                        problems.append(f'{ws.title}!{col}{k} markdown residue {bad!r}')
                if re.search(r'pending human adjudication', v, re.I):
                    problems.append(f'{ws.title}!{col}{k} boilerplate')
                if len(v) > 32767:
                    problems.append(f'{ws.title}!{col}{k} too long')
        if ws.max_row != len(exp) + 1:
            problems.append(f'{ws.title} max_row {ws.max_row} != {len(exp) + 1}')
    if len(seen) != len(set(seen)):
        problems.append('duplicate keys')
    return check(2, f'{role}: keys unique/sequential, E valid, D blank, F no, B/C clean', not problems, problems[:20])


def covers(sqref, last):
    """Exactly one single-column range spanning rows 2..last (openpyxl writes 'D2' for 'D2:D2')."""
    rs = list(sqref.ranges)
    return len(rs) == 1 and rs[0].min_col == rs[0].max_col and (rs[0].min_row, rs[0].max_row) == (2, last)


def check3_formulas_layout(catalog_json, wb, tpl, role):
    pats = B.verify_template(tpl)
    tcn = [n for n in tpl.sheetnames if n.startswith('C_')]
    tws = tpl[tcn[0]]
    problems = []
    for p in sorted({c['paper'] for c in catalog_json['candidates']}):
        ws = wb[f'C_{p}']
        last = ws.max_row
        prefix = f'{YEAR}::{p}::'
        for r in range(2, last + 1):
            for col in B.FORMULA_COLS:
                if ws[f'{col}{r}'].value != B.render(pats[col], r, last, prefix):
                    problems.append(f'{ws.title}!{col}{r} formula mismatch')
        # data validations / conditional formats: attributes as template, sqref = same column rows 2..last
        def dvsig(d):
            return (str(d.sqref).split(':')[0].rstrip('0123456789'), d.type, d.formula1, d.allow_blank,
                    d.showErrorMessage, d.showDropDown, d.showInputMessage, d.errorTitle, d.error, d.errorStyle)
        if [dvsig(d) for d in ws.data_validations.dataValidation] != [dvsig(d) for d in tws.data_validations.dataValidation]:
            problems.append(f'{ws.title} DV mismatch')
        for d in ws.data_validations.dataValidation:
            if not covers(d.sqref, last):
                problems.append(f'{ws.title} DV sqref {d.sqref} not rows 2..{last}')

        def cfsig(cf):
            return (str(cf.sqref).split(':')[0].rstrip('0123456789'),
                    [(r.type, r.formula, r.priority, r.stopIfTrue, r.dxf) for r in cf.rules])
        if [cfsig(cf) for cf in ws.conditional_formatting] != [cfsig(cf) for cf in tws.conditional_formatting]:
            problems.append(f'{ws.title} CF mismatch')
        for cf in ws.conditional_formatting:
            if not covers(cf.sqref, last):
                problems.append(f'{ws.title} CF sqref {cf.sqref}')
        for col in 'MNOPQ':
            cd = ws.column_dimensions[col]
            if not cd.hidden or cd.width != 3:
                problems.append(f'{ws.title} column {col} not hidden width 3')
        widths = {k: (v.width, v.hidden) for k, v in ws.column_dimensions.items()}
        if widths != {k: (v.width, v.hidden) for k, v in tws.column_dimensions.items()}:
            problems.append(f'{ws.title} column widths differ from template')
        if ws.freeze_panes != 'B2':
            problems.append(f'{ws.title} freeze {ws.freeze_panes}')
        if ws.auto_filter.ref != f'A1:R{last}':
            problems.append(f'{ws.title} autofilter {ws.auto_filter.ref}')
        hdr = [ws.cell(1, c).value for c in range(1, 19)]
        thdr = [tws.cell(1, c).value for c in range(1, 19)]
        thdr[7] = B.H1_NEW
        if hdr != thdr:
            problems.append(f'{ws.title} header row {hdr}')
        for r in range(1, last + 1):
            src = 1 if r == 1 else (2 if r % 2 == 0 else 3)
            if ws.row_dimensions[r].height != tws.row_dimensions[src].height:
                problems.append(f'{ws.title} row {r} height')
            for c in range(1, 19):
                if ws.cell(r, c)._style != tws.cell(src, c)._style:
                    problems.append(f'{ws.title} {ws.cell(r, c).coordinate} style')
            if r > 1 and not all(ws.cell(r, c).alignment.wrap_text for c in range(1, 19)):
                problems.append(f'{ws.title} row {r} wrap')
    return check(3, f'{role}: G-L formulas = template pattern; DV/CF rows 2..last; M-Q hidden; freeze B2; layout/styles',
                 not problems, problems[:20])


def check4_sheets_codebook(catalog_json, wbs, tpl):
    papers = sorted({c['paper'] for c in catalog_json['candidates']})
    expected = ['Setup', 'Checks', 'Papers', *[f'C_{p}' for p in papers], 'Codebook']
    ok = True
    for role, wb in wbs.items():
        ok &= check(4, f'{role}: sheet list exactly as specified', wb.sheetnames == expected, wb.sheetnames)
        st = wb['Setup']
        ok &= check(4, f'{role}: Setup B2/B3 blank, B4 role, B5 integer year',
                    st['B2'].value is None and st['B3'].value is None and st['B4'].value == role
                    and type(st['B5'].value) is int and st['B5'].value == YEAR,
                    [st[c].value for c in ('B2', 'B3', 'B4', 'B5')])
        ck = wb['Checks']
        sl = [(f'C_{p}', wb[f'C_{p}'].max_row) for p in papers]
        ok &= check(4, f'{role}: Checks B3 integer total; B4/B5 COUNTIF over exact row ranges; other cells as template',
                    type(ck['B3'].value) is int and ck['B3'].value == catalog_json['total_candidates']
                    and ck['B4'].value == B.checks_formula(sl, 'G', 'PRIMARY READY')
                    and ck['B5'].value == B.checks_formula(sl, 'D', 'discuss')
                    and all(ck[c].value == tpl['Checks'][c].value for c in
                            ('A1', 'B1', 'A2', 'B2', 'A3', 'A4', 'A5', 'A6', 'B6', 'A7', 'B7', 'A8', 'B8', 'A9', 'B9'))
                    and ck.max_row == tpl['Checks'].max_row,
                    ck['B3'].value)
        pp = list(wb['Papers'].iter_rows(values_only=True))
        want = [('Paper', 'Review worksheet', 'Candidates')] + \
            [(x['paper'], f'C_{x["paper"]}', x['candidates']) for x in catalog_json['papers']]
        ok &= check(4, f'{role}: Papers rows', pp == want and all(type(r[2]) is int for r in pp[1:]), pp)
        # Codebook values
        tv = [list(r) for r in tpl['Codebook'].iter_rows(values_only=True)]
        nv = [list(r) for r in wb['Codebook'].iter_rows(values_only=True)]
        diffs = [(i, j, tv[i][j], nv[i][j]) for i in range(max(len(tv), len(nv))) for j in range(5)
                 if (tv[i][j] if i < len(tv) else None) != (nv[i][j] if i < len(nv) else None)]
        allowed = {(B.H1_OLD, B.H1_NEW), (B.CB_SECTION_OLD, B.CB_SECTION_NEW)}
        ok &= check(4, f'{role}: Codebook values identical to template except the two fixes',
                    len(diffs) == 2 and {(d[2], d[3]) for d in diffs} == allowed, diffs)
    # Codebook / Setup / Checks / Papers XML vs template (formatting verbatim)
    tz = ZipFile(B.TEMPLATE)
    tmap = sheet_xml_map(tz)
    for role in wbs:
        z = ZipFile(wb_path(role))
        m = sheet_xml_map(z)
        t = tz.read(tmap['Codebook']).replace(B.H1_OLD.encode(), B.H1_NEW.encode()) \
            .replace(B.CB_SECTION_OLD.encode(), B.CB_SECTION_NEW.encode())
        ok &= check(4, f'{role}: Codebook sheet XML byte-identical to template after the two fixes',
                    z.read(m['Codebook']) == t)
        tc = tz.read(tmap['Setup']).replace(b'<v>2018</v>', f'<v>{YEAR}</v>'.encode()) \
            .replace(b'<t>R_A</t>', f'<t>{role}</t>'.encode())
        ok &= check(4, f'{role}: Setup sheet XML identical to template except B4/B5', z.read(m['Setup']) == tc)
        ok &= check(4, f'{role}: styles.xml and theme identical to template',
                    z.read('xl/styles.xml') == tz.read('xl/styles.xml')
                    and z.read('xl/theme/theme1.xml') == tz.read('xl/theme/theme1.xml'))
    # R_A vs R_B: every cell value
    a, b = wbs['R_A'], wbs['R_B']
    diffs = []
    for ws in a:
        wb2 = b[ws.title]
        for r in range(1, max(ws.max_row, wb2.max_row) + 1):
            for c in range(1, max(ws.max_column, wb2.max_column) + 1):
                if ws.cell(r, c).value != wb2.cell(r, c).value:
                    diffs.append(f'{ws.title}!{ws.cell(r, c).coordinate}')
    ok &= check(4, 'R_A vs R_B differ only in Setup!B4 (all cell values compared)', diffs == ['Setup!B4'], diffs)
    za, zb = ZipFile(wb_path('R_A')), ZipFile(wb_path('R_B'))
    ma = sheet_xml_map(za)
    part_diffs = [n for n in za.namelist() if za.read(n) != zb.read(n)]
    ok &= check(4, 'R_A vs R_B zip parts differ only in the Setup sheet part',
                za.namelist() == zb.namelist() and part_diffs == [ma['Setup']], part_diffs)
    return ok


def check_fidelity(tpl):
    """Builder fed the template's own 2018 content must reproduce the template's XML."""
    from openpyxl import load_workbook as lw
    import io
    pats = B.verify_template(tpl)
    wb = lw(B.TEMPLATE)
    cn = [n for n in tpl.sheetnames if n.startswith('C_')]
    proto = tpl[cn[0]]
    for n in cn:
        wb.remove(wb[n])
    new = []
    for n in cn:
        ws = tpl[n]
        cards = [{'id': ws[f'A{r}'].value.split('::')[-1], 'title': ws[f'B{r}'].value,
                  'source_locations': ws[f'C{r}'].value, 'domain': ws[f'E{r}'].value} for r in range(2, ws.max_row + 1)]
        s, _ = B.build_c_sheet(wb, proto, n, B.TEMPLATE_YEAR, n[2:], cards, pats)
        s['H1'] = B.H1_OLD
        new.append(s)
    wb._sheets = [wb['Setup'], wb['Checks'], wb['Papers'], *new, wb['Codebook']]
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    za, zb = ZipFile(B.TEMPLATE), ZipFile(buf)
    diffs = [n for n in za.namelist() if n != 'docProps/core.xml' and za.read(n) != zb.read(n)]
    return check(3, 'Fidelity: C-sheet builder reproduces the template workbook XML byte-for-byte from its own content',
                 not diffs and za.namelist() == zb.namelist(), diffs)


# --------------------------------------------------------------------------
# Recalculation scenarios (formulas package)
# --------------------------------------------------------------------------

def answers_for(domain):
    if domain == 'statistical_results':
        return {'H': 'confirmed_issue', 'I': 'substantive'}
    if domain in ('report_internal_inconsistency', 'plan_material'):
        return {'K': 'minor_presentation'}
    if domain == 'protocol_report':
        return {'L': 'version_timing'}
    raise ValueError(domain)


def calc(path):
    import formulas
    model = formulas.ExcelModel().loads(str(path)).finish()
    sol = model.calculate()
    out = {}
    fname = path.name.upper()
    for k, v in sol.items():
        m = re.match(r"^'\[(.+?)\](.+?)'!([A-Z]+\d+)$", k)
        if not m or m.group(1).upper() != fname:
            continue
        val = v.value[0][0] if hasattr(v, 'value') else v
        out[(m.group(2).upper(), m.group(3))] = val
    return out


def check5_recalc(scratch: Path, catalog_json):
    import numpy as np  # noqa: F401  (formulas dependency)
    papers = sorted({c['paper'] for c in catalog_json['candidates']})
    total = catalog_json['total_candidates']
    src = wb_path('R_A')
    rows = [(f'C_{p}', r) for p in papers for r in range(2, sum(c['paper'] == p for c in catalog_json['candidates']) + 2)]
    ok = True

    def gvals(res):
        return [res.get((s.upper(), f'G{r}')) for s, r in rows]

    def ck(res, cell):
        v = res.get(('CHECKS', cell))
        return v.item() if hasattr(v, 'item') else v

    # (a) blank
    pa = scratch / 'scenario_a_blank.xlsx'
    pa.write_bytes(src.read_bytes())
    ra = calc(pa)
    g = gvals(ra)
    ok &= check(5, '(a) blank file: every G = CHECK REVIEWER ID and Checks!B7 = ACTION NEEDED',
                all(x == 'CHECK REVIEWER ID' for x in g) and len(g) == total and ck(ra, 'B7') == 'ACTION NEEDED',
                {'G_values': sorted(set(map(str, g))), 'B7': ck(ra, 'B7'), 'B4': ck(ra, 'B4')})
    hl_problems = []
    wb0 = load_workbook(src)
    for s, r in rows:
        dom = wb0[s][f'E{r}'].value
        for col in 'HIJKL':
            active = (col in 'HIJ' and dom == 'statistical_results') or \
                     (col == 'K' and dom in ('report_internal_inconsistency', 'plan_material')) or \
                     (col == 'L' and dom == 'protocol_report')
            v = ra.get((s.upper(), f'{col}{r}'))
            v = '' if v is None or (not isinstance(v, str) and str(v) in ('', 'empty')) else str(v)
            if v != ('' if active else 'not_applicable'):
                hl_problems.append(f'{s}!{col}{r}={v!r}')
    ok &= check(5, '(a) blank file: H-L defaults are blank for the active branch and not_applicable otherwise',
                not hl_problems, hl_problems[:10])

    # (b) fully filled
    def filled():
        wb = load_workbook(src)
        wb['Setup']['B2'] = 'TEST_REVIEWER'
        for s, r in rows:
            ws = wb[s]
            ws[f'D{r}'] = 'retain'
            for col, val in answers_for(ws[f'E{r}'].value).items():
                ws[f'{col}{r}'] = val
            ws[f'F{r}'] = 'primary_done'
        return wb
    pb = scratch / 'scenario_b_filled.xlsx'
    filled().save(pb)
    rb = calc(pb)
    g = gvals(rb)
    ok &= check(5, '(b) reviewer ID + every row validly filled + primary_done: every G PRIMARY READY, B4 = B3, B7 PRIMARY READY',
                all(x == 'PRIMARY READY' for x in g) and ck(rb, 'B4') == ck(rb, 'B3') == total
                and ck(rb, 'B7') == 'PRIMARY READY',
                {'G_values': sorted(set(map(str, g))), 'B3': ck(rb, 'B3'), 'B4': ck(rb, 'B4'), 'B7': ck(rb, 'B7')})

    # (c) one invalid inactive value
    wb = filled()
    s, r = rows[0]
    dom = wb[s][f'E{r}'].value
    col, val = ('K', 'minor_presentation') if dom == 'statistical_results' else ('H', 'confirmed_issue')
    wb[s][f'{col}{r}'] = val
    pc = scratch / 'scenario_c_inactive.xlsx'
    wb.save(pc)
    rc = calc(pc)
    gc = rc.get((s.upper(), f'G{r}'))
    ok &= check(5, '(c) one row with an invalid inactive value: that row CLEAR INACTIVE VALUES; others PRIMARY READY; B7 ACTION NEEDED',
                gc == 'CLEAR INACTIVE VALUES' and ck(rc, 'B7') == 'ACTION NEEDED' and ck(rc, 'B4') == total - 1
                and sum(x == 'PRIMARY READY' for x in gvals(rc)) == total - 1,
                {'cell': f'{s}!{col}{r}={val} (domain {dom})', 'G': gc, 'B4': ck(rc, 'B4'), 'B7': ck(rc, 'B7')})

    # (d) one discuss row
    wb = filled()
    s, r = rows[-1]
    orig = load_workbook(src)[s]
    wb[s][f'D{r}'] = 'discuss'
    for col in 'HIJKL':
        wb[s][f'{col}{r}'] = orig[f'{col}{r}'].value
    pd = scratch / 'scenario_d_discuss.xlsx'
    wb.save(pd)
    rd = calc(pd)
    gd = rd.get((s.upper(), f'G{r}'))
    ok &= check(5, '(d) one discuss row: counted in B5 and still PRIMARY READY; B4 = B3; B7 PRIMARY READY',
                gd == 'PRIMARY READY' and ck(rd, 'B5') == 1 and ck(rd, 'B4') == total and ck(rd, 'B7') == 'PRIMARY READY',
                {'row': f'{s}!{r}', 'G': gd, 'B5': ck(rd, 'B5'), 'B4': ck(rd, 'B4'), 'B7': ck(rd, 'B7')})
    return ok


def check6_sources(catalog_json, papers):
    hashes = json.loads((HERE / 'output_hashes.json').read_text())
    now = B.source_hashes(papers)
    rec = {x['paper']: x['report_sha256'] for x in catalog_json['papers']}
    oh = {k.split('/')[1]: v for k, v in hashes['source_reports'].items()}
    ok = check(6, 'Source reports unchanged: current sha256 = catalog sha256 = output_hashes (recorded before and after build)',
               now == rec == oh and hashes.get('source_hashes_unchanged_during_build') is True,
               {p: now[p][:12] for p in papers})
    tpl_rel, tpl_sha = next(iter(hashes['template'].items()))
    ok &= check(6, 'Template workbook unchanged', B.sha256(B.TEMPLATE) == tpl_sha)
    wbh = {n: B.sha256(HERE / n) for n in hashes['workbooks']}
    ok &= check(6, 'Workbook hashes match output_hashes.json', wbh == hashes['workbooks'], wbh)
    return ok


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--scratch-dir', required=True)
    args = ap.parse_args()
    scratch = Path(args.scratch_dir).resolve()
    if scratch == HERE or HERE in scratch.parents:
        ap.error('scratch dir must be outside the year folder')
    scratch.mkdir(parents=True, exist_ok=True)

    papers = B.discover_papers()
    src_before = B.source_hashes(papers)
    catalog_json = json.loads((HERE / f'candidates_{YEAR}.json').read_text(encoding='utf-8'))
    tpl = load_workbook(B.TEMPLATE)
    wbs = {role: load_workbook(wb_path(role)) for role in B.ROLES}

    check1_counts(catalog_json, papers, wbs)
    for role, wb in wbs.items():
        check2_content(catalog_json, wb, role)
        check3_formulas_layout(catalog_json, wb, tpl, role)
    check_fidelity(tpl)
    check4_sheets_codebook(catalog_json, wbs, tpl)
    check5_recalc(scratch, catalog_json)
    check6_sources(catalog_json, papers)
    check(6, 'Source reports unchanged during validation', B.source_hashes(papers) == src_before)

    failed = [r for r in RESULTS if not r['passed']]
    report = {
        'version': 'v.2.4.1', 'year': YEAR, 'status': 'PASS' if not failed else 'FAIL',
        'papers': len(papers), 'candidates': catalog_json['total_candidates'],
        'passed': len(RESULTS) - len(failed), 'failed': len(failed),
        'recalculation_engine': 'formulas 1.3.4 (Python); not Microsoft Excel',
        'desktop_excel_smoke_test': 'PENDING',
        'checks': RESULTS,
    }
    (HERE / 'validation_report.json').write_text(json.dumps(report, indent=2, default=str) + '\n')
    print(report['status'], report['passed'], 'passed', report['failed'], 'failed')
    sys.exit(0 if not failed else 1)


if __name__ == '__main__':
    main()
