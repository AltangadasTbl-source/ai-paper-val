#!/usr/bin/env python3
"""Validate the v.2.4.1 2024 workbooks against ../BUILD_SPEC.md "Required checks".

Structural checks use openpyxl; recalculation checks use the `formulas` package on
copies written to a scratch directory (never into this folder). Writes validation_report.json.

Run: PYTHONDONTWRITEBYTECODE=1 <repo>/human_adjudication/human_adjudication_v_2_3_1/2019/.venv/bin/python \
     validate_v_2_4_1.py --scratch <scratch dir outside the repo>
"""
import argparse
import hashlib
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

YEAR = 2024
HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
YEAR_DIR = REPO / str(YEAR)
REPORT_REL = Path(".ai_paper_validation") / "final_report_1_6_2.md"
TEMPLATE = REPO / "human_adjudication/human_adjudication_v_2_3_1/2018/release/human_adjudication_v_2_3_1_2018_R_A.xlsx"
TPL_PATTERN_SHEET = "C_jama.2017.21906"
PLACEHOLDER = None  # jama.2024.12829 placeholder retired 2026-10-11 (1.6.2 report now exists)
PLACEHOLDER_SHEET_NOTE = "Placeholder: no workflow 1.6.2 output for this paper yet; nothing to review in v.2.4.1."
PLACEHOLDER_PAPERS_NOTE = "Placeholder: workflow 1.6.2 output pending; no candidates in this version"
DOMAINS = {"statistical_results", "report_internal_inconsistency", "plan_material", "protocol_report"}
WB_NAMES = {r: f"human_adjudication_v_2_4_1_{YEAR}_{r}.xlsx" for r in ("R_A", "R_B")}
ALLOWED_FILES = {"build_v_2_4_1.py", "validate_v_2_4_1.py", f"candidates_{YEAR}.json", "validation_report.json",
                 "output_hashes.json", "VALIDATION.md", *WB_NAMES.values()}

RESULTS = []


def check(group, name, ok, detail=None):
    RESULTS.append({"group": group, "check": name, "status": "pass" if ok else "fail", "detail": detail})
    print(("PASS " if ok else "FAIL ") + f"[{group}] {name}" + (f" :: {detail}" if (detail and not ok) else ""))
    return ok


def sha256(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def style_sig(c):
    f, fl, a, b, p = c.font, c.fill, c.alignment, c.border, c.protection
    return (f.name, f.sz, f.b, f.i, f.color.rgb if f.color is not None else None,
            fl.fill_type, fl.fgColor.rgb if fl.fill_type else None,
            a.wrap_text, a.vertical, a.horizontal,
            tuple(getattr(getattr(b, s), "style", None) for s in ("left", "right", "top", "bottom")),
            c.number_format, p.locked)


def tpl_formula(f, row, last, prefix):
    f = re.sub(r"(?<![A-Za-z$])([A-L])2(?![0-9])", lambda m: f"{m.group(1)}{row}", f)
    f = f.replace("$A$2:$A$7", f"$A$2:$A${last}")
    return f.replace('"2018::jama.2017.21906::"', '"' + prefix + '"')


def dv_sig(dv):
    return (dv.type, dv.formula1, dv.formula2, dv.allow_blank, dv.showErrorMessage, dv.showInputMessage,
            dv.errorTitle, dv.error, dv.promptTitle, dv.prompt, dv.errorStyle, dv.operator, dv.showDropDown)


def cf_list(ws):
    out = []
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for r in rules:
            dxf = r.dxf
            out.append((str(rng.sqref), r.type, tuple(r.formula), r.priority, r.stopIfTrue,
                        repr(dxf.fill.fgColor.rgb if dxf and dxf.fill is not None else None),
                        repr(dxf.font.color.rgb if dxf and dxf.font is not None and dxf.font.color is not None else None)))
    return out


# ---------------------------------------------------------------- report re-parse (independent, minimal)
def report_counts(path):
    lines = path.read_text(encoding="utf-8").split("\n")

    def table(heading):
        i = lines.index(f"## {heading}") + 1
        while not lines[i].startswith("|"):
            i += 1
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        return rows[2:]

    idx_ids = [re.match(r"^\[?\s*(C\d{3})", r[0]).group(1) for r in table("Candidate Index")]
    total = [r for r in table("Review-Domain Summary") if re.sub(r"[*`]", "", r[0]).strip() == "Total"]
    total = int(re.sub(r"[*`]", "", total[0][-1]).strip())
    heads = [(m.group(1), m.group(2)) for m in (re.match(r"^## (C\d{3}) — (.+?)(?:\s*\{#[^}]*\})?\s*$", l) for l in lines) if m]
    return idx_ids, total, heads


# ---------------------------------------------------------------- structural
def structural(paths, tpl, catalog):
    wbs = {r: openpyxl.load_workbook(p) for r, p in paths.items()}
    papers = sorted(d.name for d in YEAR_DIR.iterdir() if d.is_dir() and d.name.startswith("jama."))
    expected_sheets = ["Setup", "Checks", "Papers"] + [f"C_{p}" for p in papers] + ["Codebook"]
    tpl_c = tpl[TPL_PATTERN_SHEET]
    by_paper = {}
    for c in catalog["candidates"]:
        by_paper.setdefault(c["paper"], []).append(c)

    # template self-test of the substitution model
    bad = []
    for ws in tpl.worksheets:
        if not ws.title.startswith("C_"):
            continue
        last = ws.max_row
        prefix = f"2018::{ws.title[2:]}::"
        for r in range(2, last + 1):
            for col in "GHIJKL":
                if ws[f"{col}{r}"].value != tpl_formula(tpl_c[f"{col}2"].value, r, last, prefix):
                    bad.append(f"{ws.title}!{col}{r}")
    check("3", "substitution model reproduces every G-L formula of the 2018 template", not bad, bad[:5])

    for role, wb in wbs.items():
        g = f"{role}"
        check("4", f"{g} sheet list exactly as specified", wb.sheetnames == expected_sheets, wb.sheetnames)
        # Setup
        s = wb["Setup"]
        check("4", f"{g} Setup B2/B3 blank, B4 role, B5 integer year",
              s["B2"].value is None and s["B3"].value is None and s["B4"].value == role
              and type(s["B5"].value) is int and s["B5"].value == YEAR)
        ok = all(s.cell(r, c).value == tpl["Setup"].cell(r, c).value for r in range(1, 6) for c in (1,)) \
            and s.max_row == tpl["Setup"].max_row and s.max_column == tpl["Setup"].max_column
        check("4", f"{g} Setup labels and extent as template", ok)
        # Papers
        p = wb["Papers"]
        rows = list(p.iter_rows(min_row=1, values_only=True))
        if PLACEHOLDER:
            exp = [("Paper", "Review worksheet", "Candidates", "Note")]
            for pp in papers:
                n = 0 if pp == PLACEHOLDER else len(by_paper.get(pp, []))
                exp.append((pp, f"C_{pp}", n, PLACEHOLDER_PAPERS_NOTE if pp == PLACEHOLDER else None))
        else:
            exp = [("Paper", "Review worksheet", "Candidates")]
            exp += [(pp, f"C_{pp}", len(by_paper.get(pp, []))) for pp in papers]
        check("4", f"{g} Papers rows (Candidates integer; Note only on placeholder row)",
              rows == exp and all(type(p.cell(r, 3).value) is int for r in range(2, len(papers) + 2)), rows)
        # Checks
        ck = wb["Checks"]
        real = [pp for pp in papers if pp != PLACEHOLDER]
        total = sum(len(by_paper[pp]) for pp in real)
        b4 = "=" + "+".join(f"COUNTIF('C_{pp}'!$G$2:$G${len(by_paper[pp]) + 1},\"PRIMARY READY\")" for pp in real)
        b5 = "=" + "+".join(f"COUNTIF('C_{pp}'!$D$2:$D${len(by_paper[pp]) + 1},\"discuss\")" for pp in real)
        check("4", f"{g} Checks B3 integer total = {total}", type(ck["B3"].value) is int and ck["B3"].value == total)
        check("4", f"{g} Checks B4/B5 COUNTIF sums over exact non-placeholder ranges", ck["B4"].value == b4 and ck["B5"].value == b5)
        others = [(r, c) for r in range(1, tpl["Checks"].max_row + 1) for c in (1, 2) if (r, c) not in ((3, 2), (4, 2), (5, 2))]
        check("4", f"{g} Checks other cells identical to template",
              all(ck.cell(r, c).value == tpl["Checks"].cell(r, c).value for r, c in others)
              and ck.max_row == tpl["Checks"].max_row and ck.max_column == tpl["Checks"].max_column)
        allf = " ".join(str(c.value) for row in ck.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        if PLACEHOLDER:
            check("3", f"{g} placeholder sheet absent from Checks formulas", f"C_{PLACEHOLDER}" not in allf)
        # Codebook
        cb, tcb = wb["Codebook"], tpl["Codebook"]
        fixes = {"Statistcal_or_numerical issue": "Statistical_or_numerical issue",
                 "6  |  Candidate context": "5  |  Candidate context"}
        diffs, sdiffs, nfix = [], [], 0
        for r in range(1, max(cb.max_row, tcb.max_row) + 1):
            for c in range(1, max(cb.max_column, tcb.max_column) + 1):
                tv, v = tcb.cell(r, c).value, cb.cell(r, c).value
                if tv in fixes and v == fixes[tv]:
                    nfix += 1
                elif tv != v:
                    diffs.append(cb.cell(r, c).coordinate)
                if style_sig(tcb.cell(r, c)) != style_sig(cb.cell(r, c)):
                    sdiffs.append(cb.cell(r, c).coordinate)
        dims_ok = all(cb.column_dimensions[k].width == v.width for k, v in tcb.column_dimensions.items()) and \
            all(cb.row_dimensions[k].height == v.height for k, v in tcb.row_dimensions.items()) and cb.freeze_panes == tcb.freeze_panes
        check("4", f"{g} Codebook identical to template except the two fixes", not diffs and not sdiffs and nfix == 2 and dims_ok,
              {"value_diffs": diffs, "style_diffs": sdiffs[:5], "fixes_applied": nfix, "dims_ok": dims_ok})

        # C_ sheets
        keys_all = []
        for pp in papers:
            ws = wb[f"C_{pp}"]
            tag = f"{g} C_{pp}"
            hdr = [ws.cell(1, c).value for c in range(1, 19)]
            thdr = [tpl_c.cell(1, c).value for c in range(1, 19)]
            thdr[7] = "Statistical_or_numerical issue"
            check("3", f"{tag} header row as template with H1 fixed", hdr == thdr, hdr)
            hid = all(ws.column_dimensions[c].hidden and ws.column_dimensions[c].width == 3 for c in "MNOPQ")
            widths = all(ws.column_dimensions[k].width == v.width and bool(ws.column_dimensions[k].hidden) == bool(v.hidden)
                         for k, v in tpl_c.column_dimensions.items())
            check("3", f"{tag} M-Q hidden width 3; column widths as template; freeze B2",
                  hid and widths and ws.freeze_panes == "B2")
            check("3", f"{tag} header styles as template",
                  all(style_sig(ws.cell(1, c)) == style_sig(tpl_c.cell(1, c)) for c in range(1, 19))
                  and ws.row_dimensions[1].height == tpl_c.row_dimensions[1].height)
            if pp == PLACEHOLDER:
                cells = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
                exp_cells = {f"{openpyxl.utils.get_column_letter(c)}1": v for c, v in enumerate(thdr, 1) if v is not None}
                exp_cells["A3"] = PLACEHOLDER_SHEET_NOTE
                nform = sum(1 for v in cells.values() if isinstance(v, str) and v.startswith("="))
                check("3", f"{tag} placeholder: header + A3 note only; no formulas, DVs or CFs",
                      cells == exp_cells and nform == 0 and not ws.data_validations.dataValidation
                      and not cf_list(ws), {"cells": len(cells), "formulas": nform,
                                            "dv": len(ws.data_validations.dataValidation), "cf": len(cf_list(ws))})
                continue
            cands = by_paper[pp]
            n = len(cands)
            last = n + 1
            idx_ids, total, heads = report_counts(YEAR_DIR / pp / REPORT_REL)
            check("1", f"{tag} rows = index rows = cards = domain-summary total",
                  ws.max_row - 1 == n == len(idx_ids) == len(heads) == total,
                  {"rows": ws.max_row - 1, "catalog": n, "index": len(idx_ids), "cards": len(heads), "summary_total": total})
            keys = [ws[f"A{r}"].value for r in range(2, last + 1)]
            keys_all += keys
            ids = [k.split("::")[-1] if isinstance(k, str) else None for k in keys]
            check("2", f"{tag} keys '<year>::<paper>::Cnnn', sequential, match report IDs",
                  keys == [f"{YEAR}::{pp}::C{i:03d}" for i in range(1, n + 1)] and ids == idx_ids == [h[0] for h in heads])
            Es = [ws[f"E{r}"].value for r in range(2, last + 1)]
            check("2", f"{tag} E values valid and equal catalog domains",
                  all(e in DOMAINS for e in Es) and Es == [c["domain"] for c in cands])
            check("2", f"{tag} D blank, F 'no', M-R blank",
                  all(ws[f"D{r}"].value is None and ws[f"F{r}"].value == "no" and
                      all(ws[f"{c}{r}"].value is None for c in "MNOPQR") for r in range(2, last + 1)))
            Bs = [ws[f"B{r}"].value for r in range(2, last + 1)]
            Cs = [ws[f"C{r}"].value for r in range(2, last + 1)]
            residue = [x for x in Bs + Cs if not isinstance(x, str) or not x.strip() or any(t in x for t in ("](", "**", "<a id", "`"))
                       or re.search("pending human adjudication", x, re.I) or x != re.sub(r"\s+", " ", x).strip()]
            check("2", f"{tag} B and C non-empty, no markdown residue or boilerplate", not residue, residue[:3])
            heading_titles = [re.sub(r"[*`]", "", h[1]).strip() for h in heads]
            check("2", f"{tag} B equals card heading titles; C equals catalog locations",
                  Bs == heading_titles and Bs == [c["title"] for c in cands] and Cs == [c["source_locations"] for c in cands])
            fbad = [f"{col}{r}" for r in range(2, last + 1) for col in "GHIJKL"
                    if ws[f"{col}{r}"].value != tpl_formula(tpl_c[f"{col}2"].value, r, last, f"{YEAR}::{pp}::")]
            check("3", f"{tag} G-L formulas equal template pattern", not fbad, fbad[:5])
            dvs = ws.data_validations.dataValidation
            tdv = {str(d.sqref)[0]: dv_sig(d) for d in tpl_c.data_validations.dataValidation}
            got = {str(d.sqref): dv_sig(d) for d in dvs}
            expdv = {f"{c}2:{c}{last}": s for c, s in tdv.items()}
            check("3", f"{tag} DV sqrefs exactly rows 2..{last} with template settings", got == expdv, sorted(got))
            tcf = [(re.sub(r"\d+$", str(last), x[0]),) + x[1:] for x in cf_list(tpl_c)]
            check("3", f"{tag} CF sqrefs exactly rows 2..{last} with template rules", sorted(cf_list(ws)) == sorted(tcf))
            sty = all(style_sig(ws.cell(r, c)) == style_sig(tpl_c.cell(2 if r % 2 == 0 else 3, c)) for r in range(2, last + 1) for c in range(1, 19)) \
                and all(ws.row_dimensions[r].height == tpl_c.row_dimensions[2].height for r in range(2, last + 1))
            check("3", f"{tag} data-row styles (incl. wrap) and heights as template", sty)
            check("3", f"{tag} autofilter A1:R{last}; no cells beyond row {last} or column R",
                  ws.auto_filter.ref == f"A1:R{last}" and ws.max_row == last and ws.max_column == 18)
        check("2", f"{g} keys unique across workbook", len(keys_all) == len(set(keys_all)), len(keys_all))
        check("1", f"{g} total rows = {catalog['total_candidates']}", len(keys_all) == catalog["total_candidates"])

    # R_A vs R_B
    a, b = wbs["R_A"], wbs["R_B"]
    diffs = []
    for ws in a.worksheets:
        wb2 = b[ws.title]
        for r in range(1, max(ws.max_row, wb2.max_row) + 1):
            for c in range(1, max(ws.max_column, wb2.max_column) + 1):
                if ws.cell(r, c).value != wb2.cell(r, c).value:
                    diffs.append(f"{ws.title}!{ws.cell(r, c).coordinate}")
    check("4", "R_A vs R_B differ only in Setup!B4 (all cell values)", diffs == ["Setup!B4"], diffs[:5])
    return wbs


# ---------------------------------------------------------------- recalculation
def recalc(path, scratch):
    import formulas
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    out = {}
    fname = path.name.lower()
    for k, v in sol.items():
        m = re.match(r"^'\[(.+?)\](.+?)'!([A-Z]+\d+)$", k)
        if m and m.group(1).lower() == fname:
            val = v.value
            try:
                val = val[0][0]
            except Exception:
                pass
            out[(m.group(2).upper(), m.group(3))] = val
    return out


def scenarios(wb_path, catalog, scratch):
    by_paper = {}
    for c in catalog["candidates"]:
        by_paper.setdefault(c["paper"], []).append(c)
    rows = [(f"C_{p}", i + 2, c["domain"]) for p, cs in by_paper.items() for i, c in enumerate(cs)]
    n = len(rows)

    def G(vals, sheet, r):
        return vals.get((sheet.upper(), f"G{r}"))

    def ck(vals, cell):
        return vals.get(("CHECKS", cell))

    def fill(wb):
        wb["Setup"]["B2"] = "TEST_REVIEWER"
        for sheet, r, dom in rows:
            ws = wb[sheet]
            ws[f"D{r}"] = "retain"
            ws[f"F{r}"] = "primary_done"
            if dom == "statistical_results":
                ws[f"H{r}"] = "confirmed_issue"
                ws[f"I{r}"] = "substantive"
            elif dom in ("report_internal_inconsistency", "plan_material"):
                ws[f"K{r}"] = "minor_presentation"
            else:
                ws[f"L{r}"] = "substantive_difference"

    out = {}
    # (a) blank
    p = scratch / "a_blank.xlsx"
    shutil.copy(wb_path, p)
    v = recalc(p, scratch)
    gs = [G(v, s, r) for s, r, _ in rows]
    out["a"] = check("5", "(a) blank file: every G = CHECK REVIEWER ID and Checks!B7 = ACTION NEEDED",
                     all(x == "CHECK REVIEWER ID" for x in gs) and len(gs) == n and ck(v, "B7") == "ACTION NEEDED",
                     {"distinct_G": sorted(set(map(str, gs))), "B7": ck(v, "B7")})
    # (b) all filled
    wb = openpyxl.load_workbook(wb_path)
    fill(wb)
    p = scratch / "b_filled.xlsx"
    wb.save(p)
    v = recalc(p, scratch)
    gs = [G(v, s, r) for s, r, _ in rows]
    out["b"] = check("5", "(b) reviewer ID + valid answers + primary_done: all G PRIMARY READY, B4 = B3, B7 PRIMARY READY",
                     all(x == "PRIMARY READY" for x in gs) and ck(v, "B4") == ck(v, "B3") == n and ck(v, "B7") == "PRIMARY READY"
                     and ck(v, "B5") == 0,
                     {"distinct_G": sorted(set(map(str, gs))), "B3": ck(v, "B3"), "B4": ck(v, "B4"), "B5": ck(v, "B5"), "B7": ck(v, "B7")})
    # (c) invalid inactive value
    stat = next(x for x in rows if x[2] == "statistical_results")
    wb = openpyxl.load_workbook(wb_path)
    fill(wb)
    wb[stat[0]][f"K{stat[1]}"] = "minor_presentation"
    p = scratch / "c_inactive.xlsx"
    wb.save(p)
    v = recalc(p, scratch)
    others = [G(v, s, r) for s, r, _ in rows if (s, r) != stat[:2]]
    out["c"] = check("5", f"(c) invalid inactive value ({stat[0]}!K{stat[1]} on a statistical_results row): CLEAR INACTIVE VALUES",
                     G(v, *stat[:2]) == "CLEAR INACTIVE VALUES" and all(x == "PRIMARY READY" for x in others)
                     and ck(v, "B4") == n - 1 and ck(v, "B7") == "ACTION NEEDED",
                     {"G": G(v, *stat[:2]), "B4": ck(v, "B4"), "B7": ck(v, "B7")})
    # (d) discuss row
    tgt = rows[-1]
    wb = openpyxl.load_workbook(wb_path)
    fill(wb)
    ws = wb[tgt[0]]
    orig = openpyxl.load_workbook(wb_path)[tgt[0]]
    ws[f"D{tgt[1]}"] = "discuss"
    for col in "HIJKL":
        ws[f"{col}{tgt[1]}"] = orig[f"{col}{tgt[1]}"].value   # restore template formulas
    p = scratch / "d_discuss.xlsx"
    wb.save(p)
    v = recalc(p, scratch)
    gs = [G(v, s, r) for s, r, _ in rows]
    out["d"] = check("5", f"(d) one discuss row ({tgt[0]} row {tgt[1]}): counted in B5, still PRIMARY READY",
                     G(v, *tgt[:2]) == "PRIMARY READY" and ck(v, "B5") == 1 and all(x == "PRIMARY READY" for x in gs)
                     and ck(v, "B4") == n and ck(v, "B7") == "PRIMARY READY",
                     {"G": G(v, *tgt[:2]), "B5": ck(v, "B5"), "B4": ck(v, "B4"), "B7": ck(v, "B7")})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scratch", help="scratch directory for recalculation copies (outside the repo)")
    args = ap.parse_args()
    scratch = Path(args.scratch) if args.scratch else Path(tempfile.mkdtemp(prefix="v241_2024_"))
    scratch.mkdir(parents=True, exist_ok=True)
    if REPO in scratch.resolve().parents:
        sys.exit("scratch directory must be outside the repository")

    catalog = json.loads((HERE / f"candidates_{YEAR}.json").read_text(encoding="utf-8"))
    hashes = json.loads((HERE / "output_hashes.json").read_text(encoding="utf-8"))
    paths = {r: HERE / n for r, n in WB_NAMES.items()}

    # 6. sources unchanged
    cur = {}
    for d in sorted(YEAR_DIR.iterdir()):
        rp = d / REPORT_REL
        if d.is_dir() and rp.exists():
            cur[rp.relative_to(REPO).as_posix()] = sha256(rp)
    check("6", "source reports unchanged (current sha256 = sha256 recorded before/after build)",
          cur == hashes["source_reports"] and all(c["report_sha256"] == cur[c["report_path"]] for c in catalog["candidates"]),
          {"reports": len(cur)})
    check("6", "template unchanged", {k: sha256(REPO / k) for k in hashes["template"]} == hashes["template"])
    check("6", "workbook sha256 match output_hashes.json", {p.name: sha256(p) for p in paths.values()} == hashes["workbooks"])
    if PLACEHOLDER:
        check("6", f"{PLACEHOLDER} has no 1.6.2 report (placeholder justified)", not (YEAR_DIR / PLACEHOLDER / REPORT_REL).exists())

    tpl = openpyxl.load_workbook(TEMPLATE)
    structural(paths, tpl, catalog)

    work = scratch / "recalc"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir()
    for role in ("R_A", "R_B"):
        d = work / role
        d.mkdir()
        if role == "R_A":
            scenarios(paths[role], catalog, d)
        else:
            p = d / "blank.xlsx"
            shutil.copy(paths[role], p)
            v = recalc(p, d)
            check("5", "R_B blank recalculation: Checks!B2 = R_B and B7 = ACTION NEEDED",
                  v.get(("CHECKS", "B2")) == "R_B" and v.get(("CHECKS", "B7")) == "ACTION NEEDED")

    extra = sorted(p.name for p in HERE.iterdir() if p.name not in ALLOWED_FILES)
    check("hygiene", "year folder contains only the specified deliverables", not extra, extra)

    failed = [r for r in RESULTS if r["status"] != "pass"]
    report = {
        "version": "v.2.4.1", "year": YEAR, "checks": RESULTS, "passed": len(RESULTS) - len(failed), "failed": len(failed),
        "recalculation_engine": "Python 'formulas' package on scratch copies; not Microsoft Excel.",
        "pending": ["Desktop Excel smoke test (open, dropdowns, conditional formatting rendering, recalculation, save/reopen): PENDING"],
        "workbooks": {p.name: sha256(p) for p in paths.values()},
    }
    (HERE / "validation_report.json").write_text(json.dumps(report, indent=2, default=str) + "\n", encoding="utf-8")
    print(f"{report['passed']} passed, {report['failed']} failed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
