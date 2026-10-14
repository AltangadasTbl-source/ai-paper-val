#!/usr/bin/env python3
"""Apply one reviewer's judgments JSON to a copy of a v.2.4.1 blank workbook.

Construction helper for the R_A and R_Bulun reviews. It never modifies the
blank workbooks in ../<year>/; it writes a new copy and refuses to overwrite one.
R_Bulun is filled from the R_B blank workbook; Setup!B4 is set to R_Bulun.

Usage:
  python fill_review.py --year 2019 --role R_Bulun --judgments J.json --check
  python fill_review.py --year 2019 --role R_A --judgments J.json --out OUT.xlsx

Judgments JSON:
{
  "reviewer_id": "R_Bulun", "reviewer_name": null,
  "rows": [{"key": "2019::jama.2019.10517::C001", "D": "retain",
            "E": "statistical_results", "H": "confirmed_issue",
            "I": "substantive", "J": null, "K": null, "L": null,
            "R": "short English note", "evidence": "...", "rationale": "..."}]
}
H-L: give a code only for fields active under D/E; use null otherwise.
Inactive and unanswered-optional cells keep the template's default formula.
"""
import argparse
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent  # human_adjudication_v_2_4_1
TEMPLATE_ROLE = {"R_A": "R_A", "R_Bulun": "R_B"}
KEY_RE = re.compile(r"^\d{4}::[^:]+::C\d{3}$")

D_OPT = {"retain", "exclude", "discuss"}
E_OPT = {"statistical_results", "report_internal_inconsistency", "plan_material", "protocol_report"}
OPTS = {
    "H": {"confirmed_issue", "probable_issue", "not_issue", "indeterminate"},
    "I": {"substantive", "minor_presentation", "indeterminate"},
    "J": {"direct", "possible", "unlikely", "unknown"},
    "K": {"substantive_inconsistency", "minor_presentation", "non_inconsistency", "indeterminate"},
    "L": {"substantive_difference", "minor_presentation", "version_timing", "non_inconsistency", "indeterminate"},
}
# (required, optional) by domain when D=retain; everything else inactive
RULES = {
    "statistical_results": ({"H", "I"}, {"J"}),
    "report_internal_inconsistency": ({"K"}, set()),
    "plan_material": ({"K"}, set()),
    "protocol_report": (set(), {"L"}),
}
COL = {"D": 4, "E": 5, "F": 6, "H": 8, "I": 9, "J": 10, "K": 11, "L": 12, "R": 18}


def template_path(year, role):
    return ROOT / year / f"human_adjudication_v_2_4_1_{year}_{TEMPLATE_ROLE[role]}.xlsx"


def workbook_keys(wb):
    keys = {}
    for s in wb.sheetnames:
        if not s.startswith("C_"):
            continue
        ws = wb[s]
        for r in range(2, ws.max_row + 1):
            k = ws.cell(r, 1).value
            if k and KEY_RE.match(str(k)):  # skips placeholder-sheet notes
                if k in keys:
                    raise SystemExit(f"duplicate key in template: {k}")
                keys[k] = (s, r)
    return keys


def check_rows(data, keys):
    errs = []
    if not str(data.get("reviewer_id") or "").strip():
        errs.append("reviewer_id missing")
    seen = set()
    for i, row in enumerate(data.get("rows", [])):
        k = row.get("key")
        tag = f"row {i} ({k})"
        if k not in keys:
            errs.append(f"{tag}: key not in workbook")
            continue
        if k in seen:
            errs.append(f"{tag}: duplicate key")
        seen.add(k)
        d, e = row.get("D"), row.get("E")
        if d not in D_OPT:
            errs.append(f"{tag}: invalid D {d!r}")
        if e not in E_OPT:
            errs.append(f"{tag}: invalid E {e!r}")
            continue
        req, opt = RULES[e] if d == "retain" else (set(), set())
        for f in "HIJKL":
            v = row.get(f)
            if f in req:
                if v not in OPTS[f]:
                    errs.append(f"{tag}: required {f} missing/invalid ({v!r})")
            elif f in opt:
                if v is not None and v not in OPTS[f]:
                    errs.append(f"{tag}: optional {f} invalid ({v!r})")
            elif v is not None:
                errs.append(f"{tag}: {f} is inactive for D={d}, E={e}; must be null (got {v!r})")
        if not str(row.get("rationale") or "").strip():
            errs.append(f"{tag}: rationale missing")
    missing = sorted(set(keys) - seen)
    return errs, missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", required=True, choices=["2018", "2019", "2024", "2025"])
    ap.add_argument("--role", required=True, choices=["R_A", "R_Bulun"])
    ap.add_argument("--judgments", required=True)
    ap.add_argument("--out")
    ap.add_argument("--check", action="store_true", help="validate only; allow missing keys")
    a = ap.parse_args()

    tpl = template_path(a.year, a.role)
    data = json.loads(Path(a.judgments).read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(tpl)
    keys = workbook_keys(wb)
    errs, missing = check_rows(data, keys)
    print(f"template: {tpl.name}; keys={len(keys)}; judged={len(data.get('rows', []))}; missing={len(missing)}")
    for m in errs:
        print("ERROR", m)
    if a.check:
        for m in missing:
            print("MISSING", m)
        sys.exit(1 if errs else 0)
    if errs or missing:
        for m in missing:
            print("MISSING", m)
        sys.exit("refusing to write: fix errors / missing keys first")
    if not a.out:
        sys.exit("--out required")
    out = Path(a.out).resolve()
    if out.exists():
        sys.exit(f"refusing to overwrite {out}")
    if out.parent == tpl.parent:
        sys.exit("refusing to write next to the blank workbooks")

    setup = wb["Setup"]
    setup["B2"] = str(data["reviewer_id"]).strip()
    setup["B4"] = a.role
    if data.get("reviewer_name"):
        setup["B3"] = str(data["reviewer_name"]).strip()
    for row in data["rows"]:
        s, r = keys[row["key"]]
        ws = wb[s]
        d, e = row["D"], row["E"]
        ws.cell(r, COL["D"]).value = d
        ws.cell(r, COL["E"]).value = e
        req, opt = RULES[e] if d == "retain" else (set(), set())
        for f in "HIJKL":
            v = row.get(f)
            if v is not None and f in req | opt:
                ws.cell(r, COL[f]).value = v
            # otherwise keep template default formula (blank or not_applicable)
        note = str(row.get("R") or "").strip()
        if note:
            ws.cell(r, COL["R"]).value = note
        ws.cell(r, COL["F"]).value = "primary_done"
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    shutil.move(tmp, out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
