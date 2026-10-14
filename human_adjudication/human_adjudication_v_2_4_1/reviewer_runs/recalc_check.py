#!/usr/bin/env python3
"""Recalculate a completed workbook with the `formulas` interpreter and report
column G per candidate plus the Checks sheet. Read-only; construction aid only.

Usage (needs a venv with `formulas`, e.g. ../../human_adjudication_v_2_3_1/2019/.venv):
  python recalc_check.py COMPLETED.xlsx
"""
import collections
import sys
from pathlib import Path

import formulas
import openpyxl


def main(path):
    path = Path(path).resolve()
    wb = openpyxl.load_workbook(path)
    m = formulas.ExcelModel().loads(str(path)).finish()
    sol = m.calculate()
    book = path.name

    def val(sheet, cell):
        ref = f"'[{book}]{sheet.upper()}'!{cell}"
        v = sol.get(ref)
        if v is None:
            return None
        v = v.value
        try:
            return v[0][0]
        except Exception:
            return v

    g = collections.Counter()
    bad = []
    for s in wb.sheetnames:
        if not s.startswith("C_"):
            continue
        ws = wb[s]
        for r in range(2, ws.max_row + 1):
            k = ws.cell(r, 1).value
            if not k or "::C" not in str(k):
                continue
            v = val(s, f"G{r}")
            g[str(v)] += 1
            if v != "PRIMARY READY":
                bad.append((k, v))
    print("G:", dict(g))
    for k, v in bad:
        print("NOT READY", k, v)
    ws = wb["Checks"]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        print("Checks", r, a, "->", val("Checks", f"B{r}"))


if __name__ == "__main__":
    main(sys.argv[1])
