#!/usr/bin/env python3
"""v.2.4.1 primary-review workbook builder for one year (follows ../BUILD_SPEC.md).

Reads every <repo>/<year>/<paper>/.ai_paper_validation/final_report_1_6_2.md (read-only),
writes candidates_<year>.json, human_adjudication_v_2_4_1_<year>_R_A.xlsx / _R_B.xlsx and
output_hashes.json into this folder. The v.2.3.1 2018 R_A release workbook is the structural
template. Parsing fails loudly on anything unexpected rather than guessing.

Run: PYTHONDONTWRITEBYTECODE=1 <repo>/human_adjudication/human_adjudication_v_2_3_1/2019/.venv/bin/python build_v_2_4_1.py
"""
import hashlib, io, json, os, re, sys
from copy import copy, deepcopy
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.cell_range import MultiCellRange

YEAR = 2018
HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
TEMPLATE = ROOT / 'human_adjudication/human_adjudication_v_2_3_1/2018/release/human_adjudication_v_2_3_1_2018_R_A.xlsx'
REPORT_REL = '.ai_paper_validation/final_report_1_6_2.md'
ROLES = ['R_A', 'R_B']
DOMAINS = ['statistical_results', 'report_internal_inconsistency', 'plan_material', 'protocol_report']
RELATIONS = ['P_R', 'P_P', 'S_C', 'S_W']
OPTIONS = {'D': ['retain', 'exclude', 'discuss'], 'E': DOMAINS, 'F': ['no', 'primary_done'],
           'H': ['confirmed_issue', 'probable_issue', 'not_issue', 'indeterminate'],
           'I': ['substantive', 'minor_presentation', 'indeterminate'],
           'J': ['direct', 'possible', 'unlikely', 'unknown'],
           'K': ['substantive_inconsistency', 'minor_presentation', 'non_inconsistency', 'indeterminate'],
           'L': ['substantive_difference', 'minor_presentation', 'version_timing', 'non_inconsistency', 'indeterminate']}
H1_FIXED = 'Statistical_or_numerical issue'
CODEBOOK_FIXES = {'Statistcal_or_numerical issue': 'Statistical_or_numerical issue',
                  '6  |  Candidate context': '5  |  Candidate context'}
BOILERPLATE = 'Pending Human Adjudication'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
ANOMALIES = []          # parse notes reported in candidates json (non-fatal, source-text observations)


class ParseError(Exception):
    pass


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def fail(msg):
    raise ParseError(msg)


# ---------------------------------------------------------------- report parsing
LINK = re.compile(r'\[([^\]]*)\]\((<[^>]*>|[^)\s]*)\)')
FIELD = re.compile(r'^\*\*([^*]+?):\*\*\s*(.*)$')
ANCHOR_TAG = re.compile(r'<a\s+id="[^"]*"\s*>\s*</a>|<a\s+id="[^"]*"\s*/>', re.I)
BRACE_ANCHOR = re.compile(r'\s*\{#[^}]*\}\s*$')
CARD_HEAD = re.compile(r'^(C(\d{3}))\s+—\s+(.+)$')


def strip_boiler(s):
    s = s.replace(BOILERPLATE, '')
    return s


def clean_text(s, where, add_filenames=False):
    """Markdown -> plain text per spec. Fails on residue."""
    def repl(m):
        text, target = m.group(1).strip(), m.group(2).strip()
        if target.startswith('<') and target.endswith('>'):
            target = target[1:-1]
        path = target.split('#', 1)[0]
        fname = unquote(os.path.basename(path)) if path else ''
        if add_filenames:
            if not fname:
                ANOMALIES.append(f'{where}: link "{text}" -> "{target}" has no file name in target; text kept as is')
            elif fname not in text:
                stem = os.path.splitext(fname)[0]
                if stem and stem in text:
                    ANOMALIES.append(f'{where}: link text "{text}" contains stem but not full file name {fname}; prefixed')
                text = f'{fname} — {text}'
        return text
    out = LINK.sub(repl, s)
    out = ANCHOR_TAG.sub('', out)
    out = BRACE_ANCHOR.sub('', out)
    out = out.replace('**', '').replace('`', '')
    out = strip_boiler(out)
    out = re.sub(r'\s+', ' ', out).strip()
    for bad in ['](', '**', '<a id', '{#', 'http://', 'https://', '`']:
        if bad in out:
            fail(f'{where}: markdown/URL residue {bad!r} in {out[:200]!r}')
    if re.search(r'\[[^\]]*\]\(', out) or '<' in out and re.search(r'<[a-zA-Z/][^>]*>', out):
        fail(f'{where}: residual link/HTML in {out[:200]!r}')
    if not out:
        fail(f'{where}: empty after cleaning')
    return out


def section(lines, title_regex):
    """Return the lines of the first heading section matching title_regex (any level)."""
    rx = re.compile(title_regex)
    for i, l in enumerate(lines):
        if re.match(r'^#{1,6}\s', l) and rx.search(l):
            level = len(l) - len(l.lstrip('#'))
            out = []
            for l2 in lines[i + 1:]:
                m = re.match(r'^(#{1,6})\s', l2)
                if m and len(m.group(1)) <= level:
                    break
                out.append(l2)
            return out
    fail(f'section {title_regex!r} not found')


def tables(sec):
    """Markdown tables in a section: list of (header, rows)."""
    res, cur = [], []
    for l in sec + ['']:
        if l.startswith('|'):
            cur.append([c.strip() for c in l.strip().strip('|').split('|')])
        elif cur:
            if len(cur) < 2 or not all(re.fullmatch(r':?-{3,}:?', c) for c in cur[1]):
                fail(f'malformed table starting {cur[0]}')
            res.append((cur[0], cur[2:]))
            cur = []
    return res


def unmd(c):
    return c.replace('**', '').replace('`', '').strip()


def parse_index(lines, where):
    sec = section(lines, r'^#{1,6}\s+Candidate Index\s*$')
    tbs = tables(sec)
    if len(tbs) != 1:
        fail(f'{where}: Candidate Index has {len(tbs)} tables, expected 1')
    head, rows = tbs[0]
    h = [unmd(x).lower() for x in head]
    try:
        di, ri = h.index('review domain'), h.index('relationship type')
    except ValueError:
        fail(f'{where}: Candidate Index header lacks Review domain/Relationship type: {head}')
    out = []
    for r in rows:
        if len(r) != len(head):
            fail(f'{where}: index row width {len(r)} != {len(head)}: {r}')
        m = re.match(r'^\[?\s*(C\d{3})\b', unmd(r[0]))
        if not m:
            fail(f'{where}: index row first cell lacks Cnnn: {r[0]!r}')
        out.append({'id': m.group(1), 'domain': unmd(r[di]), 'relationship': unmd(r[ri])})
    return out


def parse_domain_summary(lines, where):
    sec = section(lines, r'^#{1,6}\s+Review-Domain Summary\s*$')
    cands = [t for t in tables(sec) if unmd(t[0][0]).lower() == 'review domain']
    if len(cands) != 1:
        fail(f'{where}: expected exactly one Review domain table in Review-Domain Summary, found {len(cands)}')
    head, rows = cands[0]
    counts, total = {}, None
    for r in rows:
        k, v = unmd(r[0]), unmd(r[-1])
        if not re.fullmatch(r'\d+', v):
            fail(f'{where}: non-integer count in domain summary row {r}')
        if k in DOMAINS:
            if k in counts:
                fail(f'{where}: duplicate domain row {k}')
            counts[k] = int(v)
        elif k.lower() == 'total':
            total = int(v)
        else:
            fail(f'{where}: unexpected domain summary row {r}')
    if set(counts) != set(DOMAINS) or total is None:
        fail(f'{where}: domain summary incomplete: {counts}, total={total}')
    if sum(counts.values()) != total:
        fail(f'{where}: domain summary rows sum {sum(counts.values())} != Total {total}')
    return counts, total


def flatten_field(lines, where):
    """Paragraphs joined by a space; consecutive list items joined by '; '. Tables/quotes fail."""
    groups, para = [], []
    def flush():
        if para:
            groups.append(('p', ' '.join(para))); para.clear()
    for l in lines:
        s = l.strip()
        if not s:
            flush(); continue
        if s.startswith('|') or s.startswith('>') or s.startswith('```'):
            fail(f'{where}: unsupported table/quote/code block in field: {s[:120]!r}')
        m = re.match(r'^(?:[-*+]|\d+[.)])\s+(.*)$', s)
        if m:
            flush(); groups.append(('li', m.group(1).strip()))
            ANOMALIES.append(f'{where}: list item flattened') if not any(a.startswith(f'{where}: list') for a in ANOMALIES) else None
        else:
            para.append(s)
    flush()
    out = ''
    for i, (kind, text) in enumerate(groups):
        if i == 0:
            out = text
        elif kind == 'li' and groups[i - 1][0] == 'li':
            out += '; ' + text
        else:
            out += ' ' + text
    return out.strip()


def parse_cards(lines, where):
    heads = []
    for i, l in enumerate(lines):
        m = re.match(r'^(#{1,6})\s+(.*)$', l)
        if not m:
            continue
        txt = BRACE_ANCHOR.sub('', ANCHOR_TAG.sub('', m.group(2))).strip()
        if re.match(r'^C\d', txt):
            if m.group(1) != '##':
                fail(f'{where}: card-like heading at wrong level line {i+1}: {l!r}')
            cm = CARD_HEAD.match(txt)
            if not cm:
                fail(f'{where}: unrecognized card heading line {i+1}: {l!r}')
            heads.append((i, cm))
    if ANCHOR_TAG.search(''.join(lines)):
        ANOMALIES.append(f'{where}: report contains <a id> anchor tags (removed from text)')
    cards = []
    for n, (i, cm) in enumerate(heads):
        body = []
        for l2 in lines[i + 1:]:
            if re.match(r'^#{1,2}\s', l2):
                break
            body.append(l2)
        # A field runs from its **Label:** line to the next **Label:** line (it may span
        # several paragraphs and bullet lists, e.g. jama.2018.0948 / 14280 / 14282).
        fields, cur = {}, None
        for l2 in body:
            fm = FIELD.match(l2)
            if fm:
                cur = fm.group(1).strip()
                fields.setdefault(cur, []).append([fm.group(2)])
            elif cur:
                fields[cur][-1].append(l2)

        def get(label_rx, name):
            keys = [k for k in fields if re.fullmatch(label_rx, k)]
            vals = [v for k in keys for v in fields[k]]
            if len(vals) != 1:
                fail(f'{where} {cm.group(1)}: field {name} found {len(vals)} times (labels {keys})')
            v = flatten_field(vals[0], f'{where} {cm.group(1)} {name}')
            if not v:
                fail(f'{where} {cm.group(1)}: field {name} empty')
            return v
        cid = cm.group(1)
        domain = unmd(get(r'Review domain', 'Review domain'))
        rel = unmd(get(r'Relationship type', 'Relationship type'))
        loc_raw = get(r'Exact source locations?', 'Exact source locations')
        stmt_raw = get(r'Candidate statement', 'Candidate statement')
        if domain not in DOMAINS:
            fail(f'{where} {cid}: invalid domain {domain!r}')
        if rel not in RELATIONS:
            fail(f'{where} {cid}: invalid relationship type {rel!r}')
        if not LINK.search(loc_raw):
            ANOMALIES.append(f'{where} {cid} Exact source locations: no file/page links in the field; text kept verbatim')
        if re.match(r'^(Yes|No)\b', loc_raw):
            ANOMALIES.append(f'{where} {cid} Exact source locations: field text begins with "{loc_raw.split()[0].rstrip(",:")}" (recheck-answer wording in source); kept verbatim')
        title = clean_text(cm.group(3), f'{where} {cid} title')
        cards.append({'id': cid, 'title': title, 'title_raw': cm.group(3),
                      'statement': clean_text(stmt_raw, f'{where} {cid} statement'),
                      'domain': domain, 'relationship_type': rel,
                      'source_locations': clean_text(loc_raw, f'{where} {cid} locations', add_filenames=True),
                      'source_locations_markdown': loc_raw})
    return cards


def parse_report(paper, path):
    where = paper
    lines = path.read_text(encoding='utf-8').splitlines()
    index = parse_index(lines, where)
    counts, total = parse_domain_summary(lines, where)
    cards = parse_cards(lines, where)
    ids = [c['id'] for c in cards]
    expected = [f'C{i:03d}' for i in range(1, len(cards) + 1)]
    if ids != expected:
        fail(f'{where}: card IDs not sequential from C001: {ids}')
    if [r['id'] for r in index] != ids:
        fail(f'{where}: Candidate Index IDs {[r["id"] for r in index]} != card IDs {ids}')
    if not (len(index) == len(cards) == total):
        fail(f'{where}: index rows {len(index)}, cards {len(cards)}, domain-summary total {total} disagree')
    for r, c in zip(index, cards):
        if r['domain'] != c['domain'] or r['relationship'] != c['relationship_type']:
            fail(f'{where} {c["id"]}: index ({r["domain"]},{r["relationship"]}) != card ({c["domain"]},{c["relationship_type"]})')
    card_counts = {d: sum(c['domain'] == d for c in cards) for d in DOMAINS}
    if card_counts != counts:
        fail(f'{where}: card domain counts {card_counts} != summary {counts}')
    return {'paper': paper, 'index_rows': len(index), 'cards': len(cards), 'summary_total': total,
            'summary_counts': counts, 'candidates': cards}


def collect():
    year_dir = ROOT / str(YEAR)
    papers = sorted(p.name for p in year_dir.iterdir() if p.is_dir() and not p.name.startswith('.'))
    out = []
    for p in papers:
        rp = year_dir / p / REPORT_REL
        if not rp.is_file():
            fail(f'missing report {rp}')
        out.append((p, rp))
    return out


# ---------------------------------------------------------------- formulas
def q(v):
    return '"' + str(v).replace('"', '""') + '"'


def valid(c, opts):
    return 'OR(' + ','.join(f'EXACT({c},{q(v)})' for v in opts) + ')'


def branch(col, e):
    if col in 'HIJ':
        return f'{e}="statistical_results"'
    if col == 'K':
        return f'OR({e}="report_internal_inconsistency",{e}="plan_material")'
    return f'{e}="protocol_report"'


def g_formula(r, last, prefix):
    a, d, e, f = (f'{c}{r}' for c in 'ADEF')
    act = {col: f'AND({d}="retain",{branch(col, e)})' for col in 'HIJKL'}
    stale = [f'AND(NOT({act[col]}),{col}{r}<>"",NOT(EXACT({col}{r},"not_applicable")))' for col in 'HIJKL']
    bad = [f'NOT({valid(d, OPTIONS["D"])})', f'NOT({valid(e, OPTIONS["E"])})']
    for col in 'HIJKL':
        opt = f'{col}{r}<>"",' if col in 'JL' else ''
        bad.append(f'AND({act[col]},{opt}NOT({valid(col + str(r), OPTIONS[col])}))')
    identity = 'LEN(TRIM(CLEAN(SUBSTITUTE(Setup!$B$2,CHAR(160)," "))))=0'
    key = f'OR({a}="",COUNTIF($A$2:$A${last},{a})<>1,LEFT({a},LEN({q(prefix)}))<>{q(prefix)})'
    return (f'=IFERROR(IF({identity},"CHECK REVIEWER ID",IF({key},"CHECK KEY",'
            f'IF(OR({",".join(stale)}),"CLEAR INACTIVE VALUES",IF(OR({",".join(bad)}),"CHECK FIELDS",'
            f'IF(AND({f}<>"",NOT({valid(f, OPTIONS["F"])})),"CHECK COMPLETE",'
            f'IF(EXACT({f},"primary_done"),"PRIMARY READY","PRIMARY DRAFT")))))),"CHECK INPUT / KEY")')


def hl_formula(col, r):
    return f'=IF(AND(OR(D{r}="retain",D{r}=""),{branch(col, f"E{r}")}),"","not_applicable")'


def sqn(s):
    return "'" + s.replace("'", "''") + "'"


def checks_formula(sheets, col, value):
    return '=' + '+'.join(f'COUNTIF({sqn(t)}!${col}$2:${col}${last},{q(value)})' for t, last in sheets)


def self_test_generators(tw):
    """Generators must reproduce every template G-L formula exactly."""
    n = 0
    for s in tw:
        if not s.title.startswith('C_'):
            continue
        last = s.max_row
        prefix = f'2018::{s.title[2:]}::'
        for r in range(2, last + 1):
            assert s[f'G{r}'].value == g_formula(r, last, prefix), (s.title, r)
            for col in 'HIJKL':
                assert s[f'{col}{r}'].value == hl_formula(col, r), (s.title, col, r)
            n += 1
    return n


# ---------------------------------------------------------------- workbook
def build_workbook(role, catalog, papers):
    w = load_workbook(TEMPLATE)
    tsheets = [s for s in w if s.title.startswith('C_')]
    proto = tsheets[0]
    total = sum(len(v) for v in catalog.values())

    setup = w['Setup']
    setup['B2'] = None; setup['B3'] = None; setup['B4'] = role; setup['B5'] = int(YEAR)

    # Papers
    ps = w['Papers']
    pstyle = {r: [copy(ps.cell(r, c)._style) for c in range(1, 4)] for r in (2, 3)}
    prow = {r: (ps.row_dimensions[r].height, copy(ps.row_dimensions[r]._style)) for r in (2, 3)}
    for r in range(2, ps.max_row + 1):
        for c in range(1, ps.max_column + 1):
            ps.cell(r, c).value = None; ps.cell(r, c).hyperlink = None
    for r in list(ps._cells):
        if r[0] >= 2:
            del ps._cells[r]
    for r in list(ps.row_dimensions):
        if r >= 2:
            del ps.row_dimensions[r]
    for r, p in enumerate(papers, 2):
        par = 2 if r % 2 == 0 else 3
        for c, v in enumerate([p, 'C_' + p, len(catalog[p])], 1):
            cell = ps.cell(r, c); cell.value = v; cell._style = copy(pstyle[par][c - 1])
        ps[f'B{r}'].hyperlink = Hyperlink(ref=f'B{r}', location=f"{sqn('C_' + p)}!A1")
        ps.row_dimensions[r].height = prow[par][0]; ps.row_dimensions[r]._style = copy(prow[par][1])

    # C_ sheets built from the first template C_ sheet as prototype
    new = []
    for i, p in enumerate(papers):
        ns = w.copy_worksheet(proto)
        ns.title = f'tmp{i}'
        for k in list(ns._cells):
            if k[0] >= 2:
                del ns._cells[k]
        for k in list(ns.row_dimensions):
            if k >= 2:
                del ns.row_dimensions[k]
        ns.views = deepcopy(proto.views)
        ns['H1'] = H1_FIXED
        rows = catalog[p]
        last = len(rows) + 1
        prefix = f'{YEAR}::{p}::'
        for r, c in enumerate(rows, 2):
            par = 2 if r % 2 == 0 else 3
            for col in range(1, 19):
                ns.cell(r, col)._style = copy(proto.cell(par, col)._style)
            ns.row_dimensions[r].height = proto.row_dimensions[par].height
            ns.row_dimensions[r]._style = copy(proto.row_dimensions[par]._style)
            ns[f'A{r}'] = c['key']; ns[f'B{r}'] = c['title']; ns[f'C{r}'] = c['source_locations']
            ns[f'E{r}'] = c['domain']; ns[f'F{r}'] = 'no'
            ns[f'G{r}'] = g_formula(r, last, prefix)
            for col in 'HIJKL':
                ns[f'{col}{r}'] = hl_formula(col, r)
        ns.auto_filter.ref = f'A1:R{last}'
        for dv in proto.data_validations.dataValidation:
            nd = deepcopy(dv)
            col = str(dv.sqref)[0]
            assert str(dv.sqref) == f'{col}2:{col}{proto.max_row}', dv.sqref
            nd.sqref = MultiCellRange(f'{col}2:{col}{last}')
            ns.data_validations.append(nd)
        for cf in proto.conditional_formatting:
            ref = str(cf.sqref); col = ref[0]
            assert ref == f'{col}2:{col}{proto.max_row}', ref
            for rule in cf.rules:
                ns.conditional_formatting.add(f'{col}2:{col}{last}', deepcopy(rule))
        new.append((ns, 'C_' + p))
    for s in tsheets:
        w.remove(s)
    for ns, title in new:
        ns.title = title
    order = [w['Setup'], w['Checks'], w['Papers']] + [ns for ns, _ in new] + [w['Codebook']]
    assert set(order) == set(w._sheets) and len(order) == len(w._sheets)
    w._sheets = order
    w.active = 0

    # Checks
    ck = w['Checks']
    ranges = [('C_' + p, len(catalog[p]) + 1) for p in papers if catalog[p]]
    ck['B3'] = int(total)
    ck['B4'] = checks_formula(ranges, 'G', 'PRIMARY READY')
    ck['B5'] = checks_formula(ranges, 'D', 'discuss')

    # Codebook: two fixes only
    cb = w['Codebook']
    for old, newv in CODEBOOK_FIXES.items():
        hits = [c for row in cb for c in row if c.value == old]
        assert len(hits) == 1, (old, len(hits))
        hits[0].value = newv
    return w


def save_unlocked(w, dest):
    buf = io.BytesIO(); w.save(buf); buf.seek(0)
    out = io.BytesIO()
    with ZipFile(buf) as src, ZipFile(out, 'w', ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == 'xl/styles.xml':
                root = ET.fromstring(content)
                for kind in ['cellXfs', 'cellStyleXfs']:
                    for xf in root.find(NS + kind):
                        xf.set('applyProtection', '1')
                        for old in list(xf.findall(NS + 'protection')):
                            xf.remove(old)
                        ET.SubElement(xf, NS + 'protection', {'locked': '0', 'hidden': '0'})
                content = ET.tostring(root, encoding='utf-8', xml_declaration=False)
            dst.writestr(item, content)
    dest.write_bytes(out.getvalue())


def main():
    ET.register_namespace('', NS[1:-1])
    reports = collect()
    before = {str(rp.relative_to(ROOT)): sha(rp) for _, rp in reports}
    tw = load_workbook(TEMPLATE)
    ntested = self_test_generators(tw)
    parsed, catalog, flat = [], {}, []
    for p, rp in reports:
        res = parse_report(p, rp)
        parsed.append(res)
        rel = str(rp.relative_to(ROOT))
        catalog[p] = []
        for c in res['candidates']:
            rec = {'year': YEAR, 'paper': p, 'id': c['id'], 'key': f'{YEAR}::{p}::{c["id"]}',
                   'title': c['title'], 'statement': c['statement'], 'domain': c['domain'],
                   'relationship_type': c['relationship_type'], 'source_locations': c['source_locations'],
                   'source_locations_markdown': c['source_locations_markdown'],
                   'report_path': rel, 'report_sha256': before[rel]}
            catalog[p].append(rec); flat.append(rec)
    papers = [p for p, _ in reports]
    assert len({r['key'] for r in flat}) == len(flat)
    files = []
    for role in ROLES:
        w = build_workbook(role, catalog, papers)
        dest = HERE / f'human_adjudication_v_2_4_1_{YEAR}_{role}.xlsx'
        save_unlocked(w, dest); files.append(dest)
    after = {str(rp.relative_to(ROOT)): sha(rp) for _, rp in reports}
    if before != after:
        fail('source report changed during build')
    doc = {'spec': 'human_adjudication_v_2_4_1/BUILD_SPEC.md', 'year': YEAR,
           'template': str(TEMPLATE.relative_to(ROOT)), 'template_sha256': sha(TEMPLATE),
           'generator_self_test_template_rows': ntested,
           'papers': [{'paper': r['paper'], 'candidates': r['cards'], 'index_rows': r['index_rows'],
                       'domain_summary_total': r['summary_total'], 'domain_counts': r['summary_counts']} for r in parsed],
           'total_candidates': len(flat), 'parse_notes': ANOMALIES, 'candidates': flat}
    (HERE / f'candidates_{YEAR}.json').write_text(json.dumps(doc, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    hashes = {'workbooks': {f.name: sha(f) for f in files}, 'source_reports': after}
    (HERE / 'output_hashes.json').write_text(json.dumps(hashes, indent=2) + '\n')
    for r in parsed:
        print(r['paper'], r['cards'], r['summary_counts'])
    print('total', len(flat)); print('parse notes', len(ANOMALIES))
    for f in files:
        print(f)


if __name__ == '__main__':
    try:
        main()
    except ParseError as e:
        sys.exit(f'BUILD FAILED: {e}')
