#!/usr/bin/env python3
"""v.2.4.1 structural + recalculation checks (BUILD_SPEC.md "Required checks" 1-6).

Never modifies the delivered workbooks: recalculation scenarios run on copies written to
--scratch, which must be outside this year folder. Writes validation_report.json and
VALIDATION.md into this folder.

Run: PYTHONDONTWRITEBYTECODE=1 <venv python> validate_v_2_4_1.py --scratch <dir outside repo year folder>
"""
import argparse, json, re, sys, warnings
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

sys.dont_write_bytecode = True
import build_v_2_4_1 as B  # noqa: E402  (parsing helpers, constants)

YEAR, HERE, ROOT = B.YEAR, B.HERE, B.ROOT
NSM = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RESULTS = []          # (check, label, status, detail)


def record(check, label, ok, detail=''):
    RESULTS.append({'check': check, 'label': label, 'status': 'PASS' if ok else 'FAIL', 'detail': detail})
    if not ok:
        print(f'FAIL [{check}] {label}: {detail}')
    return ok


def wb_path(role):
    return HERE / f'human_adjudication_v_2_4_1_{YEAR}_{role}.xlsx'


def style_sig(c):
    return (c.font.name, c.font.size, c.font.bold, c.font.color.rgb if c.font.color else None,
            c.fill.fill_type, c.fill.fgColor.rgb, c.alignment.wrap_text, c.alignment.vertical,
            c.alignment.horizontal, c.number_format, c.protection.locked, c.protection.hidden)


def dv_sig(dv):
    return (dv.type, dv.formula1, dv.formula2, dv.allow_blank, dv.showDropDown, dv.showErrorMessage,
            dv.showInputMessage, dv.errorTitle, dv.error, dv.errorStyle, dv.operator, dv.promptTitle, dv.prompt)


def cf_list(ws):
    out = []
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            d = r.dxf
            out.append((str(cf.sqref), r.type, tuple(r.formula), r.priority, r.stopIfTrue,
                        d.fill.fgColor.rgb if d and d.fill else None,
                        d.fill.bgColor.rgb if d and d.fill else None,
                        d.font.color.rgb if d and d.font and d.font.color else None))
    return sorted(out, key=lambda x: x[3])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--scratch', required=True, help='directory for recalculation copies (outside this folder)')
    args = ap.parse_args()
    scratch = Path(args.scratch).resolve()
    if scratch == HERE or HERE in scratch.parents:
        sys.exit('scratch must be outside the year folder')
    scratch.mkdir(parents=True, exist_ok=True)

    cat = json.loads((HERE / f'candidates_{YEAR}.json').read_text(encoding='utf-8'))
    hashes = json.loads((HERE / 'output_hashes.json').read_text())
    cands = cat['candidates']
    reports = B.collect()
    papers = [p for p, _ in reports]
    by_paper = {p: [c for c in cands if c['paper'] == p] for p in papers}
    tw = load_workbook(B.TEMPLATE)
    tproto = [s for s in tw if s.title.startswith('C_')][0]
    tlast = tproto.max_row
    tprefix = f'2018::{tproto.title[2:]}::'
    wbs = {role: load_workbook(wb_path(role)) for role in B.ROLES}
    w = wbs['R_A']

    # ---- 6 (part 1). source hashes: recorded at build == now
    now = {str(rp.relative_to(ROOT)): B.sha(rp) for _, rp in reports}
    record(6, 'source report sha256 recorded at build (output_hashes.json) equals current', now == hashes['source_reports'])
    record(6, 'candidates json report_sha256 equals current', all(c['report_sha256'] == now[c['report_path']] for c in cands))
    record(6, 'workbook sha256 equals output_hashes.json',
           all(B.sha(wb_path(r)) == hashes['workbooks'][wb_path(r).name] for r in B.ROLES))

    # ---- 1. per-paper counts
    per_paper = []
    for p, rp in reports:
        lines = rp.read_text(encoding='utf-8').splitlines()
        raw_heads = sum(1 for l in lines if re.match(r'^##\s+(<a\s+id="[^"]*"\s*>\s*</a>\s*)?C\d{3}\s', l))
        idx = B.parse_index(lines, p)
        counts, total = B.parse_domain_summary(lines, p)
        ws = w['C_' + p]
        rows = ws.max_row - 1
        prow = next(r for r in range(2, w['Papers'].max_row + 1) if w['Papers'][f'A{r}'].value == p)
        pc = w['Papers'][f'C{prow}'].value
        wb_dom = {d: sum(ws[f'E{r}'].value == d for r in range(2, ws.max_row + 1)) for d in B.DOMAINS}
        ok = rows == len(idx) == raw_heads == total == len(by_paper[p]) == pc and wb_dom == counts and isinstance(pc, int)
        record(1, f'{p}: rows {rows} = index {len(idx)} = cards {raw_heads} = domain total {total} = Papers {pc}; domain counts match',
               ok, f'workbook domains {wb_dom} vs summary {counts}')
        per_paper.append({'paper': p, 'rows': rows, 'index_rows': len(idx), 'cards': raw_heads,
                          'domain_summary_total': total, 'domain_counts': counts})
    total_all = sum(x['rows'] for x in per_paper)

    # ---- 2. keys / values
    all_keys = []
    for p in papers:
        ws = w['C_' + p]
        bad = []
        for i, r in enumerate(range(2, ws.max_row + 1), 1):
            c = by_paper[p][i - 1]
            key = f'{YEAR}::{p}::C{i:03d}'
            a, b_, cc, d, e, f = (ws[f'{x}{r}'].value for x in 'ABCDEF')
            all_keys.append(a)
            if a != key or c['key'] != key: bad.append(f'{r}: key {a}')
            if e not in B.DOMAINS or e != c['domain']: bad.append(f'{r}: E {e}')
            if d is not None: bad.append(f'{r}: D not blank')
            if f != 'no': bad.append(f'{r}: F {f}')
            for name, v, ref in [('B', b_, c['title']), ('C', cc, c['source_locations'])]:
                if not isinstance(v, str) or not v.strip(): bad.append(f'{r}: {name} empty')
                elif any(t in v for t in ['](', '**', '<a id', '{#', '`', B.BOILERPLATE]): bad.append(f'{r}: {name} residue')
                elif v != ref: bad.append(f'{r}: {name} != catalog')
            if ws[f'R{r}'].value is not None: bad.append(f'{r}: R not blank')
            if any(ws[f'{x}{r}'].value is not None for x in 'MNOPQ'): bad.append(f'{r}: M-Q not blank')
        record(2, f'{p}: sequential keys, valid E, D blank, F no, B/C non-empty without markdown residue, R and M-Q blank', not bad, '; '.join(bad[:5]))
    record(2, f'keys unique across workbook ({len(all_keys)})', len(set(all_keys)) == len(all_keys) == len(cands) == cat['total_candidates'])
    boiler = [(s.title, c.coordinate) for s in w for row in s for c in row if isinstance(c.value, str) and B.BOILERPLATE in c.value]
    record(2, 'no "Pending Human Adjudication" text anywhere in the workbook', not boiler, str(boiler[:3]))

    # ---- 3. formulas, DV, CF, layout vs template
    tG = tproto['G2'].value; tHL = {col: tproto[f'{col}2'].value for col in 'HIJKL'}
    # sanity: translator reproduces template's own row 3
    record(3, 'template-substitution method reproduces template row 3 formulas',
           Translator(tG, origin='G2').translate_formula('G3') == tproto['G3'].value and
           all(Translator(tHL[c], origin=f'{c}2').translate_formula(f'{c}3') == tproto[f'{c}3'].value for c in 'HIJKL'))
    tdv = {str(dv.sqref)[0]: dv_sig(dv) for dv in tproto.data_validations.dataValidation}
    tcf = cf_list(tproto)
    for p in papers:
        ws = w['C_' + p]; last = ws.max_row; prefix = f'{YEAR}::{p}::'
        bad = []
        for r in range(2, last + 1):
            exp = Translator(tG, origin='G2').translate_formula(f'G{r}')
            exp = exp.replace(f'$A$2:$A${tlast}', f'$A$2:$A${last}').replace(B.q(tprefix), B.q(prefix))
            if ws[f'G{r}'].value != exp: bad.append(f'G{r}')
            for col in 'HIJKL':
                if ws[f'{col}{r}'].value != Translator(tHL[col], origin=f'{col}2').translate_formula(f'{col}{r}'):
                    bad.append(f'{col}{r}')
        record(3, f'{p}: every G-L formula equals template pattern after substitution', not bad, str(bad[:5]))
        dvs = ws.data_validations.dataValidation
        dv_ok = len(dvs) == len(tdv) and all(str(dv.sqref) == f'{str(dv.sqref)[0]}2:{str(dv.sqref)[0]}{last}' and dv_sig(dv) == tdv[str(dv.sqref)[0]] for dv in dvs) \
            and sorted(str(dv.sqref)[0] for dv in dvs) == sorted(tdv)
        exp_cf = [(f'{x[0][0]}2:{x[0][0]}{last}',) + x[1:] for x in tcf]
        record(3, f'{p}: data validations (D,E,F,H-L) and conditional formats identical to template over rows 2..{last}',
               dv_ok and cf_list(ws) == exp_cf, f'dv_ok={dv_ok}')
        hidden = all(ws.column_dimensions[c].hidden and ws.column_dimensions[c].width == 3 for c in 'MNOPQ')
        widths = all(ws.column_dimensions[c].width == tproto.column_dimensions[c].width and
                     bool(ws.column_dimensions[c].hidden) == bool(tproto.column_dimensions[c].hidden) for c in 'ABCDEFGHIJKLMNOPQR')
        head = [ws.cell(1, c).value for c in range(1, 19)]
        thead = [tproto.cell(1, c).value for c in range(1, 19)]; thead[7] = B.H1_FIXED
        styles = all(style_sig(ws.cell(r, c)) == style_sig(tproto.cell(1 if r == 1 else (2 if r % 2 == 0 else 3), c))
                     for r in range(1, last + 1) for c in range(1, 19))
        heights = all(ws.row_dimensions[r].height == tproto.row_dimensions[1 if r == 1 else 2].height for r in range(1, last + 1))
        wrap = all(ws.cell(r, c).alignment.wrap_text for r in range(1, last + 1) for c in range(1, 19))
        record(3, f'{p}: M-Q hidden width 3; freeze B2; widths, headers (H1 fixed), styles, wrap, row heights, autofilter as template',
               hidden and widths and ws.freeze_panes == 'B2' and head == thead and styles and heights and wrap
               and ws.auto_filter.ref == f'A1:R{last}' and ws.max_column == 18 and not ws.sheet_view.showGridLines,
               f'hidden={hidden} widths={widths} head={head == thead} styles={styles} heights={heights} wrap={wrap} af={ws.auto_filter.ref}')

    # ---- 4. sheet list, Setup/Checks/Papers, Codebook, R_A vs R_B
    exp_sheets = ['Setup', 'Checks', 'Papers'] + ['C_' + p for p in sorted(papers)] + ['Codebook']
    for role, wx in wbs.items():
        record(4, f'{role}: sheet list exactly {len(exp_sheets)} sheets in spec order', wx.sheetnames == exp_sheets, str(wx.sheetnames))
        st = wx['Setup']
        record(4, f'{role}: Setup B2/B3 blank, B4={role}, B5={YEAR} (int); other Setup cells as template',
               st['B2'].value is None and st['B3'].value is None and st['B4'].value == role and st['B5'].value == YEAR
               and type(st['B5'].value) is int and all(st.cell(r, c).value == tw['Setup'].cell(r, c).value
               for r in range(1, 6) for c in (1, 2) if (r, c) not in [(4, 2), (5, 2)]) and st.max_row == 5)
    ck, tck = w['Checks'], tw['Checks']
    ranges = [('C_' + p, len(by_paper[p]) + 1) for p in papers if by_paper[p]]
    exp4 = '=' + '+'.join(f"COUNTIF('{t}'!$G$2:$G${l},\"PRIMARY READY\")" for t, l in ranges)
    exp5 = '=' + '+'.join(f"COUNTIF('{t}'!$D$2:$D${l},\"discuss\")" for t, l in ranges)
    other = all(ck.cell(r, c).value == tck.cell(r, c).value for r in range(1, 10) for c in (1, 2) if (r, c) not in [(3, 2), (4, 2), (5, 2)])
    record(4, 'Checks: B3 total (int), B4/B5 COUNTIF sums over every C_ sheet exact row range, other cells as template',
           ck['B3'].value == total_all == len(cands) and type(ck['B3'].value) is int and ck['B4'].value == exp4 and ck['B5'].value == exp5
           and other and ck.max_row == tck.max_row)
    ps = w['Papers']
    pok = [ps.cell(1, c).value for c in range(1, 4)] == ['Paper', 'Review worksheet', 'Candidates'] and ps.max_column == 3 \
        and ps.max_row == len(papers) + 1 and all(
            ps[f'A{r}'].value == p and ps[f'B{r}'].value == 'C_' + p and ps[f'C{r}'].value == len(by_paper[p])
            and type(ps[f'C{r}'].value) is int and ps[f'B{r}'].hyperlink is not None and ps[f'B{r}'].hyperlink.location == f"'C_{p}'!A1"
            for r, p in enumerate(papers, 2))
    record(4, 'Papers: Paper / Review worksheet (internal link) / Candidates (int), one row per paper', pok)
    cb, tcb = w['Codebook'], tw['Codebook']
    diffs = []
    for r in range(1, max(cb.max_row, tcb.max_row) + 1):
        for c in range(1, max(cb.max_column, tcb.max_column) + 1):
            a, t = cb.cell(r, c), tcb.cell(r, c)
            if a.value != t.value:
                diffs.append((a.coordinate, t.value, a.value))
            if style_sig(a) != style_sig(t):
                diffs.append((a.coordinate, 'style'))
    dims = all(cb.column_dimensions[k].width == tcb.column_dimensions[k].width for k in 'ABCDE') and \
        all(cb.row_dimensions[r].height == tcb.row_dimensions[r].height for r in range(1, tcb.max_row + 1))
    fixes_ok = len(diffs) == len(B.CODEBOOK_FIXES) and all(len(x) == 3 for x in diffs) and \
        sorted((o, n) for _, o, n in diffs) == sorted(B.CODEBOOK_FIXES.items())
    record(4, 'Codebook identical to template (values, styles, widths, heights) except the two specified fixes',
           fixes_ok and dims and cb.freeze_panes == tcb.freeze_panes, str(diffs))
    a, b = wbs['R_A'], wbs['R_B']
    vd = []
    for sa in a:
        sb = b[sa.title]
        for r in range(1, max(sa.max_row, sb.max_row) + 1):
            for c in range(1, max(sa.max_column, sb.max_column) + 1):
                if sa.cell(r, c).value != sb.cell(r, c).value:
                    vd.append(f'{sa.title}!{sa.cell(r, c).coordinate}')
    record(4, 'R_A vs R_B: all cell values identical except Setup!B4', vd == ['Setup!B4'], str(vd[:5]))
    # protection / calc (template-inherited properties)
    with ZipFile(wb_path('R_A')) as z, ZipFile(B.TEMPLATE) as tz:
        xfs = [xf for tag in ['cellXfs', 'cellStyleXfs'] for xf in ET.fromstring(z.read('xl/styles.xml')).find('m:' + tag, NSM)]
        unlocked = all(xf.find('m:protection', NSM) is not None and xf.find('m:protection', NSM).get('locked') == '0' for xf in xfs)
        same_styles = z.read('xl/styles.xml') == tz.read('xl/styles.xml')
        prot = any(b'sheetProtection' in z.read(n) for n in z.namelist() if n.startswith('xl/worksheets/')) or b'workbookProtection' in z.read('xl/workbook.xml')
        extra = [n for n in z.namelist() if any(x in n.lower() for x in ['externallink', 'vbaproject', 'connections', 'querytable'])]
    record(4, 'styles.xml byte-identical to template; all cell xfs unlocked; no sheet/workbook protection, external links or macros',
           same_styles and unlocked and not prot and not extra, f'same_styles={same_styles}')
    record(4, 'full automatic calculation on load; Setup active; no defined names other than autofilter',
           w.calculation.calcMode == 'auto' and w.calculation.fullCalcOnLoad and w.active.title == 'Setup' and
           all(n == '_xlnm._FilterDatabase' for s in w for n in s.defined_names) and not list(w.defined_names))

    # ---- 5. recalculation with `formulas` on copies
    import formulas
    warnings.filterwarnings('ignore')

    def calc(wx, name):
        path = scratch / name
        wx.save(path)
        sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
        def get(sheet, cell):
            k = f"'[{name}]{sheet.upper()}'!{cell}"
            v = sol[k].value[0, 0]
            return v
        return get

    def gvals(get, wx):
        return {f'{s.title}!G{r}': get(s.title, f'G{r}') for s in wx if s.title.startswith('C_') for r in range(2, s.max_row + 1)}

    def fill_all(wx):
        wx['Setup']['B2'] = 'TEST_REVIEWER'
        for s in wx:
            if not s.title.startswith('C_'):
                continue
            for r in range(2, s.max_row + 1):
                s[f'D{r}'] = 'retain'; s[f'F{r}'] = 'primary_done'; d = s[f'E{r}'].value
                if d == 'statistical_results':
                    s[f'H{r}'] = 'confirmed_issue'; s[f'I{r}'] = 'substantive'
                elif d in ('report_internal_inconsistency', 'plan_material'):
                    s[f'K{r}'] = 'substantive_inconsistency'
                # protocol_report: no required judgment; L optional left as formula
    for role in B.ROLES:
        wx = load_workbook(wb_path(role)); get = calc(wx, f'TEST_{role}_a_blank.xlsx'); g = gvals(get, wx)
        record(5, f'(a) {role} blank file: all {len(g)} G = CHECK REVIEWER ID; Checks B7 ACTION NEEDED, B4=0, B5=0, B6=B3',
               len(g) == total_all and set(g.values()) == {'CHECK REVIEWER ID'} and get('Checks', 'B7') == 'ACTION NEEDED'
               and get('Checks', 'B4') == 0 and get('Checks', 'B5') == 0 and get('Checks', 'B6') == total_all,
               str(set(g.values())))
    wx = load_workbook(wb_path('R_A')); fill_all(wx); get = calc(wx, 'TEST_R_A_b_filled.xlsx'); g = gvals(get, wx)
    record(5, '(b) reviewer ID + every row valid for its domain + primary_done: all G PRIMARY READY, B4 = B3, B7 PRIMARY READY',
           set(g.values()) == {'PRIMARY READY'} and get('Checks', 'B4') == get('Checks', 'B3') == total_all
           and get('Checks', 'B7') == 'PRIMARY READY' and get('Checks', 'B5') == 0, f"{set(g.values())} B4={get('Checks', 'B4')}")
    # (c) one invalid inactive value
    wx = load_workbook(wb_path('R_A')); fill_all(wx)
    s0 = wx['C_' + papers[0]]
    col = 'K' if s0['E2'].value == 'statistical_results' else 'H'
    s0[f'{col}2'] = 'substantive_inconsistency' if col == 'K' else 'confirmed_issue'
    get = calc(wx, 'TEST_R_A_c_inactive.xlsx'); g = gvals(get, wx)
    others = {k: v for k, v in g.items() if k != f'{s0.title}!G2'}
    record(5, f'(c) one row ({s0.title}!{col}2 inactive value) -> CLEAR INACTIVE VALUES; others ready; B4=B3-1; B7 ACTION NEEDED',
           g[f'{s0.title}!G2'] == 'CLEAR INACTIVE VALUES' and set(others.values()) == {'PRIMARY READY'}
           and get('Checks', 'B4') == total_all - 1 and get('Checks', 'B7') == 'ACTION NEEDED', g[f'{s0.title}!G2'])
    # (c, extra) one inactive value per domain present, each in its own row
    wx = load_workbook(wb_path('R_A')); fill_all(wx); seen = {}
    for s in wx:
        if s.title.startswith('C_'):
            for r in range(2, s.max_row + 1):
                d = s[f'E{r}'].value
                if d not in seen:
                    c2 = 'K' if d == 'statistical_results' else ('L' if d != 'protocol_report' else 'H')
                    s[f'{c2}{r}'] = {'K': 'minor_presentation', 'L': 'version_timing', 'H': 'not_issue'}[c2]
                    seen[d] = f'{s.title}!G{r}'
    get = calc(wx, 'TEST_R_A_c_per_domain.xlsx'); g = gvals(get, wx)
    record(5, f'(c, extra) one inactive literal per domain ({len(seen)} domains) -> CLEAR INACTIVE VALUES on exactly those rows',
           all(g[k] == 'CLEAR INACTIVE VALUES' for k in seen.values()) and
           sum(v == 'PRIMARY READY' for v in g.values()) == total_all - len(seen), str(seen))
    # (d) one discuss row
    wx = load_workbook(wb_path('R_A')); fill_all(wx)
    s1 = wx['C_' + papers[-1]]
    s1['D2'] = 'discuss'
    for c3 in 'HIJKL':
        s1[f'{c3}2'] = B.hl_formula(c3, 2)
    get = calc(wx, 'TEST_R_A_d_discuss.xlsx'); g = gvals(get, wx)
    record(5, f'(d) one discuss row ({s1.title}!D2): G PRIMARY READY, counted in B5 = 1, B4 = B3, B7 PRIMARY READY',
           g[f'{s1.title}!G2'] == 'PRIMARY READY' and get('Checks', 'B5') == 1 and get('Checks', 'B4') == total_all
           and get('Checks', 'B7') == 'PRIMARY READY', f"G={g[f'{s1.title}!G2']} B5={get('Checks', 'B5')}")

    # ---- 6 (part 2). sources unchanged after validation as well
    record(6, 'source report sha256 unchanged after validation run', {str(rp.relative_to(ROOT)): B.sha(rp) for _, rp in reports} == now)

    failed = [r for r in RESULTS if r['status'] == 'FAIL']
    report = {'year': YEAR, 'validated_at_utc': datetime.now(timezone.utc).isoformat(timespec='seconds'),
              'workbooks': {wb_path(r).name: B.sha(wb_path(r)) for r in B.ROLES},
              'template': str(B.TEMPLATE.relative_to(ROOT)), 'template_sha256': B.sha(B.TEMPLATE),
              'recalculation_engine': f'formulas {formulas.__version__} (Python), on copies in a scratch folder outside the repo year folder',
              'per_paper': per_paper, 'total_candidates': total_all,
              'checks': RESULTS, 'passed': len(RESULTS) - len(failed), 'failed': len(failed),
              'pending': ['Desktop Excel open/recalculate/edit/save/reopen smoke test and visual rendering: PENDING (not run; no Excel available).'],
              'parse_notes': cat['parse_notes']}
    (HERE / 'validation_report.json').write_text(json.dumps(report, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    write_md(report)
    print(f"{report['passed']} passed, {report['failed']} failed")
    sys.exit(1 if failed else 0)


def write_md(rep):
    L = [f'# v.2.4.1 {YEAR} validation', '',
         f'Generated by `validate_v_2_4_1.py` at {rep["validated_at_utc"]}. Builder: `build_v_2_4_1.py` (spec: `../BUILD_SPEC.md`).', '',
         '## Inputs', '',
         f'- Candidate source: `{YEAR}/<paper>/.ai_paper_validation/final_report_1_6_2.md` for {len(rep["per_paper"])} papers (sha256 in `output_hashes.json`).',
         f'- Structural template: `{rep["template"]}` (sha256 `{rep["template_sha256"]}`).', '',
         '## Candidates', '', '| Paper | Rows | statistical_results | report_internal_inconsistency | plan_material | protocol_report |', '|---|---:|---:|---:|---:|---:|']
    for p in rep['per_paper']:
        d = p['domain_counts']
        L.append(f"| {p['paper']} | {p['rows']} | {d['statistical_results']} | {d['report_internal_inconsistency']} | {d['plan_material']} | {d['protocol_report']} |")
    L += [f"| Total | {rep['total_candidates']} | | | | |", '',
          '## How the checks were run', '',
          '1. Counts: workbook rows per C_ sheet compared with the report Candidate Index rows, a raw count of `## Cnnn` card headings and the Review-Domain Summary total (and per-domain counts).',
          '2. Row content: keys sequential and unique; E in the four domains; D blank; F `no`; B and C non-empty, equal to the extracted catalog and free of `](`, `**`, `<a id`, `{#`, backticks and the "Pending Human Adjudication" boilerplate.',
          '3. Formulas: every G-L formula compared with the template\'s first C_ sheet row-2 formula translated to the row with openpyxl `Translator`, then the `$A$2:$A$<last>` range and `<year>::<paper>::` prefix substituted (an independent path from the builder\'s generator). Data validations and conditional formats compared attribute by attribute with the template, with sqrefs rows 2..last. Column widths, M-Q hidden, freeze B2, headers (H1 spelling fixed), cell styles by row parity, wrap and row heights compared with the template.',
          '4. Sheet list and order; Setup, Checks (B3 integer; B4/B5 exact COUNTIF strings) and Papers; Codebook cell-by-cell (values and styles) against the template with only the two specified fixes; R_A vs R_B every cell value; styles.xml byte-identical to the template (all xfs unlocked).',
          f'5. Recalculation: copies saved to a scratch folder and evaluated with {rep["recalculation_engine"].split(" (")[0]}. Scenarios (a) blank, (b) all rows valid + primary_done, (c) one inactive literal (plus one per domain), (d) one discuss row.',
          '6. Source report sha256 before build, after build and after validation.', '',
          '## Text conversion conventions (column C)', '',
          '- Markdown links are replaced by their link text; targets (including `<...>` targets) are dropped. When the link text lacks the target\'s file name, it is prefixed as `<file name> — <link text>`.',
          '- A field runs from its `**Exact source locations:**` label to the next `**Label:**` line; paragraphs are joined by a space and consecutive list items by `; `. `**`, backticks and anchors are removed and whitespace collapsed. Role labels and page numbers are kept as written.', '',
          '## Results', '', '| Check | Result | Item |', '|---|---|---|']
    for r in rep['checks']:
        L.append(f"| {r['check']} | {r['status']} | {r['label']} |")
    L += ['', f"Passed {rep['passed']}, failed {rep['failed']}.", '',
          '## Pending', '', '- Desktop Excel smoke test (open, recalculate, edit, save, reopen) and visual rendering: **PENDING**. Not run; `formulas` recalculation is not a substitute for Excel.', '',
          '## Parse notes', '']
    L += [f'- {n}' for n in rep['parse_notes']] or ['- none']
    L.append('')
    (HERE / 'VALIDATION.md').write_text('\n'.join(L), encoding='utf-8')


if __name__ == '__main__':
    main()
